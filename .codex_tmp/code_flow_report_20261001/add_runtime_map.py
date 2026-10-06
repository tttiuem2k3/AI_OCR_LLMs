from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
TMP=Path(r"E:\Asoft\AI_BEM\BEM_AI_PROJECT\.codex_tmp\code_flow_report_20261001")
IMG=TMP/'images'; IMG.mkdir(parents=True,exist_ok=True)
MD=BASE/'Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md'
DOCX=BASE/'Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx'
OUT=IMG/'08_ban_do_instance_db.png'
FONT=r'C:\Windows\Fonts\times.ttf'; FONT_B=r'C:\Windows\Fonts\timesbd.ttf'
NAVY='#17365D'; BLUE='#2F75B5'; TEAL='#168A8A'; GOLD='#D6A300'; ORANGE='#ED7D31'; GREEN='#70AD47'; PURPLE='#7030A0'; LINE='#71879A'; BG='#F4F7FA'
W,H=2100,1320

def ft(s,b=False): return ImageFont.truetype(FONT_B if b else FONT,s)
def center(draw,xy,text,font,fill='#1F2933',leading=5):
    x1,y1,x2,y2=xy; width=x2-x1; lines=[]
    for raw in str(text).split('\n'):
        words=raw.split(); line=''
        if not words: lines.append(''); continue
        for word in words:
            test=(line+' '+word).strip()
            if draw.textbbox((0,0),test,font=font)[2] <= width: line=test
            else:
                if line: lines.append(line)
                line=word
        if line: lines.append(line)
    hs=[max(1,draw.textbbox((0,0),x,font=font)[3]-draw.textbbox((0,0),x,font=font)[1]) for x in lines]
    total=sum(hs)+leading*(len(lines)-1); y=y1+(y2-y1-total)/2
    for line,h in zip(lines,hs):
        bb=draw.textbbox((0,0),line,font=font); tw=bb[2]-bb[0]
        draw.text((x1+(width-tw)/2,y),line,font=font,fill=fill); y+=h+leading

def box(draw,xy,title,body,color,size=25):
    x1,y1,x2,y2=xy
    draw.rounded_rectangle(xy,radius=18,fill='white',outline=color,width=4)
    draw.rounded_rectangle((x1,y1,x2,y1+52),radius=18,fill=color,outline=color)
    draw.rectangle((x1,y1+28,x2,y1+52),fill=color)
    center(draw,(x1+8,y1+5,x2-8,y1+47),title,ft(24,True),'white')
    center(draw,(x1+14,y1+62,x2-14,y2-10),body,ft(size),'#1F2933')

def arrow(draw,a,b,color=LINE,width=6,head=20,dashed=False):
    import math
    x1,y1=a; x2,y2=b
    if dashed:
        n=10
        for i in range(n):
            if i%2==0:
                t1=i/n; t2=(i+1)/n
                draw.line((x1+(x2-x1)*t1,y1+(y2-y1)*t1,x1+(x2-x1)*t2,y1+(y2-y1)*t2),fill=color,width=width)
    else: draw.line((x1,y1,x2,y2),fill=color,width=width)
    ang=math.atan2(y2-y1,x2-x1)
    p1=(x2-head*math.cos(ang-.52),y2-head*math.sin(ang-.52)); p2=(x2-head*math.cos(ang+.52),y2-head*math.sin(ang+.52))
    draw.polygon([(x2,y2),p1,p2],fill=color)

def poly_arrow(draw,points,color=LINE,width=6,head=20,dashed=False):
    for i in range(len(points)-2):
        x1,y1=points[i]; x2,y2=points[i+1]
        if dashed:
            n=8
            for j in range(n):
                if j%2==0:
                    t1=j/n; t2=(j+1)/n
                    draw.line((x1+(x2-x1)*t1,y1+(y2-y1)*t1,x1+(x2-x1)*t2,y1+(y2-y1)*t2),fill=color,width=width)
        else:
            draw.line((x1,y1,x2,y2),fill=color,width=width)
    arrow(draw,points[-2],points[-1],color,width,head,dashed)
def label(draw,xy,text,color='#5B6573'):
    draw.rounded_rectangle(xy,radius=8,fill='white',outline='#D4DEE8',width=2)
    center(draw,(xy[0]+5,xy[1]+2,xy[2]-5,xy[3]-2),text,ft(19,True),color)

im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
d.rectangle((0,0,W,105),fill=NAVY); d.text((45,18),'Bản đồ instance và dữ liệu vận hành hiện tại',font=ft(42,True),fill='white'); d.text((47,69),'Nhìn từ lúc ERP9 phát lệnh đến khi DB ghi trạng thái và kết quả',font=ft(23),fill='#DCE6F1')
lanes=[(25,350,'WEB ERP9',BLUE),(375,675,'API ERP',TEAL),(700,1280,'API AI',GOLD),(1305,1715,'AI Python',ORANGE),(1740,2075,'Cơ sở dữ liệu',PURPLE)]
for x1,x2,name,col in lanes:
    d.rounded_rectangle((x1,125,x2,1280),radius=18,fill=BG,outline=col,width=3)
    d.rectangle((x1,125,x2,185),fill=col); center(d,(x1+5,130,x2-5,180),name,ft(28,True),'white')

# WEB
box(d,(55,235,320,390),'Giao diện','Bấm Đối chiếu AI\nhoặc xem kết quả',BLUE,24)
box(d,(55,470,320,655),'ERP Controller','Lấy phiếu + chi tiết + file\nUpload file và gửi request',BLUE,23)
arrow(d,(187,390),(187,470))

# API ERP
box(d,(400,270,650,475),'Quartz chung','Có framework lịch/job\nChưa thấy job BEM cụ thể',TEAL,23)
label(d,(405,520,645,585),'Luồng thủ công không đi qua đây','#8A3E00')

# API AI
box(d,(735,220,1000,365),'Controller','Nhận HandlerFileAsync',GOLD,23)
box(d,(1030,220,1245,365),'Orchestrator','Tạo lượt chạy\nStatus = PROCESSING',GOLD,22)
box(d,(815,450,1165,610),'Queue Singleton trong RAM','Bounded Channel\nSức chứa 200 job',GOLD,24)
box(d,(815,690,1165,855),'ReadFileWorker','HostedService lấy job\nvà gọi JobExecutor',GOLD,23)
box(d,(725,945,970,1135),'OCR Client','OcrService\nTối đa 7 file/phiếu',GOLD,23)
box(d,(1010,945,1255,1135),'LLM/Rules Client','Gửi prompt trích xuất\nvà đối chiếu',GOLD,23)
arrow(d,(1000,292),(1030,292)); arrow(d,(1137,365),(1020,450)); arrow(d,(990,610),(990,690)); arrow(d,(900,855),(850,945)); arrow(d,(1080,855),(1130,945))

# AI Python
box(d,(1340,790,1680,955),'OCR instance logic','Flask route /ocr\nOCR file → text',ORANGE,23)
box(d,(1340,1000,1680,1215),'LLM + Rule Engine','Route /llms/api/ai_llms_models\nLLM trích xuất\nRules + LLM đối chiếu',ORANGE,22)
arrow(d,(970,1035),(1340,872)); arrow(d,(1255,1035),(1340,1107))
label(d,(1335,675,1685,750),'Cùng nằm trong ứng dụng Python\nSố process/replica cần xem cấu hình chạy','#8A3E00')

# DB
box(d,(1765,220,2050,360),'BEMT2003','Tạo lượt chạy\nPROCESSING',PURPLE,23)
box(d,(1765,510,2050,650),'BEMT2002','Text OCR và dữ liệu\nnguồn/tổng hợp',PURPLE,22)
box(d,(1765,735,2050,875),'BEMT2005 / 2006','Dữ liệu được trích xuất\ntừ từng chứng từ',PURPLE,22)
box(d,(1765,960,2050,1095),'BEMT2004','Kết quả từng tiêu chí\nOK / NG + giải thích',PURPLE,22)
box(d,(1765,1150,2050,1260),'BEMT2003 cập nhật','COMPLETED hoặc FAILED\nKết quả tổng + %',PURPLE,21)
arrow(d,(1245,292),(1765,290)); arrow(d,(1680,872),(1765,580)); arrow(d,(1680,1107),(1765,805)); arrow(d,(1680,1107),(1765,1025)); arrow(d,(1907,1095),(1907,1150))

# Across zone arrows
poly_arrow(d,[(320,562),(680,562),(680,292),(735,292)]); label(d,(410,585,650,635),'POST hồ sơ sang API AI')
arrow(d,(650,375),(735,260),color='#9A6A00',dashed=True); label(d,(435,205,675,255),'Tự động: cần xác nhận job BEM')

center(d,(35,1228,1710,1280),'Queue hiện là bộ nhớ trong API-AI. DB lưu trạng thái/kết quả nhưng không phải hàng đợi job trong code đã rà.',ft(22,True),'#8A3E00')
im.save(OUT)

section='''
## 1B. Hiện trạng instance, Queue và DB

![Bản đồ instance và DB](08_ban_do_instance_db.png)

### Cái gì đang nằm ở đâu

| Thành phần anh quan tâm | Khu vực code | Hiện trạng theo source |
|---|---|---|
| Quản lý Queue | API-AI | `ChannelJobQueue` được đăng ký **Singleton**, dùng `BoundedChannel` trong RAM với sức chứa 200 job. `ReadFileOrchestratorService` ghi job vào Queue; `ReadFileWorker` lấy job ra. Đây không phải DB queue. Khi có nhiều instance API-AI, về nguyên tắc mỗi process sẽ có Queue riêng; số process triển khai thực tế cần kiểm tra cấu hình IIS/service. |
| OCR instance | AI Python | OCR thực sự chạy tại route `/ocr` trong `App/main_iis.py`. `OcrService` ở API-AI chỉ là phần gửi file/gọi OCR và nhận text. API-AI có thể xử lý tối đa 7 file song song trong một phiếu. Số process OCR đang chạy thực tế cần xem cấu hình deploy Python. |
| Rule Engine instance | AI Python | Rule Engine là hàm `process_ai_llms_models_rules` trong `Rules_AI_BEM_MEIKO.py`, được gọi từ endpoint LLM. Nó nằm trong cùng ứng dụng Python, không phải một service/queue riêng. |
| LLM instance | AI Python | Endpoint `/llms/api/ai_llms_models` nhận prompt; nhánh trích xuất và nhánh đối chiếu đều gọi model LLM. Rules quyết định trường hợp xử lý bằng logic và trường hợp cần gọi LLM. Số model worker/GPU process thực tế cần xem cấu hình chạy. |
| Worker điều phối | API-AI | `ReadFileWorker` là `HostedService`; mỗi instance API-AI đăng ký một worker. Worker lấy job từ Queue rồi gọi `ReadFileBackgroundWorkflow`. |
| DB trạng thái/kết quả | DB | DB không trực tiếp chạy OCR/LLM; DB lưu dấu vết của từng lượt chạy, dữ liệu đọc được và kết quả. Các bảng liên kết theo APK của lượt chạy `BEMT2003`. |

### DB ghi gì ở từng bước

| Thời điểm | Bảng | Dữ liệu được ghi |
|---|---|---|
| Nhận yêu cầu hợp lệ | `BEMT2003` | Tạo lượt chạy mới; `StatusProcess = PROCESSING`. Sau đó job mới được đưa vào Queue. |
| OCR xong | `BEMT2002` | Lưu text OCR và các dữ liệu nguồn/tổng hợp phục vụ đối chiếu. |
| LLM trích xuất xong | `BEMT2005`, `BEMT2006` | Lưu dữ liệu master/detail đã trích từ từng chứng từ. |
| Đối chiếu từng tiêu chí xong | `BEMT2004` | Lưu tên tiêu chí, trạng thái OK/NG, mô tả/giải thích và dữ liệu liên quan. |
| Hoàn tất toàn bộ phiếu | `BEMT2003` | Cập nhật `StatusProcess = COMPLETED`, kết quả tổng OK/NG và phần trăm. |
| Luồng bị lỗi | `BEMT2003` | Cập nhật `StatusProcess = FAILED`, kết quả tổng NG để nhận biết lượt chạy không hoàn tất. |

### Mapping gọi nhau

`WEB UI` → `ERP Controller` → `API-AI Controller` → `Orchestrator` → ghi `BEMT2003 PROCESSING` → `Queue` → `Worker` → `Workflow` → gọi `OCR Python` → lưu OCR → gọi `LLM/Rules Python` → lưu dữ liệu trích xuất và từng tiêu chí → cập nhật `BEMT2003 COMPLETED/FAILED` → ERP9 đọc DB và hiển thị.

**Điểm cần hiểu đúng:** source hiện cho thấy `/ocr` và `/llms/api/ai_llms_models` là hai route trong cùng ứng dụng Python. “OCR instance”, “Rule Engine instance” và “LLM instance” là các khối xử lý logic; số process/replica thực tế không thể kết luận chỉ từ source, cần xem cấu hình IIS/service/GPU đang chạy.

'''
md=MD.read_text(encoding='utf-8')
marker='## 2. Luồng khi người dùng đối chiếu thủ công\n'
if '## 1B. Hiện trạng instance, Queue và DB' not in md:
    md=md.replace(marker,section+marker)
    MD.write_text(md,encoding='utf-8')

# Word helpers
def font_run(r,size=10,bold=False,color=None):
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); pr.append(e)
def border(cell):
    pr=cell._tc.get_or_add_tcPr(); borders=pr.first_child_found_in('w:tcBorders')
    if borders is None: borders=OxmlElement('w:tcBorders'); pr.append(borders)
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9E2F3'); borders.append(e)
def heading(doc,text):
    p=doc.add_paragraph(); p.style='Heading 1'; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(5); r=p.add_run(text); font_run(r,17,True,'17365D'); return p
def small_heading(doc,text):
    p=doc.add_paragraph(); p.style='Heading 2'; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(4); r=p.add_run(text); font_run(r,13,True,'2F75B5'); return p
def para(doc,text,bold=False):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05; r=p.add_run(text); font_run(r,10.2,bold); return p
def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(t.rows[0].cells,headers):
        shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); font_run(r,9.5,True,'FFFFFF')
    for row in rows:
        cells=t.add_row().cells
        for idx,(c,txt) in enumerate(zip(cells,row)):
            border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1); r=p.add_run(txt); font_run(r,8.8,idx==0)
    return t

doc=Document(DOCX)
if not any(p.text=='1B. Hiện trạng instance, Queue và DB' for p in doc.paragraphs):
    target=next(p for p in doc.paragraphs if p.text=='2. Khi người dùng bấm đối chiếu thủ công')
    body=doc._body._element; anchor=target._p
    before_children=list(body)
    # Create at end, then move generated elements before section 2.
    heading(doc,'1B. Hiện trạng instance, Queue và DB')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(OUT),width=Inches(7.05))
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=cp.add_run('Hình 3. Queue, OCR, Rules, LLM và DB đang nằm ở đâu trong luồng hiện tại'); font_run(r,9,False,'5B6573')
    small_heading(doc,'Cái gì đang nằm ở đâu')
    rows=[
      ('Queue','API-AI','ChannelJobQueue là Singleton trong RAM, sức chứa 200 job. Orchestrator ghi job; Worker lấy job. Không phải DB queue.'),
      ('OCR instance logic','AI Python','Route /ocr chạy OCR. OcrService trong API-AI chỉ gửi file và nhận text; tối đa 7 file song song/phiếu.'),
      ('Rule Engine','AI Python','process_ai_llms_models_rules nằm trong ứng dụng Python; chọn nhánh trích xuất/đối chiếu và áp dụng rules.'),
      ('LLM instance logic','AI Python','Endpoint /llms/api/ai_llms_models gọi model cho trích xuất hoặc đối chiếu. Số process/GPU worker cần xem cấu hình chạy.'),
      ('Worker','API-AI','ReadFileWorker là HostedService, lấy job từ Queue và gọi workflow.'),
      ('Trạng thái/kết quả','DB','BEMT2003–2006 lưu lượt chạy, OCR, dữ liệu trích xuất và kết quả từng tiêu chí.'),
    ]
    table(doc,['Thành phần','Nằm tại','Hiện trạng'],rows)
    small_heading(doc,'DB ghi gì ở từng bước')
    rows2=[
      ('Nhận yêu cầu','BEMT2003','Tạo lượt chạy; StatusProcess = PROCESSING; sau đó mới đưa job vào Queue.'),
      ('OCR xong','BEMT2002','Text OCR và dữ liệu nguồn/tổng hợp phục vụ đối chiếu.'),
      ('Trích xuất xong','BEMT2005 / BEMT2006','Dữ liệu master/detail đã đọc từ từng chứng từ.'),
      ('Đối chiếu xong','BEMT2004','Kết quả OK/NG và giải thích của từng tiêu chí.'),
      ('Hoàn tất','BEMT2003','COMPLETED, kết quả tổng OK/NG và phần trăm.'),
      ('Bị lỗi','BEMT2003','FAILED và kết quả tổng NG để nhận biết lượt chạy không hoàn tất.'),
    ]
    table(doc,['Thời điểm','Bảng','Dữ liệu lưu'],rows2)
    small_heading(doc,'Mapping gọi nhau')
    para(doc,'WEB UI → ERP Controller → API-AI Controller → Orchestrator → BEMT2003 PROCESSING → Queue → Worker → Workflow → OCR Python → LLM/Rules Python → BEMT2002/2004/2005/2006 → BEMT2003 COMPLETED hoặc FAILED → ERP9 hiển thị.',True)
    para(doc,'Lưu ý: source xác nhận cấu trúc logic nhưng không cho biết chắc môi trường đang chạy bao nhiêu process/replica OCR, LLM hoặc API-AI. Muốn chốt số instance thực tế cần kiểm tra cấu hình IIS/service và tiến trình GPU đang chạy.')
    after_children=list(body); new_elems=[e for e in after_children if e not in before_children]
    for e in new_elems: body.remove(e)
    idx=body.index(anchor)
    for e in new_elems:
        body.insert(idx,e); idx+=1

doc.save(DOCX)
print(OUT)
print(MD)
print(DOCX)

