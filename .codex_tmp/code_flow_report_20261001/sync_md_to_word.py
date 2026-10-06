from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
MD_PATH = BASE / "Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
DOCX_PATH = BASE / "Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"
IMG_DIR = BASE / "Temp" / "co_che_code_van_hanh_02102026"

IMAGE_MAP = {
    "1.": (IMG_DIR / "01_cau_truc_ma_nguon.png", "Hình 1. Trách nhiệm của ERP9, API-AI, AI Python và Database"),
    "2.": (IMG_DIR / "02_vi_du_phieu_10_file.png", "Hình 2. Ví dụ một phiếu NVL có 10 file đi qua toàn bộ luồng xử lý"),
    "3.": (IMG_DIR / "03_du_lieu_qua_tung_khu_vuc.png", "Hình 3. Dữ liệu vào, xử lý và kết quả tại từng khu vực"),
    "4.": (IMG_DIR / "04_database_luu_du_lieu.png", "Hình 4. Các lớp dữ liệu được lưu trong Database"),
    "6.": (IMG_DIR / "05_rui_ro_task_15_10.png", "Hình 5. Rủi ro chính và hướng xử lý để đạt mục tiêu 15/10"),
}

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F4F7"
ORANGE = "F4B183"
GREEN = "E2F0D9"
DARK = "1F2933"
WHITE = "FFFFFF"
BORDER = "AAB7C4"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=65, start=70, bottom=65, end=70) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, name="Times New Roman", size=9.2, bold=None, color=DARK, italic=None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str, size=9.2, color=DARK, bold=False) -> None:
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(7.0, size - 0.5), color="244062", bold=False)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color, bold=bold)


def set_paragraph_layout(paragraph, before=0, after=3, line=1.05, keep_with_next=False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_with_next


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[f"Heading {level}"]
    set_paragraph_layout(paragraph, before=5 if level == 2 else 3, after=4, line=1.0, keep_with_next=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=15 if level == 1 else 11.5, bold=True, color=NAVY if level == 1 else BLUE)


def add_body(doc: Document, text: str, kind="plain") -> None:
    paragraph = doc.add_paragraph()
    if kind == "bullet":
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        add_inline(paragraph, "• " + text, 9.1)
    elif kind == "number":
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        add_inline(paragraph, text, 9.1)
    elif kind == "callout":
        paragraph.paragraph_format.left_indent = Inches(0.14)
        paragraph.paragraph_format.right_indent = Inches(0.14)
        set_cell_like_paragraph(paragraph, GREEN)
        add_inline(paragraph, text, 9.3, color="385723", bold=True)
    else:
        add_inline(paragraph, text, 9.2)
    set_paragraph_layout(paragraph, after=2.5, line=1.08)


def set_cell_like_paragraph(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "70AD47")
    borders.append(bottom)
    p_pr.append(borders)


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def normalize_row(headers: list[str], cells: list[str]) -> list[str]:
    count = len(headers)
    if len(cells) == count:
        return cells
    if headers and headers[0] == "Khu vực":
        if cells and cells[0] == "ERP9 WEB" and len(cells) >= 6:
            return [cells[0], " | ".join(cells[1:4]), cells[4], " | ".join(cells[5:])]
        if cells and cells[0] == "Rules + LLM đối chiếu" and len(cells) >= 7:
            return [cells[0], " | ".join(cells[1:5]), cells[5], " | ".join(cells[6:])]
    if len(cells) > count:
        return cells[: count - 1] + [" | ".join(cells[count - 1 :])]
    return cells + [""] * (count - len(cells))


def suggested_widths(headers: list[str]) -> list[float]:
    count = len(headers)
    first = headers[0] if headers else ""
    if first == "STT":
        return [0.45, 2.8, 1.75, 5.0]
    if first == "Khu vực":
        return [1.35, 3.15, 3.1, 3.25]
    if first == "Bảng":
        return [1.0, 1.85, 4.2, 3.8]
    if first == "Bước":
        return [1.45, 1.1, 8.3]
    if first == "Nội dung":
        return [1.7, 1.5, 7.65]
    if first == "Trường hợp":
        return [2.15, 4.15, 4.55]
    if first == "Owner":
        return [1.05, 1.45, 5.75, 1.0, 1.0]
    if first == "Thành phần":
        return [1.05, 1.55, 1.2, 1.8, 2.85, 1.15, 0.8]
    return [10.85 / max(1, count)] * count


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    widths = suggested_widths(headers)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    keep_row_together(header_row)
    header_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell, 70, 55, 70, 55)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_layout(paragraph, after=0, line=1.0)
        add_inline(paragraph, header, 7.2 if len(headers) <= 5 else 6.5, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_shading(cell, WHITE if row_index % 2 == 0 else LIGHT_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell, 60, 55, 60, 55)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 and len(headers) <= 5 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_layout(paragraph, after=0, line=1.0)
            size = 7.1 if len(headers) <= 4 else (6.7 if len(headers) == 5 else 6.1)
            add_inline(paragraph, value, size=size)
    paragraph = doc.add_paragraph()
    set_paragraph_layout(paragraph, after=1)


def add_picture(doc: Document, image_path: Path, caption: str) -> None:
    if not image_path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_layout(paragraph, before=1, after=2, line=1.0)
    paragraph.add_run().add_picture(str(image_path), width=Inches(10.6))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_layout(cap, after=4, line=1.0)
    run = cap.add_run(caption)
    set_run_font(run, size=8.2, italic=True, color="5B6573")


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=8, color="667085")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.15)
    section.footer_distance = Inches(0.18)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9.2)


def parse_and_build(md_text: str) -> Document:
    lines = md_text.splitlines()
    doc = Document()
    configure_document(doc)
    title_done = False
    current_h2 = 0
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            title = line[2:].strip()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_layout(paragraph, before=80, after=10, line=1.0, keep_with_next=True)
            run = paragraph.add_run(title.upper())
            set_run_font(run, size=24, bold=True, color=NAVY)
            sub = doc.add_paragraph()
            sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_layout(sub, after=8, line=1.0)
            run = sub.add_run("Báo cáo hiện trạng mã nguồn, luồng dữ liệu và phương án ổn định vận hành")
            set_run_font(run, size=12.5, bold=True, color=BLUE)
            date_p = doc.add_paragraph()
            date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_layout(date_p, after=10)
            run = date_p.add_run("Cập nhật ngày 02/10/2026")
            set_run_font(run, size=10.5, color="667085")
            title_done = True
            index += 1
            continue
        if line.startswith("## "):
            if title_done and current_h2 > 0:
                doc.add_page_break()
            current_h2 += 1
            heading = line[3:].strip()
            add_heading(doc, heading, 1)
            for prefix, (image, caption) in IMAGE_MAP.items():
                if heading.startswith(prefix):
                    add_picture(doc, image, caption)
                    break
            index += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\|\s*:?-+", lines[index + 1].strip()):
            headers = split_row(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(normalize_row(headers, split_row(lines[index])))
                index += 1
            add_table(doc, headers, rows)
            continue
        number_match = re.match(r"^(\d+\.\s+)(.+)$", line)
        if number_match:
            add_body(doc, line, "number")
        elif line.startswith("- "):
            add_body(doc, line[2:].strip(), "bullet")
        elif line.startswith("=>") or line.startswith("\=>"):
            add_body(doc, line.lstrip("\\"), "callout")
        else:
            add_body(doc, line, "plain")
        index += 1
    return doc


def main() -> None:
    md_text = MD_PATH.read_text(encoding="utf-8")
    document = parse_and_build(md_text)
    document.core_properties.title = "Cơ chế vận hành đối chiếu AI BEM"
    document.core_properties.subject = "Đồng bộ từ nội dung MD ngày 02/10/2026"
    document.core_properties.author = "ASOFT"
    document.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
