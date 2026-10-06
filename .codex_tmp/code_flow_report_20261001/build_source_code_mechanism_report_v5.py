from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMAGE_DIR = BASE / "Temp" / "co_che_code_van_hanh_02102026"
IMAGE_DIR.mkdir(parents=True, exist_ok=True)
MD_PATH = BASE / "Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
DOCX_PATH = BASE / "Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"

FONT = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"
NAVY = "#17365D"
ERP = "#2F75B5"
API = "#D89500"
PY = "#C65911"
DB = "#7030A0"
GREEN = "#548235"
RED = "#C00000"
GRAY = "#667085"
DARK = "#1F2933"
BG = "#F5F7FA"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def text_size(draw, value, font_value):
    box = draw.textbbox((0, 0), str(value), font=font_value)
    return box[2] - box[0], box[3] - box[1]


def wrap(draw, value, font_value, max_width):
    lines = []
    for paragraph in str(value).split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if not current or text_size(draw, candidate, font_value)[0] <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines


def draw_text_box(draw, rect, value, size=24, bold=False, fill=DARK, align="center", valign="middle", gap=6):
    x1, y1, x2, y2 = rect
    fnt = font(size, bold)
    lines = wrap(draw, value, fnt, max(20, x2 - x1 - 28))
    heights = [max(1, text_size(draw, line, fnt)[1]) for line in lines]
    total_h = sum(heights) + gap * max(0, len(lines) - 1)
    if valign == "top":
        y = y1 + 14
    else:
        y = y1 + (y2 - y1 - total_h) / 2
    for line, height in zip(lines, heights):
        width = text_size(draw, line, fnt)[0]
        if align == "left":
            x = x1 + 18
        elif align == "right":
            x = x2 - width - 18
        else:
            x = x1 + (x2 - x1 - width) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += height + gap


def canvas(title, subtitle, width=3200, height=1800):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 150), fill=NAVY)
    draw.text((60, 28), title, font=font(48, True), fill="white")
    draw.text((62, 92), subtitle, font=font(25), fill="#DDE7F3")
    return image, draw


def rounded(draw, rect, fill="white", outline="#CDD5DF", width=3, radius=28):
    draw.rounded_rectangle(rect, radius=radius, fill=fill, outline=outline, width=width)

def arrow(draw, start, end, color=NAVY, width=10, head=24):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    if x1 == x2:
        direction = 1 if y2 > y1 else -1
        draw.polygon(
            [
                (x2, y2),
                (x2 - head, y2 - direction * head * 1.5),
                (x2 + head, y2 - direction * head * 1.5),
            ],
            fill=color,
        )
    elif y1 == y2:
        direction = 1 if x2 > x1 else -1
        draw.polygon(
            [
                (x2, y2),
                (x2 - direction * head * 1.5, y2 - head),
                (x2 - direction * head * 1.5, y2 + head),
            ],
            fill=color,
        )


def card(draw, rect, title, body, color, title_size=25, body_size=22, number=None):
    x1, y1, x2, y2 = rect
    rounded(draw, rect, "white", color, 4, 26)
    draw.rounded_rectangle((x1, y1, x2, y1 + 78), radius=26, fill=color, outline=color)
    draw.rectangle((x1, y1 + 44, x2, y1 + 78), fill=color)
    if number is not None:
        draw.ellipse((x1 + 18, y1 + 14, x1 + 62, y1 + 58), fill="white")
        draw_text_box(draw, (x1 + 18, y1 + 14, x1 + 62, y1 + 58), str(number), 20, True, color)
        title_rect = (x1 + 72, y1 + 8, x2 - 18, y1 + 68)
    else:
        title_rect = (x1 + 16, y1 + 8, x2 - 16, y1 + 68)
    draw_text_box(draw, title_rect, title, title_size, True, "white")
    draw_text_box(draw, (x1 + 20, y1 + 92, x2 - 20, y2 - 20), body, body_size, False, DARK, "center", "middle", 7)

def compact_card(draw, rect, title, body, color, title_size=28, body_size=23, number=None):
    x1, y1, x2, y2 = rect
    rounded(draw, rect, "white", color, 4, 24)
    draw.rounded_rectangle((x1, y1, x2, y1 + 72), radius=24, fill=color, outline=color)
    draw.rectangle((x1, y1 + 40, x2, y1 + 72), fill=color)
    if number is not None:
        draw.ellipse((x1 + 18, y1 + 14, x1 + 58, y1 + 54), fill="white")
        draw_text_box(draw, (x1 + 18, y1 + 14, x1 + 58, y1 + 54), str(number), 18, True, color)
        title_rect = (x1 + 68, y1 + 7, x2 - 18, y1 + 65)
    else:
        title_rect = (x1 + 16, y1 + 7, x2 - 16, y1 + 65)
    draw_text_box(draw, title_rect, title, title_size, True, "white")
    draw_text_box(draw, (x1 + 24, y1 + 96, x2 - 24, y2 - 26), body, body_size, False, DARK, "left", "top", 7)


def pill(draw, rect, text, color, size=22):
    rounded(draw, rect, "#FFFFFF", color, 3, 22)
    draw_text_box(draw, rect, text, size, True, color)


def table_header(draw, x, y, widths, headers, colors, h=72):
    cur = x
    for w, header, color in zip(widths, headers, colors):
        draw.rounded_rectangle((cur, y, cur + w, y + h), radius=18, fill=color, outline=color)
        draw.rectangle((cur, y + 35, cur + w, y + h), fill=color)
        draw_text_box(draw, (cur + 8, y + 6, cur + w - 8, y + h - 6), header, 24, True, "white")
        cur += w + 18


def row_card(draw, x, y, widths, values, colors, h=142, sizes=None):
    cur = x
    if sizes is None:
        sizes = [20] * len(values)
    for w, value, color, size in zip(widths, values, colors, sizes):
        rounded(draw, (cur, y, cur + w, y + h), "white", color, 3, 18)
        draw_text_box(draw, (cur + 10, y + 10, cur + w - 10, y + h - 10), value, size, False, DARK, "center", "middle", 5)
        cur += w + 18


def make_source_structure():
    image, draw = canvas(
        "01. CẤU TRÚC MÃ NGUỒN VÀ TRÁCH NHIỆM",
        "Không hiển thị đường dẫn source; chỉ thể hiện mỗi phần làm gì trong luồng đối chiếu AI"
    )
    compact_card(draw, (100, 230, 960, 610), "ERP9 - WEB", "- Nhận thao tác đối chiếu thủ công/lịch tự động.\n- Lấy master, detail và file đính kèm.\n- Gửi request sang API-AI.\n- Đọc DB để hiển thị kết quả.", ERP, 29, 24, 1)
    compact_card(draw, (1170, 230, 2030, 610), "API-AI", "- Nhận ReadFileRequest từ ERP9.\n- Tạo BEMT2003 = PROCESSING.\n- Đưa ReadFileJob vào Queue.\n- Điều phối OCR, LLM, Rules và ghi DB.", API, 29, 24, 2)
    compact_card(draw, (2240, 230, 3100, 610), "AI - PYTHON", "- OCR đọc file và trả text.\n- LLM trích xuất OCR thành JSON.\n- Rule Engine chọn chứng từ cần dùng.\n- LLM đối chiếu và trả OK/NG.", PY, 29, 24, 3)

    arrow(draw, (960, 420), (1170, 420), NAVY, 8, 22)
    arrow(draw, (2030, 420), (2240, 420), NAVY, 8, 22)
    draw_text_box(draw, (100, 680, 3100, 735), "ERP9 API: trong luồng đã rà, chưa thấy là bước trung gian của đối chiếu BEM-AI; WEB ERP9 gọi trực tiếp API-AI.", 25, True, ERP)

    draw_text_box(draw, (130, 835, 3070, 890), "Database là nơi các phần trên cùng đọc/ghi dữ liệu vận hành", 30, True, DB)
    table_header(draw, 145, 960, [550, 550, 550, 550, 550], ["BEMT2003", "BEMT2002", "BEMT2005", "BEMT2006", "BEMT2004"], [DB, DB, DB, DB, DB], 74)
    row_card(draw, 145, 1060, [550, 550, 550, 550, 550], ["Lượt chạy, trạng thái, kết quả tổng", "Tên file, mã file, text OCR", "Loại/header chứng từ", "FieldID01...FieldID100", "OK/NG từng tiêu chí"], [DB, DB, DB, DB, DB], 128, [20, 20, 20, 20, 20])

    rounded(draw, (330, 1320, 2870, 1580), "#F7FAFC", "#CBD5E1", 3, 28)
    draw_text_box(draw, (370, 1350, 2830, 1550), "Cách đọc sơ đồ: ERP9 chỉ gom hồ sơ và hiển thị kết quả. API-AI là lớp điều phối job/queue và ghi DB. AI Python là nơi chạy OCR, LLM trích xuất và Rules/LLM đối chiếu. DB giữ trạng thái, dữ liệu OCR, dữ liệu đã trích xuất và kết quả từng tiêu chí.", 27, False, DARK)
    return image


def make_manual_flow():
    image, draw = canvas(
        "02. VÍ DỤ ĐỐI CHIẾU THỦ CÔNG 1 PHIẾU NVL CÓ 10 FILE",
        "Đọc theo số thứ tự: thành phần xử lý, dữ liệu đi qua và kết quả đạt được"
    )
    # file summary
    draw_text_box(draw, (90, 190, 3110, 245), "Phiếu ví dụ: NVL/09/2026/0011 | 10 file: PO, Ringi, 2 Commercial Invoice, 2 tờ khai, 2 Packing List, Statement, Bill | Tổng yêu cầu: 3,250,000 JPY", 27, True, NAVY)
    headers = ["ERP9\nkhông đổi", "API-AI\nđiều phối", "Python AI\nOCR/LLM/Rules", "Database\nlưu trạng thái & kết quả"]
    widths = [650, 760, 910, 650]
    colors = [ERP, API, PY, DB]
    table_header(draw, 90, 300, widths, headers, colors, 78)
    x = [90, 90 + 650 + 18, 90 + 650 + 18 + 760 + 18, 90 + 650 + 18 + 760 + 18 + 910 + 18]
    h = 142
    steps = [
        (1, 0, 420, "Người dùng bấm\nĐối chiếu AI", "ERP9 lấy phiếu ĐNTT, 2 dòng chi tiết và 10 file đính kèm."),
        (2, 1, 420, "API-AI nhận hồ sơ", "Nhận thông tin phiếu, dòng ĐNTT và danh sách 10 file từ ERP9."),
        (3, 1, 610, "Tạo lượt chạy", "Tạo BEMT2003 = PROCESSING để nhận diện lần xử lý này."),
        (4, 1, 800, "Queue lấy job", "Job nằm trong Queue RAM; worker đơn lấy job theo thứ tự."),
        (5, 2, 990, "OCR 10 file", "Gọi /ocr một lần với 10 file; Python trả text OCR theo từng file."),
        (6, 3, 990, "Lưu OCR", "BEMT2002 lưu FileName, AttachID, APK_File và RawContent."),
        (7, 2, 1180, "LLM trích xuất", "Mỗi file gọi prompt đọc file; trả loại chứng từ và JSON dữ liệu."),
        (8, 3, 1180, "Lưu dữ liệu chứng từ", "BEMT2005 lưu header/loại chứng từ; BEMT2006 lưu từng trường trích xuất. Tên file gốc vẫn ở BEMT2002."),
        (9, 2, 1370, "Rules + LLM đối chiếu", "Nhận dữ liệu phiếu + dòng ĐNTT + dataFiles; chạy 9 tiêu chí NVL."),
        (10, 3, 1370, "Kết quả cuối", "BEMT2004 lưu OK/NG từng tiêu chí; BEMT2003 = COMPLETED. ERP9 đọc DB để hiển thị."),
    ]
    for idx, col, y, title, body in steps:
        rect = (x[col], y, x[col] + widths[col], y + h)
        rounded(draw, rect, "white", colors[col], 3, 20)
        draw.ellipse((rect[0] + 14, rect[1] + 16, rect[0] + 58, rect[1] + 60), fill=colors[col])
        draw_text_box(draw, (rect[0] + 14, rect[1] + 16, rect[0] + 58, rect[1] + 60), str(idx), 19, True, "white")
        draw_text_box(draw, (rect[0] + 70, rect[1] + 14, rect[2] - 12, rect[1] + 56), title, 22, True, colors[col])
        draw_text_box(draw, (rect[0] + 18, rect[1] + 62, rect[2] - 18, rect[3] - 12), body, 20, False, DARK, "center", "middle", 5)
    # result box
    rounded(draw, (420, 1580, 2780, 1750), "#EEF7EE", GREEN, 4, 24)
    draw_text_box(draw, (440, 1598, 2760, 1732), "Kết quả ví dụ: tiêu chí Số tiền OK vì 2 Commercial Invoice = 1,500,000 + 1,750,000 JPY khớp với 2 tờ khai và Ringi. STATEMENT chỉ kiểm tra thêm, không cộng trùng. Nếu có tiêu chí NG, DB lưu tiêu chí NG + file/lý do để ERP9 hiển thị.", 24, True, GREEN)
    return image


def make_data_contract_flow():
    image, draw = canvas(
        "03. DỮ LIỆU ĐI QUA TỪ ERP9 ĐẾN KẾT QUẢ AI",
        "Mỗi bước thể hiện rõ dữ liệu nhận vào, xử lý thực tế và dữ liệu trả ra"
    )
    headers = ["Khu vực", "Dữ liệu nhận vào", "Xử lý chính", "Dữ liệu trả ra / lưu lại"]
    widths = [390, 840, 850, 840]
    table_header(draw, 70, 205, widths, headers, [NAVY, NAVY, NAVY, NAVY], 74)
    rows = [
        (ERP, "ERP9 WEB", "Phiếu master: VoucherNo, loại DNTT, NCC, Currency, Deadline...\nDòng chi tiết: Description, InvoiceNo, RequestAmount, InvoiceDate, RingiNo.\nFile: APK, AttachID, AttachName, AttachURL.", "Lấy master + detail + file.\nThủ công: 1 phiếu.\nTự động: nhiều phiếu, giới hạn 5 yêu cầu ERP gửi đồng thời.", "ReadFileRequest JSON: UserId/UserName, BEMF2000ViewModel, BEMF2001ViewModels, AttachFiles, dữ liệu thông báo OOT9002/OOT9003."),
        (API, "API-AI nhận request", "ReadFileRequest và loại DNTT PaymentRequestTypeID.", "Validate; xóa dữ liệu lần chạy trước; lấy danh sách prompt; tạo BEMT2003 = PROCESSING.", "Trả ERP9 ngay: đã nhận yêu cầu, đang xử lý nền.\nTạo ReadFileJob gồm BEMT2003APK + Request + PromptContents."),
        (API, "Queue + Worker", "ReadFileJob: JobType, BEMT2003APK, Request, PromptContents.", "Queue RAM tối đa 200 job; 1 reader/worker lấy từng job theo thứ tự và mở workflow xử lý nền.", "Một job được chuyển sang OCR. Nếu service restart, job đang chờ trong RAM có rủi ro mất."),
        (PY, "OCR", "Danh sách AttachFiles với FilePath, FileName, AttachID, APK_File.", "Nhận dạng MIME; đọc tối đa 7 file song song. Ảnh/PDF gọi OCR Python; Word/Excel đọc bằng bộ xử lý tương ứng. Timeout gọi OCR: 10 phút.", "TextMerged và từng ResultReadFileModel: FileName, FilePath, TextContent, NumberOrder, AttachID, APK_File, HasErrorReadFile."),
        (PY, "LLM trích xuất", "Prompt BEM_AGENT_READFILE + FileName + result = text OCR của từng file.", "LLM nhận diện loại chứng từ và chuẩn hóa text OCR thành JSON sections; mapping trường động theo ONT1041.", "aiSectionCompares/dataFiles; BEMT2005 lưu header/loại chứng từ; BEMT2006 lưu các trường đã trích xuất; BEMT2002 lưu OCR."),
        (PY, "Rules + LLM đối chiếu", "datas = dữ liệu master; details = Description, InvoiceNo, RequestAmount, InvoiceDate, RingiNo; dataFiles = dữ liệu trích xuất; prompt từng tiêu chí.", "Sắp xếp dataFiles theo SectionType; chạy lần lượt các prompt tiêu chí NVL; Rule Engine chọn/chuẩn hóa chứng từ, LLM kết luận.", "Mỗi tiêu chí trả CriteriaName, CriteriaStatus, Description, FileName, DueDateAI; lưu BEMT2004 và dữ liệu tổng hợp prompt vào BEMT2002."),
        (DB, "DB + ERP9 hiển thị", "Danh sách BEMT2004 của tất cả tiêu chí.", "Tính số tiêu chí đạt; tạo Percentage; nếu có tiêu chí khác OK thì Status = NG; cập nhật COMPLETED hoặc FAILED; gửi thông báo.", "BEMT2003 lưu Status, Percentage, TextConditionFail, TextContentOCR. ERP9 đọc DB và hiển thị kết quả/lý do cho người dùng."),
    ]
    y = 305
    for color, component, input_data, processing, output_data in rows:
        row_card(draw, 70, y, widths, [component, input_data, processing, output_data], [color, color, color, color], 186, [20, 18, 18, 18])
        y += 207
    return image

def make_responsibility_db_state():
    image, draw = canvas(
        "04. DATABASE LƯU DỮ LIỆU ĐỐI CHIẾU NHƯ THẾ NÀO",
        "Một lượt chạy, file OCR, loại chứng từ, trường trích xuất và kết quả tiêu chí được lưu tách riêng"
    )

    card(
        draw,
        (350, 200, 2850, 480),
        "BEMT2003 | 1 LƯỢT CHẠY AI CỦA PHIẾU",
        "APK_BEMT2000: phiếu ĐNTT | StatusProcess: PROCESSING / COMPLETED / FAILED | Status: OK / NG | Percentage: tỷ lệ tiêu chí đạt | TextConditionFail: lỗi tổng | CreateDate / LastModifyDate: thời điểm chạy",
        NAVY,
        28,
        22,
    )

    draw.line((1600, 480, 1600, 555), fill=NAVY, width=9)
    draw.line((540, 555, 2670, 555), fill=NAVY, width=9)
    arrow(draw, (540, 555), (540, 650), NAVY, 9, 22)
    arrow(draw, (1600, 555), (1600, 650), NAVY, 9, 22)
    arrow(draw, (2670, 555), (2670, 650), NAVY, 9, 22)

    card(
        draw,
        (100, 650, 980, 1420),
        "BEMT2002 | 1 DÒNG / FILE OCR",
        "Lưu danh tính file và nội dung OCR.\n\n- APK_BEMT2003: thuộc lượt chạy nào\n- FileName: tên file gốc\n- AttachID / APK_File: mã file đính kèm\n- DataSourceType = OCR\n- RawContent: text OCR của file\n- ConsolidatedContent: nội dung đã tổng hợp khi cần\n\nVí dụ: CI_SSK-MV-2026-02-001.pdf + text OCR Invoice 1,500,000 JPY.",
        DB,
        25,
        20,
    )
    card(
        draw,
        (1160, 650, 2040, 1110),
        "BEMT2005 | HEADER / LOẠI CHỨNG TỪ",
        "- APK_BEMT2003: thuộc lượt chạy\n- SectionType: COMMERCIALINVOICE / PO / RINGI...\n- SectionOrder / SectionTitle\n- TotalAmount / TotalCurrency\n- Signature / PromptSystem\n\nKhông có cột FileName cố định; tên file gốc vẫn nằm ở BEMT2002.",
        DB,
        25,
        20,
    )
    arrow(draw, (1600, 1110), (1600, 1210), DB, 9, 22)
    card(
        draw,
        (1160, 1210, 2040, 1650),
        "BEMT2006 | CÁC TRƯỜNG ĐÃ TRÍCH XUẤT",
        "- APK_BEMT2005: thuộc section nào\n- OrderNo / Description\n- FieldID01 ... FieldID100: giá trị theo mapping ONT1041\n\nVí dụ mapping lưu: InvoiceNo, InvoiceDate, SupplierName, Amount, Currency, Incoterm...\nFileName chỉ nằm trong FieldID khi ONT1041 có cấu hình map FileName.",
        DB,
        25,
        20,
    )
    card(
        draw,
        (2220, 650, 3100, 1420),
        "BEMT2004 | KẾT QUẢ TỪNG TIÊU CHÍ",
        "Lưu kết quả sau khi Rules + LLM đối chiếu.\n\n- APK_BEMT2003: thuộc lượt chạy\n- CriteriaID / CriteriaName\n- CriteriaStatus: OK / NG\n- Description: kết luận hoặc lý do NG\n- FileName: file liên quan đến tiêu chí\n- DueDateAI / PromptSystem\n\nVí dụ: Số tiền = OK vì 2 Invoice cộng đúng 3,250,000 JPY.",
        DB,
        25,
        20,
    )

    rounded(draw, (100, 1690, 3100, 1770), "#FFF7E6", API, 3, 20)
    draw_text_box(
        draw,
        (120, 1695, 3080, 1765),
        "Lưu ý truy vết: code dùng OCR của từng file để tạo BEMT2005/BEMT2006, nhưng BEMT2005 hiện không có khóa trực tiếp tới BEMT2002. Muốn biết chắc section/field sinh từ file nào nên bổ sung APK_BEMT2002 hoặc FileID/FileHash vào dữ liệu trích xuất.",
        22,
        True,
        API,
    )
    return image


def make_risk_task():
    image, draw = canvas(
        "06. RỦI RO THÁNG 09 VÀ PHƯƠNG ÁN CHỐT 15/10",
        "Tập trung sửa API-AI, Python AI và DB; không mở rộng ngoài mục tiêu vận hành"
    )
    headers = ["Rủi ro chính", "Tác động", "Phương án chốt", "Owner", "Deadline"]
    widths = [620, 610, 910, 360, 300]
    colors = [RED, API, GREEN, NAVY, DB]
    table_header(draw, 80, 230, widths, headers, colors, 74)
    rows = [
        ("Queue RAM, 1 worker tuần tự", "Dễ backlog; restart có rủi ro mất job chờ.", "Persistent Queue DB + JobID + trạng thái + ưu tiên.", "API/DB", "06/10"),
        ("PROCESSING quá chung", "Không biết phiếu đang OCR, LLM hay Rule.", "Thêm StepLog: queued, OCR, extract, rule, done, failed.", "API/DB", "08/10"),
        ("File đổi vẫn có thể chạy lại cả bộ", "Tốn OCR/LLM, chậm khi hồ sơ nhiều file.", "FileID + Hash + Version; file không đổi thì dùng lại OCR/extract.", "API/DB", "10/10"),
        ("Lỗi OCR/LLM chưa retry chuẩn", "Một lỗi tạm thời làm cả phiếu fail hoặc khó xử lý lại.", "Retry giới hạn theo bước; báo rõ file lỗi, bước lỗi, lý do lỗi.", "API/Python", "10/10"),
        ("LLM làm cả logic xác định", "Tốn GPU, kết quả thiếu ổn định.", "Rule Engine xử lý logic rõ; LLM chỉ xử lý phần cần suy luận.", "Python", "12/10"),
        ("Thiếu dashboard vận hành", "Khó biết backlog, p95, số job fail, retry, phiếu gần hạn.", "Dashboard Queue/OCR/LLM/Rule + cảnh báo kỳ thanh toán.", "API/DB", "15/10"),
    ]
    y = 330
    for row in rows:
        row_card(draw, 80, y, widths, row, colors, 150, [20, 20, 20, 20, 20])
        y += 172
    rounded(draw, (140, 1460, 3060, 1710), "#F4F9F4", GREEN, 4, 28)
    draw_text_box(draw, (170, 1480, 3030, 1690), "Baseline đề xuất: ERP9 giữ nguyên vai trò nhận thao tác/hiển thị kết quả. API-AI chịu trách nhiệm điều phối job và ghi DB. Python chịu trách nhiệm OCR, LLM, Rule Engine. DB là nơi lưu trạng thái, dữ liệu, version, lỗi và thời gian để vận hành ổn định trước 15/10.", 29, True, GREEN)
    return image


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def write_cell(cell, value, size=8.0, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(value))
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color="000000", after=3):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, text):
    return add_text(doc, text, 14, True, WD_ALIGN_PARAGRAPH.LEFT, "17365D", 4)


def picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(10.7))
    add_text(doc, caption, 8.5, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 2)


def add_table(doc, title, headers, rows, widths=None, font_size=7.2):
    add_text(doc, title, 10.2, True, WD_ALIGN_PARAGRAPH.LEFT, "17365D", 2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "17365D")
        write_cell(cell, header, 7.8, True, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_shading(cell, "F3F7FC" if index % 2 == 0 else "FFFFFF")
            write_cell(cell, value, font_size, False, "000000", WD_ALIGN_PARAGRAPH.LEFT)
    return table


EXAMPLE_FILES = [
    ("1", "PO_PO-2026-0901.pdf", "PO", "PO, NCC, tiền tệ, điều kiện giao hàng, điều khoản thanh toán"),
    ("2", "RINGI_RG-2026-088.pdf", "RINGI", "Ringi, NCC, tổng tiền phê duyệt, trạng thái duyệt"),
    ("3", "CI_SSK-MV-2026-02-001.pdf", "COMMERCIALINVOICE", "Invoice 1,500,000 JPY, ngày hóa đơn, NCC, Incoterm"),
    ("4", "CI_SSK-MV-2026-02-002.pdf", "COMMERCIALINVOICE", "Invoice 1,750,000 JPY, ngày hóa đơn, NCC, Incoterm"),
    ("5", "TK_107333888810.pdf", "CUSTOMSHEET", "Tờ khai gắn invoice 001, tiền, ngày hoàn thành kiểm tra"),
    ("6", "TK_107333888811.pdf", "CUSTOMSHEET", "Tờ khai gắn invoice 002, tiền, ngày hoàn thành kiểm tra"),
    ("7", "PACKINGLIST_SSK-MV-2026-02-001.pdf", "PACKINGLIST", "Dữ liệu đóng gói, không phải tiêu chí chính NVL hiện tại"),
    ("8", "PACKINGLIST_SSK-MV-2026-02-002.pdf", "PACKINGLIST", "Dữ liệu đóng gói, không phải tiêu chí chính NVL hiện tại"),
    ("9", "STATEMENT_09-2026.pdf", "STATEMENT", "Bảng kê 2 invoice, dùng kiểm tra thêm, không cộng trùng"),
    ("10", "BL_ABC123456.pdf", "BILL", "Thông tin vận chuyển, hiện không phải chứng từ bắt buộc của 9 prompt NVL"),
]

DATA_FLOW_ROWS = [
    ("ERP9 WEB", "Master: VoucherNo, loại DNTT, NCC, Currency, Deadline... | Detail: Description, InvoiceNo, RequestAmount, InvoiceDate, RingiNo | File: APK, AttachID, AttachName, AttachURL", "Gom dữ liệu phiếu. Thủ công gửi 1 phiếu; tự động có SemaphoreSlim(5), tối đa 5 request ERP gửi đồng thời.", "ReadFileRequest gồm UserId/UserName, master, detail, AttachFiles, OOT9002/OOT9003."),
    ("API-AI nhận request", "ReadFileRequest + PaymentRequestTypeID", "Validate; xóa dữ liệu đối chiếu trước; đọc prompt theo loại DNTT; tạo BEMT2003 = PROCESSING.", "Trả ERP9 ngay: đã nhận, xử lý nền. Tạo ReadFileJob: JobType, BEMT2003APK, Request, PromptContents."),
    ("Queue + Worker", "ReadFileJob", "Queue RAM capacity 200, SingleReader = true; worker lấy lần lượt từng job để chạy workflow.", "Job được chuyển OCR. Job chỉ nằm trong RAM; restart có nguy cơ mất job chờ."),
    ("OCR", "AttachFiles: đường dẫn file, tên file, AttachID, APK file", "Nhận dạng MIME. Tối đa 7 file được đọc song song; ảnh/PDF gọi OCR Python; timeout gọi OCR 10 phút.", "TextMerged + ResultReadFileModel mỗi file: FileName, FilePath, TextContent, NumberOrder, AttachID, APK_File, HasErrorReadFile."),
    ("LLM trích xuất", "BEM_AGENT_READFILE + FileName + result = text OCR từng file", "LLM nhận diện loại chứng từ và trả JSON sections; mapping trường động theo ONT1041.", "dataFiles/aiSectionCompares; BEMT2002 lưu OCR; BEMT2005 lưu header/loại chứng từ; BEMT2006 lưu trường trích xuất."),
    ("Rules + LLM đối chiếu", "datas = master | details = Description, InvoiceNo, RequestAmount, InvoiceDate, RingiNo | dataFiles = section/field đã trích xuất | prompt từng tiêu chí", "Sắp xếp dataFiles theo SectionType; Rule Engine chọn/chuẩn hóa chứng từ; LLM kết luận theo từng prompt tiêu chí.", "Mỗi tiêu chí trả CriteriaName, CriteriaStatus, Description, FileName, DueDateAI; API-AI lưu BEMT2004 và prompt dữ liệu tổng hợp vào BEMT2002."),
    ("DB + ERP9 hiển thị", "Danh sách BEMT2004 của toàn bộ tiêu chí", "Tính Percentage; có tiêu chí không OK thì Status = NG; chốt COMPLETED hoặc FAILED; gửi thông báo.", "BEMT2003 lưu Status, Percentage, TextConditionFail, TextContentOCR. ERP9 đọc DB để hiển thị kết quả/lý do."),
]

DB_ROWS = [
    ("BEMT2003", "Một lượt chạy AI của phiếu", "APK_BEMT2000, AttachID, AttachName, StatusProcess, Status, Percentage, TextContentOCR, TextContentAI, TextConditionFail, CreateDate, LastModifyDate", "Là bản ghi tổng của lần chạy: PROCESSING/COMPLETED/FAILED và OK/NG. Chưa tách trạng thái từng bước OCR/Extract/Rule/LLM."),
    ("BEMT2002", "Một file OCR hoặc dữ liệu nguồn", "APK_BEMT2003, FileName, AttachID, APK_File, DataSourceType, RawContent, ConsolidatedContent, Version", "Tên file gốc và text OCR nằm tại đây. Ví dụ FileName = CI_SSK-MV-2026-02-001.pdf; RawContent = text OCR của Invoice."),
    ("BEMT2005", "Header/section chứng từ do LLM nhận diện", "APK_BEMT2003, SectionType, SectionOrder, SectionTitle, TotalAmount, TotalCurrency, Signature, PromptSystem", "Lưu loại chứng từ như COMMERCIALINVOICE/PO/RINGI và thông tin tổng. Không có cột FileName cố định."),
    ("BEMT2006", "Các trường chi tiết LLM trích xuất", "APK_BEMT2005, APK_BEMT2003, OrderNo, Description, FieldID01...FieldID100", "Mỗi FieldID có ý nghĩa theo mapping ONT1041, ví dụ InvoiceNo, InvoiceDate, NCC, Amount, Currency, Incoterm. FileName chỉ được lưu nếu mapping có cấu hình."),
    ("BEMT2004", "Kết quả đối chiếu từng tiêu chí", "APK_BEMT2003, CriteriaID, CriteriaName, CriteriaStatus, Description, FileName, DueDateAI, PromptSystem", "Lưu OK/NG, kết luận/lý do NG và file liên quan của Số tiền, Invoice, NCC, Ringi, Deadline..."),
]

DB_EXAMPLE_ROWS = [
    ("1. Tạo lượt chạy", "BEMT2003", "StatusProcess = PROCESSING; gắn với phiếu NVL/09/2026/0011."),
    ("2. Lưu OCR file", "BEMT2002", "FileName = CI_SSK-MV-2026-02-001.pdf; RawContent chứa Invoice No, ngày, NCC, 1,500,000 JPY, Incoterm..."),
    ("3. Lưu loại chứng từ", "BEMT2005", "SectionType = COMMERCIALINVOICE; TotalAmount = 1,500,000; TotalCurrency = JPY."),
    ("4. Lưu trường trích xuất", "BEMT2006", "Các FieldID theo ONT1041 giữ InvoiceNo, InvoiceDate, SupplierName, Amount, Currency, Incoterm..."),
    ("5. Lưu kết quả tiêu chí", "BEMT2004", "CriteriaName = Số tiền; CriteriaStatus = OK; Description nêu 2 Invoice cộng đúng 3,250,000 JPY."),
    ("6. Chốt kết quả tổng", "BEMT2003", "StatusProcess = COMPLETED; Status = OK/NG; Percentage cập nhật theo toàn bộ tiêu chí."),
]

RESPONSIBILITY_ROWS = [
    ("Queue Management", "API-AI", "Queue RAM giới hạn 200 job; 1 reader/worker xử lý tuần tự; Queue đầy thì yêu cầu chờ; job fail chưa tự retry theo RetryCount."),
    ("OCR instance", "Python AI", "OCR engine lazy-load ở lần gọi đầu khi cần và dùng lại đến khi chuyển tài nguyên hoặc service restart; API-AI lưu kết quả theo file vào BEMT2002."),
    ("LLM instance", "Python AI", "LLM engine được giữ và dùng lại giữa các request; Python trim/cleanup và retry 1 lần khi CUDA OOM; lỗi tiếp thì request fail."),
    ("Rule Engine", "Python AI", "Hàm Rules chạy theo request, không phải GPU instance riêng; chọn chứng từ cần dùng trước khi kết luận hoặc gọi LLM."),
    ("Trạng thái và dữ liệu", "Database", "BEMT2003 lưu lượt chạy; BEMT2002 lưu file + OCR; BEMT2005 lưu header/loại chứng từ; BEMT2006 lưu trường trích xuất; BEMT2004 lưu kết quả tiêu chí."),
]

TASK_ROWS = [
    ("API/DB", "Persistent Queue", "Thay Queue RAM bằng bảng Job có JobID, Version, Priority, Deadline, RetryCount.", "Cao", "06/10"),
    ("API/DB", "StepLog trạng thái", "Tách queued, OCR, extract, rule, done, failed; lưu thời gian bắt đầu/kết thúc từng bước.", "Cao", "08/10"),
    ("API/DB", "File Hash/Version", "File không đổi thì dùng lại OCR/extract; file đổi thì chỉ xử lý lại file đó.", "Cao", "10/10"),
    ("API/Python", "Retry + lỗi rõ", "Retry giới hạn theo bước; báo tên file lỗi, bước lỗi, lý do lỗi.", "Cao", "10/10"),
    ("Python", "Rule trước LLM", "Logic xác định chạy bằng Rule/code; LLM chỉ dùng phần cần suy luận.", "Trung bình", "12/10"),
    ("API/DB", "Dashboard vận hành", "Theo dõi backlog, p95, số job lỗi, retry, phiếu gần hạn và năng lực xử lý.", "Trung bình", "15/10"),
]


def markdown_content():
    file_rows = "\n".join(f"| {a} | `{b}` | `{c}` | {d} |" for a, b, c, d in EXAMPLE_FILES)
    data_flow_rows = "\n".join(f"| {a} | {b} | {c} | {d} |" for a, b, c, d in DATA_FLOW_ROWS)
    db_rows = "\n".join(f"| `{a}` | {b} | `{c}` | {d} |" for a, b, c, d in DB_ROWS)
    db_example_rows = "\n".join(f"| {a} | `{b}` | {c} |" for a, b, c in DB_EXAMPLE_ROWS)
    task_rows = "\n".join(f"| {a} | {b} | {c} | {d} | {e} |" for a, b, c, d, e in TASK_ROWS)
    return f"""# Cơ chế vận hành đối chiếu AI BEM

Ngày cập nhật: 02/10/2026  
Mục tiêu: nhìn vào là hiểu Một phiếu NVL 10 file đi qua ERP9, API-AI, Python AI và Database như thế nào; API/Python/DB cần chốt gì để ổn định trước 15/10.

## 1. Cấu trúc mã nguồn và trách nhiệm hiện tại

- ERP9 WEB: nhận thao tác đối chiếu thủ công/lịch tự động, lấy phiếu - dòng ĐNTT - file đính kèm, gửi yêu cầu sang API-AI và hiển thị kết quả.
- ERP9 API: trong luồng BEM-AI hiện đã rà, chưa thấy đứng giữa WEB và API-AI để xử lý request đối chiếu; WEB ERP9 gọi trực tiếp API-AI.
- API-AI: nhận yêu cầu, tạo lượt chạy, quản lý Queue, gọi Python OCR/LLM/Rules và ghi dữ liệu vào DB.
- AI Python: OCR file, LLM trích xuất, Rule Engine chọn chứng từ, LLM đối chiếu và trả OK/NG.
- Database: BEMT2003 lưu lượt chạy; BEMT2002 lưu file và OCR; BEMT2005 lưu header/loại chứng từ; BEMT2006 lưu các trường trích xuất; BEMT2004 lưu kết quả từng tiêu chí.
- Không thay đổi ERP9 trong phạm vi task. Tập trung cải tiến API-AI, Python AI và Database để đạt kết quả ổn định trước 15/10.

## 2. Ví dụ đối chiếu thủ công một phiếu NVL có 10 file

Phiếu ví dụ: `NVL/09/2026/0011`, tổng yêu cầu `3,250,000 JPY`, NCC `SAKURA MATERIAL CO., LTD.`, PO `PO-2026-0901`, Ringi `RG-2026-088`.

| STT | File | AI nhận diện | Dữ liệu chính cần đọc |
|---|---|---|---|
{file_rows}

### Luồng xử lý

1. ERP9 nhận thao tác người dùng bấm đối chiếu thủ công và gom thông tin phiếu + 10 file.
2. API-AI tạo lượt chạy `BEMT2003 = PROCESSING` và đưa job vào Queue RAM.
3. Python OCR đọc 10 file qua `/ocr`, trả text OCR theo từng file.
4. API-AI lưu OCR vào `BEMT2002` gồm `FileName`, `AttachID`, `APK_File`, `RawContent`.
5. Python LLM trích xuất từng file qua `BEM_AGENT_READFILE`, trả JSON `sections`.
6. API-AI lưu header/loại chứng từ vào `BEMT2005` và các trường chi tiết vào `BEMT2006.FieldID01...FieldID100`. Tên file gốc vẫn nằm ở `BEMT2002.FileName`.
7. Python Rule Engine + LLM đối chiếu 9 tiêu chí NVL.
8. API-AI lưu kết quả từng tiêu chí vào `BEMT2004` và cập nhật tổng vào `BEMT2003`.
9. ERP9 đọc kết quả để hiển thị cho người dùng.

Kết quả ví dụ: tiêu chí Số tiền OK vì 2 invoice `1,500,000 + 1,750,000 = 3,250,000 JPY` khớp tờ khai và Ringi. `STATEMENT` chỉ kiểm tra thêm, không cộng trùng.

## 3. Dữ liệu đi qua từng khu vực code

| Khu vực | Dữ liệu nhận vào | Xử lý chính | Dữ liệu trả ra hoặc lưu lại |
|---|---|---|---|
{data_flow_rows}

Điểm cần lưu ý:
- ERP9 nhận phản hồi ngay sau khi API-AI tạo job; đây chưa phải kết quả đối chiếu cuối.
- Queue chỉ giữ job trong RAM. Worker lấy một job tại một thời điểm, nhưng OCR bên trong một phiếu có thể đọc tối đa 7 file song song.
- Prompt trích xuất chỉ nhận `FileName` và text OCR `result` của từng file.
- Prompt đối chiếu nhận ba nhóm chính: `datas` là master, `details` là 5 trường dòng ĐNTT và `dataFiles` là dữ liệu chứng từ đã trích xuất.
- Kết quả cuối của từng tiêu chí là `CriteriaName`, `CriteriaStatus`, `Description`, `FileName`, `DueDateAI`; kết quả tổng được tính lại vào `BEMT2003`.

## 4. Database lưu dữ liệu gì

| Bảng | Lưu gì | Trường chính | Ghi chú |
|---|---|---|---|
{db_rows}

### Ví dụ một file Invoice đi qua Database

| Bước | Bảng | Dữ liệu ví dụ |
|---|---|---|
{db_example_rows}

Lưu ý: code dùng OCR của từng file để tạo `BEMT2005/BEMT2006`, nhưng `BEMT2005` hiện không có khóa trực tiếp tới `BEMT2002`. Vì vậy tên file gốc nằm ở `BEMT2002.FileName`; nếu cần truy vết chắc chắn section/field sinh từ file nào thì nên bổ sung `APK_BEMT2002` hoặc `FileID/FileHash` vào dữ liệu trích xuất.

## 5. Trách nhiệm API/Python/DB và trạng thái hiện tại

- API-AI: tạo lượt chạy, quản lý Queue, gọi OCR/LLM/Rules, ghi DB.
- Python OCR: engine được tạo/lazy-load trong Python process và dùng lại giữa các request cho đến khi chuyển tài nguyên hoặc service restart.
- Python LLM: engine được giữ và dùng lại giữa các request; có trim/cleanup và retry 1 lần khi CUDA OOM.
- Python Rule Engine: hàm xử lý theo request, không phải GPU instance riêng; chọn đúng nhóm chứng từ trước khi kết luận.
- Database: lưu trạng thái, OCR, dữ liệu trích xuất, kết quả từng tiêu chí, lỗi và thời điểm.

| Nội dung | Thành phần | Hiện trạng |
|---|---|---|
{chr(10).join(f"| {a} | {b} | {c} |" for a, b, c in RESPONSIBILITY_ROWS)}

- Hiện DB có `PROCESSING`, `COMPLETED`, `FAILED` tại `BEMT2003.StatusProcess`.
- `PROCESSING` hiện đang bao trùm cả chờ Queue, đang OCR, đang LLM và đang Rule nên chưa biết phiếu đang kẹt ở bước nào.
- Queue hiện là RAM Queue, giới hạn 200 job, 1 reader/worker xử lý tuần tự; khi đầy thì yêu cầu chờ. API-AI chưa có Persistent Queue/Retry history chuẩn.
- Python chỉ retry cục bộ cho một số lỗi LLM/OOM; job fail ở API-AI chưa tự đưa lại Queue theo RetryCount.
- Cần bổ sung JobID, Version, Priority, Deadline, RetryCount, StepLog và thời gian bắt đầu/kết thúc từng bước.

## 6. Rủi ro chính dẫn đến fail như tháng 09

- Queue trong RAM có rủi ro mất job chờ khi service restart.
- Chưa chống trùng/chưa quản lý version tốt khi người dùng sửa phiếu hoặc thay file.
- Lỗi OCR/LLM chưa retry chuẩn và chưa báo rõ file lỗi/bước lỗi.
- GPU tải cao, LLM là điểm nghẽn; không nên tăng song song nếu chưa đo throughput/p95/VRAM.
- Nhiều logic xác định vẫn có thể đi qua LLM, làm chậm và thiếu ổn định.

## 7. Task chốt để đạt mục tiêu 15/10

Task baseline để không lặp lại fail tháng 09.

| Owner | Hạng mục | Việc cần làm | Ưu tiên | Deadline |
|---|---|---|---|---|
{task_rows}

## 8. Baseline chốt

- ERP9 không đổi trong phạm vi task này.
- API-AI chịu trách nhiệm điều phối job và ghi DB.
- Python AI chịu trách nhiệm OCR, LLM và Rule Engine.
- DB là nơi lưu trạng thái, dữ liệu, version, lỗi và thời gian để vận hành ổn định.
- Không mở rộng sang nội dung AI ngoài mục tiêu chốt baseline vận hành API/Python/DB trước 15/10.
"""


def build_doc(images):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.36)
    section.right_margin = Inches(0.36)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_text(doc, "CƠ CHẾ VẬN HÀNH ĐỐI CHIẾU AI BEM", 21, True, WD_ALIGN_PARAGRAPH.CENTER, "17365D", 1)
    add_text(doc, "Một phiếu NVL 10 file đi qua ERP9, API-AI, Python AI và Database như thế nào | Cập nhật 02/10/2026", 10.2, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 4)
    add_text(doc, "Kết luận: ERP9 giữ nguyên; phạm vi chốt là API-AI, Python AI và DB để ổn định vận hành trước 15/10.", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, "17365D", 4)

    heading(doc, "1. Cấu trúc mã nguồn và trách nhiệm hiện tại")
    picture(doc, images[0], "Hình 1. ERP9 WEB/API, API-AI, AI Python và Database làm gì trong luồng đối chiếu")

    doc.add_page_break()
    heading(doc, "2. Ví dụ đối chiếu thủ công một phiếu NVL có 10 file")
    picture(doc, images[1], "Hình 2. Phiếu ví dụ đi qua ERP9, API-AI, Python AI và Database")
    add_table(doc, "Danh sách 10 file trong ví dụ", ["STT", "File", "AI nhận diện", "Dữ liệu chính cần đọc"], EXAMPLE_FILES, font_size=6.8)

    doc.add_page_break()
    heading(doc, "3. Dữ liệu đi qua từng khu vực code")
    picture(doc, images[2], "Hình 3. Dữ liệu nhận vào, xử lý và trả ra tại ERP9, API-AI, Queue, OCR, LLM, Rules và Database")
    add_table(doc, "Dữ liệu đầu vào và đầu ra của từng khu vực", ["Khu vực", "Dữ liệu nhận vào", "Xử lý chính", "Dữ liệu trả ra hoặc lưu lại"], DATA_FLOW_ROWS, font_size=6.8)

    doc.add_page_break()
    heading(doc, "4. Database lưu dữ liệu gì")
    picture(doc, images[3], "Hình 4. BEMT2003, BEMT2002, BEMT2005, BEMT2006 và BEMT2004 lưu dữ liệu ở các lớp riêng")
    add_table(doc, "Database lưu dữ liệu gì", ["Bảng", "Lưu gì", "Trường chính", "Ghi chú"], DB_ROWS, font_size=6.8)
    add_table(doc, "Ví dụ một file Invoice đi qua Database", ["Bước", "Bảng", "Dữ liệu ví dụ"], DB_EXAMPLE_ROWS, font_size=7.0)
    add_text(doc, "Lưu ý: tên file gốc nằm tại BEMT2002.FileName. BEMT2005 hiện không có khóa trực tiếp tới BEMT2002; để truy vết chắc chắn section/field sinh từ file nào cần bổ sung APK_BEMT2002 hoặc FileID/FileHash.", 8.6, True, WD_ALIGN_PARAGRAPH.LEFT, "C65911", 3)

    doc.add_page_break()
    heading(doc, "5. Trách nhiệm API Python DB và trạng thái hiện tại")
    add_table(doc, "Queue, OCR, LLM, Rule Engine và DB hiện hoạt động ra sao", ["Nội dung", "Thành phần", "Hiện trạng"], RESPONSIBILITY_ROWS, font_size=7.2)
    for line in [
        "BEMT2003.StatusProcess hiện có PROCESSING, COMPLETED, FAILED; PROCESSING chưa cho biết phiếu đang chờ Queue, OCR, Extract, Rule hay LLM.",
        "Queue là RAM Queue: sức chứa 200 job, 1 reader/worker xử lý tuần tự; Queue đầy thì yêu cầu phải chờ.",
        "Python chỉ retry cục bộ cho một số lỗi LLM/OOM; API-AI chưa có RetryCount và Job History chuẩn để chạy lại job fail.",
        "Cần bổ sung JobID, Version, Priority, Deadline, RetryCount, StepLog cùng thời gian bắt đầu/kết thúc từng bước.",
    ]:
        add_text(doc, "- " + line, 9.0, False, WD_ALIGN_PARAGRAPH.LEFT, "000000", 1)

    doc.add_page_break()
    heading(doc, "6. Rủi ro tháng 09")
    picture(doc, images[4], "Hình 5. Rủi ro chính và phương án thực hiện tập trung vào API/Python/DB")
    heading(doc, "7. Task chốt để đạt mục tiêu 15/10")
    add_table(doc, "Task baseline để không lặp lại fail tháng 09", ["Owner", "Hạng mục", "Việc cần làm", "Ưu tiên", "Deadline"], TASK_ROWS, font_size=6.8)

    heading(doc, "8. Baseline chốt")
    for line in [
        "ERP9 không đổi trong phạm vi task này; chỉ giữ vai trò nhận thao tác và hiển thị kết quả.",
        "API-AI chịu trách nhiệm tạo job, điều phối Queue, gọi Python và ghi kết quả vào DB.",
        "Python AI chịu trách nhiệm OCR, LLM trích xuất, Rule Engine và LLM đối chiếu.",
        "DB phải là nơi lưu trạng thái, dữ liệu, version, lỗi, retry và thời gian để vận hành ổn định.",
        "Kết quả báo cáo này dùng làm baseline triển khai, không mở rộng sang nội dung AI ngoài mục tiêu vận hành API/Python/DB trước 15/10.",
    ]:
        add_text(doc, "- " + line, 9.4, False, WD_ALIGN_PARAGRAPH.LEFT, "000000", 1)
    doc.save(DOCX_PATH)


def main():
    image_specs = [
        ("01_cau_truc_ma_nguon.png", make_source_structure()),
        ("02_vi_du_phieu_10_file.png", make_manual_flow()),
        ("03_du_lieu_qua_tung_khu_vuc.png", make_data_contract_flow()),
        ("04_database_luu_du_lieu.png", make_responsibility_db_state()),
        ("05_rui_ro_task_15_10.png", make_risk_task()),
    ]
    images = []
    for filename, img in image_specs:
        path = IMAGE_DIR / filename
        img.save(path, quality=95)
        images.append(path)
    MD_PATH.write_text(markdown_content(), encoding="utf-8")
    build_doc(images)
    print(f"Created: {DOCX_PATH}")
    print(f"Created: {MD_PATH}")
    for path in images:
        print(f"Image: {path}")


if __name__ == "__main__":
    main()
