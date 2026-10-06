from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMG = BASE / "Temp" / "code_flow_acceptance_20261001"
IMG.mkdir(parents=True, exist_ok=True)
MD = BASE / "Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
DOCX = BASE / "Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"

FONT = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"

NAVY = "#17365D"
BLUE = "#2F75B5"
GOLD = "#D6A300"
ORANGE = "#ED7D31"
PURPLE = "#7030A0"
TEAL = "#168A8A"
GREEN = "#70AD47"
RED = "#C00000"
GRAY = "#6B7280"
DARK = "#1F2933"
LIGHT = "#F7FAFC"
LINE = "#6E7F91"


def font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def text_size(draw, text, ft):
    box = draw.textbbox((0, 0), text, font=ft)
    return box[2] - box[0], box[3] - box[1]


def wrap(draw, text, ft, width):
    lines = []
    for para in str(text).split("\n"):
        if not para.strip():
            lines.append("")
            continue
        current = ""
        for word in para.split():
            trial = (current + " " + word).strip()
            if text_size(draw, trial, ft)[0] <= width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines


def draw_center(draw, rect, text, ft, fill=DARK, leading=6):
    x1, y1, x2, y2 = rect
    lines = wrap(draw, text, ft, x2 - x1)
    heights = [max(1, text_size(draw, line, ft)[1]) for line in lines]
    total = sum(heights) + leading * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = text_size(draw, line, ft)[0]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=ft, fill=fill)
        y += height + leading


def header(draw, title, subtitle, w):
    draw.rectangle((0, 0, w, 112), fill=NAVY)
    draw.text((48, 18), title, font=font(42, True), fill="white")
    draw.text((50, 72), subtitle, font=font(23), fill="#DCE6F1")


def rounded_box(draw, rect, title, body, color, body_size=23, fill="white"):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, 20, fill=fill, outline=color, width=4)
    draw.rounded_rectangle((x1, y1, x2, y1 + 56), 20, fill=color, outline=color)
    draw.rectangle((x1, y1 + 30, x2, y1 + 56), fill=color)
    draw_center(draw, (x1 + 12, y1 + 5, x2 - 12, y1 + 50), title, font(23, True), "white")
    draw_center(draw, (x1 + 18, y1 + 70, x2 - 18, y2 - 18), body, font(body_size), DARK)


def arrow(draw, start, end, color=LINE, width=6, head=20, dashed=False):
    import math

    x1, y1 = start
    x2, y2 = end
    if dashed:
        segments = 14
        for i in range(segments):
            if i % 2 == 0:
                xa = x1 + (x2 - x1) * i / segments
                ya = y1 + (y2 - y1) * i / segments
                xb = x1 + (x2 - x1) * (i + 1) / segments
                yb = y1 + (y2 - y1) * (i + 1) / segments
                draw.line((xa, ya, xb, yb), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    p1 = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
    p2 = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
    draw.polygon([(x2, y2), p1, p2], fill=color)


def new_canvas(title, subtitle, w=2200, h=1320):
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    header(draw, title, subtitle, w)
    return img, draw


def make_overview():
    img, draw = new_canvas(
        "Bản đồ vận hành AI BEM",
        "Luồng 2 hàng trái sang phải: ERP9 ngoài phạm vi sửa; API-AI điều phối; Python xử lý; DB lưu trạng thái và kết quả",
    )
    draw.rounded_rectangle((45, 150, 2155, 1170), 26, fill="#F8FBFD", outline="#CBD5E1", width=3)
    draw_center(draw, (90, 170, 2110, 230), "Luồng xử lý chính", font(32, True), NAVY)

    top_y1, top_y2 = 300, 470
    bot_y1, bot_y2 = 680, 850
    top = [
        ((90, top_y1, 350, top_y2), "WEB ERP9", "Người dùng hoàn tất phiếu\nhoặc bấm đối chiếu", GRAY),
        ((430, top_y1, 690, top_y2), "ERP/API", "Lấy phiếu + file\ngửi sang API-AI", GRAY),
        ((770, top_y1, 1030, top_y2), "API-AI", "Tạo lượt chạy\nBEMT2003=PROCESSING", GOLD),
        ((1110, top_y1, 1370, top_y2), "Queue RAM", "ChannelJobQueue\nTối đa 200 job", GOLD),
        ((1450, top_y1, 1710, top_y2), "Worker", "Lấy job\ngọi chuỗi xử lý", GOLD),
    ]
    bottom = [
        ((1450, bot_y1, 1710, bot_y2), "OCR", "Python /ocr\nFile → Text", ORANGE),
        ((1110, bot_y1, 1370, bot_y2), "LLM trích xuất", "Text OCR → dữ liệu\nInvoice, NCC, Amount...", ORANGE),
        ((770, bot_y1, 1030, bot_y2), "Rules + LLM", "Đối chiếu tiêu chí\nOK/NG + giải thích", ORANGE),
        ((430, bot_y1, 690, bot_y2), "Database", "Lưu OCR, trích xuất\nkết quả, trạng thái", PURPLE),
        ((90, bot_y1, 350, bot_y2), "ERP9 hiển thị", "Đọc DB\nhiển thị kết quả", GRAY),
    ]
    for rect, title, body, color in top + bottom:
        rounded_box(draw, rect, title, body, color, 20)
    for i in range(len(top) - 1):
        arrow(draw, (top[i][0][2], 385), (top[i + 1][0][0], 385), LINE, 6, 20)
    arrow(draw, (1580, top_y2), (1580, bot_y1), ORANGE, 7, 22)
    for i in range(len(bottom) - 1):
        arrow(draw, (bottom[i][0][0], 765), (bottom[i + 1][0][2], 765), LINE, 6, 20)

    # DB details row
    db_boxes = [
        ((470, 960, 740, 1055), "BEMT2003", "Lượt chạy, trạng thái, %"),
        ((800, 960, 1070, 1055), "BEMT2002", "Text OCR, file"),
        ((1130, 960, 1400, 1055), "BEMT2005/2006", "Dữ liệu trích xuất"),
        ((1460, 960, 1730, 1055), "BEMT2004", "Kết quả tiêu chí"),
    ]
    draw_center(draw, (120, 972, 420, 1045), "DB lưu theo từng bước", font(25, True), PURPLE)
    for rect, title, body in db_boxes:
        rounded_box(draw, rect, title, body, PURPLE, 16)
    draw_center(
        draw,
        (300, 1260, 1900, 1310),
        "Chốt phạm vi: không sửa ERP9; chỉ ổn định API-AI, AI Python và DB để tránh lặp fail tháng 09.",
        font(25, True),
        "#8A3E00",
    )
    img.save(IMG / "01_ban_do_van_hanh_acceptance.png")


def make_manual_flow():
    img, draw = new_canvas(
        "Luồng đối chiếu thủ công một phiếu",
        "ERP9 giữ nguyên; người dùng chủ động chọn một phiếu và API-AI xử lý phần còn lại",
        h=1120,
    )
    draw.rounded_rectangle((55, 145, 2145, 1025), 24, fill="#F8FBFD", outline="#CBD5E1", width=3)
    draw.rounded_rectangle((75, 170, 590, 235), 16, fill=GRAY)
    draw_center(draw, (85, 177, 580, 228), "ERP9 hiện tại — không thay đổi", font(25, True), "white")

    steps = [
        ((100, 340, 395, 520), "1. Người dùng", "Mở phiếu DNTT\nBấm Đối chiếu AI", GRAY),
        ((475, 340, 770, 520), "2. ERP9", "Kiểm tra có file\nLấy phiếu + file đính kèm", GRAY),
        ((850, 340, 1145, 520), "3. ERP9/API", "Tạo CompareFileRequest\nPOST một request sang API-AI", GRAY),
        ((1225, 340, 1520, 520), "4. API-AI", "Tạo BEMT2003\nĐưa job vào Queue RAM", GOLD),
        ((1600, 340, 1895, 520), "5. AI Python", "OCR → LLM trích xuất\nRules + LLM đối chiếu", ORANGE),
    ]
    for rect, title, body, color in steps:
        rounded_box(draw, rect, title, body, color, 21)
    for index in range(len(steps) - 1):
        arrow(draw, (steps[index][0][2], 430), (steps[index + 1][0][0], 430), LINE, 6, 20)

    rounded_box(draw, (1225, 680, 1520, 850), "6. Database", "Lưu OCR, dữ liệu trích xuất\nkết quả tiêu chí, trạng thái", PURPLE, 20)
    rounded_box(draw, (1600, 680, 1895, 850), "7. ERP9", "Đọc DB\nHiển thị kết quả cho người dùng", GRAY, 20)
    arrow(draw, (1748, 520), (1372, 680), PURPLE, 6, 20)
    arrow(draw, (1520, 765), (1600, 765), PURPLE, 6, 20)
    draw.rounded_rectangle((105, 900, 2095, 985), 16, fill="#FFF8E1", outline=GOLD, width=3)
    draw_center(
        draw,
        (130, 912, 2070, 972),
        "Lưu ý hiện trạng ERP9: khi người dùng cập nhật phiếu hoặc xóa file đính kèm, ERP9 gọi UpdateResultAI để cập nhật/xóa kết quả AI cũ; phiếu cần được đưa vào đối chiếu lại.",
        font(22, True),
        "#8A3E00",
    )
    img.save(IMG / "01a_luong_thu_cong_mot_phieu.png")


def make_automatic_flow():
    img, draw = new_canvas(
        "Luồng đối chiếu tự động nhiều phiếu",
        "ERP9 giữ nguyên; một lần chạy tự động chọn nhiều phiếu rồi gửi từng phiếu sang API-AI",
        h=1160,
    )
    draw.rounded_rectangle((55, 145, 2145, 1060), 24, fill="#F8FBFD", outline="#CBD5E1", width=3)
    draw.rounded_rectangle((75, 170, 680, 235), 16, fill=GRAY)
    draw_center(draw, (85, 177, 670, 228), "ERP9 hiện tại — không thay đổi", font(25, True), "white")

    top = [
        ((100, 335, 420, 530), "1. Lịch ERP9", "Gọi endpoint\nUpdateResultCompareAI\nGiờ chạy theo cấu hình vận hành", GRAY),
        ((510, 335, 830, 530), "2. ERP9", "Lấy danh sách phiếu cần đối chiếu\nGetDataBEMT2000Compare", GRAY),
        ((920, 335, 1240, 530), "3. ERP9", "Lấy chi tiết + file đính kèm\nBỏ qua phiếu không có file", GRAY),
        ((1330, 335, 1650, 530), "4. ERP9", "Tối đa 5 tác vụ song song\nMỗi phiếu POST sang API-AI", GRAY),
        ((1740, 335, 2040, 530), "5. API-AI", "Mỗi request tạo lượt chạy\nĐưa job vào Queue RAM", GOLD),
    ]
    for rect, title, body, color in top:
        rounded_box(draw, rect, title, body, color, 20)
    for index in range(len(top) - 1):
        arrow(draw, (top[index][0][2], 432), (top[index + 1][0][0], 432), LINE, 6, 20)

    rounded_box(draw, (1685, 700, 2095, 880), "Queue API-AI", "ChannelJobQueue trong RAM\nTối đa 200 job\nSingleReader: Worker lấy lần lượt", GOLD, 22)
    rounded_box(draw, (1110, 700, 1520, 880), "AI Python", "OCR → LLM trích xuất\nRules + LLM đối chiếu\nTrả kết quả hoặc ErrorCode", ORANGE, 22)
    rounded_box(draw, (535, 700, 945, 880), "Database", "Lưu trạng thái từng lượt\nOCR, trích xuất\nkết quả OK/NG, %", PURPLE, 22)
    arrow(draw, (1890, 530), (1890, 700), GOLD, 7, 22)
    arrow(draw, (1685, 790), (1520, 790), LINE, 7, 22)
    arrow(draw, (1110, 790), (945, 790), LINE, 7, 22)
    draw.rounded_rectangle((160, 940, 2040, 1015), 16, fill="#FFF8E1", outline=GOLD, width=3)
    draw_center(
        draw,
        (180, 950, 2020, 1005),
        "Điểm cần theo dõi: ERP9 có thể gửi tối đa 5 phiếu cùng lúc, nhưng API-AI dùng Queue RAM và SingleReader; nhiều phiếu sẽ xếp hàng chờ Worker. Đây là phạm vi cần ổn định ở API-AI/Python/DB, không sửa ERP9.",
        font(20, True),
        "#8A3E00",
    )
    img.save(IMG / "01b_luong_tu_dong_nhieu_phieu.png")


def make_responsibility():
    img, draw = new_canvas(
        "Trách nhiệm từng thành phần",
        "Tập trung vào Queue, OCR, Rule Engine, LLM và DB lưu trạng thái/kết quả",
        h=1180,
    )
    cols = [
        ("API-AI", 80, 690, GOLD, "Nhận request\nTạo BEMT2003\nĐưa job vào Queue RAM\nWorker gọi Python\nGhi trạng thái trước/sau mỗi bước\nRetry theo lỗi tạm thời"),
        ("AI Python", 795, 1405, ORANGE, "OCR instance: đọc file → text\nLLM instance: trích xuất dữ liệu\nRule Engine: đối chiếu theo tiêu chí\nTrả OK/NG, %, ErrorCode\nGiới hạn tải OCR/LLM"),
        ("Database", 1510, 2120, PURPLE, "BEMT2003: lượt chạy/trạng thái\nBEMT2002: OCR/text/file\nBEMT2005/2006: dữ liệu trích xuất\nBEMT2004: kết quả tiêu chí\nStage/retry/time/error cần bổ sung"),
    ]
    for title, x1, x2, color, body in cols:
        rounded_box(draw, (x1, 190, x2, 760), title, body, color, 28, "#FFFFFF")
    arrow(draw, (690, 475), (795, 475), LINE, 8, 26)
    arrow(draw, (1405, 475), (1510, 475), LINE, 8, 26)
    draw.rounded_rectangle((140, 860, 2060, 1060), 24, fill="#F8FBFD", outline=BLUE, width=3)
    draw_center(
        draw,
        (180, 885, 2020, 1035),
        "Nguyên tắc baseline: DB là nguồn trạng thái chính. Queue RAM chỉ là bộ tăng tốc; nếu service restart, API-AI phải phục hồi job chưa hoàn tất từ DB.",
        font(34, True),
        NAVY,
    )
    img.save(IMG / "02_trach_nhiem_api_python_db.png")


def make_status():
    img, draw = new_canvas(
        "Trạng thái xử lý một phiếu",
        "Hiện tại mới có trạng thái tổng; baseline cần biết phiếu đang ở đúng bước nào",
        h=1180,
    )
    draw.rounded_rectangle((60, 165, 2140, 430), 24, fill="#F8FBFD", outline=BLUE, width=3)
    draw_center(draw, (90, 180, 320, 230), "HIỆN TẠI", font(28, True), BLUE)
    current = [
        ((220, 260, 540, 380), "PROCESSING", "Đang xử lý nhưng DB không biết đang ở bước nào", BLUE),
        ((930, 260, 1250, 380), "COMPLETED", "Có kết quả tổng OK/NG + %", GREEN),
        ((1640, 260, 1960, 380), "FAILED", "Chỉ có khi luồng bắt lỗi và ghi thất bại", RED),
    ]
    for rect, title, body, color in current:
        rounded_box(draw, rect, title, body, color, 19)
    arrow(draw, (540, 320), (930, 320), LINE, 6, 20)
    arrow(draw, (1250, 320), (1640, 320), LINE, 6, 20, dashed=True)
    draw_center(draw, (595, 352, 1605, 410), "Thiếu: chờ Queue / đang OCR / đang LLM / đang Rules / retry", font(24, True), "#8A3E00")

    draw.rounded_rectangle((60, 500, 2140, 1045), 24, fill="#F5FAF3", outline=GREEN, width=3)
    draw_center(draw, (90, 515, 470, 570), "BASELINE CẦN CHỐT", font(28, True), GREEN)
    states = [
        ("QUEUED", "Chờ Worker"),
        ("OCR_RUNNING", "Đang đọc file"),
        ("LLM_EXTRACTING", "Đang trích xuất"),
        ("RULE_CHECKING", "Đang đối chiếu"),
        ("COMPLETED", "Hoàn tất"),
    ]
    x = 110
    for index, (title, body) in enumerate(states):
        color = GREEN if title == "COMPLETED" else TEAL
        rounded_box(draw, (x, 650, x + 350, 800), title, body, color, 22)
        if index < len(states) - 1:
            arrow(draw, (x + 350, 725), (x + 390, 725), LINE, 6, 18)
        x += 390
    rounded_box(draw, (770, 900, 1150, 1005), "RETRY_WAIT", "Lỗi tạm thời\nChờ chạy lại tối đa 2 lần", GOLD, 20)
    rounded_box(draw, (1390, 900, 1770, 1005), "FAILED", "Hết retry hoặc lỗi dữ liệu/hồ sơ", RED, 20)
    arrow(draw, (945, 800), (945, 900), GOLD, 5, 18)
    arrow(draw, (1150, 952), (1390, 952), RED, 5, 18)
    img.save(IMG / "03_trang_thai_baseline.png")


def make_db_map():
    img, draw = new_canvas(
        "DB lưu trạng thái và kết quả gì",
        "Mục tiêu: mở DB lên biết phiếu đang ở bước nào, lỗi gì, retry ra sao và kết quả nằm ở đâu",
        h=1180,
    )
    boxes = [
        ((80, 200, 520, 450), "BEMT2003", "Lượt chạy chính\nStatusProcess\nStatus OK/NG\nPercentage\nTextContentOCR/AI\nTextConditionFail", PURPLE),
        ((610, 200, 1050, 450), "BEMT2002", "File/OCR\nRawContent\nConsolidatedContent\nDataSourceType\nFileName", PURPLE),
        ((1140, 200, 1580, 450), "BEMT2005/2006", "Dữ liệu trích xuất\nMaster/detail\nThông tin chứng từ\nNCC, Invoice, Amount...", PURPLE),
        ((1670, 200, 2110, 450), "BEMT2004", "Kết quả từng tiêu chí\nCriteriaStatus\nDescription\nPromptSystem\nDueDateAI", PURPLE),
    ]
    for rect, title, body, color in boxes:
        rounded_box(draw, rect, title, body, color, 22)
    for start, end in [((520, 325), (610, 325)), ((1050, 325), (1140, 325)), ((1580, 325), (1670, 325))]:
        arrow(draw, start, end, PURPLE, 6, 20)
    draw.rounded_rectangle((100, 650, 2100, 1000), 24, fill="#FFF8E1", outline=GOLD, width=4)
    draw_center(draw, (135, 675, 2065, 750), "Cần bổ sung để vận hành ổn định", font(34, True), "#8A3E00")
    needed = [
        "ProcessStage: QUEUED / OCR_RUNNING / LLM_EXTRACTING / RULE_CHECKING / RETRY_WAIT",
        "RetryCount, NextRetryAt, StartedAt, FinishedAt, LastHeartbeatAt",
        "ErrorCode, ErrorMessage, FailedStep để biết lỗi tại Queue/OCR/LLM/Rules",
        "EngineSource ở BEMT2004: RULE / LLM / MIXED để tách lỗi rules và lỗi model",
    ]
    y = 780
    for line in needed:
        draw.text((170, y), "• " + line, font=font(28), fill=DARK)
        y += 50
    img.save(IMG / "04_db_luu_gi.png")


def make_risks():
    img, draw = new_canvas(
        "Rủi ro có thể làm fail như tháng 09",
        "Chỉ ghi các rủi ro đã thấy từ source/log; xử lý trước khi chốt baseline 15/10",
        h=1180,
    )
    risks = [
        ((90, 180, 1040, 390), "P0  Phiếu treo PROCESSING", "Có nhánh return trước khi ghi FAILED: ngày nghỉ lỗi hoặc OCR không trả đủ dữ liệu."),
        ((1160, 180, 2110, 390), "P0  Queue RAM chưa bền", "Queue nằm trong RAM; restart có thể mất job chờ nếu DB không đủ trạng thái để phục hồi."),
        ((90, 500, 1040, 710), "P1  OCR biến động", "Log app_2: 591 lượt OCR, median 2,99s nhưng max 205,04s; 8 lượt ≥ 60s."),
        ((1160, 500, 2110, 710), "P1  LLM dùng chung GPU", "Một LLM engine tái sử dụng và gen_lock tuần tự hóa generate; yêu cầu dài làm phiếu sau chờ."),
        ((625, 820, 1575, 1030), "P0  Thiếu retry/error chuẩn", "Lỗi tạm thời, lỗi dữ liệu và lỗi hồ sơ chưa tách rõ; khó chạy lại đúng phiếu."),
    ]
    for rect, title, body in risks:
        rounded_box(draw, rect, title, body, RED if title.startswith("P0") else ORANGE, 24, "#FFF8F8" if title.startswith("P0") else "#FFF4EA")
    img.save(IMG / "05_rui_ro_chinh.png")


def make_solution():
    img, draw = new_canvas(
        "Phương án chốt API Python DB",
        "Giữ ERP9 nguyên trạng; chốt baseline vận hành để đạt mục tiêu 15/10",
        h=1180,
    )
    cols = [
        ("API-AI", 90, 670, GOLD, "1. Ghi stage trước/sau mỗi bước\n2. Queue RAM + phục hồi từ DB\n3. Retry ≤ 2 cho lỗi tạm thời\n4. Một RunID không chạy trùng\n5. Mọi lỗi phải ghi FAILED/ErrorCode"),
        ("AI Python", 810, 1390, ORANGE, "1. Tái sử dụng OCR/LLM engine\n2. Health check OCR và LLM\n3. Timeout riêng từng bước\n4. Trả ErrorCode có cấu trúc\n5. Giới hạn tải: OCR song song thấp, LLM giữ gen_lock"),
        ("Database", 1530, 2110, PURPLE, "1. BEMT2003 là nguồn trạng thái chính\n2. Bổ sung stage/retry/time/error\n3. Lưu nguồn kết quả RULE/LLM/MIXED\n4. Có query phục hồi job sau restart\n5. Có số liệu đo thời gian từng bước"),
    ]
    for title, x1, x2, color, body in cols:
        rounded_box(draw, (x1, 210, x2, 850), title, body, color, 27)
    arrow(draw, (670, 530), (810, 530), LINE, 8, 26)
    arrow(draw, (1390, 530), (1530, 530), LINE, 8, 26)
    draw_center(
        draw,
        (240, 930, 1960, 1055),
        "Điều kiện đạt: restart không mất job • không còn PROCESSING treo • biết phiếu đang ở bước nào • lỗi cuối có mã/thời gian • test fail/retry/tải đạt.",
        font(32, True),
        NAVY,
    )
    img.save(IMG / "06_phuong_an_chot.png")


def make_timeline():
    img, draw = new_canvas(
        "Kế hoạch chốt baseline đến 15/10",
        "Ưu tiên P0 trước: trạng thái, Queue DB, retry, test fail/restart/tải",
        h=1020,
    )
    dates = [("02–04/10", 240), ("05–07/10", 560), ("08–10/10", 880), ("11–13/10", 1200), ("14/10", 1520), ("15/10", 1790)]
    for label, x in dates:
        draw.rounded_rectangle((x, 150, x + 260, 215), 12, fill=NAVY)
        draw_center(draw, (x, 150, x + 260, 215), label, font(22, True), "white")
    rows = [
        ("API-AI", 300, GOLD, [(240, 520, "Fix treo trạng thái"), (560, 840, "Queue DB + retry"), (880, 1460, "Tích hợp stage"), (1520, 1760, "Sửa UAT")]),
        ("Python", 450, ORANGE, [(240, 840, "ErrorCode + health"), (880, 1160, "Timeout + giới hạn tải"), (1200, 1460, "Tích hợp"), (1520, 1760, "Sửa UAT")]),
        ("DB", 600, PURPLE, [(240, 520, "Chốt field"), (560, 1160, "Migration + recovery query"), (1200, 1460, "Đối soát dữ liệu")]),
        ("QA/Lead", 750, TEAL, [(240, 840, "Kịch bản test"), (880, 1460, "Test fail/restart/tải"), (1520, 1760, "UAT"), (1790, 2050, "Chốt baseline")]),
    ]
    for name, y, color, tasks in rows:
        draw.rounded_rectangle((40, y, 180, y + 78), 12, fill=color)
        draw_center(draw, (40, y, 180, y + 78), name, font(22, True), "white")
        for x1, x2, label in tasks:
            draw.rounded_rectangle((x1, y + 5, x2, y + 73), 12, fill=color)
            draw_center(draw, (x1 + 10, y + 5, x2 - 10, y + 73), label, font(19, True), "white")
    draw_center(draw, (240, 895, 1960, 970), "15/10: chỉ nghiệm thu khi trạng thái rõ, retry hoạt động, restart không mất job và bộ test tháng 09 chạy đạt.", font(29, True), RED)
    img.save(IMG / "07_ke_hoach_1510.png")


def make_all_images():
    make_overview()
    make_manual_flow()
    make_automatic_flow()
    make_responsibility()
    make_status()
    make_db_map()
    make_risks()
    make_solution()
    make_timeline()


def md_text():
    return """# Phân tích và chốt phương án ổn định xử lý AI BEM

Ngày báo cáo: 01/10/2026  
Mục tiêu nghiệm thu: đến **15/10/2026** có kết quả xử lý ổn định, không lặp lại tình trạng fail như hai kỳ tháng 09.

## 1. Phạm vi chốt

- Không thay đổi ERP9.
- Chỉ xem xét/cải tiến API-AI, AI Python và Database.
- Không mở rộng sang Rules/Training/nghiệp vụ AI ngoài mục tiêu ổn định vận hành.
- Kết quả báo cáo này dùng làm baseline triển khai đến 15/10.

## 2. Luồng end-to-end hiện tại

![Bản đồ vận hành](Temp/code_flow_acceptance_20261001/01_ban_do_van_hanh_acceptance.png)

- ERP9 hiện có vai trò khởi tạo yêu cầu, lấy dữ liệu phiếu/file và hiển thị kết quả; không nằm trong phạm vi sửa.
- API-AI tạo lượt chạy, đưa job vào Queue RAM, Worker lấy job và gọi AI Python.
- AI Python xử lý OCR, LLM trích xuất, Rule Engine + LLM đối chiếu.
- DB lưu lượt chạy, OCR, dữ liệu trích xuất, kết quả tiêu chí và trạng thái cuối.

### 2.1. Đối chiếu thủ công từng phiếu

![Luồng thủ công](Temp/code_flow_acceptance_20261001/01a_luong_thu_cong_mot_phieu.png)

- Người dùng bấm **Đối chiếu AI** trên một phiếu; ERP9 kiểm tra file đính kèm, tạo request và gọi API-AI cho đúng phiếu đó.
- Khi có kết quả, ERP9 đọc DB để hiển thị lại cho người dùng.
- Khi người dùng cập nhật phiếu hoặc xóa file đính kèm, ERP9 hiện gọi `UpdateResultAI` để cập nhật/xóa kết quả AI cũ; phiếu cần được đối chiếu lại.

### 2.2. Đối chiếu tự động nhiều phiếu

![Luồng tự động](Temp/code_flow_acceptance_20261001/01b_luong_tu_dong_nhieu_phieu.png)

- Lịch ERP9 gọi `UpdateResultCompareAI`; endpoint lấy danh sách phiếu cần đối chiếu, lấy chi tiết/file và bỏ qua phiếu không có file.
- ERP9 đang giới hạn tối đa 5 tác vụ gọi API-AI song song; mỗi phiếu vẫn trở thành một job riêng trong Queue RAM API-AI.
- Queue API-AI có `SingleReader`, nên job được Worker lấy lần lượt. Nhiều phiếu có thể chờ Queue trước khi OCR/LLM chạy.
- Thời điểm chạy cụ thể thuộc cấu hình vận hành/lịch ERP9; báo cáo không đề xuất sửa ERP9.

## 3. API / Python / DB đang chịu trách nhiệm gì

![Trách nhiệm](Temp/code_flow_acceptance_20261001/02_trach_nhiem_api_python_db.png)

| Thành phần | Thuộc phần | Hiện trạng |
|---|---|---|
| Queue Management | API-AI | `ChannelJobQueue` là Singleton trong RAM, tối đa 200 job, `SingleReader = true`; chưa có persistent queue/recovery đầy đủ sau restart. |
| OCR instance | AI Python | API-AI gọi route `/ocr`; OCR engine được giữ trong process Python và tái sử dụng. |
| LLM instance | AI Python | `_llm_engine`, `_llm_registry` được khởi tạo/tái sử dụng; `_gen_lock` làm thao tác generate chạy tuần tự trên cùng model/GPU. |
| Rule Engine | AI Python | `process_ai_llms_models_rules` xử lý logic đối chiếu và gọi LLM khi cần. |
| DB kết quả/trạng thái | Database | `BEMT2003`, `BEMT2002`, `BEMT2004`, `BEMT2005/2006` lưu trạng thái, OCR, dữ liệu trích xuất và kết quả tiêu chí. |

## 4. Trạng thái chính của một phiếu

![Trạng thái](Temp/code_flow_acceptance_20261001/03_trang_thai_baseline.png)

Hiện DB chủ yếu có `PROCESSING`, `COMPLETED`, `FAILED`. Baseline cần bổ sung stage vận hành: `QUEUED`, `OCR_RUNNING`, `LLM_EXTRACTING`, `RULE_CHECKING`, `RETRY_WAIT` để biết phiếu đang nằm ở bước nào.

## 5. DB/table/field đang lưu gì

![DB](Temp/code_flow_acceptance_20261001/04_db_luu_gi.png)

| Bảng/field | Đang lưu | Cần bổ sung để vận hành ổn định |
|---|---|---|
| `BEMT2003.StatusProcess` | `PROCESSING / COMPLETED / FAILED` | `ProcessStage` chi tiết: Queue/OCR/LLM/Rules/Retry. |
| `BEMT2003.Status`, `Percentage` | Kết quả tổng OK/NG và % | Cho biết kết quả được tạo sau bước nào. |
| `BEMT2003.TextContentOCR`, `TextContentAI`, `TextConditionFail` | Nội dung OCR/AI và điều kiện fail | Chuẩn hóa `ErrorCode`, `ErrorMessage`, `FailedStep`. |
| `BEMT2002` | OCR, nội dung tổng hợp, file | Lưu thời điểm bắt đầu/kết thúc OCR và trạng thái OCR. |
| `BEMT2005/BEMT2006` | Dữ liệu trích xuất chứng từ | Lưu trạng thái/thời điểm trích xuất theo lượt. |
| `BEMT2004` | Kết quả từng tiêu chí | Phân biệt nguồn kết quả `RULE / LLM / MIXED`. |
| Chưa có field riêng | Retry và heartbeat | `RetryCount`, `NextRetryAt`, `StartedAt`, `FinishedAt`, `LastHeartbeatAt`. |

## 6. Cơ chế Queue hiện tại

- API-AI nhận request và tạo lượt chạy `BEMT2003`.
- Job được đưa vào `ChannelJobQueue` trong RAM, tối đa 200 job.
- Worker lấy job từ Queue và xử lý theo chuỗi OCR → LLM trích xuất → Rules/LLM đối chiếu → lưu DB.
- Rủi ro hiện tại: Queue nằm trong RAM nên restart có thể mất job đang chờ; retry chưa có trạng thái bền vững trong DB.
- Baseline: DB là nguồn trạng thái chính, Queue RAM chỉ là bộ tăng tốc; khi restart phải phục hồi các job chưa hoàn tất từ DB.

## 7. Rủi ro chính có thể dẫn đến fail như tháng 09

![Rủi ro](Temp/code_flow_acceptance_20261001/05_rui_ro_chinh.png)

- Phiếu có thể treo `PROCESSING` nếu luồng return trước khi ghi `FAILED`.
- Queue RAM chưa bền khi service restart.
- Retry/lỗi chưa được phân loại rõ giữa lỗi tạm thời, lỗi dữ liệu và lỗi hồ sơ.
- OCR có độ trễ đuôi dài: `app_2.log` có 591 lượt OCR, median 2,99 giây, max 205,04 giây, 8 lượt từ 60 giây trở lên.
- LLM dùng chung GPU và có `gen_lock`; request dài làm phiếu sau phải chờ.

## 8. Phương án chốt API/Python/DB

![Phương án](Temp/code_flow_acceptance_20261001/06_phuong_an_chot.png)

- Giữ `StatusProcess = PROCESSING / COMPLETED / FAILED` để tương thích ERP9.
- Bổ sung `ProcessStage`, retry, thời gian, heartbeat và mã lỗi để vận hành.
- API-AI chịu trách nhiệm ghi trạng thái, retry và phục hồi job từ DB.
- AI Python trả kết quả/lỗi có cấu trúc cho OCR, LLM, Rule Engine.
- DB là nguồn trạng thái chính; Queue RAM chỉ hỗ trợ xử lý nhanh.
- Chỉ tăng tải sau khi test fail/restart/timeout đạt.

## 9. Task owner deadline

![Kế hoạch](Temp/code_flow_acceptance_20261001/07_ke_hoach_1510.png)

| Thành phần | Hiện trạng | API/Python/DB | Vấn đề | Phương án | Người thực hiện | Ưu tiên | Deadline |
|---|---|---|---|---|---|---|---|
| Trạng thái run | Có nhánh return không ghi FAILED | API-AI | Treo PROCESSING | Mọi lối thoát ghi stage/final status/ErrorCode | DEV API-AI | P0 | 04/10 |
| Queue | RAM Channel 200 | API-AI/DB | Restart mất job chờ | DB là nguồn trạng thái; startup recovery | DEV API-AI + DB | P0 | 07/10 |
| Retry | Chưa bền vững | API-AI/DB | Lỗi tạm thời phải chạy tay | Retry ≤ 2; `RetryCount/NextRetryAt` | DEV API-AI + DB | P0 | 07/10 |
| OCR | Nhiều file có thể làm chậm | Python/API-AI | Độ trễ đuôi dài | Giới hạn tải, timeout, health check, ErrorCode | DEV Python + API-AI | P1 | 10/10 |
| LLM | Một engine, có `gen_lock` | Python | Phiếu dài làm phiếu sau chờ | Timeout, metric thời gian/token/lỗi, giữ lock an toàn | DEV Python | P1 | 10/10 |
| Rule Engine | Kết quả chung | Python/DB | Khó tách lỗi Rules hay LLM | Lưu `EngineSource = RULE/LLM/MIXED` | DEV Python + DB | P1 | 10/10 |
| DB vận hành | Thiếu field | DB/API-AI | Không đo được stage/thời gian | Bổ sung stage/retry/time/heartbeat/error | DEV DB + API-AI | P0 | 07/10 |
| Kiểm thử | Chưa có baseline fail/restart | QA | Nguy cơ lặp fail tháng 09 | Test restart/OCR/LLM/timeout/retry/tải/case T09 | QA ASOFT | P0 | 13/10 |
| UAT/baseline | Chưa chốt | API/Python/DB/QA | Dễ mở rộng ngoài mục tiêu | UAT 14/10; Go/No-Go 15/10 | Tech Lead + owner | P0 | 15/10 |

## 10. Điều kiện nghiệm thu cuối

- Không thay đổi ERP9.
- Biết được mỗi phiếu đang ở stage nào.
- Không còn `PROCESSING` treo khi luồng dừng.
- Restart API-AI không làm mất job chưa xử lý.
- Retry lỗi tạm thời tối đa 2 lần; lỗi cuối có mã và thời gian.
- Test đạt OCR fail, LLM fail, timeout, restart, retry và tải.
- Bộ đại diện hai kỳ tháng 09 chạy đạt trước khi chốt baseline ngày 15/10.
- Sau khi chốt chỉ triển khai nội dung trong baseline; không mở rộng sang Rules/Training/nghiệp vụ AI ngoài task này.
"""


def set_cell(cell, text, bold=False, size=10, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc, headers, rows, font_size=8.5):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, header_text in enumerate(headers):
        set_cell(tbl.rows[0].cells[idx], header_text, True, font_size, "FFFFFF")
        shade(tbl.rows[0].cells[idx], "17365D")
    for row in rows:
        cells = tbl.add_row().cells
        for idx, value in enumerate(row):
            set_cell(cells[idx], str(value), False, font_size)
    return tbl


def para(doc, text, bullet=False, bold=False, color=None):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(11)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(17 if level == 1 else 14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    return p


def image(doc, path, caption, width=8.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in c.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = RGBColor(91, 101, 115)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(0.55)
    sec.right_margin = Inches(0.55)
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Báo cáo chốt phương án ổn định xử lý AI BEM")
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    subtitle = doc.add_paragraph("Phạm vi: API-AI, AI Python và Database. Không thay đổi ERP9. Mục tiêu nghiệm thu: 15/10/2026.")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading(doc, "1. Kết luận phạm vi")
    for item in [
        "Không thay đổi ERP9; ERP9 chỉ đưa phiếu vào và hiển thị kết quả.",
        "Chỉ cải tiến API-AI, AI Python và Database để vận hành ổn định.",
        "Mục tiêu: đến 15/10/2026 xử lý đạt và không lặp lại tình trạng fail như hai kỳ tháng 09.",
        "Không mở rộng sang Rules/Training/nghiệp vụ AI ngoài baseline vận hành.",
    ]:
        para(doc, item, True)

    heading(doc, "2. Luồng end-to-end hiện tại")
    image(doc, IMG / "01_ban_do_van_hanh_acceptance.png", "Hình 1. Luồng từ ERP9 đưa phiếu vào đến DB lưu kết quả và ERP9 hiển thị")

    heading(doc, "2.1. Đối chiếu thủ công từng phiếu", 2)
    image(doc, IMG / "01a_luong_thu_cong_mot_phieu.png", "Hình 2. Vai trò hiện tại của ERP9 trong luồng người dùng bấm đối chiếu một phiếu")
    for item in [
        "Người dùng bấm Đối chiếu AI trên một phiếu; ERP9 kiểm tra file đính kèm, lấy thông tin phiếu và gửi một request sang API-AI.",
        "Kết quả cuối được lưu DB; ERP9 đọc DB để hiển thị lại cho người dùng.",
        "Khi cập nhật phiếu hoặc xóa file đính kèm, ERP9 hiện gọi UpdateResultAI để cập nhật/xóa kết quả AI cũ; phiếu cần được đưa vào chạy lại.",
    ]:
        para(doc, item, True)

    heading(doc, "2.2. Đối chiếu tự động nhiều phiếu", 2)
    image(doc, IMG / "01b_luong_tu_dong_nhieu_phieu.png", "Hình 3. Vai trò hiện tại của ERP9 trong luồng tự động nhiều phiếu")
    for item in [
        "Lịch ERP9 gọi endpoint UpdateResultCompareAI; ERP9 lấy danh sách phiếu cần đối chiếu, chi tiết và file đính kèm.",
        "ERP9 giới hạn tối đa 5 tác vụ gọi API-AI song song; mỗi phiếu được gửi thành một request riêng.",
        "API-AI có Queue RAM và SingleReader nên job được Worker lấy lần lượt; nhiều phiếu có thể xếp hàng chờ trước khi OCR/LLM chạy.",
        "Thời điểm chạy cụ thể là cấu hình vận hành/lịch ERP9. Báo cáo mô tả hiện trạng, không đề xuất sửa ERP9.",
    ]:
        para(doc, item, True)

    heading(doc, "3. API Python DB chịu trách nhiệm gì")
    image(doc, IMG / "02_trach_nhiem_api_python_db.png", "Hình 4. Trách nhiệm vận hành của API-AI, AI Python và Database")
    table(
        doc,
        ["Thành phần", "Phần", "Hiện trạng"],
        [
            ("Queue Management", "API-AI", "ChannelJobQueue nằm trong RAM, tối đa 200 job, SingleReader=true; chưa có recovery đầy đủ sau restart."),
            ("OCR instance", "AI Python", "Route /ocr đọc file thành text; engine được giữ trong process và tái sử dụng."),
            ("LLM instance", "AI Python", "LLM engine/registry tái sử dụng; gen_lock làm generate tuần tự trên cùng model/GPU."),
            ("Rule Engine", "AI Python", "process_ai_llms_models_rules đối chiếu tiêu chí và gọi LLM khi cần."),
            ("Kết quả/trạng thái", "Database", "BEMT2003/2002/2004/2005/2006 lưu lượt chạy, OCR, trích xuất và kết quả."),
        ],
        8.8,
    )

    doc.add_page_break()
    heading(doc, "4. Trạng thái xử lý một phiếu")
    image(doc, IMG / "03_trang_thai_baseline.png", "Hình 5. Trạng thái hiện tại và trạng thái baseline cần bổ sung")

    heading(doc, "5. DB đang lưu gì và còn thiếu gì")
    image(doc, IMG / "04_db_luu_gi.png", "Hình 6. DB lưu trạng thái, OCR, trích xuất và kết quả tiêu chí")
    table(
        doc,
        ["Bảng/field", "Đang lưu", "Cần bổ sung"],
        [
            ("BEMT2003.StatusProcess", "PROCESSING / COMPLETED / FAILED", "ProcessStage: QUEUED, OCR_RUNNING, LLM_EXTRACTING, RULE_CHECKING, RETRY_WAIT."),
            ("BEMT2003.Status, Percentage", "Kết quả tổng OK/NG và %", "Cho biết kết quả tạo sau bước nào."),
            ("TextContentOCR/AI/Fail", "Nội dung OCR/AI và điều kiện fail", "ErrorCode, ErrorMessage, FailedStep."),
            ("BEMT2002", "OCR, nội dung tổng hợp, file", "Thời điểm bắt đầu/kết thúc OCR và trạng thái OCR."),
            ("BEMT2005/2006", "Dữ liệu trích xuất chứng từ", "Trạng thái/thời điểm trích xuất theo lượt."),
            ("BEMT2004", "Kết quả từng tiêu chí", "EngineSource = RULE / LLM / MIXED."),
            ("Chưa có field riêng", "Retry và heartbeat", "RetryCount, NextRetryAt, StartedAt, FinishedAt, LastHeartbeatAt."),
        ],
        8.4,
    )

    doc.add_page_break()
    heading(doc, "6. Cơ chế Queue hiện tại")
    for item in [
        "API-AI nhận request, tạo lượt chạy BEMT2003 và đưa job vào ChannelJobQueue trong RAM.",
        "Worker lấy job từ Queue và gọi chuỗi OCR → LLM trích xuất → Rule Engine/LLM đối chiếu → lưu DB.",
        "Rủi ro: Queue nằm trong RAM nên restart có thể mất job chờ nếu DB không đủ trạng thái phục hồi.",
        "Baseline: DB là nguồn trạng thái chính; Queue RAM chỉ là bộ tăng tốc, restart phải phục hồi job chưa hoàn tất từ DB.",
    ]:
        para(doc, item, True)

    heading(doc, "7. Rủi ro có thể gây fail như tháng 09")
    image(doc, IMG / "05_rui_ro_chinh.png", "Hình 7. Các rủi ro chính đã xác nhận từ source/log")

    heading(doc, "8. Phương án chốt API Python DB")
    image(doc, IMG / "06_phuong_an_chot.png", "Hình 8. Baseline tập trung API-AI, AI Python và Database")
    for item in [
        "Giữ StatusProcess hiện có để tương thích ERP9, nhưng bổ sung ProcessStage để vận hành.",
        "API-AI ghi stage, retry, thời gian, heartbeat và error vào DB.",
        "AI Python trả kết quả/lỗi có cấu trúc cho OCR, LLM và Rule Engine.",
        "DB là nguồn trạng thái chính; Queue RAM chỉ hỗ trợ xử lý nhanh.",
        "Chỉ tăng tải sau khi test restart, timeout, OCR fail, LLM fail và retry đạt.",
    ]:
        para(doc, item, True)

    doc.add_page_break()
    heading(doc, "9. Kế hoạch đến 15/10")
    image(doc, IMG / "07_ke_hoach_1510.png", "Hình 9. Thứ tự thực hiện để chốt baseline 15/10")

    sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.top_margin = Inches(0.42)
    sec.bottom_margin = Inches(0.42)
    sec.left_margin = Inches(0.42)
    sec.right_margin = Inches(0.42)
    heading(doc, "10. Bảng task owner deadline")
    table(
        doc,
        ["Thành phần", "Hiện trạng", "API/Python/DB", "Vấn đề", "Phương án", "Người thực hiện", "P", "Deadline"],
        [
            ("Trạng thái run", "Có nhánh return không ghi FAILED", "API-AI", "Treo PROCESSING", "Mọi lối thoát ghi stage/final status/ErrorCode", "DEV API-AI", "P0", "04/10"),
            ("Queue", "RAM Channel 200", "API-AI/DB", "Restart mất job chờ", "DB là nguồn trạng thái; startup recovery", "DEV API-AI + DB", "P0", "07/10"),
            ("Retry", "Chưa bền vững", "API-AI/DB", "Lỗi tạm thời phải chạy tay", "Retry ≤ 2; RetryCount/NextRetryAt", "DEV API-AI + DB", "P0", "07/10"),
            ("OCR", "Nhiều file có thể làm chậm", "Python/API-AI", "Độ trễ đuôi dài", "Giới hạn tải, timeout, health check, ErrorCode", "DEV Python + API-AI", "P1", "10/10"),
            ("LLM", "Một engine, có gen_lock", "Python", "Phiếu dài làm phiếu sau chờ", "Timeout, metric thời gian/token/lỗi", "DEV Python", "P1", "10/10"),
            ("Rule Engine", "Kết quả chung", "Python/DB", "Khó tách lỗi Rules hay LLM", "EngineSource = RULE/LLM/MIXED", "DEV Python + DB", "P1", "10/10"),
            ("DB vận hành", "Thiếu field", "DB/API-AI", "Không đo stage/thời gian", "Stage/retry/time/heartbeat/error", "DEV DB + API-AI", "P0", "07/10"),
            ("Kiểm thử", "Chưa có baseline fail/restart", "QA", "Nguy cơ lặp fail tháng 09", "Test restart/OCR/LLM/timeout/retry/tải/case T09", "QA ASOFT", "P0", "13/10"),
            ("UAT/baseline", "Chưa chốt", "API/Python/DB/QA", "Dễ mở rộng ngoài mục tiêu", "UAT 14/10; Go/No-Go 15/10", "Tech Lead + owner", "P0", "15/10"),
        ],
        7.3,
    )
    heading(doc, "11. Điều kiện nghiệm thu cuối")
    for item in [
        "Không thay đổi ERP9.",
        "Biết được mỗi phiếu đang ở stage nào.",
        "Không còn PROCESSING treo khi luồng dừng.",
        "Restart API-AI không làm mất job chưa xử lý.",
        "Retry lỗi tạm thời tối đa 2 lần; lỗi cuối có mã và thời gian.",
        "Test đạt OCR fail, LLM fail, timeout, restart, retry và tải.",
        "Bộ đại diện hai kỳ tháng 09 chạy đạt trước khi chốt baseline ngày 15/10.",
        "Sau khi chốt chỉ triển khai nội dung trong baseline; không mở rộng Rules/Training/nghiệp vụ AI ngoài task.",
    ]:
        para(doc, item, True)
    doc.save(DOCX)


if __name__ == "__main__":
    make_all_images()
    MD.write_text(md_text(), encoding="utf-8")
    build_docx()
    print(MD)
    print(DOCX)
