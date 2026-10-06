from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMAGE_DIR = BASE / "Temp" / "co_che_code_ngan_gon_01102026"
IMAGE_DIR.mkdir(parents=True, exist_ok=True)
MD_PATH = BASE / "Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
DOCX_PATH = BASE / "Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"

FONT = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"
NAVY, ERP, API_AI, AI, DB = "#17365D", "#2F75B5", "#C58A00", "#C65911", "#7030A0"
DARK, LINE = "#1F2933", "#4B5563"


def f(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def text_size(draw, value, font_value):
    box = draw.textbbox((0, 0), str(value), font=font_value)
    return box[2] - box[0], box[3] - box[1]


def wrap(draw, value, font_value, max_width):
    lines = []
    for paragraph in str(value).split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if not current or text_size(draw, candidate, font_value)[0] <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines


def centered(draw, rect, value, font_value, fill=DARK, gap=6):
    x1, y1, x2, y2 = rect
    lines = wrap(draw, value, font_value, max(20, x2 - x1 - 24))
    heights = [max(1, text_size(draw, line, font_value)[1]) for line in lines]
    y = y1 + ((y2 - y1 - sum(heights) - gap * max(0, len(lines) - 1)) / 2)
    for line, height in zip(lines, heights):
        width = text_size(draw, line, font_value)[0]
        draw.text((x1 + ((x2 - x1 - width) / 2), y), line, font=font_value, fill=fill)
        y += height + gap


def left_wrapped(draw, rect, value, font_value, fill=DARK, gap=7):
    x1, y1, x2, y2 = rect
    y = y1
    for raw_line in str(value).split("\n"):
        if raw_line == "":
            y += text_size(draw, "A", font_value)[1] + gap
            continue
        for line in wrap(draw, raw_line, font_value, max(20, x2 - x1)):
            if y > y2:
                return
            draw.text((x1, y), line, font=font_value, fill=fill)
            y += max(1, text_size(draw, line, font_value)[1]) + gap


def canvas(title, subtitle, width=2800, height=1500):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 140), fill=NAVY)
    draw.text((58, 25), title, font=f(44, True), fill="white")
    draw.text((60, 84), subtitle, font=f(23), fill="#DEEAF6")
    return image, draw


def box(draw, rect, title, body, color, body_size=24):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=26, fill="white", outline=color, width=5)
    draw.rounded_rectangle((x1, y1, x2, y1 + 72), radius=26, fill=color, outline=color)
    draw.rectangle((x1, y1 + 40, x2, y1 + 72), fill=color)
    centered(draw, (x1 + 14, y1 + 8, x2 - 14, y1 + 63), title, f(25, True), "white")
    centered(draw, (x1 + 22, y1 + 92, x2 - 22, y2 - 20), body, f(body_size), DARK, 7)


def tag(draw, rect, value, color, size=18):
    draw.rounded_rectangle(rect, radius=14, fill="#F8FBFF", outline=color, width=2)
    centered(draw, rect, value, f(size, True), color, 4)


def right_arrow(draw, x1, y, x2, color=LINE, width=6, tip=23):
    draw.line((x1, y, x2 - tip, y), fill=color, width=width)
    draw.polygon([(x2, y), (x2 - tip, y - tip * 0.64), (x2 - tip, y + tip * 0.64)], fill=color)


def left_arrow(draw, x1, y, x2, color=LINE, width=6, tip=23):
    right_arrow(draw, x2, y, x1, color, width, tip)


def down_arrow(draw, x, y1, y2, color=LINE, width=6, tip=23):
    draw.line((x, y1, x, y2 - tip), fill=color, width=width)
    draw.polygon([(x, y2), (x - tip * 0.64, y2 - tip), (x + tip * 0.64, y2 - tip)], fill=color)


def card(draw, rect, number, title, body, color, body_size=21):
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=26, fill="white", outline=color, width=5)
    draw.rounded_rectangle((x1, y1, x2, y1 + 74), radius=26, fill=color, outline=color)
    draw.rectangle((x1, y1 + 42, x2, y1 + 74), fill=color)
    draw.ellipse((x1 + 20, y1 + 13, x1 + 66, y1 + 59), fill="white")
    centered(draw, (x1 + 22, y1 + 15, x1 + 64, y1 + 57), str(number), f(21, True), color, 0)
    centered(draw, (x1 + 84, y1 + 8, x2 - 16, y1 + 66), title, f(23, True), "white")
    left_wrapped(draw, (x1 + 34, y1 + 112, x2 - 30, y2 - 24), body, f(body_size), DARK, 8)


def make_source_structure():
    image, draw = canvas("01. CẤU TRÚC SOURCE", "Đọc theo thứ tự 1 đến 4; mỗi khu vực có trách nhiệm riêng")
    card(draw, (100, 270, 730, 1260), 1, "ERP9  WEB VÀ API", "Vị trí: 01.ERP9\nCode: BEMF2000Controller.cs\n\n- Nhận thao tác thủ công hoặc lịch tự động.\n- Lấy phiếu ĐNTT và file đính kèm.\n- Gọi API-AI qua HandlerFile.\n- Đọc kết quả để hiển thị UI.\n\n03.SERVICES: chưa thấy trực tiếp trong flow BEM-AI đã rà.", ERP, 24)
    card(draw, (780, 270, 1410, 1260), 2, "API-AI  .NET", "Vị trí: 06.API_AI\n\nCode chính:\n- FileDataHandlerController\n- ReadFileOrchestratorService\n- ChannelJobQueue / ReadFileWorker\n\nTạo lượt chạy, đưa job vào Queue, gọi AI Python, lưu và đọc Database.", API_AI, 24)
    card(draw, (1460, 270, 2090, 1260), 3, "AI BEM  PYTHON", "Vị trí: BEM_AI_PROJECT/App\n\nCode chính:\n- main_iis.py: /ocr, /api/ai_llms_models\n- Rules_AI_BEM_MEIKO.py\n- LLMs_BE/engine.py\n\nNhận file/text, OCR, trích xuất dữ liệu, áp dụng Rules và LLM đối chiếu.", AI, 24)
    card(draw, (2140, 270, 2770, 1260), 4, "DATABASE", "BEMT2003: một lượt chạy AI\nBEMT2002: tên file và text OCR\nBEMT2005: loại chứng từ\nBEMT2006: dữ liệu cột đã trích xuất\nBEMT2004: kết quả từng tiêu chí\n\nAPI-AI là nơi ghi/đọc dữ liệu vận hành.", DB, 24)
    return image


def make_processing_flow():
    image, draw = canvas("02. LUỒNG XỬ LÝ MỘT PHIẾU", "Sáu bước theo thứ tự; ô sau nhận dữ liệu đầu ra của ô trước")
    cards = [
        ((120, 250, 1360, 590), 1, "ERP9 NHẬN PHIẾU VÀ FILE", "Code: CallAPICompareFileOCRAsync hoặc UpdateResultCompareAI\n\nNhận: thao tác thủ công / lịch tự động, phiếu ĐNTT, file đính kèm.\nTrả sang API-AI: CompareFileRequest.", ERP),
        ((1440, 250, 2680, 590), 2, "API-AI TẠO LƯỢT CHẠY", "Code: HandlerFileAsync → ReadFileOrchestratorService\n\nTạo BEMT2003 với StatusProcess = PROCESSING.\nĐưa ReadFileJob vào ChannelJobQueue.", API_AI),
        ((120, 655, 1360, 995), 3, "QUEUE LẤY JOB NỀN", "Code: ChannelJobQueue → ReadFileWorker\n\nNhận: ReadFileJob.\nChuyển file và thông tin lượt chạy sang workflow OCR / LLM.", API_AI),
        ((1440, 655, 2680, 995), 4, "OCR ĐỌC FILE", "Code: AI Python endpoint /ocr\n\nNhận: từng file đính kèm.\nTrả: text OCR; API-AI lưu theo file vào BEMT2002.", AI),
        ((120, 1060, 1360, 1400), 5, "LLM TRÍCH XUẤT DỮ LIỆU", "Code: /api/ai_llms_models\n\nNhận: text OCR.\nTrả: loại chứng từ + dữ liệu cột; API-AI lưu BEMT2005 / BEMT2006.", AI),
        ((1440, 1060, 2680, 1400), 6, "RULES VÀ LLM ĐỐI CHIẾU", "Code: CompareAsync + Rules_AI_BEM_MEIKO.py\n\nNhận: dữ liệu đã trích xuất + Rules.\nTrả: OK/NG từng tiêu chí vào BEMT2004; kết quả tổng cập nhật BEMT2003.", AI),
    ]
    for rect, number, title, body, color in cards:
        card(draw, rect, number, title, body, color, 24)
    return image


def make_database_structure():
    image, draw = canvas("03. DATABASE LƯU DỮ LIỆU GÌ", "BEMT2003 là mã lượt chạy; các bảng dưới đều liên kết qua APK_BEMT2003")
    card(draw, (190, 220, 2610, 500), 1, "BEMT2003  LƯỢT CHẠY AI", "Khóa phiếu: APK_BEMT2000  |  File đầu vào: AttachID, AttachName\nTrạng thái: StatusProcess  |  Kết quả tổng: Status, Percentage\nNội dung tổng: TextContentOCR, TextContentAI  |  Lỗi: TextConditionFail", DB, 24)
    card(draw, (60, 650, 710, 1390), 2, "BEMT2002  FILE VÀ OCR", "Tên file: FileName\nMã file: APK_File, AttachID\nNguồn: DataSourceType\n\nOCR thô: RawContent\nOCR đã gom/chuẩn hóa: ConsolidatedContent\n\nDùng để biết file nào đã OCR và OCR đọc ra gì.", DB, 20)
    card(draw, (750, 650, 1400, 1390), 3, "BEMT2005  LOẠI CHỨNG TỪ", "Loại chứng từ: SectionType\nTên phần/chứng từ: SectionTitle\nThứ tự: SectionOrder\n\nTổng tiền: TotalAmount\nTiền tệ: TotalCurrency\nChữ ký: Signature\nPrompt dùng: PromptSystem", DB, 20)
    card(draw, (1440, 650, 2090, 1390), 4, "BEMT2006  CỘT ĐÃ TRÍCH XUẤT", "Thuộc chứng từ: APK_BEMT2005\nDòng: OrderNo, Description\n\nGiá trị AI đọc: FieldID01 ... FieldID100\nÝ nghĩa từng FieldID do cấu hình ONT1041 map.\n\nTên file không có cột cố định; khi cần lưu sẽ map vào một FieldID theo cấu hình.", DB, 19)
    card(draw, (2130, 650, 2780, 1390), 5, "BEMT2004  KẾT QUẢ ĐỐI CHIẾU", "Tiêu chí: CriteriaID, CriteriaName\nKết quả: CriteriaStatus (OK/NG)\nGiải thích: Description\n\nFile dùng đối chiếu: FileName\nHạn AI xác định: DueDateAI\nPrompt dùng: PromptSystem", DB, 20)
    return image


def set_cell_shading(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def write_cell(cell, value, size=10.0, bold=False, color="000000"):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def write_cell_left(cell, value, size=8.0, bold=False, color="000000"):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(document, value, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, color=None, after=3):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return paragraph


def heading(document, value):
    return add_text(document, value, 14, True, WD_ALIGN_PARAGRAPH.LEFT, "17365D", 4)


def picture(document, path, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(path), width=Inches(10.65))
    add_text(document, caption, 9, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 1)


def database_table(document):
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Cần tìm", "Bảng", "Trường chính", "Hiểu đơn giản"]):
        set_cell_shading(cell, "17365D")
        write_cell(cell, value, 9.8, True, "FFFFFF")
    rows = [
        ("Phiếu đang chạy tới đâu", "BEMT2003", "APK_BEMT2000, StatusProcess, Status, Percentage, TextConditionFail", "Một bản ghi = một lượt chạy AI của phiếu."),
        ("Tên file và OCR", "BEMT2002", "FileName, APK_File, AttachID, RawContent, ConsolidatedContent", "Biết file nào OCR và OCR đọc ra text gì."),
        ("Loại chứng từ", "BEMT2005", "SectionType, SectionTitle, SectionOrder, TotalAmount, TotalCurrency", "Header của nhóm dữ liệu/chứng từ AI đã trích xuất."),
        ("Dữ liệu các cột đã trích xuất", "BEMT2006", "OrderNo, Description, FieldID01...FieldID100", "Mỗi FieldID là một cột dữ liệu do ONT1041 mapping."),
        ("Kết quả đối chiếu", "BEMT2004", "CriteriaName, CriteriaStatus, Description, FileName, DueDateAI", "Kết quả OK/NG và lý do của từng tiêu chí."),
    ]
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_shading(cell, "F3F7FC" if row_index % 2 == 0 else "FFFFFF")
            write_cell(cell, value, 8.6)


EXAMPLE_FILES = [
    ("1", "PO_PO-2026-0901.pdf", "PO", "ContractNo/PO No = PO-2026-0901; RingiNo = RG-2026-088; SupplierName = SAKURA MATERIAL CO., LTD.; Currency = JPY; Amount = 3,250,000; PaymentTerm = TT 30 days; DeliveryTerm = CIF HAI PHONG", "BEMT2002 lưu FileName + OCR; BEMT2005 lưu SectionType=PO; BEMT2006 lưu các FieldID tương ứng."),
    ("2", "RINGI_RG-2026-088.pdf", "RINGI", "RingiNo = RG-2026-088; SupplierName = SAKURA MATERIAL CO., LTD.; Currency = JPY; Amount = 3,250,000; ApprovalLast = đã duyệt", "BEMT2005 SectionType=RINGI; BEMT2006 lưu RingiNo, NCC, số tiền, tiền tệ, người/ngày duyệt."),
    ("3", "CI_SSK-MV-2026-02-001.pdf", "COMMERCIALINVOICE", "VoucherNo = SSK-MV-2026/02-001; VoucherDate = 05/09/2026; SupplierName = SAKURA MATERIAL CO., LTD.; Currency = JPY; Amount = 1,500,000; DeliveryTerm = CIF HAI PHONG", "BEMT2006 lưu một detail theo VoucherNo."),
    ("4", "CI_SSK-MV-2026-02-002.pdf", "COMMERCIALINVOICE", "VoucherNo = SSK-MV-2026/02-002; VoucherDate = 06/09/2026; Currency = JPY; Amount = 1,750,000", "BEMT2006 lưu detail invoice thứ hai."),
    ("5", "TK_107333888810.pdf", "CUSTOMSHEET", "DeclarationNo = 107333888810; VoucherNo = SSK-MV-2026/02-001; DeliveryTerm = CIF; Currency = JPY; Amount = 1,500,000; ClearanceDate = 10/09/2026", "BEMT2006 lưu số tờ khai, invoice liên quan, trị giá, ngày hoàn thành kiểm tra."),
    ("6", "TK_107333888811.pdf", "CUSTOMSHEET", "DeclarationNo = 107333888811; VoucherNo = SSK-MV-2026/02-002; DeliveryTerm = CIF; Currency = JPY; Amount = 1,750,000; ClearanceDate = 10/09/2026", "BEMT2006 lưu tờ khai thứ hai."),
    ("7", "PACKINGLIST_SSK-MV-2026-02-001.pdf", "PACKINGLIST", "Danh sách đóng gói của SSK-MV-2026/02-001; không phải tiêu chí chính trong 9 prompt NVL hiện tại", "Có thể lưu trích xuất, nhưng prompt đối chiếu NVL thường không bắt buộc dùng."),
    ("8", "PACKINGLIST_SSK-MV-2026-02-002.pdf", "PACKINGLIST", "Danh sách đóng gói của SSK-MV-2026/02-002; không phải tiêu chí chính trong 9 prompt NVL hiện tại", "Có thể lưu trích xuất, nhưng prompt đối chiếu NVL thường không bắt buộc dùng."),
    ("9", "STATEMENT_09-2026.pdf", "STATEMENT", "VoucherNo = SSK-MV-2026/02-001, SSK-MV-2026/02-002; Currency = JPY; Amount = 3,250,000", "Là chứng từ thay thế/đối chiếu thêm trong nhóm Invoice/Statement nếu có."),
    ("10", "BL_ABC123456.pdf", "BILL", "BillNo = ABC123456; thông tin vận chuyển; không phải tiêu chí chính của NVL hiện tại", "Có thể lưu SectionType=BILL; Rules NVL hiện tại không dùng làm chứng từ bắt buộc."),
]


EXAMPLE_STEPS = [
    ("1", "WEB ERP9", "BEMF2000Controller.CallAPICompareFileOCRAsync hoặc UpdateResultCompareAI", "Phiếu NVL/09/2026/0011, 2 dòng ĐNTT, 10 file đính kèm", "Tạo CompareFileRequest gồm BEMF2000ViewModel, BEMF2001ViewModels, AttachFiles; POST /CoreAI/FileDataHandler/HandlerFile."),
    ("2", "API ERP9 / 03.SERVICES", "Không thấy route BEM-AI trực tiếp trong source đã rà", "Không có payload riêng trong flow đang rà", "WEB ERP9 dùng BaseApi để gọi thẳng API-AI; 03.SERVICES không biến đổi dữ liệu trong ví dụ này."),
    ("3", "API-AI Controller", "FileDataHandlerController.HandlerFileAsync", "CompareFileRequest từ ERP9", "Gọi ReadFileOrchestratorService.HandleAsync."),
    ("4", "API-AI Orchestrator", "ReadFileOrchestratorService.HandleAsync", "Thông tin phiếu, file, prompt đối chiếu NVL", "Tạo BEMT2003: APK_BEMT2000 = phiếu, StatusProcess = PROCESSING; đưa ReadFileJob vào ChannelJobQueue."),
    ("5", "API-AI Queue", "ChannelJobQueue → ReadFileWorker → ReadFileBackgroundWorkflow", "ReadFileJob của lượt chạy BEMT2003", "Worker lấy job nền và bắt đầu OCR → trích xuất → đối chiếu."),
    ("6", "AI Python OCR", "OcrService.ReadAsync gọi http://192.168.0.134:4444/ocr", "10 file chứng từ", "OCR trả TextMerged và Results. API-AI tạo 10 dòng BEMT2002 với DataSourceType=OCR, FileName, AttachID, RawContent = text OCR."),
    ("7", "AI Python LLM trích xuất", "FormatOCRText → SendPromptWithSumaryResultAsync → /api/ai_llms_models", "Từng file: FileName + text OCR + Prompt_ReadFile", "AI trả JSON sections. ProcessInfomationFileAsync lưu BEMT2005/BEMT2006: BEMT2005.SectionType = PO/RINGI/CUSTOMSHEET...; BEMT2006.FieldID01...FieldID100 = dữ liệu cột đã trích xuất theo ONT1041."),
    ("8", "API-AI chuẩn bị dữ liệu đối chiếu", "BuildAiSectionComparesAsync", "Các dictionary trả về sau trích xuất", "Tạo dataFiles gồm SectionType, field đã trích xuất và FileName. FileName có trong dictionary dùng cho prompt; DB BEMT2006 chỉ lưu tên file nếu ONT1041 map FileName vào một FieldID."),
    ("9", "AI Python Rules/LLM đối chiếu", "BuildCriteriaListAsync → CompareAsync → /api/ai_llms_models", "details ĐNTT + dataFiles + 9 prompt NVL", "Gọi từng prompt đã cấu hình: Số tiền, NCC, Số hóa đơn, Ngày hóa đơn, Loại tiền, Điều kiện giao hàng, Deadline, Ngày hoàn thành kiểm tra, Chữ ký/con dấu."),
    ("10", "AI Python Rules Engine", "Rules_AI_BEM_MEIKO.py", "DnttType=Nguyên vật liệu; FormationID=DATCOC_TRATRUOC; Installment=1; CriterionName", "Resolve rule NVL. Ví dụ Số tiền cần CUSTOMSHEET + RINGI + một trong INVOICE/COMMERCIALINVOICE/STATEMENT; lọc đúng chứng từ rồi trả CriteriaStatus, FileName, Description."),
    ("11", "Database kết quả", "ParseCriteriaResult + SaveData + UpdateCompareResult", "JSON tiêu chí AI trả về", "Lưu BEMT2004 từng tiêu chí; lưu thêm BEMT2002 DataSourceType=CriteriaSynthesis; cập nhật BEMT2003 Status, Percentage, TextConditionFail, StatusProcess=COMPLETED/FAILED."),
]


EXAMPLE_DATA_LIFECYCLE = [
    ("1. Gửi OCR", "10 file vật lý", "API-AI gửi 10 file dạng multipart tới POST /ocr. OCR trả về một danh sách 10 ResultReadFileModel, mỗi phần tử có FileName, AttachID, APK_File và TextContent."),
    ("2. Lưu OCR", "Text OCR của từng file", "API-AI chuẩn bị 10 dòng BEMT2002: DataSourceType=OCR, FileName, AttachID, APK_File, RawContent=text OCR. Các dòng này gắn với cùng APK_BEMT2003 của lượt chạy."),
    ("3. Gọi prompt trích xuất", "Từng FileName + RawContent", "Mỗi file được gọi Prompt BEM_AGENT_READFILE một lần. Template nhận đúng 2 biến: FileName và result (text OCR). AI trả JSON sections."),
    ("4. Lưu dữ liệu trích xuất", "sections.master và sections.details[]", "API-AI lưu SectionType, tổng tiền, loại tiền, chữ ký vào BEMT2005; map các cột detail vào BEMT2006.FieldID01...FieldID100 theo cấu hình ONT1041."),
    ("5. Chuẩn bị đối chiếu", "Dữ liệu phiếu + dòng ĐNTT + dữ liệu 10 file", "API-AI tạo dataFiles từ kết quả trích xuất và giữ FileName trong dictionary để prompt biết chứng từ nào cung cấp dữ liệu."),
    ("6. Gọi prompt đối chiếu", "datas + details + dataFiles", "Với NVL, mỗi prompt tiêu chí được gọi lần lượt. Rules Python chọn nhóm chứng từ phải dùng theo DnttType, FormationID, Installment, CriterionName."),
    ("7. Lưu kết quả", "JSON criteria của từng tiêu chí", "API-AI lưu BEMT2004: CriteriaName, CriteriaStatus, Description, FileName, DueDateAI; sau cùng cập nhật BEMT2003 là COMPLETED và kết quả tổng OK/NG, Percentage."),
]


EXAMPLE_DB_ROWS = [
    ("BEMT2003", "01 dòng cho lượt chạy", "APK_BEMT2000 = phiếu; StatusProcess: PROCESSING -> COMPLETED; Status = OK; Percentage = 100%; TextContentOCR là text OCR đã gộp."),
    ("BEMT2002", "10 dòng OCR; thêm tối đa 09 dòng tổng hợp tiêu chí", "Ví dụ OCR file CI: FileName=CI_SSK-MV-2026-02-001.pdf; AttachID/APK_File; DataSourceType=OCR; RawContent=text OCR. Dòng tổng hợp có DataSourceType=CriteriaSynthesis và APK_BEMT2004."),
    ("BEMT2005", "Khoảng 10 header nếu mỗi file trả 01 section", "Số dòng thực tế phụ thuộc JSON sections. Ví dụ Commercial Invoice: SectionType=COMMERCIALINVOICE; TotalAmount=1,500,000; TotalCurrency=JPY; PromptSystem là prompt hệ thống đã dùng."),
    ("BEMT2006", "01 hoặc nhiều dòng chi tiết cho mỗi BEMT2005", "Ví dụ invoice: VoucherNo, VoucherDate, SupplierName, DeliveryTerm, Currency, Amount được map vào các FieldID theo ONT1041. Không cố định số FieldID vì phụ thuộc cấu hình."),
    ("BEMT2004", "09 dòng nếu cả 9 tiêu chí NVL đều trả kết quả", "Ví dụ Số tiền: CriteriaStatus=OK; Description nêu các giá trị khớp; FileName chỉ có khi AI cần chỉ ra chứng từ liên quan; DueDateAI dùng cho tiêu chí Deadline."),
]


EXAMPLE_PROMPT_ROWS = [
    ("Prompt trích xuất", "BEM_AGENT_READFILE trong ONT1042", "FileName = CI_SSK-MV-2026-02-001.pdf\nresult = text OCR của đúng file này", "JSON sections: SectionType=COMMERCIALINVOICE; total=1,500,000 JPY; detail có InvoiceNo, ngày hóa đơn, NCC, Incoterm, tiền tệ, số tiền."),
    ("Prompt Số tiền", "Prompt_WareHouse_Amount", "datas = thông tin phiếu NVL\ndetails = 2 dòng ĐNTT: invoice, tiền yêu cầu, Ringi\ndataFiles = dữ liệu 10 file đã trích xuất + FileName", "JSON criteria: CriteriaName=Số tiền; CriteriaStatus=OK/NG; Description; FileName; DueDateAI."),
]


def example_section_markdown():
    files_rows = "\n".join(
        f"| {no} | `{file}` | `{section}` | {extract} | {storage} |"
        for no, file, section, extract, storage in EXAMPLE_FILES
    )
    step_rows = "\n".join(
        f"| {no} | {area} | `{code}` | {input_data} | {output_data} |"
        for no, area, code, input_data, output_data in EXAMPLE_STEPS
    )
    return f"""

## 4. Ví dụ mẫu: một phiếu NVL có 10 file chứng từ

Ví dụ minh họa theo cấu trúc code và DB hiện tại: phiếu `NVL/09/2026/0011` có tổng tiền yêu cầu `3,250,000 JPY`, nhà cung cấp `SAKURA MATERIAL CO., LTD.`, Ringi `RG-2026-088`, PO `PO-2026-0901`, gồm 10 file đính kèm.

### 4.1. Mười file đi qua OCR và trích xuất như thế nào

| STT | File | SectionType AI nhận diện | Dữ liệu chính sau trích xuất | DB lưu ở đâu |
|---|---|---|---|---|
{files_rows}

### 4.2. Luồng chi tiết từ đầu đến cuối

| Bước | Khu vực | Code/API chính | Dữ liệu vào | Dữ liệu ra và nơi lưu |
|---|---|---|---|---|
{step_rows}

### 4.3. Dữ liệu đi qua AI và Database như thế nào

- `Prompt_ReadFile` chỉ làm nhiệm vụ đọc OCR và trả JSON `sections`: `master.SectionType`, `SectionTitle`, `TotalAmount`, `TotalCurrency`, `Signature`, và `details[]`.
- API-AI chuyển `sections` thành `BEMT2005` và `BEMT2006`. `BEMT2005` giữ loại chứng từ; `BEMT2006` giữ từng dòng dữ liệu đã trích xuất qua `FieldID01...FieldID100`.
- Khi đối chiếu, API-AI truyền vào prompt đối chiếu các nhóm dữ liệu chính: `details` của phiếu ĐNTT và `dataFiles` là dữ liệu đã trích xuất từ chứng từ.
- Với bộ Nguyên vật liệu, `Rules_AI_BEM_MEIKO.py` kiểm tra rule theo `DnttType / FormationID / Installment / CriterionName` để biết tiêu chí đó cần chứng từ nào.
- Ví dụ tiêu chí `Số tiền`: rule NVL yêu cầu có `CUSTOMSHEET`, `RINGI` và một trong `INVOICE / COMMERCIALINVOICE / STATEMENT`. Trường hợp này AI chọn 2 Commercial Invoice đúng số hóa đơn của ĐNTT, đối chiếu với 2 tờ khai và Ringi. `STATEMENT` chỉ kiểm tra thêm, không cộng trùng. Tổng đúng là `1,500,000 + 1,750,000 = 3,250,000 JPY`, khớp ĐNTT và Ringi thì lưu `BEMT2004.CriteriaStatus = OK`.
- Nếu một tiêu chí NG, `BEMT2004.FileName` lưu file liên quan đến lỗi và `Description` lưu lý do; `BEMT2003.Percentage` giảm theo số tiêu chí NG.

### 4.4. Prompt thực nhận dữ liệu gì

**a. Prompt trích xuất cho file `CI_SSK-MV-2026-02-001.pdf`**

API-AI lấy prompt `BEM_AGENT_READFILE` trong `ONT1042`, sau đó truyền hai biến vào template: `FileName` và `result` (text OCR). Nội dung prompt sau render có dạng:

```text
Tên file: CI_SSK-MV-2026-02-001.pdf
Dữ liệu OCR: COMMERCIAL INVOICE ... Invoice No SSK-MV-2026/02-001 ...
Supplier SAKURA MATERIAL CO., LTD. ... CIF HAI PHONG ... Total JPY 1,500,000
```

AI Python trả JSON `sections`, ví dụ rút gọn:

```json
{{
  "sections": [{{
    "master": {{"SectionType": "COMMERCIALINVOICE", "TotalAmount": 1500000, "TotalCurrency": "JPY"}},
    "details": [{{"OrderNo": "1", "VoucherNo": "SSK-MV-2026/02-001", "VoucherDate": "05/09/2026", "SupplierName": "SAKURA MATERIAL CO., LTD.", "DeliveryTerm": "CIF HAI PHONG", "Currency": "JPY", "Amount": 1500000}}]
  }}]
}}
```

API-AI lưu `master` vào `BEMT2005` và map từng field của `details` vào `BEMT2006.FieldID01...FieldID100` theo cấu hình `ONT1041`. Không nên ghi cố định “VoucherNo luôn là FieldID01”, vì FieldID thực tế phụ thuộc mapping đang cấu hình.

**b. Prompt đối chiếu tiêu chí Số tiền**

API-AI dùng `Prompt_WareHouse_Amount`; template nhận `datas` (thông tin phiếu), `details` (dòng ĐNTT) và `dataFiles` (dữ liệu chứng từ đã trích xuất). Ví dụ nội dung render rút gọn:

```text
PromptType: Đối chiếu | DnttType: Nguyên vật liệu
FormationID: DATCOC_TRATRUOC | Installment: 1 | CriterionName: Số tiền

ĐNTT: SSK-MV-2026/02-001 = 1,500,000 JPY; SSK-MV-2026/02-002 = 1,750,000 JPY; Ringi = RG-2026-088
Tổng tiền yêu cầu: 3,250,000 JPY

COMMERCIALINVOICE: SSK-MV-2026/02-001 = 1,500,000 JPY | file CI_SSK-MV-2026-02-001.pdf
COMMERCIALINVOICE: SSK-MV-2026/02-002 = 1,750,000 JPY | file CI_SSK-MV-2026-02-002.pdf
CUSTOMSHEET: SSK-MV-2026/02-001 = 1,500,000 JPY | file TK_107333888810.pdf
CUSTOMSHEET: SSK-MV-2026/02-002 = 1,750,000 JPY | file TK_107333888811.pdf
RINGI: 3,250,000 JPY | file RINGI_RG-2026-088.pdf
```

AI trả đúng một JSON `criteria`. Nếu toàn bộ giá trị khớp, API-AI lưu ví dụ: `BEMT2004.CriteriaName = Số tiền`, `CriteriaStatus = OK`, `FileName = ""`, `Description = "Số tiền đã hoàn toàn khớp với nhau."`.

### 4.5. Dữ liệu còn lại trong DB sau khi ví dụ hoàn tất

| Bảng | Số dòng dự kiến | Dữ liệu ví dụ |
|---|---|---|
{chr(10).join(f"| `{table}` | {count} | {data} |" for table, count, data in EXAMPLE_DB_ROWS)}
"""


def add_example_table(document, title, headers, rows, font_size=7.2):
    add_text(document, title, 10.5, True, WD_ALIGN_PARAGRAPH.LEFT, "17365D", 2)
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "17365D")
        write_cell_left(cell, value, 8.0, True, "FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_shading(cell, "F3F7FC" if row_index % 2 == 0 else "FFFFFF")
            write_cell_left(cell, value, font_size)


def add_example_section(document):
    document.add_page_break()
    heading(document, "4. Ví dụ mẫu: một phiếu NVL có 10 file chứng từ")
    add_text(
        document,
        "Ví dụ minh họa theo cấu trúc code và DB hiện tại: phiếu NVL/09/2026/0011 có tổng tiền yêu cầu 3,250,000 JPY, nhà cung cấp SAKURA MATERIAL CO., LTD., Ringi RG-2026-088, PO PO-2026-0901 và 10 file đính kèm.",
        9.2,
        False,
        WD_ALIGN_PARAGRAPH.LEFT,
        "333333",
        3,
    )
    add_example_table(
        document,
        "4.1. Mười file đi qua OCR và trích xuất",
        ["STT", "File", "SectionType", "Dữ liệu chính sau trích xuất", "DB lưu ở đâu"],
        EXAMPLE_FILES,
        6.5,
    )
    document.add_page_break()
    heading(document, "4.2. Luồng chi tiết từ đầu đến cuối")
    add_example_table(
        document,
        "Một lượt chạy đi qua WEB ERP9, API-AI, AI Python và DB như sau",
        ["Bước", "Khu vực", "Code/API chính", "Dữ liệu vào", "Dữ liệu ra và nơi lưu"],
        EXAMPLE_STEPS,
        6.4,
    )
    document.add_page_break()
    heading(document, "4.3. Dữ liệu đi qua AI và Database như thế nào")
    add_example_table(
        document,
        "Ví dụ chỉ minh họa dữ liệu của phiếu NVL/09/2026/0011; số APK thực tế do Database tự sinh",
        ["Giai đoạn", "Dữ liệu được dùng", "AI/API-AI làm gì và lưu ở đâu"],
        EXAMPLE_DATA_LIFECYCLE,
        7.2,
    )
    document.add_page_break()
    heading(document, "4.4. Prompt thực nhận dữ liệu gì")
    add_example_table(
        document,
        "Prompt không tự đọc file vật lý: OCR đọc file trước; LLM chỉ nhận text OCR hoặc dữ liệu đã trích xuất",
        ["Lần gọi", "Prompt", "Biến được truyền", "Kết quả AI trả về"],
        EXAMPLE_PROMPT_ROWS,
        7.4,
    )
    add_text(document, "Ví dụ dữ liệu prompt trích xuất: FileName = CI_SSK-MV-2026-02-001.pdf; result = 'COMMERCIAL INVOICE ... Invoice No SSK-MV-2026/02-001 ... Supplier SAKURA MATERIAL CO., LTD. ... CIF HAI PHONG ... Total JPY 1,500,000'.", 8.2, False, WD_ALIGN_PARAGRAPH.LEFT, "333333", 2)
    add_text(document, "Ví dụ prompt Số tiền: details có 2 dòng ĐNTT 1,500,000 + 1,750,000 JPY; dataFiles có 2 Commercial Invoice, 2 tờ khai, Ringi và STATEMENT. AI chọn 2 Commercial Invoice đúng số hóa đơn để tính, dùng tờ khai và Ringi để đối chiếu; STATEMENT chỉ kiểm tra thêm, không cộng trùng. Khi tổng khớp 3,250,000 JPY, AI trả CriteriaStatus = OK.", 8.2, False, WD_ALIGN_PARAGRAPH.LEFT, "333333", 3)
    add_example_table(
        document,
        "4.5. Dữ liệu còn lại trong Database khi ví dụ hoàn tất",
        ["Bảng", "Số dòng dự kiến", "Dữ liệu ví dụ"],
        EXAMPLE_DB_ROWS,
        7.2,
    )
    add_text(document, "Cách đọc nhanh: BEMT2002 cho biết OCR đọc được gì; BEMT2005/BEMT2006 cho biết LLM trích xuất được gì; BEMT2004 cho biết từng tiêu chí OK/NG; BEMT2003 cho biết kết quả tổng của lượt chạy.", 8.5, True, WD_ALIGN_PARAGRAPH.LEFT, "17365D", 0)


def markdown_content():
    return """# Cơ chế mã nguồn đối chiếu AI BEM

Ngày cập nhật: 02/10/2026  
Mục tiêu: nhìn nhanh là hiểu source nào làm gì và Database lưu dữ liệu AI ở đâu.

## 1. Cấu trúc source

```text
ERP9 (WEB + API)                         API AI .NET                         AI BEM Python                         Database
└─ BEMF2000Controller.cs                 ├─ FileDataHandlerController        ├─ App/main_iis.py                   ├─ BEMT2003: lượt chạy
   • Nhận thao tác / lịch tự động         ├─ ReadFileOrchestratorService       │  • /ocr                             ├─ BEMT2002: OCR
   • Lấy phiếu và file                    ├─ ChannelJobQueue / ReadFileWorker  │  • /api/ai_llms_models              ├─ BEMT2005 / 2006: trích xuất
   • Gọi API AI                           └─ Tạo run, Queue, gọi AI, ghi DB   ├─ Rules_AI_BEM_MEIKO.py             └─ BEMT2004: đối chiếu
                                                                          └─ LLMs_BE/engine.py
```

- Source đã rà cho thấy WEB ERP9 gọi trực tiếp API-AI qua `HandlerFile`; chưa thấy `03.SERVICES` đứng trực tiếp trong flow BEM-AI hiện tại.
- ERP9 chỉ là nơi tạo yêu cầu, gửi file và hiển thị kết quả; API-AI và AI Python là hai khu vực xử lý chính.

## 2. Luồng một phiếu

1. ERP9 nhận lệnh thủ công hoặc lịch tự động, lấy phiếu ĐNTT và file đính kèm.
2. API-AI tạo `BEMT2003 = PROCESSING`, đưa job vào `ChannelJobQueue` và xử lý nền.
3. AI Python OCR file, trả text; API-AI lưu text OCR vào `BEMT2002`.
4. AI Python dùng LLM trích xuất text thành JSON chứng từ; API-AI lưu loại chứng từ vào `BEMT2005` và dữ liệu từng cột vào `BEMT2006`.
5. AI Python áp dụng Rules và LLM để đối chiếu; API-AI lưu kết quả từng tiêu chí vào `BEMT2004`, cập nhật kết quả tổng vào `BEMT2003`; ERP9 đọc để hiển thị.

## 3. Database lưu gì

| Cần tìm | Bảng | Trường chính | Hiểu đơn giản |
|---|---|---|---|
| Phiếu đang chạy tới đâu | `BEMT2003` | `APK_BEMT2000`, `StatusProcess`, `Status`, `Percentage`, `TextConditionFail` | Một bản ghi = một lượt chạy AI của phiếu. |
| Tên file và OCR | `BEMT2002` | `FileName`, `APK_File`, `AttachID`, `RawContent`, `ConsolidatedContent` | Biết file nào OCR và OCR đọc ra text gì. |
| Loại chứng từ | `BEMT2005` | `SectionType`, `SectionTitle`, `SectionOrder`, `TotalAmount`, `TotalCurrency` | Header của nhóm dữ liệu/chứng từ AI đã trích xuất. |
| Dữ liệu các cột đã trích xuất | `BEMT2006` | `OrderNo`, `Description`, `FieldID01`...`FieldID100` | Mỗi `FieldID` là một cột dữ liệu do `ONT1041` mapping. |
| Kết quả đối chiếu | `BEMT2004` | `CriteriaName`, `CriteriaStatus`, `Description`, `FileName`, `DueDateAI` | Kết quả OK/NG và lý do của từng tiêu chí. |

**Cách đọc nhanh:** `BEMT2002.FileName` là tên file OCR; `BEMT2005.SectionType/SectionTitle` là loại hoặc nhóm chứng từ; `BEMT2006.FieldID01...FieldID100` là dữ liệu các cột đã trích xuất; `BEMT2004.CriteriaStatus/Description` là kết quả đối chiếu.
""" + example_section_markdown()


def main():
    images = []
    for filename, image in [
        ("01_cau_truc_source.png", make_source_structure()),
        ("02_luong_mot_phieu.png", make_processing_flow()),
        ("03_db_luu_du_lieu.png", make_database_structure()),
    ]:
        path = IMAGE_DIR / filename
        image.save(path, quality=95)
        images.append(path)
    MD_PATH.write_text(markdown_content(), encoding="utf-8")

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.36)
    section.right_margin = Inches(0.36)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    add_text(document, "CƠ CHẾ MÃ NGUỒN ĐỐI CHIẾU AI BEM", 21, True, WD_ALIGN_PARAGRAPH.CENTER, "17365D", 1)
    add_text(document, "Bản tóm tắt cấu trúc source, luồng xử lý và Database | Cập nhật 02/10/2026", 10.5, False, WD_ALIGN_PARAGRAPH.CENTER, "666666", 5)
    add_text(document, "Kết luận: ERP9 tạo yêu cầu và đọc kết quả; API-AI điều phối Queue và ghi DB; AI Python OCR, trích xuất và đối chiếu.", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER, "17365D", 5)
    heading(document, "1. Cấu trúc source")
    picture(document, images[0], "Hình 1. Bốn khu vực source và trách nhiệm chính")
    document.add_page_break()
    heading(document, "2. Luồng xử lý một phiếu")
    picture(document, images[1], "Hình 2. Từ thao tác ERP9 đến khi có kết quả hiển thị")
    add_text(document, "Quy ước: OCR trả text; LLM trích xuất trả JSON dữ liệu chứng từ; Rules/LLM đối chiếu trả OK/NG và lý do từng tiêu chí.", 10, False, WD_ALIGN_PARAGRAPH.CENTER, "555555", 1)
    document.add_page_break()
    heading(document, "3. Database lưu dữ liệu theo từng lượt chạy")
    picture(document, images[2], "Hình 3. Field chính lưu tên file, loại chứng từ, dữ liệu trích xuất và kết quả đối chiếu")
    add_text(document, "Lưu ý: BEMT2003 là điểm neo của một lần chạy. Khi chạy lại, hệ thống tạo lượt chạy mới để lưu dữ liệu và kết quả của lần đó.", 9.5, False, WD_ALIGN_PARAGRAPH.LEFT, "555555", 0)
    add_example_section(document)
    document.save(DOCX_PATH)
    print(f"Created: {DOCX_PATH}")
    print(f"Created: {MD_PATH}")
    for path in images:
        print(f"Image: {path}")


if __name__ == "__main__":
    main()
