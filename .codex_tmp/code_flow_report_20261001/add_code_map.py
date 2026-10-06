from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
md_path=base/'Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md'
docx_path=base/'Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx'
section_md='''
## 1A. Muốn xem code thì mở file nào

| Khu vực | File và hàm chính | Code đang xử lý gì |
|---|---|---|
| WEB – giao diện | `BEMF2002.js`: `hanldeCompareOCR` (dòng 770), `CallAPICompareFileOCRAsync` (dòng 796) | Nhận thao tác bấm đối chiếu; kiểm tra phiếu đã có kết quả cũ hay chưa; nếu có thì hỏi người dùng có chạy lại hay không; gọi controller ERP. |
| WEB – backend ERP | `BEMF2000Controller.cs`: `CallAPICompareFileOCRAsync` (dòng 2267), `GetAttachFileModels` (dòng 3573), `CompareWithAI` (dòng 3526) | Lấy dữ liệu phiếu, chi tiết và file đính kèm; upload file sang API AI; đóng gói `CompareFileRequest` và gửi API AI. |
| API ERP Services | `CoreStartup.cs`: đăng ký Quartz (dòng 119, 188); `AutomationJobBusiness.cs` | Cung cấp cơ chế job/lịch tổng quát. Chưa thấy code BEM cụ thể chọn phiếu, giờ chạy hoặc số phiếu mỗi đợt trong source đã rà. Luồng thủ công không đi qua khu vực này. |
| API AI – nhận và xếp hàng | `FileDataHandlerController.cs`: `HandlerFileAsync` (dòng 31); `ReadFileOrchestratorService.cs` (dòng 71, 88); `ChannelJobQueue.cs` (dòng 10, 21, 24) | Nhận request; tạo bản ghi lượt chạy `BEMT2003` trạng thái `PROCESSING`; đưa job đối chiếu vào Queue. Queue có sức chứa cấu hình 200 job. |
| API AI – xử lý nền/OCR/lưu DB | `ReadFileWorker.cs`: `ExecuteAsync` (dòng 22); `ReadFileJobExecutor.cs`: `ExecuteAsync` (dòng 18); `ReadFileBackgroundWorkflow.cs`: `RunReadFileFlowAsync` (dòng 474), `SaveCompareResultAsync` (dòng 568); `OcrService.cs`: `ReadAsync` (dòng 45) | Worker lấy job khỏi Queue; lựa chọn chạy tuần tự/song song theo cấu hình; gọi OCR; gọi các prompt trích xuất và đối chiếu; lưu kết quả từng tiêu chí vào `BEMT2004` và cập nhật kết quả tổng/lượt chạy tại `BEMT2003`. OCR tối đa 7 file song song trong một phiếu. |
| AI BEM – OCR | `App/main_iis.py`: route `/ocr` (dòng 374) | Nhận file từ API AI, chạy OCR và trả text để các bước sau đọc được nội dung chứng từ. |
| AI BEM – LLM/Rules | `App/main_iis.py`: route `/llms/api/ai_llms_models` (dòng 1636), gọi Rules (dòng 2128); `Rules_AI_BEM_MEIKO.py`: `process_ai_llms_models_rules` (dòng 3800) | Nhận prompt và text OCR; tách nhánh **Trích xuất** hoặc **Đối chiếu**; trích trường dữ liệu từ chứng từ, lọc chứng từ liên quan theo tiêu chí, áp dụng Rules/logic đặc biệt hoặc gọi LLM; trả kết quả về API AI. |

**Đường đi thực tế khi bấm thủ công:** `BEMF2002.js` → `BEMF2000Controller.cs` → `FileDataHandlerController.cs` → `ReadFileOrchestratorService.cs` → `ChannelJobQueue` → `ReadFileWorker` → `ReadFileBackgroundWorkflow` → `OcrService` / `main_iis.py` → `Rules_AI_BEM_MEIKO.py` → lưu DB → ERP9 hiển thị.

'''
md=md_path.read_text(encoding='utf-8')
marker='## 2. Luồng khi người dùng đối chiếu thủ công\n'
if '## 1A. Muốn xem code thì mở file nào' not in md:
    md=md.replace(marker,section_md+marker)
    md_path.write_text(md,encoding='utf-8')

def set_font(run,size=11,bold=False,color=None):
    run.font.name='Times New Roman'; run._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); run._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); tcPr.append(e)
def border(cell):
    tcPr=cell._tc.get_or_add_tcPr(); b=tcPr.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9E2F3'); b.append(e)
def add_heading(doc,text):
    p=doc.add_paragraph(); p.style='Heading 1'; p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(6); r=p.add_run(text); set_font(r,17,True,'17365D')
def add_para(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.05; r=p.add_run(text); set_font(r,10.5)

doc=Document(docx_path)
if not any(p.text=='1A. Muốn xem code thì mở file nào' for p in doc.paragraphs):
    # This table is placed at the end to avoid breaking the existing illustrated structure.
    add_heading(doc,'Phụ lục A. Muốn xem code thì mở file nào')
    add_para(doc,'Bảng dưới đây chỉ đúng khu vực code, hàm chính và việc xử lý. Có thể mở file theo dòng tham khảo để kiểm tra chi tiết.')
    rows=[
      ('WEB – giao diện','BEMF2002.js: hanldeCompareOCR (770); CallAPICompareFileOCRAsync (796)','Nhận nút bấm; kiểm tra có kết quả cũ; hỏi xác nhận chạy lại; gọi controller ERP.'),
      ('WEB – backend ERP','BEMF2000Controller.cs: CallAPICompareFileOCRAsync (2267); GetAttachFileModels (3573); CompareWithAI (3526)','Lấy phiếu, chi tiết, file; upload file sang API AI; gửi CompareFileRequest.'),
      ('API ERP Services','CoreStartup.cs: Quartz (119, 188); AutomationJobBusiness.cs','Cơ chế lịch/job chung. Chưa thấy code BEM cụ thể chọn phiếu, giờ chạy hoặc số phiếu mỗi đợt.'),
      ('API AI – Queue','FileDataHandlerController.cs: HandlerFileAsync (31); ReadFileOrchestratorService.cs (71, 88); ChannelJobQueue.cs (10, 21, 24)','Nhận request; tạo BEMT2003 PROCESSING; đưa job vào Queue 200 job.'),
      ('API AI – xử lý','ReadFileWorker.cs (22, 31, 43); ReadFileJobExecutor.cs (18); ReadFileBackgroundWorkflow.cs (474, 568); OcrService.cs (45, 81)','Lấy job; gọi OCR/trích xuất/đối chiếu; lưu BEMT2004 và cập nhật BEMT2003. OCR tối đa 7 file song song/phiếu.'),
      ('AI BEM – OCR','App/main_iis.py: /ocr (374)','Nhận file và trả text OCR.'),
      ('AI BEM – LLM/Rules','App/main_iis.py: /llms/api/ai_llms_models (1636, 2128); Rules_AI_BEM_MEIKO.py: process_ai_llms_models_rules (3800)','Tách Trích xuất/Đối chiếu; gọi LLM, áp dụng Rules, lọc chứng từ và trả kết quả.'),
    ]
    table=doc.add_table(rows=1,cols=3); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    headers=['Khu vực','Mở code tại','Code xử lý gì']
    for c,h in zip(table.rows[0].cells,headers):
        shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); set_font(r,10,True,'FFFFFF')
    for a,b,cval in rows:
        cells=table.add_row().cells
        for cell,txt in zip(cells,[a,b,cval]):
            border(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); r=p.add_run(txt); set_font(r,9.2,cell==cells[0])
    add_para(doc,'Đường đi thủ công: BEMF2002.js → BEMF2000Controller.cs → FileDataHandlerController.cs → ReadFileOrchestratorService.cs → ChannelJobQueue → ReadFileWorker → ReadFileBackgroundWorkflow → OcrService / main_iis.py → Rules_AI_BEM_MEIKO.py → lưu DB → ERP9 hiển thị.')
doc.save(docx_path)
print('Updated:',md_path)
print('Updated:',docx_path)
