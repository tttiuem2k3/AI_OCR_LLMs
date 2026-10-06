from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMG_DIR=BASE/'Temp'/'code_flow_simplified_20261001'; IMG_DIR.mkdir(parents=True,exist_ok=True)
MD=BASE/'Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md'
DOCX=BASE/'Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx'
FONT=r'C:\Windows\Fonts\times.ttf'; FONT_B=r'C:\Windows\Fonts\timesbd.ttf'
NAVY='#17365D'; BLUE='#2F75B5'; TEAL='#168A8A'; GOLD='#D6A300'; ORANGE='#ED7D31'; PURPLE='#7030A0'; LINE='#6B8299'; DARK='#1F2933'; SOFT='#F6F9FC'

def f(size,bold=False): return ImageFont.truetype(FONT_B if bold else FONT,size)
def center(d,xy,text,font,fill=DARK,leading=6):
    x1,y1,x2,y2=xy; lines=[]
    for raw in str(text).split('\n'):
        words=raw.split(); line=''
        if not words: lines.append(''); continue
        for w in words:
            test=(line+' '+w).strip()
            if d.textbbox((0,0),test,font=font)[2] <= x2-x1: line=test
            else:
                if line: lines.append(line)
                line=w
        if line: lines.append(line)
    heights=[max(1,d.textbbox((0,0),s,font=font)[3]-d.textbbox((0,0),s,font=font)[1]) for s in lines]
    y=y1+((y2-y1)-(sum(heights)+leading*(len(lines)-1)))/2
    for s,h in zip(lines,heights):
        b=d.textbbox((0,0),s,font=font); d.text((x1+(x2-x1-(b[2]-b[0]))/2,y),s,font=font,fill=fill); y+=h+leading

def box(d,xy,title,body,color,body_size=26):
    x1,y1,x2,y2=xy
    d.rounded_rectangle(xy,22,fill='white',outline=color,width=4)
    d.rounded_rectangle((x1,y1,x2,y1+58),22,fill=color,outline=color)
    d.rectangle((x1,y1+31,x2,y1+58),fill=color)
    center(d,(x1+8,y1+6,x2-8,y1+51),title,f(27,True),'white')
    center(d,(x1+18,y1+68,x2-18,y2-15),body,f(body_size),DARK)

def arrow(d,a,b,color=LINE,width=7,head=24,dashed=False):
    import math
    x1,y1=a; x2,y2=b
    if dashed:
        n=12
        for i in range(n):
            if i%2==0:
                t=i/n; u=(i+1)/n; d.line((x1+(x2-x1)*t,y1+(y2-y1)*t,x1+(x2-x1)*u,y1+(y2-y1)*u),fill=color,width=width)
    else: d.line((x1,y1,x2,y2),fill=color,width=width)
    ang=math.atan2(y2-y1,x2-x1)
    p1=(x2-head*math.cos(ang-.52),y2-head*math.sin(ang-.52)); p2=(x2-head*math.cos(ang+.52),y2-head*math.sin(ang+.52))
    d.polygon([(x2,y2),p1,p2],fill=color)

def title_canvas(title,subtitle,w,h):
    im=Image.new('RGB',(w,h),'white'); d=ImageDraw.Draw(im)
    d.rectangle((0,0,w,105),fill=NAVY); d.text((48,18),title,font=f(41,True),fill='white'); d.text((50,70),subtitle,font=f(22),fill='#DCE6F1')
    return im,d

# Diagram 1: system map, no overlapping lines
im,d=title_canvas('Bản đồ vận hành hiện tại','Hai cách kích hoạt cùng đi vào API-AI; Queue, OCR, LLM/Rules và DB nằm ở các khu vực như bên dưới',1900,960)
box(d,(55,195,375,390),'WEB ERP9','Người dùng bấm\nĐối chiếu AI\n\nSau cùng đọc DB\nvà hiển thị kết quả',BLUE,27)
box(d,(55,530,375,725),'API ERP','Có Quartz/lịch chung\n\nChưa xác nhận được\njob BEM cụ thể',TEAL,27)
box(d,(540,300,935,640),'API-AI','1. Nhận request\n2. Tạo BEMT2003 = PROCESSING\n3. Queue trong RAM: 200 job\n4. Worker lấy job xử lý',GOLD,27)
box(d,(1085,300,1440,640),'AI Python','OCR: file → text\n\nLLM: trích xuất dữ liệu\n\nRules + LLM: đối chiếu',ORANGE,27)
box(d,(1580,300,1845,640),'DB','BEMT2003: lượt chạy\nBEMT2002: OCR\nBEMT2005/2006: trích xuất\nBEMT2004: tiêu chí',PURPLE,23)
arrow(d,(375,292),(540,420)); arrow(d,(375,627),(540,520),dashed=True)
arrow(d,(935,470),(1085,470)); arrow(d,(1440,470),(1580,470))
center(d,(425,205,520,270),'Thủ công',f(20,True),'#4E667B')
center(d,(400,650,530,710),'Tự động\n(cần xác nhận job BEM)',f(18,True),'#8A3E00')
box(d,(640,735,1260,920),'Điểm cần nhớ','Queue nằm trong RAM của API-AI. DB chỉ lưu trạng thái và kết quả; DB không phải nơi xếp hàng job.',NAVY,23)
im.save(IMG_DIR/'01_ban_do_van_hanh.png')

# Diagram 2: one voucher state
im,d=title_canvas('Một phiếu được xử lý và ghi DB như thế nào','Mỗi mũi tên là một bước; DB cho biết phiếu đang chạy tới đâu và đã có kết quả gì',1900,910)
steps=[
 ('1. Nhận lệnh','WEB/API-AI nhận phiếu',BLUE),
 ('2. PROCESSING','BEMT2003 tạo lượt chạy',PURPLE),
 ('3. Queue','Chờ Worker trong RAM',GOLD),
 ('4. OCR','AI Python đọc file\nBEMT2002 lưu text',ORANGE),
 ('5. Trích xuất','LLM đọc dữ liệu\nBEMT2005/2006 lưu',ORANGE),
 ('6. Đối chiếu','Rules + LLM\nBEMT2004 lưu tiêu chí',ORANGE),
 ('7. Kết thúc','BEMT2003 cập nhật\nCOMPLETED hoặc FAILED',PURPLE),
]
x=38
for i,(t,b,c) in enumerate(steps):
    box(d,(x,270,x+230,605),t,b,c,23)
    if i < len(steps)-1: arrow(d,(x+230,438),(x+260,438),LINE,6,22)
    x+=260
center(d,(140,700,1760,805),'Nếu COMPLETED: DB có kết quả tổng OK/NG, %, dữ liệu trích xuất và kết quả từng tiêu chí. Nếu FAILED: lượt chạy không hoàn tất, cần xem lỗi để chạy lại có kiểm soát.',f(27),DARK)
im.save(IMG_DIR/'02_trang_thai_mot_phieu.png')

md='''# Cơ chế vận hành AI BEM

Ngày cập nhật: 01/10/2026  
Mục tiêu: nhìn một lần là biết phần nào xử lý việc gì, Queue nằm ở đâu và DB lưu gì.

## 1. Hiện trạng đang chạy

![Bản đồ vận hành](Temp/code_flow_simplified_20261001/01_ban_do_van_hanh.png)

- **WEB ERP9:** người dùng bấm đối chiếu, gửi hồ sơ và hiển thị kết quả từ DB.
- **API ERP:** source chỉ cho thấy Quartz/lịch chung; chưa xác nhận được code BEM cụ thể chọn phiếu và gửi phiếu tự động.
- **API-AI:** nhận yêu cầu, tạo lượt chạy, xếp Queue và điều phối xử lý nền.
- **AI Python:** OCR đọc file; LLM trích xuất dữ liệu; Rules + LLM đối chiếu.
- **DB:** lưu từng lượt chạy, OCR, dữ liệu trích xuất và kết quả đối chiếu.

## 2. Năm thành phần cần quan tâm

| Thành phần | Nằm tại | Hiện trạng |
|---|---|---|
| Queue | API-AI | `ChannelJobQueue` là Queue trong RAM, Singleton, sức chứa 200 job. `ReadFileWorker` lấy job từ Queue. Không phải DB queue. |
| OCR | AI Python | Route `/ocr` xử lý file thành text. `OcrService` ở API-AI chỉ gọi OCR; tối đa 7 file song song trong một phiếu. |
| Rule Engine | AI Python | Hàm `process_ai_llms_models_rules` chọn nhánh trích xuất/đối chiếu và áp dụng Rules. Không phải service riêng. |
| LLM | AI Python | Route `/llms/api/ai_llms_models` gọi LLM để trích xuất hoặc đối chiếu. |
| DB | SQL DB | Lưu theo APK của lượt chạy `BEMT2003`; không trực tiếp xử lý OCR/LLM. |

## 3. Một phiếu được xử lý và ghi DB như thế nào

![Trạng thái một phiếu](Temp/code_flow_simplified_20261001/02_trang_thai_mot_phieu.png)

| Khi nào | DB lưu tại | Ý nghĩa |
|---|---|---|
| Bắt đầu | `BEMT2003` | Tạo lượt chạy với `StatusProcess = PROCESSING`. |
| OCR xong | `BEMT2002` | Text OCR và dữ liệu nguồn/tổng hợp. |
| Trích xuất xong | `BEMT2005`, `BEMT2006` | Dữ liệu đọc từ chứng từ. |
| Đối chiếu xong | `BEMT2004` | Kết quả từng tiêu chí OK/NG và giải thích. |
| Hoàn tất | `BEMT2003` | `COMPLETED`, kết quả tổng OK/NG và %. |
| Lỗi | `BEMT2003` | `FAILED`, nhận biết lượt chạy không hoàn tất. |

## 4. Mapping code cần xem

| Khu vực | File/hàm chính | Việc xử lý |
|---|---|---|
| WEB | `BEMF2002.js`: `hanldeCompareOCR`, `CallAPICompareFileOCRAsync`; `BEMF2000Controller.cs`: `CallAPICompareFileOCRAsync`, `GetAttachFileModels`, `CompareWithAI` | Bấm đối chiếu; lấy phiếu và file; gửi request sang API-AI. |
| API | `CoreStartup.cs`; `AutomationJobBusiness.cs` | Cấu hình Quartz/lịch chung; chưa thấy job BEM cụ thể trong source đã rà. |
| API-AI | `FileDataHandlerController.cs`; `ReadFileOrchestratorService.cs`; `ChannelJobQueue.cs`; `ReadFileWorker.cs`; `ReadFileBackgroundWorkflow.cs`; `OcrService.cs` | Tạo run, Queue, Worker, gọi OCR/LLM, lưu DB. |
| AI Python | `App/main_iis.py`; `App/Rules_AI_BEM_MEIKO.py` | Endpoint OCR/LLM, xử lý Rules và gọi LLM. |

## Kết luận

- Queue đang ở **API-AI/RAM**, DB không phải Queue.
- OCR, Rule Engine và LLM đều nằm trong **AI Python**; số process/replica thực tế cần kiểm tra cấu hình deploy/GPU.
- DB cho biết lượt chạy đang `PROCESSING`, đã `COMPLETED` hay `FAILED`, đồng thời lưu toàn bộ OCR/trích xuất/tiêu chí để ERP9 hiển thị lại.
'''
MD.write_text(md,encoding='utf-8')

# Word helpers
def set_font(r,size=11,bold=False,color=None):
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade(cell,color):
    p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),color); p.append(e)
def border(cell):
    p=cell._tc.get_or_add_tcPr(); b=p.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders'); p.append(b)
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9E2F3'); b.append(e)
def heading(doc,text):
    p=doc.add_paragraph(); p.style='Heading 1'; p.paragraph_format.space_before=Pt(11); p.paragraph_format.space_after=Pt(5); r=p.add_run(text); set_font(r,16,True,'17365D')
def para(doc,text,bullet=False):
    p=doc.add_paragraph(style='List Bullet' if bullet else None); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05; r=p.add_run(text); set_font(r,10.5)
def image(doc,path,caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(7.05))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); set_font(r,9,False,'5B6573')
def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(t.rows[0].cells,headers):
        shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); set_font(r,9.4,True,'FFFFFF')
    for row in rows:
        cells=t.add_row().cells
        for idx,(c,txt) in enumerate(zip(cells,row)):
            border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1); r=p.add_run(txt); set_font(r,8.8,idx==0)
    return t

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.6); sec.right_margin=Inches(.6)
for style in ['Normal','Heading 1','Heading 2','Title']:
    st=doc.styles[style]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
doc.styles['Normal'].font.size=Pt(10.5)
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=h.add_run('AI BEM | Cơ chế vận hành'); set_font(r,9,False,'5B6573')
fo=sec.footer.paragraphs[0]; fo.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=fo.add_run('01/10/2026'); set_font(r,9,False,'5B6573')

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(45); r=p.add_run('CƠ CHẾ VẬN HÀNH AI BEM'); set_font(r,26,True,'17365D')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('WEB  |  API  |  API-AI  |  AI Python  |  DB'); set_font(r,14,False,'168A8A')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(18); r=p.add_run('Bản rút gọn để theo dõi nơi xử lý và nơi lưu kết quả'); set_font(r,11,False,'5B6573')
heading(doc,'Kết luận ngắn')
for x in ['Queue nằm trong RAM của API-AI, không phải trong DB.', 'OCR, Rule Engine và LLM nằm trong AI Python.', 'DB lưu tiến độ của từng lượt chạy: PROCESSING, COMPLETED hoặc FAILED; đồng thời lưu OCR, trích xuất và từng tiêu chí.', 'Luồng tự động có Quartz ở API ERP, nhưng chưa xác nhận được job BEM cụ thể trong source đã rà.']:
    para(doc,x,True)
heading(doc,'1. Bản đồ vận hành hiện tại')
image(doc,IMG_DIR/'01_ban_do_van_hanh.png','Hình 1. Mỗi khu vực đang chịu trách nhiệm phần nào')
heading(doc,'2. Năm thành phần cần quan tâm')
table(doc,['Thành phần','Nằm tại','Hiện trạng'],[
('Queue','API-AI','ChannelJobQueue trong RAM; Singleton; tối đa 200 job; Worker lấy job để xử lý.'),
('OCR','AI Python','Route /ocr đọc file thành text. API-AI chỉ gọi OCR; tối đa 7 file song song/phiếu.'),
('Rule Engine','AI Python','process_ai_llms_models_rules chọn trích xuất/đối chiếu và áp dụng Rules.'),
('LLM','AI Python','Route /llms/api/ai_llms_models gọi LLM để trích xuất hoặc đối chiếu.'),
('DB','SQL DB','Lưu lượt chạy, OCR, dữ liệu trích xuất, kết quả từng tiêu chí và kết quả tổng.'),
])
heading(doc,'3. Một phiếu được xử lý và ghi DB như thế nào')
image(doc,IMG_DIR/'02_trang_thai_mot_phieu.png','Hình 2. Trạng thái và dữ liệu được lưu theo từng bước')
table(doc,['Thời điểm','DB lưu tại','Ý nghĩa'],[
('Bắt đầu','BEMT2003','Tạo lượt chạy; StatusProcess = PROCESSING.'),
('OCR xong','BEMT2002','Text OCR và dữ liệu nguồn/tổng hợp.'),
('Trích xuất xong','BEMT2005, BEMT2006','Dữ liệu đọc từ chứng từ.'),
('Đối chiếu xong','BEMT2004','Kết quả từng tiêu chí OK/NG và giải thích.'),
('Hoàn tất','BEMT2003','COMPLETED, kết quả tổng OK/NG và %.'),
('Lỗi','BEMT2003','FAILED, nhận biết lượt chạy không hoàn tất.'),
])
heading(doc,'4. Muốn xem code thì mở file nào')
table(doc,['Khu vực','File/hàm chính','Xử lý'],[
('WEB','BEMF2002.js; BEMF2000Controller.cs','Bấm đối chiếu; lấy phiếu/file; gửi request sang API-AI.'),
('API','CoreStartup.cs; AutomationJobBusiness.cs','Quartz/lịch chung; chưa thấy job BEM cụ thể.'),
('API-AI','FileDataHandlerController; Orchestrator; Queue; Worker; Workflow; OcrService','Tạo run, Queue, Worker, gọi OCR/LLM, lưu DB.'),
('AI Python','main_iis.py; Rules_AI_BEM_MEIKO.py','Endpoint OCR/LLM; Rules và LLM xử lý trích xuất/đối chiếu.'),
])
para(doc,'Lưu ý: source cho biết cách các khối gọi nhau. Số process/replica OCR, LLM và API-AI đang chạy thực tế cần kiểm tra cấu hình IIS, service Python và GPU.',True)
doc.save(DOCX)
print(MD); print(DOCX); print(IMG_DIR/'01_ban_do_van_hanh.png'); print(IMG_DIR/'02_trang_thai_mot_phieu.png')

