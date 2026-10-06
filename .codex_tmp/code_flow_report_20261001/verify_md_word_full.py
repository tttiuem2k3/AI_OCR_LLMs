from __future__ import annotations
import re, zipfile
from pathlib import Path
from docx import Document

md_path=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md')
doc_path=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx')

with zipfile.ZipFile(doc_path) as archive:
    bad=archive.testzip()
    media=[x for x in archive.namelist() if x.startswith('word/media/')]
    print('DOCX_ZIP=', 'OK' if bad is None else f'BAD:{bad}')
    print('EMBEDDED_MEDIA=', len(media))

doc=Document(doc_path)
paragraph_text='\n'.join(p.text for p in doc.paragraphs)
table_text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
word_text=paragraph_text+'\n'+table_text

def norm(s:str)->str:
    s=s.replace('`','').replace('**','').replace('“','"').replace('”','"')
    s=re.sub(r'\s+',' ',s).strip().lower()
    return s

word_norm=norm(word_text)
missing=[]
checked=0
for raw in md_path.read_text(encoding='utf-8').splitlines():
    line=raw.strip()
    if not line or re.fullmatch(r'\|?\s*(?::?-+:?\s*\|\s*)+(?::?-+:?\s*)?\|?',line):
        continue
    if line.startswith('#'):
        fragments=[line.lstrip('#').strip()]
    elif line.startswith('|'):
        fragments=[x.strip() for x in line.strip('|').split('|') if x.strip()]
    else:
        line=re.sub(r'^-\s+','',line)
        fragments=[line]
    for fragment in fragments:
        fragment_norm=norm(fragment)
        if len(fragment_norm)<5:
            continue
        checked+=1
        if fragment_norm not in word_norm:
            missing.append(fragment)
print('PARAGRAPHS=',len(doc.paragraphs),'TABLES=',len(doc.tables),'INLINE_IMAGES=',len(doc.inline_shapes))
print('MD_FRAGMENTS_CHECKED=',checked,'MISSING=',len(missing))
for item in missing[:20]:
    print('MISSING_TEXT:',item)
headers=[]
for t in doc.tables:
    headers.append([c.text.strip() for c in t.rows[0].cells])
print('TABLE_HEADERS:')
for i,h in enumerate(headers,1): print(i,h)
required=[
'1. Cấu trúc mã nguồn và trách nhiệm hiện tại',
'2. Ví dụ đối chiếu thủ công một phiếu NVL có 10 file',
'3. Dữ liệu đi qua từng khu vực code',
'4. Database lưu dữ liệu gì',
'5. Trách nhiệm API/Python/DB và trạng thái hiện tại',
'6. Rủi ro chính dẫn đến fail như tháng 09',
'7. Task chốt để đạt mục tiêu 15/10',
'8. Baseline chốt triển khai',
]
for value in required:
    print(('OK ' if norm(value) in word_norm else 'MISSING ')+value)
if bad is not None or len(media)!=5 or len(doc.inline_shapes)!=5 or len(doc.tables)!=8 or missing:
    raise SystemExit(1)
