from docx import Document
from pathlib import Path
p=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx")
d=Document(p)
print('paragraphs=',len(d.paragraphs),'tables=',len(d.tables),'images=',len(d.inline_shapes))
for i, paragraph in enumerate(d.paragraphs):
    if paragraph.text.strip() and ('Heading' in paragraph.style.name or paragraph.style.name=='Title'):
        print(f'H{i+1}: {paragraph.text}')
for i, table in enumerate(d.tables,1):
    print(f'T{i}: {len(table.rows)}x{len(table.columns)} | {" || ".join(cell.text.replace(chr(10)," ") for cell in table.rows[0].cells)}')
text='\n'.join(x.text for x in d.paragraphs)
for required in [
    'Thủ công, tự động và chạy lại',
    'Bảng chốt theo tiêu chí nghiệm thu',
    'Baseline chốt triển khai',
    'Queue RAM tối đa 200 job',
    'Restart service không được làm mất job chờ',
    'Kết quả Rules + LLM hiện được chốt chung tại BEMT2004',
]:
    print(('OK ' if required in text or any(required in cell.text for table in d.tables for row in table.rows for cell in row.cells) else 'MISSING ') + required)
