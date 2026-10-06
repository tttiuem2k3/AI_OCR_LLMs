from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_acceptance_report_v2 import (
    DARK,
    GRAY,
    LINE,
    arrow,
    draw_center,
    font,
    rounded_box,
)


BASE = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMG = BASE / "Temp" / "code_flow_acceptance_20261001_v3"
IMG.mkdir(parents=True, exist_ok=True)
MD = BASE / "Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
DOCX = BASE / "Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"

NAVY = "#17365D"
ERP = "#2F75B5"
API_AI = "#D6A300"
AI = "#ED7D31"
DB = "#7030A0"
GREEN = "#70AD47"
RED = "#C00000"
TEAL = "#168A8A"
LIGHT = "#F7FAFC"


def canvas(title, subtitle, width=2400, height=1350):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 120), fill=NAVY)
    draw.text((50, 20), title, font=font(43, True), fill="white")
    draw.text((52, 76), subtitle, font=font(23), fill="#DCE6F1")
    return image, draw


def legend(draw, y=1250):
    items = [("WEB + API ERP9", ERP), ("API-AI", API_AI), ("AI Python", AI), ("Database", DB)]
    start_x = 525
    for label, color in items:
        draw.rounded_rectangle((start_x, y, start_x + 300, y + 58), 12, fill=color)
        draw_center(draw, (start_x + 10, y + 5, start_x + 290, y + 53), label, font(21, True), "white")
        start_x += 335


def make_flow_diagram():
    image, draw = canvas(
        "Hai luồng đưa phiếu vào AI BEM",
        "ERP9 được mô tả đúng hiện trạng nhưng không nằm trong phạm vi sửa; màu sắc thống nhất theo khu vực code",
    )
    draw.rounded_rectangle((45, 155, 2355, 690), 24, fill="#F7FAFD", outline=ERP, width=3)
    draw.rounded_rectangle((45, 735, 2355, 1210), 24, fill="#FBFCFD", outline=TEAL, width=3)
    draw_center(draw, (75, 170, 520, 225), "A. THỦ CÔNG — 1 PHIẾU", font(29, True), ERP)
    draw_center(draw, (75, 750, 610, 805), "B. TỰ ĐỘNG — NHIỀU PHIẾU", font(29, True), TEAL)

    manual = [
        ((95, 300, 410, 510), "WEB ERP9", "Người dùng mở phiếu\nvà bấm Đối chiếu AI", ERP),
        ((475, 300, 790, 510), "API ERP9", "Kiểm tra file\nlấy phiếu + đính kèm\ntạo CompareFileRequest", ERP),
        ((855, 300, 1170, 510), "API-AI", "Tạo lượt chạy\nBEMT2003=PROCESSING\nđưa job vào Queue", API_AI),
        ((1235, 300, 1550, 510), "AI Python", "OCR → LLM trích xuất\nRules + LLM đối chiếu", AI),
        ((1615, 300, 1930, 510), "Database", "Lưu OCR, trích xuất\nkết quả, trạng thái", DB),
        ((1995, 300, 2305, 510), "WEB ERP9", "Đọc DB và hiển thị\nkết quả cho người dùng", ERP),
    ]
    for rect, title, body, color in manual:
        rounded_box(draw, rect, title, body, color, 20)
    for index in range(len(manual) - 1):
        arrow(draw, (manual[index][0][2], 405), (manual[index + 1][0][0], 405), LINE, 7, 22)
    draw_center(
        draw,
        (140, 555, 2250, 640),
        "Khi sửa phiếu hoặc xóa file đính kèm, ERP9 gọi UpdateResultAI để cập nhật/xóa kết quả cũ; phiếu cần đối chiếu lại.",
        font(24, True),
        "#8A3E00",
    )

    automatic = [
        ((95, 865, 410, 1075), "LỊCH ERP9", "Gọi endpoint\nUpdateResultCompareAI\ntheo cấu hình vận hành", ERP),
        ((475, 865, 790, 1075), "API ERP9", "Lấy danh sách phiếu\nchi tiết + file\nbỏ phiếu không có file", ERP),
        ((855, 865, 1170, 1075), "API ERP9", "Tối đa 5 tác vụ\ngửi từng phiếu\nsang API-AI", ERP),
        ((1235, 865, 1550, 1075), "API-AI", "Queue RAM tối đa 200\nSingleReader\nWorker lấy lần lượt", API_AI),
        ((1615, 865, 1930, 1075), "AI Python", "OCR → LLM → Rules\ntrả kết quả/ErrorCode", AI),
        ((1995, 865, 2305, 1075), "Database", "Lưu trạng thái\nvà kết quả từng phiếu", DB),
    ]
    for rect, title, body, color in automatic:
        rounded_box(draw, rect, title, body, color, 20)
    for index in range(len(automatic) - 1):
        arrow(draw, (automatic[index][0][2], 970), (automatic[index + 1][0][0], 970), LINE, 7, 22)

    legend(draw)
    image.save(IMG / "01_hai_luong_erp_api_ai_db.png")


def make_technical_diagram():
    image, draw = canvas(
        "API-AI xử lý một phiếu như thế nào",
        "Một hàng xử lý chính: API-AI điều phối → AI Python OCR/LLM/Rules → DB lưu trạng thái và kết quả",
    )

    # Main horizontal pipeline. No crossing arrows: each card contains the API-AI action,
    # AI Python response, and DB storage for that exact step.
    card_w = 345
    card_h = 765
    gap = 35
    x0 = 78
    y0 = 265
    y1 = y0 + card_h

    steps = [
        {
            "title": "1. Nhận yêu cầu",
            "header": API_AI,
            "api": "Nhận phiếu từ ERP9\nXóa kết quả cũ\nTạo lượt chạy mới",
            "ai": "Chưa gọi AI Python",
            "db": "BEMT2003\nStatusProcess = PROCESSING",
        },
        {
            "title": "2. Xếp hàng",
            "header": API_AI,
            "api": "Đưa job vào Queue RAM\nWorker lấy phiếu xử lý",
            "ai": "Chưa gọi AI Python",
            "db": "Queue hiện chưa lưu DB\nRủi ro mất job khi restart",
        },
        {
            "title": "3. OCR từng file",
            "header": AI,
            "api": "Gửi file đính kèm sang /ocr\nNhận text OCR từng file",
            "ai": "OCR instance\nPDF/ảnh → text OCR",
            "db": "BEMT2002.RawContent\nAttachID, FileName, APK_File",
        },
        {
            "title": "4. LLM trích xuất",
            "header": AI,
            "api": "Gửi text OCR + prompt\nChuẩn hóa dữ liệu chứng từ",
            "ai": "LLM instance trả về:\nNCC, Invoice, Amount,\nPO/Ringi/tờ khai...",
            "db": "BEMT2002.CriteriaSynthesis\nBEMT2005/2006: cần xác nhận",
        },
        {
            "title": "5. Rules/LLM đối chiếu",
            "header": AI,
            "api": "Gọi từng prompt tiêu chí\nqua CompareAsync",
            "ai": "Rules Engine + LLM\nTrả OK/NG + lý do\ntừng tiêu chí",
            "db": "BEMT2004\nCriteriaStatus, Description,\nPromptSystem",
        },
        {
            "title": "6. Trả kết quả",
            "header": API_AI,
            "api": "Tổng hợp kết quả\nCập nhật kết quả cuối",
            "ai": "Không gọi thêm\nChỉ tổng hợp kết quả đã có",
            "db": "BEMT2003\nCOMPLETED/FAILED\nStatus, %, TextConditionFail",
        },
    ]

    def section(rect, label, body, color, body_size=15):
        x1, yy1, x2, yy2 = rect
        draw.rounded_rectangle(rect, 13, fill="white", outline=color, width=3)
        draw.rounded_rectangle((x1 + 12, yy1 + 12, x2 - 12, yy1 + 45), 11, fill=color)
        draw_center(draw, (x1 + 16, yy1 + 14, x2 - 16, yy1 + 43), label, font(15, True), "white")
        draw_center(draw, (x1 + 16, yy1 + 52, x2 - 16, yy2 - 10), body, font(body_size), DARK, leading=5)

    for idx, step in enumerate(steps):
        x = x0 + idx * (card_w + gap)
        draw.rounded_rectangle((x, y0, x + card_w, y1), 22, fill="#F8FAFC", outline="#CBD5E1", width=3)
        draw.rounded_rectangle((x, y0, x + card_w, y0 + 68), 22, fill=step["header"], outline=step["header"])
        draw.rectangle((x, y0 + 40, x + card_w, y0 + 68), fill=step["header"])
        draw_center(draw, (x + 12, y0 + 8, x + card_w - 12, y0 + 62), step["title"], font(22, True), "white")

        section((x + 18, y0 + 94, x + card_w - 18, y0 + 275), "API-AI thực hiện", step["api"], API_AI, 15)
        section((x + 18, y0 + 305, x + card_w - 18, y0 + 500), "AI Python nhận / trả", step["ai"], AI if idx in (2, 3, 4) else GRAY, 15)
        section((x + 18, y0 + 530, x + card_w - 18, y0 + 735), "DB lưu", step["db"], DB, 15)

        # Detached triangle connectors between cards: visible direction, no font glyph and no crossing lines.
        if idx < len(steps) - 1:
            mid_x = x + card_w + gap / 2
            mid_y = y0 + 385
            draw.polygon(
                [(mid_x - 9, mid_y - 18), (mid_x + 13, mid_y), (mid_x - 9, mid_y + 18)],
                fill=LINE,
            )

    def erp_chip(rect, title, body):
        x1, yy1, x2, yy2 = rect
        draw.rounded_rectangle(rect, 16, fill="white", outline=ERP, width=3)
        draw.rounded_rectangle((x1, yy1, x2, yy1 + 44), 16, fill=ERP, outline=ERP)
        draw.rectangle((x1, yy1 + 25, x2, yy1 + 44), fill=ERP)
        draw_center(draw, (x1 + 14, yy1 + 5, x2 - 14, yy1 + 40), title, font(19, True), "white")
        draw_center(draw, (x1 + 14, yy1 + 52, x2 - 14, yy2 - 10), body, font(15), DARK)

    # ERP input/output cards are outside the pipeline so ERP9 role is visible but not a change scope.
    erp_chip((78, 138, 520, 235), "ERP9 gửi vào", "Phiếu + chi tiết + file đính kèm")
    erp_chip((1850, 138, 2325, 235), "ERP9 đọc ra", "Đọc DB để hiển thị OK/NG, %, lý do")
    draw_center(draw, (560, 160, 1815, 215), "ERP9 không thay đổi: chỉ gửi yêu cầu và đọc kết quả đã lưu trong DB", font(22, True), ERP)

    # Compact baseline footer.
    draw.rounded_rectangle((78, 1080, 2325, 1164), 18, fill="#F3F7FF", outline=ERP, width=3)
    draw_center(
        draw,
        (105, 1092, 2298, 1154),
        "Trạng thái cần chuẩn hóa trong DB: QUEUED → OCR_RUNNING → LLM_EXTRACTING → RULE_CHECKING → COMPLETED / RETRY_WAIT / FAILED",
        font(23, True),
        NAVY,
    )
    draw.rounded_rectangle((78, 1188, 2325, 1238), 14, fill="#FFF8E1", outline=API_AI, width=3)
    draw_center(
        draw,
        (100, 1195, 2303, 1231),
        "Cần bổ sung: ProcessStage, RetryCount, thời gian bắt đầu/kết thúc, Heartbeat, ErrorCode, ErrorMessage, FailedStep, EngineSource",
        font(19, True),
        "#8A3E00",
    )
    legend(draw)
    image.save(IMG / "02_trach_nhiem_trang_thai.png")

def make_baseline_diagram():
    image, draw = canvas(
        "Rủi ro và phương án chốt đến 15/10",
        "Chỉ tập trung lỗi vận hành API-AI, AI Python và Database; không mở rộng sang nội dung AI ngoài phạm vi",
    )
    draw.rounded_rectangle((55, 155, 1155, 1175), 24, fill="#FFF8F8", outline=RED, width=3)
    draw.rounded_rectangle((1245, 155, 2345, 1175), 24, fill="#F5FAF3", outline=GREEN, width=3)
    draw_center(draw, (90, 180, 1120, 245), "HIỆN TRẠNG / RỦI RO", font(31, True), RED)
    draw_center(draw, (1280, 180, 2310, 245), "BASELINE CẦN CHỐT", font(31, True), GREEN)

    risks = [
        ("API-AI", "Có nhánh dừng trước khi ghi FAILED → phiếu có thể treo PROCESSING", API_AI),
        ("API-AI", "Queue nằm trong RAM → restart có thể mất job chờ", API_AI),
        ("API-AI/DB", "Chưa có retry bền vững và chưa phân loại lỗi tạm thời/lỗi dữ liệu", API_AI),
        ("AI Python", "OCR có độ trễ đuôi dài; LLM dùng chung GPU và gen_lock", AI),
        ("Database", "Chưa biết phiếu đang Queue/OCR/LLM/Rules; thiếu thời gian và ErrorCode chuẩn", DB),
    ]
    y = 300
    for area, text, color in risks:
        draw.rounded_rectangle((105, y, 1105, y + 135), 18, fill="white", outline=color, width=4)
        draw.rounded_rectangle((105, y, 310, y + 135), 18, fill=color, outline=color)
        draw.rectangle((285, y, 310, y + 135), fill=color)
        draw_center(draw, (120, y + 10, 295, y + 125), area, font(22, True), "white")
        draw_center(draw, (340, y + 12, 1075, y + 123), text, font(22), DARK)
        y += 165

    solutions = [
        ("04/10", "API-AI", "Mọi lối thoát ghi stage/final status/ErrorCode", API_AI),
        ("07/10", "API-AI + DB", "DB là nguồn trạng thái; phục hồi Queue sau restart; retry ≤ 2", API_AI),
        ("10/10", "AI Python", "Health check, timeout từng bước, giới hạn tải OCR/LLM", AI),
        ("10/10", "Database", "Bổ sung ProcessStage, retry, time, heartbeat, EngineSource", DB),
        ("13/10", "QA", "Test OCR/LLM fail, timeout, restart, retry, tải và case tháng 09", TEAL),
        ("15/10", "Tech Lead", "UAT và chốt baseline; không mở rộng ngoài phạm vi", GREEN),
    ]
    y = 285
    for date, owner, text, color in solutions:
        draw.rounded_rectangle((1295, y, 2295, y + 125), 18, fill="white", outline=color, width=4)
        draw.rounded_rectangle((1295, y, 1490, y + 125), 18, fill=color, outline=color)
        draw.rectangle((1465, y, 1490, y + 125), fill=color)
        draw_center(draw, (1310, y + 8, 1475, y + 117), date, font(24, True), "white")
        draw_center(draw, (1520, y + 6, 1810, y + 119), owner, font(21, True), color)
        draw_center(draw, (1825, y + 8, 2265, y + 117), text, font(20), DARK)
        y += 145

    draw_center(
        draw,
        (1265, 1168, 2320, 1240),
        "Điều kiện chốt: không mất job khi restart • không còn PROCESSING treo • biết phiếu ở bước nào • lỗi có mã/thời gian • test tháng 09 đạt",
        font(22, True),
        NAVY,
    )
    image.save(IMG / "03_rui_ro_phuong_an_1510.png")


def build_images():
    make_flow_diagram()
    make_technical_diagram()
    make_baseline_diagram()


def md_content():
    return """# Báo cáo chốt phương án ổn định xử lý AI BEM

Ngày báo cáo: 01/10/2026  
Mục tiêu: đến **15/10/2026** xử lý ổn định và không lặp lại tình trạng fail như hai kỳ tháng 09.

1. Phạm vi và kết luận
- ERP9 không thay đổi, nhưng báo cáo mô tả rõ vai trò hiện tại của WEB và API ERP9.
- Chỉ cải tiến API-AI, AI Python và Database.
- DB được chốt là nguồn trạng thái chính; Queue RAM chỉ là bộ tăng tốc.
- Không mở rộng sang Rules/Training/nghiệp vụ AI ngoài mục tiêu vận hành.

2. Luồng hiện tại

![Hai luồng](Temp/code_flow_acceptance_20261001_v3/01_hai_luong_erp_api_ai_db.png)

- Thủ công: người dùng bấm một phiếu; WEB/API ERP9 lấy phiếu + file và gọi API-AI.
- Tự động: ERP9 gọi `UpdateResultCompareAI`, lấy nhiều phiếu, gửi tối đa 5 tác vụ song song; API-AI nhận từng phiếu vào Queue RAM.
- Khi sửa phiếu hoặc xóa file, ERP9 gọi `UpdateResultAI` để cập nhật/xóa kết quả cũ; phiếu cần chạy lại.

3. Trách nhiệm kỹ thuật và trạng thái

![Trách nhiệm](Temp/code_flow_acceptance_20261001_v3/02_trach_nhiem_trang_thai.png)

- API-AI xóa kết quả cũ, tạo `BEMT2003=PROCESSING`, đưa job vào Queue và trả ngay trạng thái “đang xử lý nền” cho ERP9.
- Worker gọi OCR `/ocr`, nhận text OCR theo từng file; `RawContent`, tên file và liên kết file được gom để lưu `BEMT2002`.
- API-AI gửi text OCR sang LLM trích xuất (`FormatOCRText`), nhận dữ liệu chứng từ đã chuẩn hóa để dùng khi đối chiếu.
- API-AI chạy từng prompt qua `CompareAsync`; Rules/LLM trả OK/NG và giải thích từng tiêu chí để lưu `BEMT2004`.
- Cuối luồng, API-AI tính kết quả tổng, %, điều kiện fail và cập nhật `BEMT2003=COMPLETED`; nếu bắt được lỗi thì cập nhật `FAILED`.
- `BEMT2005/2006` có cấu trúc master/detail, nhưng workflow đang rà chưa thấy gọi lưu trực tiếp; cần xác nhận luồng nào đang ghi hai bảng này.
- Trạng thái baseline cần theo dõi: `QUEUED → OCR_RUNNING → LLM_EXTRACTING → RULE_CHECKING → COMPLETED`; lỗi tạm thời vào `RETRY_WAIT`, lỗi cuối vào `FAILED`.
- OCR/LLM engine được tái sử dụng trong process Python; LLM dùng `_gen_lock` để tuần tự hóa generate trên cùng GPU. API-AI cần lưu timeout/error theo từng bước.

4. DB cần lưu

| Bảng/nhóm | Hiện có | Cần bổ sung |
|---|---|---|
| `BEMT2003` | Lượt chạy, `PROCESSING/COMPLETED/FAILED`, OK/NG, % | `ProcessStage`, retry, thời gian, heartbeat, `ErrorCode/ErrorMessage/FailedStep` |
| `BEMT2002` | OCR, nội dung tổng hợp, file | Thời gian và trạng thái OCR |
| `BEMT2005/2006` | Có cấu trúc lưu dữ liệu trích xuất master/detail | Chưa thấy `ReadFileBackgroundWorkflow` ghi trực tiếp; cần xác nhận luồng thực tế |
| `BEMT2004` | Kết quả từng tiêu chí | `EngineSource = RULE/LLM/MIXED` |

5. Rủi ro, phương án và deadline

![Baseline](Temp/code_flow_acceptance_20261001_v3/03_rui_ro_phuong_an_1510.png)

| Ưu tiên | Phần | Công việc | Owner | Deadline |
|---|---|---|---|---|
| P0 | API-AI | Không để nhánh dừng mà không ghi FAILED/stage/ErrorCode | DEV API-AI | 04/10 |
| P0 | API-AI/DB | DB là nguồn trạng thái; phục hồi Queue; retry tối đa 2 lần | DEV API-AI + DB | 07/10 |
| P1 | AI Python | Health check, timeout và giới hạn tải OCR/LLM | DEV Python | 10/10 |
| P0 | Database | Bổ sung stage/retry/time/heartbeat/error | DEV DB + API-AI | 10/10 |
| P0 | Kiểm thử | Test fail, timeout, restart, retry, tải và case tháng 09 | QA ASOFT | 13/10 |
| P0 | UAT | Chốt baseline, không mở rộng ngoài phạm vi | Tech Lead + owner | 15/10 |

6. Điều kiện nghiệm thu
- Không thay đổi ERP9.
- Biết mỗi phiếu đang ở bước nào.
- Không còn `PROCESSING` treo sau khi luồng dừng.
- Restart API-AI không mất job chưa xử lý.
- Retry lỗi tạm thời tối đa 2 lần; lỗi cuối có mã và thời gian.
- Test OCR/LLM fail, timeout, restart, retry, tải và bộ đại diện tháng 09 đạt.
- Kết quả được chốt làm baseline triển khai; không mở rộng sang nội dung AI ngoài phạm vi.
"""


def set_cell(cell, text, bold=False, size=9, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows, size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, True, size, "FFFFFF", WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[index], NAVY.replace("#", ""))
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if index in (0, len(row) - 1) else WD_ALIGN_PARAGRAPH.LEFT
            set_cell(cells[index], value, False, size, None, align)
    return table


def add_heading(doc, text, size=16):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(10.5)


def add_image(doc, path, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(9.2))
    caption_paragraph = doc.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.runs[0]
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = RGBColor(91, 101, 115)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Báo cáo chốt phương án ổn định xử lý AI BEM")
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(23, 54, 93)
    subtitle = doc.add_paragraph("Phạm vi cải tiến: API-AI, AI Python và Database. ERP9 giữ nguyên. Mục tiêu nghiệm thu: 15/10/2026.")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "1. Phạm vi và kết luận")
    for item in [
        "ERP9 không thay đổi, nhưng báo cáo mô tả rõ vai trò hiện tại của WEB và API ERP9.",
        "Chỉ cải tiến API-AI, AI Python và Database.",
        "DB là nguồn trạng thái chính; Queue RAM chỉ là bộ tăng tốc.",
        "Không mở rộng sang Rules/Training/nghiệp vụ AI ngoài mục tiêu vận hành.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Luồng thủ công và tự động")
    add_image(doc, IMG / "01_hai_luong_erp_api_ai_db.png", "Hình 1. Hai luồng hiện tại; WEB và API ERP9 dùng cùng màu")
    for item in [
        "Thủ công: ERP9 lấy một phiếu và file rồi gửi một request sang API-AI.",
        "Tự động: endpoint UpdateResultCompareAI lấy nhiều phiếu, gửi tối đa 5 tác vụ song song; API-AI xếp từng phiếu vào Queue RAM.",
        "Khi sửa phiếu/xóa file, ERP9 gọi UpdateResultAI để cập nhật/xóa kết quả AI cũ; phiếu cần chạy lại.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "3. Chuỗi xử lý API-AI và dữ liệu lưu DB")
    add_image(doc, IMG / "02_trach_nhiem_trang_thai.png", "Hình 2. API-AI gọi OCR, LLM, Rules/LLM và lưu dữ liệu/kết quả vào DB")
    for item in [
        "API-AI xóa kết quả cũ, tạo BEMT2003 = PROCESSING, đưa job vào Queue và trả ngay trạng thái đang xử lý nền cho ERP9.",
        "Worker gọi OCR /ocr và nhận text OCR theo từng file; API-AI lưu RawContent/tên file/liên kết file vào BEMT2002.",
        "API-AI gửi text OCR sang LLM trích xuất; dữ liệu chuẩn hóa được dùng cho các lần đối chiếu tiêu chí.",
        "Từng prompt gọi CompareAsync; Rules/LLM trả OK/NG + giải thích để lưu BEMT2004. Sau đó BEMT2003 lưu kết quả tổng, % và trạng thái COMPLETED/FAILED.",
"BEMT2005/2006 có cấu trúc master/detail, nhưng workflow đang rà chưa thấy gọi lưu trực tiếp; cần xác nhận luồng thực tế.",
        "Trạng thái baseline: QUEUED → OCR_RUNNING → LLM_EXTRACTING → RULE_CHECKING → COMPLETED; lỗi tạm thời vào RETRY_WAIT, lỗi cuối vào FAILED.",
        "OCR/LLM engine tái sử dụng trong Python; LLM giữ gen_lock trên GPU. API-AI cần ghi timeout/error theo từng bước vào DB.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. DB cần lưu")
    add_table(
        doc,
        ["Bảng/nhóm", "Hiện có", "Cần bổ sung"],
        [
            ("BEMT2003", "Lượt chạy, PROCESSING/COMPLETED/FAILED, OK/NG, %", "ProcessStage, retry, thời gian, heartbeat, ErrorCode/ErrorMessage/FailedStep"),
            ("BEMT2002", "OCR, nội dung tổng hợp, file", "Thời gian và trạng thái OCR"),
            ("BEMT2005/2006", "Có cấu trúc master/detail", "Workflow đang rà chưa thấy gọi lưu trực tiếp; cần xác nhận luồng thực tế"),
            ("BEMT2004", "Kết quả từng tiêu chí", "EngineSource = RULE/LLM/MIXED"),
        ],
        8.7,
    )

    doc.add_page_break()
    add_heading(doc, "5. Rủi ro và phương án đến 15/10")
    add_image(doc, IMG / "03_rui_ro_phuong_an_1510.png", "Hình 3. Rủi ro đã xác nhận và baseline thực hiện đến 15/10")

    landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.45)
    landscape.bottom_margin = Inches(0.45)
    landscape.left_margin = Inches(0.45)
    landscape.right_margin = Inches(0.45)
    add_heading(doc, "6. Danh sách task chốt")
    add_table(
        doc,
        ["P", "Phần", "Công việc", "Owner", "Deadline"],
        [
            ("P0", "API-AI", "Mọi lối thoát phải ghi stage/final status/ErrorCode", "DEV API-AI", "04/10"),
            ("P0", "API-AI/DB", "DB là nguồn trạng thái; phục hồi Queue; retry tối đa 2 lần", "DEV API-AI + DB", "07/10"),
            ("P1", "AI Python", "Health check, timeout và giới hạn tải OCR/LLM", "DEV Python", "10/10"),
            ("P0", "Database", "Bổ sung stage/retry/time/heartbeat/error và EngineSource", "DEV DB + API-AI", "10/10"),
            ("P0", "Kiểm thử", "Test fail, timeout, restart, retry, tải và case tháng 09", "QA ASOFT", "13/10"),
            ("P0", "UAT", "Chốt baseline, không mở rộng ngoài phạm vi", "Tech Lead + owner", "15/10"),
        ],
        8.6,
    )
    add_heading(doc, "7. Điều kiện nghiệm thu")
    for item in [
        "Không thay đổi ERP9; chỉ mô tả đúng hiện trạng.",
        "Biết mỗi phiếu đang ở bước nào; không còn PROCESSING treo.",
        "Restart API-AI không mất job chưa xử lý.",
        "Retry lỗi tạm thời tối đa 2 lần; lỗi cuối có mã và thời gian.",
        "Test OCR/LLM fail, timeout, restart, retry, tải và bộ đại diện tháng 09 đạt.",
        "Kết quả được chốt làm baseline; không mở rộng sang nội dung AI ngoài phạm vi.",
    ]:
        add_bullet(doc, item)
    doc.save(DOCX)


if __name__ == "__main__":
    build_images()
    MD.write_text(md_content(), encoding="utf-8")
    build_docx()
    print(MD)
    print(DOCX)




