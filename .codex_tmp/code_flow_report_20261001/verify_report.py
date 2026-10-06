from pathlib import Path
from zipfile import ZipFile
from docx import Document
base=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
md=base/"Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md"
docx=base/"Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx"
assert md.exists() and md.stat().st_size > 4000
assert docx.exists() and docx.stat().st_size > 500000
content=md.read_text(encoding='utf-8')
checks=['Luồng khi người dùng đối chiếu thủ công','Luồng khi lịch tự động chạy nhiều phiếu','Điều API AI làm trong một lượt chạy','Điều AI BEM làm với hồ sơ','hàng đợi 200 job','Điểm cần xác nhận']
missing=[p for p in checks if p not in content]
assert not missing, missing
d=Document(docx)
assert len(d.inline_shapes)==8, len(d.inline_shapes)
assert len(d.tables)==1, len(d.tables)
headings=[x.text for x in d.paragraphs if x.style.name.startswith('Heading')]
for h in ['1. Bốn khu vực code làm gì','2. Khi người dùng bấm đối chiếu thủ công','3. Khi lịch tự động gửi nhiều phiếu','7. Thủ công và tự động khác nhau thế nào']:
    assert h in headings, h
with ZipFile(docx) as z:
    pngs=[n for n in z.namelist() if n.startswith('word/media/') and n.endswith('.png')]
    assert len(pngs)==7, pngs
print('PASS')
print('MD bytes:',md.stat().st_size)
print('DOCX bytes:',docx.stat().st_size)
print('Images embedded:',len(pngs))
print('Sections:',len(d.sections),'Paragraphs:',len(d.paragraphs),'Tables:',len(d.tables),'Inline shapes:',len(d.inline_shapes))
