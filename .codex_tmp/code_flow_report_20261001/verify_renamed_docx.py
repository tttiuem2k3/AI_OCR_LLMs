from pathlib import Path
from docx import Document
import zipfile
path = Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Hien_trang_doi_chieu_AI_BEM.docx')
with zipfile.ZipFile(path) as archive:
    bad = archive.testzip()
    media = [name for name in archive.namelist() if name.startswith('word/media/')]
doc = Document(path)
text = '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
print('DOCX_ZIP=', 'OK' if bad is None else f'BAD:{bad}')
print('SIZE=', path.stat().st_size)
print('PARAGRAPHS=', len(doc.paragraphs), 'TABLES=', len(doc.tables), 'IMAGES=', len(doc.inline_shapes), 'MEDIA=', len(media))
for required in ['Cấu trúc mã nguồn', 'Database lưu dữ liệu', 'Task chốt', 'Baseline']:
    print(('OK ' if required in text else 'MISSING ') + required)
if bad is not None:
    raise SystemExit(1)
