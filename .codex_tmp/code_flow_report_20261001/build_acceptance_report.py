from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
IMG=BASE/'Temp'/'code_flow_acceptance_20261001'; IMG.mkdir(parents=True,exist_ok=True)
SOURCE_MAP=BASE/'Temp'/'code_flow_simplified_20261001'/'01_ban_do_van_hanh.png'
MD=BASE/'Noi_dung_co_che_code_doi_chieu_AI_BEM_01102026.md'
DOCX=BASE/'Bao_cao_co_che_code_doi_chieu_AI_BEM_01102026.docx'
FONT=r'C:\Windows\Fonts\times.ttf'; FONT_B=r'C:\Windows\Fonts\timesbd.ttf'
NAVY='#17365D'; BLUE='#2F75B5'; TEAL='#168A8A'; GOLD='#D6A300'; ORANGE='#ED7D31'; RED='#C00000'; GREEN='#70AD47'; PURPLE='#7030A0'; GRAY='#5B6573'; LINE='#768A9E'; DARK='#1F2933'

def f(n,b=False): return ImageFont.truetype(FONT_B if b else FONT,n)
def center(d,xy,text,font,fill=DARK,leading=5):
    x1,y1,x2,y2=xy; lines=[]
    for raw in str(text).split('\n'):
        words=raw.split(); line=''
        if not words: lines.append(''); continue
        for word in words:
            t=(line+' '+word).strip()
            if d.textbbox((0,0),t,font=font)[2] <= x2-x1: line=t
            else:
                if line: lines.append(line)
                line=word
        if line: lines.append(line)
    hs=[max(1,d.textbbox((0,0),s,font=font)[3]-d.textbbox((0,0),s,font=font)[1]) for s in lines]
    y=y1+((y2-y1)-(sum(hs)+leading*(len(lines)-1)))/2
    for s,h in zip(lines,hs):
        b=d.textbbox((0,0),s,font=font); d.text((x1+(x2-x1-(b[2]-b[0]))/2,y),s,font=font,fill=fill); y+=h+leading

def box(d,xy,title,body,color,sz=22,fill='white'):
    x1,y1,x2,y2=xy; d.rounded_rectangle(xy,18,fill=fill,outline=color,width=4); d.rounded_rectangle((x1,y1,x2,y1+52),18,fill=color,outline=color); d.rectangle((x1,y1+28,x2,y1+52),fill=color); center(d,(x1+7,y1+5,x2-7,y1+47),title,f(23,True),'white'); center(d,(x1+12,y1+62,x2-12,y2-10),body,f(sz),DARK)
def arrow(d,a,b,color=LINE,w=7,head=22,dashed=False):
    import math
    x1,y1=a; x2,y2=b
    if dashed:
        n=10
        for i in range(n):
            if i%2==0:
                t=i/n; u=(i+1)/n; d.line((x1+(x2-x1)*t,y1+(y2-y1)*t,x1+(x2-x1)*u,y1+(y2-y1)*u),fill=color,width=w)
    else: d.line((x1,y1,x2,y2),fill=color,width=w)
    ang=math.atan2(y2-y1,x2-x1); p1=(x2-head*math.cos(ang-.52),y2-head*math.sin(ang-.52)); p2=(x2-head*math.cos(ang+.52),y2-head*math.sin(ang+.52)); d.polygon([(x2,y2),p1,p2],fill=color)
def canvas(title,subtitle,w=1900,h=920):
    im=Image.new('RGB',(w,h),'white'); d=ImageDraw.Draw(im); d.rectangle((0,0,w,105),fill=NAVY); d.text((48,18),title,font=f(41,True),fill='white'); d.text((50,70),subtitle,font=f(22),fill='#DCE6F1'); return im,d

# Status current vs baseline
im,d=canvas('Trạng thái phiếu hiện tại và baseline 15/10','Tách rõ trạng thái đang có trong code và trạng thái cần bổ sung để vận hành ổn định')
d.rounded_rectangle((40,140,1860,410),20,fill='#F8FBFD',outline=BLUE,width=3); center(d,(60,150,330,190),'HIỆN TẠI',f(26,True),BLUE)
box(d,(150,230,510,365),'PROCESSING','Tạo trước khi đưa job vào Queue',BLUE,23)
box(d,(790,230,1120,365),'COMPLETED','Có kết quả tổng\nOK/NG + %',GREEN,23)
box(d,(1390,230,1720,365),'FAILED','Luồng bắt được lỗi\nvà cập nhật thất bại',RED,23)
arrow(d,(510,298),(790,298)); arrow(d,(1120,298),(1390,298),dashed=True)
center(d,(525,325,1380,390),'DB chưa có trạng thái riêng cho: chờ Queue • đang OCR • đang trích xuất • đang Rules/LLM • retry',f(23,True),'#8A3E00')
d.rounded_rectangle((40,455,1860,865),20,fill='#F5FAF3',outline=GREEN,width=3); center(d,(60,465,500,510),'BASELINE CẦN CHỐT',f(26,True),GREEN)
states=[('QUEUED','Chờ Worker'),('OCR_RUNNING','Đang đọc file'),('LLM_EXTRACTING','Đang trích xuất'),('RULE_CHECKING','Đang đối chiếu'),('COMPLETED','Hoàn tất')]
x=75
for i,(t,b) in enumerate(states):
    box(d,(x,555,x+300,700),t,b,GREEN if i==4 else TEAL,20)
    if i<len(states)-1: arrow(d,(x+300,627),(x+340,627),LINE,6,18)
    x+=340
box(d,(650,745,1030,830),'RETRY_WAIT','Lỗi tạm thời: chờ chạy lại\nTối đa 2 lần',GOLD,20)
box(d,(1220,745,1580,830),'FAILED','Hết retry hoặc lỗi dữ liệu',RED,20)
arrow(d,(840,700),(840,745),GOLD,5,18); arrow(d,(1030,787),(1220,787),RED,5,18)
im.save(IMG/'02_status_current_baseline.png')

# Risks
im,d=canvas('Rủi ro chính cần xử lý trước 15/10','Ưu tiên tính ổn định; không mở rộng sang nội dung AI ngoài phạm vi API, Python và DB')
box(d,(70,170,895,430),'P0  Phiếu có thể treo PROCESSING','Hai nhánh đang return trước khi ghi FAILED:\n• Gửi dữ liệu ngày nghỉ lỗi\n• OCR không trả dữ liệu',RED,28,fill='#FFF7F7')
box(d,(1005,170,1830,430),'P0  Queue và retry chưa bền vững','Queue nằm trong RAM; restart có thể mất job chờ.\nWorker chỉ log lỗi; chưa có RetryCount / NextRetryAt / cơ chế phục hồi.',RED,26,fill='#FFF7F7')
box(d,(70,515,895,820),'P1  Thời gian OCR biến động','Mẫu app_2.log: 591 lượt OCR\nMedian: 2,99 giây • P95: 17,54 giây\nMax: 205,04 giây • 8 lượt ≥ 60 giây',ORANGE,27,fill='#FFF9F2')
# mini bars
vals=[('Median',2.99,TEAL),('P95',17.54,GOLD),('Max',205.04,RED)]; base_y=770
for idx,(lab,val,col) in enumerate(vals):
    y=625+idx*52; d.text((525,y),lab,font=f(19,True),fill=DARK); length=int(260*min(val/205.04,1)); d.rounded_rectangle((620,y,620+max(8,length),y+25),8,fill=col)
box(d,(1005,515,1830,820),'P1  LLM dùng chung GPU','Python giữ một _llm_engine và tái sử dụng.\n_gen_lock cho phép một lệnh generate tại một thời điểm.\nLệnh dài sẽ làm các yêu cầu sau phải chờ.',ORANGE,26,fill='#FFF9F2')
im.save(IMG/'03_risk_summary.png')

# Baseline architecture
im,d=canvas('Phương án chốt API Python DB','Giữ nguyên ERP9; BEMT2003 là nguồn trạng thái chính, Queue RAM chỉ là bộ tăng tốc')
box(d,(65,190,580,700),'API-AI','• Ghi trạng thái trước/sau mỗi bước\n• Queue RAM + phục hồi từ DB khi restart\n• Retry tối đa 2 lần cho lỗi tạm thời\n• Không retry lỗi thiếu/sai hồ sơ\n• Một RunID không chạy trùng',GOLD,27,fill='#FFFBEB')
box(d,(695,190,1205,700),'AI Python','• Tái sử dụng OCR/LLM engine\n• Health check cho OCR và LLM\n• Trả ErrorCode có cấu trúc\n• Timeout riêng từng bước\n• Ổn định trước: 1 phiếu/Worker, OCR tối đa 2 file song song, LLM giữ gen_lock',ORANGE,25,fill='#FFF7F0')
box(d,(1320,190,1835,700),'Database','Giữ StatusProcess tương thích ERP:\nPROCESSING / COMPLETED / FAILED\n\nBổ sung ProcessStage, RetryCount, StartedAt, FinishedAt, LastHeartbeatAt, ErrorCode, ErrorMessage.\n\nBEMT2004 bổ sung nguồn RULE / LLM / MIXED.',PURPLE,24,fill='#FAF5FF')
arrow(d,(580,445),(695,445)); arrow(d,(1205,445),(1320,445))
center(d,(175,760,1725,850),'Điều kiện đạt: restart không mất job • không còn PROCESSING treo • biết phiếu đang ở bước nào • lỗi tạm thời tự retry • lỗi cuối có mã và thời gian rõ ràng',f(28,True),NAVY)
im.save(IMG/'04_baseline_solution.png')

# Timeline
im,d=canvas('Kế hoạch chốt baseline đến 15/10','Owner theo nhóm thực hiện; ưu tiên P0 trước, kiểm thử fail/restart trước khi nghiệm thu',1900,830)
cols=[('02–04/10',180,430),('05–07/10',450,700),('08–10/10',720,970),('11–13/10',990,1240),('14/10',1260,1480),('15/10',1500,1740)]
for label,x1,x2 in cols:
    d.rounded_rectangle((x1,150,x2,205),10,fill=NAVY); center(d,(x1+3,153,x2-3,202),label,f(21,True),'white')
rows=[('API-AI',235,GOLD,[('Fix treo trạng thái',0,0),('Queue DB + retry',1,1),('Tích hợp trạng thái',2,3),('Sửa lỗi UAT',4,4)]),('Python',355,ORANGE,[('ErrorCode + health',0,1),('Timeout + giới hạn tải',2,2),('Tích hợp',3,3),('Sửa lỗi UAT',4,4)]),('DB',475,PURPLE,[('Chốt field',0,0),('Migration + recovery query',1,2),('Đối soát dữ liệu',3,3)]),('QA / Lead',595,TEAL,[('Kịch bản test',0,1),('Test fail/restart/tải',2,3),('UAT',4,4),('Chốt baseline',5,5)])]
for name,y,col,tasks in rows:
    d.rounded_rectangle((35,y,155,y+72),10,fill=col); center(d,(40,y+5,150,y+67),name,f(20,True),'white')
    for label,c1,c2 in tasks:
        x1=cols[c1][1]; x2=cols[c2][2]; d.rounded_rectangle((x1,y,x2,y+72),12,fill=col,outline='white',width=3); center(d,(x1+7,y+5,x2-7,y+67),label,f(18,True),'white')
center(d,(150,720,1750,790),'15/10: chỉ nghiệm thu khi trạng thái rõ, retry hoạt động, restart không mất job và bộ test đại diện tháng 09 chạy đạt.',f(26,True),RED)
im.save(IMG/'05_timeline_1510.png')

md='''# Phân tích và chốt phương án ổn định xử lý AI BEM

Ngày báo cáo: 01/10/2026  
Mục tiêu nghiệm thu: đến **15/10/2026** có kết quả xử lý ổn định và không lặp lại tình trạng fail như hai kỳ tháng 09.

## 1. Phạm vi chốt

- Không thay đổi ERP9.
- Chỉ cải tiến API-AI, AI Python và Database.
- Không mở rộng sang nghiệp vụ AI/Rules/Training ngoài mục tiêu ổn định vận hành.

## 2. Luồng end-to-end hiện tại

![Luồng hiện tại](Temp/code_flow_simplified_20261001/01_ban_do_van_hanh.png)

- WEB/ERP9 chỉ là điểm đưa phiếu vào và hiển thị kết quả; không nằm trong phạm vi sửa.
- API ERP có Quartz/lịch chung nhưng chưa xác nhận được job BEM cụ thể trong source đã rà.
- API-AI tạo lượt chạy, quản lý Queue RAM và Worker.
- AI Python xử lý OCR, LLM trích xuất và Rules + LLM đối chiếu.
- DB lưu lượt chạy, OCR, dữ liệu trích xuất và kết quả từng tiêu chí.

## 3. Thành phần và instance hiện tại

| Thành phần | Thuộc phần | Hiện trạng |
|---|---|---|
| Queue Management | API-AI | `ChannelJobQueue` là Singleton trong RAM; bounded 200 job; `SingleReader = true`; Queue đầy thì request chờ. `ReadFileWorker` lấy từng job. Chưa có persistent queue hoặc recovery sau restart. |
| OCR instance | AI Python | `_engine` được giữ trong process Python và tái sử dụng; có thể lazy-load theo cấu hình. API-AI gọi route `/ocr`; mỗi phiếu đang cho phép tối đa 7 file OCR song song. |
| LLM instance | AI Python | `_llm_engine` và `_llm_registry` khởi tạo ở startup và tái sử dụng. `_gen_lock` tuần tự hóa thao tác generate trên cùng model/GPU. |
| Rule Engine | AI Python | `process_ai_llms_models_rules` được gọi theo từng request; không phải service/instance riêng. Rules quyết định xử lý logic và trường hợp cần gọi LLM. |
| DB kết quả | Database | `BEMT2003` là lượt chạy; `BEMT2002` lưu OCR/tổng hợp; `BEMT2005/2006` lưu dữ liệu trích xuất; `BEMT2004` lưu kết quả từng tiêu chí. |

## 4. Trạng thái hiện tại và trạng thái cần bổ sung

![Trạng thái](Temp/code_flow_acceptance_20261001/02_status_current_baseline.png)

Hiện tại DB chủ yếu phản ánh `PROCESSING`, `COMPLETED`, `FAILED`. Chưa biết một phiếu `PROCESSING` đang chờ Queue, đang OCR, đang LLM hay đang Rules; chưa có trạng thái retry riêng.

## 5. DB đang lưu gì và còn thiếu gì

| Bảng/field hiện có | Đang lưu | Phần còn thiếu để vận hành |
|---|---|---|
| `BEMT2003.StatusProcess` | `PROCESSING / COMPLETED / FAILED` | Chưa có bước xử lý chi tiết. |
| `BEMT2003.Status`, `Percentage` | Kết quả tổng OK/NG và % | Không cho biết kết quả được tạo ở bước nào. |
| `BEMT2003.TextContentOCR`, `TextContentAI`, `TextConditionFail` | Nội dung OCR/AI và điều kiện fail | Chưa chuẩn hóa ErrorCode/ErrorMessage. |
| `BEMT2002.RawContent`, `ConsolidatedContent`, `DataSourceType` | OCR và nội dung tổng hợp | Chưa lưu thời điểm bắt đầu/kết thúc OCR. |
| `BEMT2005/BEMT2006` | Dữ liệu trích xuất chứng từ | Chưa có trạng thái trích xuất theo lần gọi. |
| `BEMT2004.CriteriaStatus`, `Description`, `PromptSystem` | Kết quả tiêu chí | Chưa phân biệt kết quả từ RULE, LLM hay MIXED. |
| Chưa có field riêng | Retry và tiến độ | Cần `ProcessStage`, `RetryCount`, `StartedAt`, `FinishedAt`, `LastHeartbeatAt`, `ErrorCode`, `ErrorMessage`. |

## 6. Rủi ro có thể gây fail tháng 09

![Rủi ro](Temp/code_flow_acceptance_20261001/03_risk_summary.png)

- Có hai nhánh `return` trước khi cập nhật `FAILED`: gửi ngày nghỉ lỗi và OCR không trả dữ liệu. Phiếu có thể giữ `PROCESSING` dù không còn chạy.
- Queue nằm trong RAM nên restart service có thể mất các job đang chờ; DB không có dữ liệu để tự dựng lại Queue.
- Chưa có retry tự động có giới hạn; lỗi tạm thời và lỗi dữ liệu đang chưa được phân loại rõ.
- OCR có độ trễ đuôi dài: 591 lượt trong `app_2.log`, median 2,99 giây nhưng max 205,04 giây; 8 lượt từ 60 giây trở lên.
- LLM tái sử dụng một engine/GPU và khóa generate; an toàn VRAM hơn nhưng yêu cầu dài sẽ làm các phiếu sau chờ.
- Chưa đủ incident log để khẳng định một nguyên nhân duy nhất cho hai kỳ tháng 09; các rủi ro trên là điểm code/log đã xác nhận và phải loại bỏ trong baseline.

## 7. Phương án chốt

![Baseline](Temp/code_flow_acceptance_20261001/04_baseline_solution.png)

- Giữ `StatusProcess = PROCESSING / COMPLETED / FAILED` để tương thích ERP9.
- Bổ sung `ProcessStage = QUEUED / OCR_RUNNING / LLM_EXTRACTING / RULE_CHECKING / RETRY_WAIT`.
- Dùng `BEMT2003` làm nguồn trạng thái chính; Queue RAM chỉ là bộ tăng tốc. Khi API-AI restart, Worker nạp lại các run chưa hoàn tất từ DB.
- Retry tối đa 2 lần cho timeout/network/HTTP 5xx/service unavailable. Không retry hồ sơ thiếu file, request sai hoặc lỗi nghiệp vụ.
- Mỗi lỗi phải ghi `ErrorCode`, `ErrorMessage`, `RetryCount`, thời điểm bắt đầu/kết thúc và heartbeat.
- Giai đoạn ổn định: một phiếu trên Worker; OCR tối đa 2 file song song; LLM giữ `_gen_lock`. Chỉ tăng tải sau khi test đạt.
- Python trả lỗi có cấu trúc cho OCR/LLM/Rules; API-AI chịu trách nhiệm chuyển lỗi thành trạng thái DB.

## 8. Task owner deadline

![Timeline](Temp/code_flow_acceptance_20261001/05_timeline_1510.png)

| Thành phần | Hiện trạng | Phần thực hiện | Vấn đề | Phương án | Owner | Ưu tiên | Deadline |
|---|---|---|---|---|---|---|---|
| Trạng thái run | Có thể return mà không ghi FAILED | API-AI | Phiếu treo PROCESSING | Mọi lối thoát phải cập nhật stage/final status và ErrorCode | DEV API-AI | P0 | 04/10 |
| Queue | RAM Channel 200 job | API-AI + DB | Restart mất job chờ | DB là source of truth; khôi phục QUEUED/RETRY_WAIT khi startup | DEV API-AI + DB | P0 | 07/10 |
| Retry | Chưa có retry bền vững | API-AI + DB | Lỗi tạm thời phải chạy tay | Retry tối đa 2, lưu RetryCount/NextRetryAt; lỗi dữ liệu không retry | DEV API-AI + DB | P0 | 07/10 |
| OCR | Tối đa 7 file song song/phiếu | Python + API-AI | Độ trễ biến động, nguy cơ dồn tải | Config mặc định 2 file/phiếu; timeout/error code; health check | DEV Python + API-AI | P1 | 10/10 |
| LLM | Một engine, generate có lock | Python | Phiếu dài làm các phiếu sau chờ | Giữ lock; timeout; health; metric thời gian/token/error | DEV Python | P1 | 10/10 |
| Rules | Hàm trong Python, kết quả chung | Python + DB | Khó tách Rule/LLM khi lỗi | Ghi EngineSource RULE/LLM/MIXED và stage RULE_CHECKING | DEV Python + DB | P1 | 10/10 |
| Timestamp/lỗi | Chưa có field vận hành riêng | DB | Không đo được thời gian từng bước | Thêm StartedAt/FinishedAt/Heartbeat/ErrorCode/ErrorMessage | DEV DB + API-AI | P0 | 07/10 |
| Kiểm thử | Chưa có baseline fail/restart | QA | Có thể lặp lại lỗi tháng 09 | Test restart, OCR fail, LLM fail, timeout, retry, tải và case tháng 09 | QA/ASOFT | P0 | 13/10 |
| UAT/baseline | Chưa chốt | API/Python/DB/QA | Tiếp tục mở rộng ngoài mục tiêu | Chỉ fix lỗi nghiệm thu; chạy UAT 14/10; Go/No-Go 15/10 | Tech Lead + các owner | P0 | 15/10 |

## 9. Tiêu chí nghiệm thu cuối

- Không thay đổi ERP9.
- Một phiếu luôn xác định được đang ở stage nào.
- Không còn run `PROCESSING` treo sau khi luồng đã dừng.
- Restart API-AI không làm mất job chưa xử lý.
- Lỗi tạm thời tự retry tối đa 2 lần; lỗi cuối có mã, nội dung và thời gian.
- Có kết quả test OCR fail, LLM fail, timeout, restart và tải.
- Chạy lại bộ đại diện hai kỳ tháng 09 đạt kết quả thống nhất trước khi chốt baseline ngày 15/10.
- Sau khi chốt, chỉ triển khai task trong baseline; không mở rộng sang Rules/Training/nghiệp vụ AI khác trong task này.
'''
MD.write_text(md,encoding='utf-8')

# DOCX helpers
def set_font(r,size=10.5,bold=False,color=None):
    r.font.name='Times New Roman'; r._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman'); r._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman'); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade(c,color):
    pr=c._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),color); pr.append(e)
def border(c):
    pr=c._tc.get_or_add_tcPr(); b=pr.first_child_found_in('w:tcBorders')
    if b is None: b=OxmlElement('w:tcBorders'); pr.append(b)
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9E2F3'); b.append(e)
def heading(doc,text):
    p=doc.add_paragraph(); p.style='Heading 1'; p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5); r=p.add_run(text); set_font(r,16,True,'17365D')
def para(doc,text,bullet=False,bold=False,color=None):
    p=doc.add_paragraph(style='List Bullet' if bullet else None); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.04; r=p.add_run(text); set_font(r,10.2,bold,color); return p
def image(doc,path,caption,width=7.05):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width)); c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=c.add_run(caption); set_font(r,9,False,'5B6573')
def table(doc,headers,rows,font_size=8.6):
    t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(t.rows[0].cells,headers):
        shade(c,'17365D'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); set_font(r,font_size,True,'FFFFFF')
    for row in rows:
        cells=t.add_row().cells
        for idx,(c,txt) in enumerate(zip(cells,row)):
            border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(txt); set_font(r,font_size-0.2,idx==0)
    return t

doc=Document(); s=doc.sections[0]; s.top_margin=Inches(.5); s.bottom_margin=Inches(.5); s.left_margin=Inches(.58); s.right_margin=Inches(.58)
for stn in ['Normal','Heading 1','Heading 2','Title']:
    st=doc.styles[stn]; st.font.name='Times New Roman'; st._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
doc.styles['Normal'].font.size=Pt(10.2)
h=s.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=h.add_run('AI BEM | Baseline ổn định API Python DB'); set_font(r,9,False,'5B6573')
fo=s.footer.paragraphs[0]; fo.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=fo.add_run('Ngày báo cáo 01/10/2026'); set_font(r,9,False,'5B6573')

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(42); r=p.add_run('PHÂN TÍCH VÀ CHỐT PHƯƠNG ÁN'); set_font(r,24,True,'17365D')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('ỔN ĐỊNH XỬ LÝ AI BEM'); set_font(r,27,True,'17365D')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16); r=p.add_run('Phạm vi API-AI  |  AI Python  |  Database'); set_font(r,15,False,'168A8A')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Mục tiêu hoàn tất và nghiệm thu ngày 15/10/2026'); set_font(r,13,True,'C00000')
heading(doc,'Phạm vi và kết luận cần chốt')
for x in ['Không thay đổi ERP9; ERP9 chỉ là điểm đưa phiếu vào và hiển thị kết quả.', 'Chỉ cải tiến API-AI, AI Python và Database; không mở rộng sang Rules/Training/nghiệp vụ AI ngoài task.', 'Hiện DB chưa cho biết phiếu đang ở Queue, OCR, LLM hay Rules; Queue chưa phục hồi được sau restart và chưa có retry bền vững.', 'Phương án chốt: DB là nguồn trạng thái chính; Queue RAM là bộ tăng tốc; Python tái sử dụng engine có giới hạn tải và lỗi có cấu trúc.']:
    para(doc,x,True)

doc.add_page_break(); heading(doc,'1. Luồng end-to-end hiện tại')
image(doc,SOURCE_MAP,'Hình 1. ERP9 vào/ra; phần sửa chỉ nằm ở API-AI, AI Python và DB')

doc.add_page_break(); heading(doc,'2. Thành phần và instance hiện tại')
table(doc,['Thành phần','Thuộc phần','Cách tạo và tái sử dụng','Cách gọi','Xử lý lỗi hiện tại'],[
('Queue','API-AI','ChannelJobQueue Singleton trong RAM, tối đa 200 job, một reader.','Orchestrator enqueue; Worker dequeue.','Queue đầy thì chờ; restart không có recovery từ DB.'),
('OCR','AI Python','_engine giữ trong process; lazy/eager theo cấu hình.','API-AI OcrService gọi /ocr; tối đa 7 file song song/phiếu.','Python trả HTTP lỗi; API-AI catch nhưng có nhánh OCR rỗng return chưa ghi FAILED.'),
('LLM','AI Python','_llm_engine/_llm_registry khởi tạo startup và tái sử dụng; _gen_lock tuần tự generate.','Endpoint /llms/api/ai_llms_models.','Có xử lý OOM/cleanup; request dài làm request sau chờ.'),
('Rule Engine','AI Python','Hàm process_ai_llms_models_rules, không phải service riêng.','Được endpoint LLM gọi theo từng prompt.','Lỗi trả về endpoint; API-AI chịu trách nhiệm ghi trạng thái cuối.'),
('DB','Database','Mỗi BEMT2003 là một lượt chạy.','API-AI ghi BEMT2002/3/4/5/6.','Chưa có stage chi tiết, retry và timestamp vận hành riêng.'),
],8.4)
heading(doc,'3. Trạng thái hiện tại và baseline')
image(doc,IMG/'02_status_current_baseline.png','Hình 2. Phân biệt trạng thái đang có và trạng thái cần bổ sung')

doc.add_page_break(); heading(doc,'4. DB đang lưu gì và còn thiếu gì')
table(doc,['Bảng hoặc field','Hiện đang lưu','Thiếu để vận hành ổn định'],[
('BEMT2003.StatusProcess','PROCESSING / COMPLETED / FAILED','ProcessStage chi tiết: QUEUED, OCR_RUNNING, LLM_EXTRACTING, RULE_CHECKING, RETRY_WAIT.'),
('BEMT2003.Status, Percentage','Kết quả tổng OK/NG và %','Không cho biết bước tạo ra kết quả.'),
('TextContentOCR / TextContentAI / TextConditionFail','Nội dung OCR/AI và điều kiện fail','ErrorCode và ErrorMessage chuẩn hóa.'),
('BEMT2002','RawContent, ConsolidatedContent, DataSourceType','Thời gian OCR bắt đầu/kết thúc và trạng thái lần gọi.'),
('BEMT2005 / BEMT2006','Dữ liệu trích xuất chứng từ','Trạng thái và thời gian trích xuất.'),
('BEMT2004','CriteriaStatus, Description, PromptSystem','EngineSource = RULE / LLM / MIXED.'),
('Chưa có','Retry và heartbeat','RetryCount, NextRetryAt, StartedAt, FinishedAt, LastHeartbeatAt.'),
],8.6)
heading(doc,'5. Rủi ro chính gây fail')
image(doc,IMG/'03_risk_summary.png','Hình 3. Các rủi ro có bằng chứng từ source và log')
para(doc,'Chưa đủ incident log để kết luận một nguyên nhân duy nhất cho hai kỳ tháng 09. Baseline phải loại bỏ toàn bộ failure mode đã xác nhận ở trên.',False,True,'8A3E00')

doc.add_page_break(); heading(doc,'6. Phương án chốt API Python DB')
image(doc,IMG/'04_baseline_solution.png','Hình 4. Baseline không thay đổi ERP9')
for x in ['Giữ StatusProcess PROCESSING/COMPLETED/FAILED để tương thích ERP9; bổ sung ProcessStage cho vận hành.', 'BEMT2003 là nguồn trạng thái chính; API-AI khôi phục run QUEUED/RETRY_WAIT khi restart.', 'Retry tối đa 2 lần cho timeout, network, HTTP 5xx và service unavailable; không retry lỗi hồ sơ/dữ liệu.', 'Ổn định trước: một phiếu trên Worker, OCR tối đa 2 file song song, LLM giữ gen_lock; chỉ tăng tải sau test.', 'Python trả ErrorCode có cấu trúc; API-AI chịu trách nhiệm ghi stage, error, retry và thời gian vào DB.']:
    para(doc,x,True)
heading(doc,'7. Kế hoạch đến 15/10')
image(doc,IMG/'05_timeline_1510.png','Hình 5. Thứ tự xử lý P0, tích hợp, kiểm thử và chốt baseline')

# Landscape task matrix
sec=doc.add_section(WD_SECTION_START.NEW_PAGE); sec.orientation=WD_ORIENT.LANDSCAPE; sec.page_width,sec.page_height=sec.page_height,sec.page_width; sec.top_margin=Inches(.42); sec.bottom_margin=Inches(.42); sec.left_margin=Inches(.42); sec.right_margin=Inches(.42)
heading(doc,'8. Danh sách task owner deadline')
table(doc,['Thành phần','Hiện trạng','Phần','Vấn đề','Phương án','Owner','P','Deadline'],[
('Trạng thái run','Có nhánh return không ghi FAILED','API-AI','Treo PROCESSING','Mọi lối thoát ghi stage/final status/ErrorCode','DEV API-AI','P0','04/10'),
('Queue','RAM Channel 200','API-AI/DB','Restart mất job chờ','DB source of truth; startup recovery','DEV API-AI + DB','P0','07/10'),
('Retry','Chưa bền vững','API-AI/DB','Lỗi tạm thời chạy tay','Retry ≤2; RetryCount/NextRetryAt','DEV API-AI + DB','P0','07/10'),
('OCR','7 file song song/phiếu','Python/API-AI','Độ trễ đuôi dài','Default 2; timeout; health; ErrorCode','DEV Python + API-AI','P1','10/10'),
('LLM','Một engine, gen_lock','Python','Phiếu dài làm phiếu sau chờ','Giữ lock; timeout; metric thời gian/token/error','DEV Python','P1','10/10'),
('Rules','Kết quả chung','Python/DB','Khó tách nguồn lỗi','EngineSource RULE/LLM/MIXED','DEV Python + DB','P1','10/10'),
('DB vận hành','Thiếu field','DB/API-AI','Không đo stage/thời gian','Stage/retry/time/heartbeat/error fields','DEV DB + API-AI','P0','07/10'),
('Kiểm thử','Chưa có baseline fail/restart','QA','Nguy cơ lặp fail tháng 09','Test restart/OCR/LLM/timeout/retry/tải/case T09','QA ASOFT','P0','13/10'),
('UAT','Chưa chốt','API/Python/DB','Mở rộng ngoài mục tiêu','UAT 14/10; Go/No-Go 15/10','Tech Lead + owner','P0','15/10'),
],7.4)
heading(doc,'9. Điều kiện nghiệm thu cuối')
for x in ['Không thay đổi ERP9.', 'Biết được mỗi phiếu đang ở stage nào.', 'Không còn PROCESSING treo sau khi luồng dừng.', 'Restart API-AI không làm mất job chưa xử lý.', 'Retry lỗi tạm thời tối đa 2 lần; lỗi cuối có mã và thời gian.', 'Test đạt các trường hợp OCR fail, LLM fail, timeout, restart, retry và tải.', 'Bộ đại diện hai kỳ tháng 09 chạy đạt trước khi chốt baseline ngày 15/10.', 'Sau khi chốt chỉ triển khai nội dung trong baseline; không mở rộng Rules/Training trong task này.']:
    para(doc,x,True)
doc.save(DOCX)
print(MD); print(DOCX)
