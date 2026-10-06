from openpyxl import load_workbook
from pathlib import Path
p=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\DATA_BEM AI_MEIKO_30092026.xlsx')
wb=load_workbook(p,data_only=False,read_only=True,keep_links=False)
for ws in wb.worksheets:
    hits=[]
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            v=cell.value
            if v is not None and ('Độ chính xác' in str(v) or 'AVERAGE' in str(v) or 'COUNTIF' in str(v) or 'Kết quả tháng 09' in str(v)):
                hits.append((cell.coordinate,str(v)[:220]))
    if hits:
        print('---',ws.title,len(hits))
        for h in hits[:80]: print(h)
