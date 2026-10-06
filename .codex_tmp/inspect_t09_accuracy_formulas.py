from openpyxl import load_workbook
from pathlib import Path
from collections import Counter
p=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\DATA_BEM AI_MEIKO_30092026.xlsx')
wb=load_workbook(p,data_only=False,read_only=True,keep_links=False)
ws=wb['Kết quả tháng 09 và xử lý']
vals=[]
for r,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
    if len(row)>=19 and row[7]: vals.append((r,row[7],row[16],row[18]))
print('rows',len(vals),'types',Counter(type(x[3]).__name__ for x in vals))
print('formulas',sum(isinstance(x[3],str) and x[3].startswith('=') for x in vals))
for x in vals:
    if ('AI trả lời đúng' in str(x[2]) or 'AI trả lời tốt' in str(x[2]) or 'AI đọc đúng' in str(x[2])) and x[3] in (0,0.0,'0%','0.0%'):
        print('ZERO',x); 
