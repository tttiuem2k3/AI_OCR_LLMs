import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
WORK = ROOT / "Accuracy_Recheck_T09_20261002"
BOOK = ROOT / "DATA_BEM AI_MEIKO_30092026.xlsx"
POS = ("AI đọc đúng", "AI trả lời đúng", "AI trả lời đúng theo Rules", "AI trả lời tốt")
NEG = ("AI đọc sai", "AI trả lời sai", "AI đối chiếu sai")

def load_json(name):
    return json.loads((WORK / name).read_text(encoding="utf-8-sig"))

def pct(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value) / 100 if float(value) > 1.0001 else float(value)
    number = float(str(value).replace("%", "").replace(",", ".").strip())
    return number / 100 if number > 1.0001 else number

def positive_only(text):
    text = str(text or "")
    return any(label in text for label in POS) and not any(label in text for label in NEG)

db = load_json("db_t09_latest_20261002.json")["Records"]
db_by = {str(row["VoucherNo"]): row for row in db}
special = {row["voucher"]: row for row in load_json("special_review_decisions_t09_20261002.json")}
missing = load_json("missing_0298_db.json")

wb = load_workbook(BOOK, read_only=False, data_only=False)
ws = wb["Kết quả tháng 09 và xử lý"]
row_by = {str(ws.cell(row, 8).value): row for row in range(3, ws.max_row + 1) if ws.cell(row, 8).value}

insert_missing = missing["VoucherNo"] not in row_by
insert_row_values = [
    None,
    missing["CreateDate"],
    "Nguyên vật liệu",
    missing["Nam"],
    missing["Thang"],
    missing["CurrencyID"],
    missing["AdvanceUserID"],
    missing["VoucherNo"],
    "OK" if missing["IsVoucherCompleted"] == 1 else "NG",
    missing["SoFile"],
    None,
    None,
    "OK" if missing.get("ApprovingLevel") == 6 else "NG",
    None,
    None,
    "- Phiếu chưa có kết quả AI OK/NG mới nhất trên DB.\n- Chưa đủ căn cứ để đánh giá AI đọc/đối chiếu đúng hay sai.",
    "- Không có kết quả AI",
    "- Cần chạy AI trước khi nghiệm thu phiếu này.",
    0,
    "- DB mới nhất chưa có lần chạy AI hợp lệ.",
]

def adjusted_row(row):
    return row + 1 if insert_missing and row >= 292 else row

updates = []
accuracy_only = []
for voucher, original_row in row_by.items():
    dbrow = db_by.get(voucher)
    if not dbrow:
        continue
    row = adjusted_row(original_row)
    if voucher in special:
        item = special[voucher]
        updates.append({
            "voucher": voucher,
            "row": row,
            "values": [item["p"], item["q"], item["r"], item["s"], item["t"]],
            "fill": item.get("fill"),
            "mode": "review",
        })
        continue
    process = str(dbrow.get("StatusProcess") or "").upper()
    status = str(dbrow.get("Status") or "").upper()
    issue = ws.cell(original_row, 17).value
    if process == "COMPLETED" and status in ("OK", "NG") and positive_only(issue) and pct(ws.cell(original_row, 19).value) != 1.0:
        updates.append({"voucher": voucher, "row": row, "values": [1], "fill": "FF77BC65", "mode": "accuracy"})
        accuracy_only.append(voucher)

stt = [[row - 2] for row in range(3, ws.max_row + 2 if insert_missing else ws.max_row + 1)]
plan = {
    "book": str(BOOK),
    "output": str(BOOK),
    "insertMissing": insert_missing,
    "insertRow": 292,
    "sourceRange": "A292:T965",
    "targetRange": "A293:T966",
    "insertRowValues": insert_row_values,
    "sttRange": f"A3:A{966 if insert_missing else ws.max_row}",
    "sttValues": stt,
    "updates": updates,
    "counts": {"review": sum(u["mode"] == "review" for u in updates), "accuracy": len(accuracy_only), "inserted": 1 if insert_missing else 0},
}
(WORK / "artifact_update_plan_t09_accuracy_20261002.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(plan["counts"], ensure_ascii=False, indent=2))
