from pathlib import Path
from collections import Counter
from decimal import Decimal, InvalidOperation
import json, re
from openpyxl import load_workbook

root=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09')
book=root/'DATA_BEM AI_MEIKO_30092026.xlsx'
db_path=root/'Accuracy_Recheck_T09_20261002'/'db_t09_latest_20261002.json'
out_dir=root/'Accuracy_Recheck_T09_20261002'
db=json.loads(db_path.read_text(encoding='utf-8-sig'))['Records']

def pct(v):
    if v is None or str(v).strip()=='': return None
    try:
        n=float(str(v).replace('%','').strip())
        return n/100 if n>1.00001 else n
    except: return None

def sval(v): return None if v is None else str(v).strip().upper()

wb=load_workbook(book,data_only=True,read_only=True,keep_links=False)
ws=wb['Kết quả tháng 09 và xử lý']
headers=None
for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=10,values_only=True),1):
    if any(str(v or '').strip()=='Độ chính xác của AI' for v in row):
        headers={str(v).strip():i+1 for i,v in enumerate(row) if v is not None}; header_row=ridx;break
cols={'voucher':headers['Số DNTT'],'ai':headers['Kết quả phiếu (OK / NG)'],'pct':headers['Kết quả phiếu (%)'],'issue':headers['Loại vấn đề'],'acc':headers['Độ chính xác của AI'],'note':headers['Ghi chú'],'type':headers['Loại DNTT'],'files':headers['Số file đính kèm']}
wb_rows={}
for r,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),header_row+1):
    v=row[cols['voucher']-1]
    if not v: continue
    wb_rows[str(v).strip()]={'row':r,'voucher':str(v).strip(),'ai':row[cols['ai']-1],'pct':row[cols['pct']-1],'issue':row[cols['issue']-1],'acc':row[cols['acc']-1],'note':row[cols['note']-1],'type':row[cols['type']-1],'files':row[cols['files']-1]}
db_rows={str(x['VoucherNo']).strip():x for x in db}
missing_in_book=sorted(set(db_rows)-set(wb_rows))
missing_in_db=sorted(set(wb_rows)-set(db_rows))
status_mismatch=[]; pct_mismatch=[]; changed_after_note=[]
for voucher,w in wb_rows.items():
    d=db_rows.get(voucher)
    if not d: continue
    wb_status=sval(w['ai']); db_status=sval(d['Status'])
    if wb_status!=db_status: status_mismatch.append({'voucher':voucher,'row':w['row'],'workbook_status':w['ai'],'db_status':d['Status'],'workbook_pct':w['pct'],'db_pct':d['Percentage'],'db_process':d['StatusProcess'],'db_run':d['RunCreateDate'],'issue':w['issue'],'note':w['note']})
    wp=pct(w['pct']); dp=pct(d['Percentage'])
    if wp is not None and dp is not None and abs(wp-dp)>0.00011: pct_mismatch.append({'voucher':voucher,'row':w['row'],'workbook_status':w['ai'],'db_status':d['Status'],'workbook_pct':w['pct'],'db_pct':d['Percentage'],'db_process':d['StatusProcess'],'db_run':d['RunCreateDate'],'issue':w['issue'],'note':w['note']})
    note=str(w['note'] or '')
    m=re.search(r'(\d{2}/\d{2}/2026)',note)
    # Report if current DB run date strictly after explicit last review date in note (DD/MM/YYYY)
    if m and d.get('RunCreateDate'):
        from datetime import datetime
        try:
            review=datetime.strptime(m.group(1),'%d/%m/%Y').date()
            run_date=str(d['RunCreateDate'])[:10]
            run=datetime.fromisoformat(run_date).date()
            if run>review: changed_after_note.append({'voucher':voucher,'row':w['row'],'review_date':str(review),'run_date':str(run),'workbook_status':w['ai'],'db_status':d['Status'],'db_pct':d['Percentage'],'issue':w['issue']})
        except: pass

pos_labels=('AI đọc đúng','AI trả lời đúng','AI trả lời đúng theo Rules','AI trả lời tốt')
neg_labels=('AI đọc sai','AI trả lời sai','AI đối chiếu sai')
positive_zero=[]; positive_below_100=[]
for v,w in wb_rows.items():
    issue=str(w['issue'] or '')
    pos=any(label.lower() in issue.lower() for label in pos_labels)
    neg=any(label.lower() in issue.lower() for label in neg_labels)
    accuracy=pct(w['acc'])
    d=db_rows.get(v)
    entry={**w,'db_status':d.get('Status') if d else None,'db_pct':d.get('Percentage') if d else None,'db_process':d.get('StatusProcess') if d else None,'db_run':d.get('RunCreateDate') if d else None,'has_negative_label':neg}
    if pos and accuracy==0: positive_zero.append(entry)
    if pos and not neg and accuracy is not None and accuracy<.9999: positive_below_100.append(entry)

result={
 'workbook_count':len(wb_rows),'db_count':len(db_rows),'missing_in_workbook':missing_in_book,'missing_in_db':missing_in_db,
 'status_mismatch_count':len(status_mismatch),'percentage_mismatch_count':len(pct_mismatch),'db_run_after_note_count':len(changed_after_note),
 'positive_zero_count':len(positive_zero),'positive_below_100_count':len(positive_below_100),
 'status_mismatch':status_mismatch,'percentage_mismatch':pct_mismatch,'db_run_after_note':changed_after_note,'positive_zero':positive_zero,'positive_below_100':positive_below_100,
 'db_status':Counter(str(x.get('Status')) for x in db),'db_process':Counter(str(x.get('StatusProcess')) for x in db),
}
(out_dir/'workbook_vs_db_t09_20261002.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print('WORKBOOK',len(wb_rows),'DB',len(db_rows))
print('MISSING_IN_WORKBOOK',len(missing_in_book),missing_in_book)
print('MISSING_IN_DB',len(missing_in_db),missing_in_db)
print('STATUS_MISMATCH',len(status_mismatch))
print('PCT_MISMATCH',len(pct_mismatch))
print('RUN_AFTER_NOTE',len(changed_after_note))
print('POSITIVE_ZERO',len(positive_zero),'POSITIVE_BELOW_100',len(positive_below_100))
print('DB_STATUS',result['db_status'])
print('DB_PROCESS',result['db_process'])
print('STATUS_SAMPLES')
for item in status_mismatch[:60]: print(json.dumps(item,ensure_ascii=False,default=str))
print('PCT_SAMPLES')
for item in pct_mismatch[:60]: print(json.dumps(item,ensure_ascii=False,default=str))
print('POS_ZERO_SAMPLES')
for item in positive_zero[:60]: print(json.dumps(item,ensure_ascii=False,default=str))
print('OUT',out_dir/'workbook_vs_db_t09_20261002.json')
