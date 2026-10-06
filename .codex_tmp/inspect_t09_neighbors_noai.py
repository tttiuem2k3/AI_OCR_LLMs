from openpyxl import load_workbook
from pathlib import Path
p=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\DATA_BEM AI_MEIKO_30092026.xlsx')
wb=load_workbook(p,data_only=True,read_only=True,keep_links=False); ws=wb['Kết quả tháng 09 và xử lý']
for r,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
    v=row[7] if len(row)>7 else None
    if v and ('NVL/09/2026/029' in str(v) or str(row[10]).upper() not in ('OK','NG')):
        if 'NVL/09/2026/029' in str(v) or r<20:
            print(r, [row[i] if i<len(row) else None for i in range(20)])
print('NO_AI_EXAMPLES')
n=0
for r,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
    v=row[7] if len(row)>7 else None
    if v and str(row[10]).upper() not in ('OK','NG'):
        print(r, v, 'K=',row[10],'L=',row[11],'P-T=',row[15:20]);n+=1
        if n>=30:break
