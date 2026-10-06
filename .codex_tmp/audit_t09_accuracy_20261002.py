from pathlib import Path
from collections import Counter, defaultdict
import json, re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

book=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\DATA_BEM AI_MEIKO_30092026.xlsx')
sheet='Kết quả tháng 09 và xử lý'
wb_formula=load_workbook(book,data_only=False,read_only=False)
wb_value=load_workbook(book,data_only=True,read_only=True)
ws=wb_formula[sheet]
wv=wb_value[sheet]
print('SHEETS', wb_formula.sheetnames)
print('DIM', ws.max_row, ws.max_column, 'STATE', ws.sheet_state, 'AUTO_FILTER', ws.auto_filter.ref)
header_row=None
headers={}
for r in range(1,min(ws.max_row,20)+1):
    vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
    if any(str(v or '').strip()=='Độ chính xác của AI' for v in vals):
        header_row=r
        headers={str(v).strip():c for c,v in enumerate(vals,1) if v is not None}
        break
print('HEADER_ROW',header_row)
for c in range(1,ws.max_column+1):
    print(c,get_column_letter(c),repr(ws.cell(header_row,c).value))

positive=('AI đọc đúng','AI trả lời đúng','AI trả lời đúng theo Rules','AI trả lời tốt')
negative=('AI đọc sai','AI trả lời sai','AI đối chiếu sai')
col_voucher=headers.get('Số DNTT') or headers.get('Số phiếu') or 8
col_issue=headers.get('Loại vấn đề') or 17
col_acc=headers.get('Độ chính xác của AI') or 19
col_ai=headers.get('Kết quả phiếu (OK / NG)') or headers.get('Kết quả AI') or 11
col_pct=headers.get('Kết quả phiếu (%)') or 12
col_type=headers.get('Loại DNTT') or 3
col_files=headers.get('Số file đính kèm') or 10
col_p=16; col_r=18; col_t=20
records=[]
for r in range(header_row+1,ws.max_row+1):
    voucher=wv.cell(r,col_voucher).value
    if not voucher: continue
    issue=str(wv.cell(r,col_issue).value or '')
    acc=wv.cell(r,col_acc).value
    f=ws.cell(r,col_acc).value
    has_pos=any(x.lower() in issue.lower() for x in positive)
    has_neg=any(x.lower() in issue.lower() for x in negative)
    zero=(acc==0 or acc==0.0 or str(acc).strip() in ('0','0%','0.0%'))
    records.append({
        'row':r,'voucher':voucher,'type':wv.cell(r,col_type).value,'ai':wv.cell(r,col_ai).value,
        'ai_pct':wv.cell(r,col_pct).value,'files':wv.cell(r,col_files).value,'issue':issue,'acc':acc,'formula':f,
        'p':wv.cell(r,col_p).value,'r':wv.cell(r,col_r).value,'t':wv.cell(r,col_t).value,
        'pos':has_pos,'neg':has_neg,'zero':zero,
    })
print('RECORDS',len(records))
print('ACC TYPES',Counter(type(x['formula']).__name__ for x in records))
print('ACC VALUES TOP',Counter(x['acc'] for x in records).most_common(20))
print('ISSUE EMPTY',sum(not x['issue'] for x in records),'ACC EMPTY',sum(x['acc'] is None for x in records))
problem=[x for x in records if x['pos'] and x['zero']]
print('POSITIVE_WITH_ZERO',len(problem))
print('  POS_ONLY_ZERO',sum(x['pos'] and not x['neg'] and x['zero'] for x in records))
print('  MIXED_POS_NEG_ZERO',sum(x['pos'] and x['neg'] and x['zero'] for x in records))
print('  BY_AI',Counter(str(x['ai']) for x in problem))
print('  BY_TYPE',Counter(str(x['type']) for x in problem))
print('SAMPLES')
for x in problem[:80]:
    print(json.dumps(x,ensure_ascii=False,default=str))

# Other logical conflicts.
conflicts={
    'positive_acc_lt_1': [x for x in records if x['pos'] and not x['neg'] and isinstance(x['acc'],(int,float)) and x['acc']<1],
    'negative_acc_1': [x for x in records if x['neg'] and x['acc']==1],
    'no_ai_but_accuracy': [x for x in records if str(x['ai']).upper() not in ('OK','NG') and isinstance(x['acc'],(int,float)) and x['acc']>0],
    'has_ai_zero_no_negative': [x for x in records if str(x['ai']).upper() in ('OK','NG') and x['zero'] and not x['neg']],
}
for name,items in conflicts.items():
    print('CONFLICT',name,len(items))
    for x in items[:20]: print(' ',x['row'],x['voucher'],repr(x['ai']),repr(x['issue']),x['acc'])

out=book.parent/'Accuracy_Recheck_T09_20261002'
out.mkdir(exist_ok=True)
(out/'audit_current_workbook.json').write_text(json.dumps({'headers':headers,'records':records,'conflict_counts':{k:len(v) for k,v in conflicts.items()}},ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print('AUDIT',out/'audit_current_workbook.json')
