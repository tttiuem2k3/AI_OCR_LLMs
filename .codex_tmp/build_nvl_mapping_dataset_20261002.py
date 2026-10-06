import json
import re
import shutil
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
WORK = ROOT / "Mapping_DNTT_Files_NVL_Work_20261002"
OUT = ROOT / "Mapping_DNTT_Files_NVL_20261002"
ATTACH_OUT = OUT / "Files_dinh_kem"
EXCEL_OUT = OUT / "Bo_du_lieu_mapping_DNTT_Files_NVL_20261002.xlsx"
VOUCHERS = ["NVL/09/2026/0045", "NVL/09/2026/0082", "NVL/09/2026/0529"]

OUT.mkdir(parents=True, exist_ok=True)
ATTACH_OUT.mkdir(parents=True, exist_ok=True)


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


db = load_json(WORK / "db_mapping_3_nvl.json")["tables"]
raw = load_json(WORK / "db_raw_and_section_mapping.json")["tables"]

headers = {x["VoucherNo"]: x for x in db["headers"]}
lines = defaultdict(list)
criteria = defaultdict(list)
sections = defaultdict(list)
fields = defaultdict(list)
attachments = defaultdict(list)
section_fields = defaultdict(list)
for row in db["lines"]:
    lines[row["VoucherNo"]].append(row)
for row in db["criteria"]:
    criteria[row["VoucherNo"]].append(row)
for row in db["sections"]:
    sections[row["VoucherNo"]].append(row)
for row in db["fields"]:
    fields[row["VoucherNo"]].append(row)
for row in db["attachments"]:
    attachments[row["VoucherNo"]].append(row)
for row in raw["section_fields"]:
    section_fields[row["VoucherNo"]].append(row)


def slug(voucher):
    return voucher.replace("/", "-")


def norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def norm_key(text):
    return re.sub(r"[^A-Za-z0-9]", "", str(text or "")).upper()


def number_tokens(text):
    tokens = []
    for token in re.findall(r"\d+", str(text or "")):
        tokens.append(token)
        stripped = token.lstrip("0") or "0"
        if stripped != token:
            tokens.append(stripped)
    return sorted(set(tokens), key=lambda x: (-len(x), x))


def split_codes(text):
    return [x.strip() for x in re.split(r"[;,\n]+", str(text or "")) if x.strip()]


def money_float(value):
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", ""))
    except Exception:
        return None


def find_source_folder(voucher):
    h = headers[voucher]
    candidates = []
    for kind in ("DNTT_FINAL", "DNTT_TEMP"):
        path = ROOT / "Attached" / kind / "2026" / "9" / str(h.get("CurrencyID") or "") / str(h.get("AdvanceUserID") or "") / slug(voucher)
        if path.exists():
            candidates.append(path)
    if not candidates:
        return None
    db_names = {str(x.get("AttachName") or "").lower() for x in attachments[voucher]}
    return max(candidates, key=lambda p: sum(1 for f in p.iterdir() if f.is_file() and f.name.lower() in db_names))


def page_count(path):
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf" and PdfReader:
            return len(PdfReader(str(path), strict=False).pages), "PDF pages"
        if suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            return len(wb.worksheets), "Excel sheets"
        if suffix in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"):
            return 1, "image file"
        return 1, "file"
    except Exception as exc:
        return None, f"count error: {exc!r}"


def classify_file(name):
    upper = name.upper()
    if upper.startswith("IV") or "INVOICE" in upper or upper.startswith("VAT"):
        return "Invoice/VAT"
    if upper.startswith("COM"):
        return "Commercial Invoice"
    if upper.startswith("PL") or "PACK" in upper:
        return "Packing list"
    if upper.startswith("PO"):
        return "PO"
    if "RINGI" in upper:
        return "Ringi"
    if "TOKHAI" in upper or "HQ" in upper:
        return "Tờ khai"
    if upper.startswith("CT") or "CONTRACT" in upper:
        return "Contract"
    if "INSPEC" in upper:
        return "Inspection"
    return "Khác"


physical_rows = []
copied_summary = {}
for voucher in VOUCHERS:
    src = find_source_folder(voucher)
    dest = ATTACH_OUT / slug(voucher)
    if dest.exists():
        shutil.rmtree(dest)
    dest.mkdir(parents=True, exist_ok=True)
    disk = {f.name.lower(): f for f in src.iterdir() if f.is_file()} if src else {}
    total_pages = 0
    missing = 0
    for idx, att in enumerate(attachments[voucher], 1):
        name = str(att.get("AttachName") or "")
        source = disk.get(name.lower())
        exists = source is not None
        if exists:
            target = dest / name
            shutil.copy2(source, target)
            pages, page_note = page_count(target)
            if isinstance(pages, int):
                total_pages += pages
        else:
            target = None
            pages, page_note = None, "missing physical file"
            missing += 1
        physical_rows.append({
            "VoucherNo": voucher,
            "FileOrder": idx,
            "AttachID": att.get("AttachID"),
            "AttachName": name,
            "DocType_by_name": classify_file(name),
            "AttachCreateDate": att.get("AttachCreateDate"),
            "AttachModifyDate": att.get("AttachModifyDate"),
            "AttachBytes": att.get("AttachBytes"),
            "AttachHash": att.get("AttachHash"),
            "PhysicalFound": "Có" if exists else "Không",
            "PhysicalSourceFolder": str(src) if src else "",
            "CopiedPath": str(target) if target else "",
            "PageOrSheetCount": pages,
            "PageCountNote": page_note,
        })
    copied_summary[voucher] = {"source": str(src) if src else "", "dest": str(dest), "files": len(attachments[voucher]), "missing": missing, "pages": total_pages}


field_groups = defaultdict(dict)
section_meta = {}
for row in section_fields:
    key = (row["VoucherNo"], row["SectionAPK"], row["ExtractRowAPK"])
    section_meta[key] = {
        "VoucherNo": row["VoucherNo"],
        "SectionAPK": row["SectionAPK"],
        "SectionType": row["SectionType"],
        "SectionOrder": row["SectionOrder"],
        "SectionTitle": row["SectionTitle"],
        "TotalAmount": row["TotalAmount"],
        "TotalCurrency": row["TotalCurrency"],
        "Signature": row["Signature"],
        "ExtractRowAPK": row["ExtractRowAPK"],
        "ExtractRowNo": row["ExtractRowNo"],
        "ExtractDescription": row["ExtractDescription"],
    }
    display = row.get("DisplayName") or row.get("FieldID")
    value = norm(row.get("FieldValue"))
    if display and value:
        if display in field_groups[key]:
            if value not in str(field_groups[key][display]).split(" | "):
                field_groups[key][display] += " | " + value
        else:
            field_groups[key][display] = value

compact_section_rows = []
for key, meta in section_meta.items():
    vals = field_groups[key]
    compact_section_rows.append({
        **meta,
        "Tên file": vals.get("Tên file", ""),
        "Số hóa đơn": vals.get("Số hóa đơn", ""),
        "Ngày hóa đơn": vals.get("Ngày hóa đơn", ""),
        "Số tiền": vals.get("Số tiền", ""),
        "Loại tiền": vals.get("Loại tiền", ""),
        "Tên nhà cung cấp": vals.get("Tên nhà cung cấp", ""),
        "Số PO/Contract": vals.get("Số PO/Contract", vals.get("Số hợp đồng", "")),
        "Số Ringi": vals.get("Số Ringi", ""),
        "Điều kiện giao hàng": vals.get("Điều kiện giao hàng", ""),
        "Điều kiện thanh toán": vals.get("Điều kiện thanh toán", ""),
    })


def section_rows_for(voucher):
    return [r for r in compact_section_rows if r["VoucherNo"] == voucher]


def files_by_names(voucher):
    return [r["AttachName"] for r in physical_rows if r["VoucherNo"] == voucher]


manual_rows = []
for voucher in VOUCHERS:
    sec_rows = section_rows_for(voucher)
    file_names = files_by_names(voucher)
    for line in lines[voucher]:
        invoice = norm(line.get("InvoiceNo"))
        inv_tokens = number_tokens(invoice)
        po_codes = split_codes(line.get("ContractNo"))
        amount = money_float(line.get("RequestAmount"))
        currency = norm(line.get("CurrencyID"))

        matched_files = set()
        evidence = []
        roles = defaultdict(list)

        for name in file_names:
            key = norm_key(name)
            role = classify_file(name)
            if invoice and any(token and token in key for token in inv_tokens):
                matched_files.add(name)
                roles[role].append(name)
                evidence.append(f"Tên file khớp Invoice {invoice}: {name}")
            for po in po_codes:
                if norm_key(po) and norm_key(po) in key:
                    matched_files.add(name)
                    roles[role].append(name)
                    evidence.append(f"Tên file khớp PO {po}: {name}")

        for sec in sec_rows:
            fname = norm(sec.get("Tên file"))
            if not fname:
                continue
            sec_inv = norm(sec.get("Số hóa đơn"))
            sec_po = norm(sec.get("Số PO/Contract"))
            sec_amount = money_float(sec.get("Số tiền")) or money_float(sec.get("TotalAmount"))
            sec_currency = norm(sec.get("Loại tiền") or sec.get("TotalCurrency"))
            matched = False
            if invoice and sec_inv and (norm_key(invoice) in norm_key(sec_inv) or norm_key(sec_inv) in norm_key(invoice) or any(token in norm_key(sec_inv) for token in inv_tokens)):
                matched = True
                evidence.append(f"DB extracted Invoice {sec_inv} từ file {fname}")
            if any(norm_key(po) and norm_key(po) in norm_key(sec_po) for po in po_codes):
                matched = True
                evidence.append(f"DB extracted PO {sec_po} từ file {fname}")
            if amount is not None and sec_amount is not None and currency and currency == sec_currency and abs(amount - sec_amount) < 0.01:
                matched = True
                evidence.append(f"DB extracted amount {sec_amount:g} {sec_currency} khớp dòng ĐNTT từ file {fname}")
            if matched:
                matched_files.add(fname)
                roles[classify_file(fname)].append(fname)

        manual_rows.append({
            "VoucherNo": voucher,
            "DNTT_LineNo": line.get("OrderNo"),
            "DNTT_InvoiceNo": invoice,
            "DNTT_InvoiceDate": line.get("InvoiceDate"),
            "DNTT_PO_Contract": norm(line.get("ContractNo")),
            "DNTT_Amount": line.get("RequestAmount"),
            "DNTT_Currency": currency,
            "DocumentGroupID_Target": f"{slug(voucher)}__LINE_{line.get('OrderNo')}__INV_{invoice or 'NA'}",
            "ManualTargetMapping": "DNTT line -> Invoice/VAT + Customs + PO/Ringi/Contract/Inspection files that share InvoiceNo, PO/ContractNo, amount, supplier or payment term.",
            "MatchedFiles_All": "\n".join(sorted(matched_files)),
            "Matched_Invoice_VAT": "\n".join(sorted(set(roles.get("Invoice/VAT", [])))),
            "Matched_CommercialInvoice": "\n".join(sorted(set(roles.get("Commercial Invoice", [])))),
            "Matched_PackingList": "\n".join(sorted(set(roles.get("Packing list", [])))),
            "Matched_Customs": "\n".join(sorted(set(roles.get("Tờ khai", [])))),
            "Matched_PO": "\n".join(sorted(set(roles.get("PO", [])))),
            "Matched_Ringi": "\n".join(sorted(set(roles.get("Ringi", [])))),
            "Matched_Contract_Inspection": "\n".join(sorted(set(roles.get("Contract", []) + roles.get("Inspection", [])))),
            "MappingEvidence": "\n".join(dict.fromkeys(evidence)),
            "MappingGap_NeedForAI": "Cần lưu sẵn DNTT_LineID/InvoiceNo/PONo -> AttachID + DocRole + ExtractedKey để AI không phải tự dò lại từ tên file và text OCR.",
        })

target_schema_rows = [
    {"Field": "PaymentVoucherID", "Meaning": "APK/Số DNTT", "WhyNeeded": "Khóa gốc để gom toàn bộ file và kết quả AI của một phiếu."},
    {"Field": "PaymentLineID", "Meaning": "Dòng ĐNTT/OrderNo", "WhyNeeded": "Một phiếu NVL có thể có nhiều Invoice/PO; cần map theo từng dòng."},
    {"Field": "DocumentGroupID", "Meaning": "Nhóm chứng từ của một lần thanh toán", "WhyNeeded": "Gom Invoice/VAT/PO/Ringi/tờ khai/PL theo cùng nghiệp vụ."},
    {"Field": "AttachID + AttachName", "Meaning": "File vật lý trong ERP", "WhyNeeded": "Tránh AI đọc nhầm file cùng tên hoặc file ngoài nhóm."},
    {"Field": "DocRole", "Meaning": "Invoice, VAT, PO, Ringi, Customs, PackingList, Contract, Inspection", "WhyNeeded": "AI biết file dùng để kiểm tra tiêu chí nào."},
    {"Field": "MatchKey", "Meaning": "InvoiceNo, PO/ContractNo, RingiNo, Amount, Currency, Supplier, Date", "WhyNeeded": "Khóa nối DNTT với dữ liệu trích xuất từ từng file."},
    {"Field": "MappingConfidence", "Meaning": "High/Medium/Need review", "WhyNeeded": "Tách mapping chắc chắn khỏi mapping cần người xác nhận."},
    {"Field": "CriteriaScope", "Meaning": "Tiêu chí dùng file này: Số tiền, NCC, Invoice, PO/Ringi, Deadline, Incoterm...", "WhyNeeded": "Giảm số file AI phải đọc lại khi check từng tiêu chí."},
]


def write_sheet(wb, title, rows, headers=None):
    ws = wb.create_sheet(title)
    if headers is None:
        keys = []
        for row in rows:
            for key in row.keys():
                if key not in keys:
                    keys.append(key)
        headers = keys
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    format_sheet(ws)
    return ws


def safe_value(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return value


def format_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, column_cells in enumerate(ws.columns, 1):
        max_len = 0
        for cell in column_cells[:200]:
            text = str(cell.value or "")
            max_len = max(max_len, min(80, max(len(line) for line in text.splitlines()) if text else 0))
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(55, max_len + 2))
    if ws.max_row > 1 and ws.max_column > 0:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=re.sub(r"[^A-Za-z0-9_]", "_", ws.title)[:25] + "Tbl", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        try:
            ws.add_table(table)
        except Exception:
            pass


summary_rows = []
for voucher in VOUCHERS:
    h = headers[voucher]
    c = copied_summary[voucher]
    summary_rows.append({
        "VoucherNo": voucher,
        "NCC": h.get("AdvanceUserID"),
        "SupplierName": h.get("SupplierName"),
        "Currency": h.get("CurrencyID"),
        "DNTT_Lines": len(lines[voucher]),
        "AttachedFiles_DB": c["files"],
        "PhysicalFilesCopied": c["files"] - c["missing"],
        "TotalPagesOrSheets": c["pages"],
        "AI_Status": h.get("AIStatus"),
        "AI_Percentage": h.get("AIPercentage"),
        "AI_RunDate": h.get("RunCreateDate"),
        "WhySelected": "NVL, nhiều hơn 30 file đính kèm, AI OK 100%, dùng tốt để phân tích mapping DNTT -> files.",
        "CopiedFolder": c["dest"],
    })

db_mapping_rows = []
for row in compact_section_rows:
    db_mapping_rows.append(row)

long_field_rows = []
for row in raw["section_fields"]:
    long_field_rows.append(row)

run_rows = []
for row in raw["runs"]:
    run_rows.append({k: (str(v)[:32000] if isinstance(v, str) else v) for k, v in row.items()})

wb = Workbook()
wb.remove(wb.active)
write_sheet(wb, "00_Tong_quan", summary_rows)
write_sheet(wb, "01_DNTT_Header_DB", list(headers.values()))
write_sheet(wb, "02_DNTT_Lines_DB", [row for voucher in VOUCHERS for row in lines[voucher]])
write_sheet(wb, "03_Files_DinhKem", physical_rows)
write_sheet(wb, "04_DB_Section_Mapping", db_mapping_rows)
write_sheet(wb, "05_DB_Field_Long", long_field_rows)
write_sheet(wb, "06_AI_Criteria_DB", [row for voucher in VOUCHERS for row in criteria[voucher]])
write_sheet(wb, "07_Manual_Target_Mapping", manual_rows)
write_sheet(wb, "08_Target_Schema_For_AI", target_schema_rows)
write_sheet(wb, "09_Raw_Run_Text", run_rows)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            cell.value = safe_value(cell.value)
    if ws.title in ("07_Manual_Target_Mapping", "04_DB_Section_Mapping", "05_DB_Field_Long"):
        ws.sheet_view.zoomScale = 80

wb.save(EXCEL_OUT)

audit = {
    "createdAt": datetime.now().isoformat(timespec="seconds"),
    "excel": str(EXCEL_OUT),
    "attachmentsFolder": str(ATTACH_OUT),
    "vouchers": VOUCHERS,
    "summary": summary_rows,
    "rowCounts": {
        "headers": len(headers),
        "lines": sum(len(lines[v]) for v in VOUCHERS),
        "criteria": sum(len(criteria[v]) for v in VOUCHERS),
        "sections": len(compact_section_rows),
        "fieldsLong": len(long_field_rows),
        "attachments": len(physical_rows),
        "manualMapping": len(manual_rows),
    },
}
(OUT / "audit_mapping_dataset_20261002.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps(audit, ensure_ascii=False, indent=2, default=str))
