from pathlib import Path
from collections import Counter
import json, re, sys, time
from openpyxl import load_workbook

book=Path(r'E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\DATA_BEM AI_MEIKO_30092026.xlsx')
sheet='Kết quả tháng 09 và xử lý'
print('START', book, flush=True)
t=time.time()
wb=load_workbook(book,data_only=True,read_only=True,keep_links=False)
print('LOADED', round(time.time()-t,2), wb.sheetnames, flush=True)
ws=wb[sheet]
print('DIM', ws.max_row, ws.max_column, flush=True)
header_row=None; headers={}; header_values=[]
for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
    if any(str(v or '').strip()=='Độ chính xác của AI' for v in row):
        header_row=ridx; header_values=list(row); headers={str(v).strip():i+1 for i,v in enumerate(row) if v is not None}; break
print('HEADER_ROW', header_row, flush=True)
for i,v in enumerate(header_values,1): print(i,repr(v), flush=True)

# fallbacks by observed workbook convention
col_voucher=headers.get('Số DNTT') or 8
col_issue=headers.get('Loại vấn đề') or 17
col_acc=headers.get('Độ chính xác của AI') or 19
col_ai=headers.get('Kết quả phiếu (OK / NG)') or headers.get('Kết quả AI') or 11
col_pct=headers.get('Kết quả phiếu (%)') or 12
col_type=headers.get('Loại DNTT') or 3
col_files=headers.get('Số file đính kèm') or 10
col_p=16; col_r=18; col_t=20
print('COLS', {'voucher':col_voucher,'issue':col_issue,'acc':col_acc,'ai':col_ai,'pct':col_pct,'type':col_type,'files':col_files}, flush=True)
pos_labels=('AI đọc đúng','AI trả lời đúng','AI trả lời đúng theo Rules','AI trả lời tốt')
neg_labels=('AI đọc sai','AI trả lời sai','AI đối chiếu sai')
records=[]
for ridx,row in enumerate(ws.iter_rows(min_row=header_row+1,values_only=True),header_row+1):
    def val(c): return row[c-1] if c-1 < len(row) else None
    voucher=val(col_voucher)
    if not voucher: continue
    issue=str(val(col_issue) or '')
    issue_low=issue.lower()
    acc=val(col_acc)
    has_pos=any(x.lower() in issue_low for x in pos_labels)
    has_neg=any(x.lower() in issue_low for x in neg_labels)
    zero=(acc==0 or acc==0.0 or str(acc).strip() in ('0','0%','0.0%'))
    records.append({'row':ridx,'voucher':voucher,'type':val(col_type),'ai':val(col_ai),'ai_pct':val(col_pct),'files':val(col_files),'issue':issue,'acc':acc,'p':val(col_p),'r':val(col_r),'t':val(col_t),'pos':has_pos,'neg':has_neg,'zero':zero})
print('RECORDS',len(records), flush=True)
print('ACC_VALUES_TOP',Counter(x['acc'] for x in records).most_common(30), flush=True)
problem=[x for x in records if x['pos'] and x['zero']]
print('POSITIVE_WITH_ZERO',len(problem), flush=True)
print('POS_ONLY_ZERO',sum(x['pos'] and not x['neg'] and x['zero'] for x in records), flush=True)
print('MIXED_ZERO',sum(x['pos'] and x['neg'] and x['zero'] for x in records), flush=True)
print('BY_AI',Counter(str(x['ai']) for x in problem), flush=True)
print('BY_TYPE',Counter(str(x['type']) for x in problem), flush=True)
for x in problem[:80]: print('PROBLEM', json.dumps(x,ensure_ascii=False,default=str), flush=True)
conflict_counts={
 'positive_acc_lt_1': sum(x['pos'] and not x['neg'] and isinstance(x['acc'],(int,float)) and x['acc']<1 for x in records),
 'negative_acc_1': sum(x['neg'] and x['acc']==1 for x in records),
 'no_ai_but_accuracy': sum(str(x['ai']).upper() not in ('OK','NG') and isinstance(x['acc'],(int,float)) and x['acc']>0 for x in records),
 'has_ai_zero_no_negative': sum(str(x['ai']).upper() in ('OK','NG') and x['zero'] and not x['neg'] for x in records),
}
print('CONFLICT_COUNTS',conflict_counts, flush=True)
out=book.parent/'Accuracy_Recheck_T09_20261002'
out.mkdir(exist_ok=True)
(out/'audit_current_workbook_fast.json').write_text(json.dumps({'headers':headers,'records':records,'conflict_counts':conflict_counts},ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print('OUT', out/'audit_current_workbook_fast.json', flush=True)
