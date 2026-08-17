from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import re, unicodedata, collections, hashlib, difflib, json
from datetime import datetime
ROOT=Path(r'E:\Asoft\AI_BEM\Du_lieu_AI_BEM')
REPORT=ROOT/'Bao_cao_Knowledge_Base_BEM_AI_MEIKO_v11_Rut_Gon_Times_New_Roman.xlsx'
KNOW=ROOT/'Kho_Thuong_Thuc_BEM_AI_MEIKO.xlsx'
TRAIN=max(ROOT.glob('BEM AI_Training_MEIKO_v*.xlsx'),key=lambda p:p.stat().st_mtime)
RULE=next(ROOT.glob('*25R025*Meiko*.xlsx'))
FONT='Times New Roman';BLUE='1F4E78';LIGHT='D9EAF7';YELLOW='FFF2CC';GREEN='E2F0D9';RED='FCE4D6';WHITE='FFFFFF';thin=Side(style='thin',color='D9E2F3')
def s(v):
    if v is None:return ''
    if isinstance(v,float) and abs(v-round(v))<1e-9:return str(int(round(v)))
    return str(v).replace('\n',' ').strip()
def nrm(t):
    t=s(t).lower();t=''.join(c for c in unicodedata.normalize('NFD',t) if unicodedata.category(c)!='Mn');return t.replace('đ','d')
def fnum(v):
    try:return float(v)
    except:return None
def pct(v):
    try:return f'{float(v)*100:.2f}%'
    except:return s(v)
def title(ws,text,cols):
    ws.sheet_view.showGridLines=False;ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=cols)
    c=ws.cell(1,1,text);c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name=FONT,size=16,bold=True,color=WHITE);c.alignment=Alignment(horizontal='center',vertical='center');ws.row_dimensions[1].height=28
def section(ws,row,text,cols):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=cols)
    c=ws.cell(row,1,text);c.fill=PatternFill('solid',fgColor=LIGHT);c.font=Font(name=FONT,size=12,bold=True,color=BLUE);c.alignment=Alignment(vertical='center');ws.row_dimensions[row].height=22
def write_table(ws,row,col,heads,rows,name=None,widths=None):
    for j,h in enumerate(heads,col):
        c=ws.cell(row,j,h);c.fill=PatternFill('solid',fgColor=BLUE);c.font=Font(name=FONT,bold=True,color=WHITE);c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    for i,data in enumerate(rows,row+1):
        for j,h in enumerate(heads,col):
            v=data.get(h,'') if isinstance(data,dict) else (data[j-col] if j-col<len(data) else '')
            c=ws.cell(i,j,v);c.font=Font(name=FONT,size=11);c.alignment=Alignment(vertical='top',wrap_text=True);c.border=Border(bottom=thin)
    if name and rows:
        ref=f'{get_column_letter(col)}{row}:{get_column_letter(col+len(heads)-1)}{row+len(rows)}'
        t=Table(displayName=name,ref=ref);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
    if widths:
        for idx,w in widths.items():ws.column_dimensions[get_column_letter(col+idx)].width=w
    return row+len(rows)+2
def note(ws,row,text,fill,cols):
    ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=cols);c=ws.cell(row,1,text);c.fill=PatternFill('solid',fgColor=fill);c.font=Font(name=FONT,italic=True,color='7F6000');c.alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=38
