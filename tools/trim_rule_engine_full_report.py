from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
REPORT = ROOT / "Bao_cao_Rule_Engine_BEM_AI_06102026.docx"
SOURCE = ROOT / "Archive" / "Bao_cao_Rule_Engine_BEM_AI_06102026_backup_before_easy_20261006_131018.docx"
ARCHIVE = ROOT / "Archive"


REMOVE_RANGES = [
    ("3. Khai báo load và thực thi Rule", "4. Cấu trúc RuleResult và quy ước trạng thái"),
    ("4. Cấu trúc RuleResult và quy ước trạng thái", "5. POC 1: Loại tiền"),
    ("5. POC 1: Loại tiền", "7. Minh chứng kiểm thử POC"),
    ("8. Ranh giới Expression Python handler và LLM", "9. Cấu trúc triển khai 9 Rule không cần nghiên cứu lại"),
    ("12. Ví dụ POC chi tiết theo một phiếu thực tế giả lập", "13. Quy ước gọi LLM khi Rule Engine chưa đủ cơ sở"),
    ("15.2. Ví dụ RuleResult được lưu", "16. Cách Rule Engine được cấu hình và chạy"),
]


RENUMBER = {
    "7. Minh chứng kiểm thử POC": "3. Minh chứng kiểm thử POC",
    "9. Cấu trúc triển khai 9 Rule không cần nghiên cứu lại": "4. Cấu trúc triển khai 9 Rule không cần nghiên cứu lại",
    "10. Nội dung đề nghị chốt": "5. Nội dung đề nghị chốt",
    "11. Sơ đồ kiến trúc và cách hoạt động": "6. Sơ đồ kiến trúc và cách hoạt động",
    "13. Quy ước gọi LLM khi Rule Engine chưa đủ cơ sở": "7. Quy ước gọi LLM khi Rule Engine chưa đủ cơ sở",
    "14. Chín prompt NVL thay đổi như thế nào": "8. Chín prompt NVL thay đổi như thế nào",
    "15. Ví dụ minh họa các quy ước": "9. Ví dụ minh họa các quy ước",
    "15.1. Ví dụ quy ước OK NG REVIEW N A": "9.1. Ví dụ quy ước OK NG REVIEW N A",
    "16. Cách Rule Engine được cấu hình và chạy": "10. Cách Rule Engine được cấu hình và chạy",
    "16.1. Cấu hình Rule Engine nằm ở đâu": "10.1. Cấu hình Rule Engine nằm ở đâu",
    "16.2. Một lần chạy đi qua hàm nào": "10.2. Một lần chạy đi qua hàm nào",
    "16.3. Dữ liệu POC đầy đủ từ đầu vào đến đầu ra": "10.3. Dữ liệu POC đầy đủ từ đầu vào đến đầu ra",
    "16.4. Khi nào Rule Engine không được chạy tiếp": "10.4. Khi nào Rule Engine không được chạy tiếp",
    "16.5. Khi triển khai thật sẽ cấu hình như thế nào": "10.5. Khi triển khai thật sẽ cấu hình như thế nào",
    "17. Ví dụ dữ liệu tháng 09 đi qua Rule Engine": "11. Ví dụ dữ liệu tháng 09 đi qua Rule Engine",
    "17.1. Dữ liệu đầu vào của phiếu": "11.1. Dữ liệu đầu vào của phiếu",
    "17.2. Mapping đúng và Mapping sai khác nhau ở đâu": "11.2. Mapping đúng và Mapping sai khác nhau ở đâu",
    "17.3. Context mà rule-engine thực sự nhận": "11.3. Context mà rule-engine thực sự nhận",
    "17.4. Khi bổ sung một Rule mới sẽ làm như thế nào": "11.4. Khi bổ sung một Rule mới sẽ làm như thế nào",
}


def paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def top_level_paragraphs(doc):
    result = []
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            result.append((element, paragraph_text(element)))
    return result


def remove_range(doc, start_text: str, end_text: str):
    body = doc.element.body
    elements = list(body)
    start_index = next(
        (i for i, element in enumerate(elements) if paragraph_text(element) == start_text),
        None,
    )
    end_index = next(
        (i for i, element in enumerate(elements) if paragraph_text(element) == end_text),
        None,
    )
    if start_index is None or end_index is None or start_index >= end_index:
        raise ValueError(f"Không tìm thấy khoảng cần xóa: {start_text} -> {end_text}")
    for element in elements[start_index:end_index]:
        if element.tag != qn("w:sectPr"):
            body.remove(element)


def rename_paragraphs(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        replacement = RENUMBER.get(text)
        if replacement is None:
            continue
        if paragraph.runs:
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(replacement)


def set_times_new_roman(doc):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            if run._element.rPr is None:
                run._element.get_or_add_rPr()
            for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                run._element.rPr.rFonts.set(qn(key), "Times New Roman")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        if run._element.rPr is None:
                            run._element.get_or_add_rPr()
                        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                            run._element.rPr.rFonts.set(qn(key), "Times New Roman")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    if REPORT.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(REPORT, ARCHIVE / f"Bao_cao_Rule_Engine_BEM_AI_06102026_backup_before_trim_{stamp}.docx")

    doc = Document(SOURCE)
    for start_text, end_text in REMOVE_RANGES:
        remove_range(doc, start_text, end_text)
    rename_paragraphs(doc)
    set_times_new_roman(doc)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    main()
