from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_rule_engine_report as base

REPORT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_POC_Rule_Engine_BEM_AI_06102026.docx")
EVIDENCE_DIR = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Minh_chung_POC_Rule_Engine_06102026\So_do_Rule_Engine")
APP_CODE = Path(r"E:\Asoft\AI_BEM\BEM_AI_PROJECT\App\rule_engine_poc.py")
MD = Path(r"E:\Asoft\AI_BEM\BEM_AI_PROJECT\docs\POC_rule_engine_NVL_20261006.md")
FONT = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"


def img_font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT, size)


def text_center(draw, xy, text, font, fill="#000000"):
    x1, y1, x2, y2 = xy
    words = text.split()
    lines, current = [], ""
    max_width = x2 - x1 - 24
    for word in words:
        candidate = (current + " " + word).strip()
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + max(0, len(lines)-1) * 5
    y = y1 + (y2 - y1 - total) / 2
    for line, h in zip(lines, heights):
        w = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x1 + (x2-x1-w)/2, y), line, font=font, fill=fill)
        y += h + 5


def node(draw, xy, fill, title, detail, title_fill):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline="#1F4E78", width=3)
    draw.rounded_rectangle((x1, y1, x2, y1+58), radius=18, fill=title_fill)
    draw.rectangle((x1, y1+38, x2, y1+58), fill=title_fill)
    text_center(draw, (x1+8, y1+5, x2-8, y1+53), title, img_font(25, True), "#FFFFFF")
    text_center(draw, (x1+15, y1+72, x2-15, y2-15), detail, img_font(21), "#000000")


def arrow(draw, start, end, color="#1F4E78"):
    draw.line([start, end], fill=color, width=6)
    x1, y1 = start; x2, y2 = end
    if x2 >= x1:
        pts = [(x2, y2), (x2-20, y2-13), (x2-20, y2+13)]
    else:
        pts = [(x2, y2), (x2+20, y2-13), (x2+20, y2+13)]
    draw.polygon(pts, fill=color)


def make_practical_diagram():
    im = Image.new("RGB", (1800, 980), "white")
    d = ImageDraw.Draw(im)
    d.text((50, 30), "Rule Engine POC: cấu hình → dữ liệu → chạy → kết quả", font=img_font(39, True), fill="#1F4E78")
    d.line((50, 92, 1750, 92), fill="#1F4E78", width=3)
    nodes = [
        ((60, 185, 380, 620), "1. CẤU HÌNH", "RULE_VERSION = nvl-poc-2026-10-06\n\nexpression =\nhas_required_evidence\nand has_required_values\nand not evidence_conflict\nand all_amounts_match", "#EAF3F8", "#1F4E78"),
        ((465, 185, 785, 620), "2. DỮ LIỆU", "DNTT = 1.500 USD\nInvoice 01 = 900\nInvoice 02 = 600\nRingi = 1.500\nTờ khai = 1.500\n\nĐầu vào đã có Mapping", "#D9EAF7", "#1F4E78"),
        ((870, 185, 1190, 620), "3. HANDLER", "Chuẩn hóa số tiền\nCộng Invoice:\n900 + 600 = 1.500\n\nconflict = False\nrequired = True\n\nTạo context", "#E2F0D9", "#548235"),
        ((1275, 185, 1595, 620), "4. ENGINE", "rule_engine.Rule(\n  expression\n).matches(context)\n\nKết quả expression:\nTrue", "#E2F0D9", "#548235"),
        ((1660, 185, 1780, 620), "5. RESULT", "SOTIEN\nOK\n\nLý do\nEvidence\nVersion", "#E2F0D9", "#548235"),
    ]
    for xy, title, detail, fill, head in nodes:
        node(d, xy, fill, title, detail, head)
    for i in range(len(nodes)-1):
        arrow(d, (nodes[i][0][2]+10, 400), (nodes[i+1][0][0]-10, 400))
    d.rounded_rectangle((190, 715, 1610, 875), radius=18, fill="#FFF2CC", outline="#BF9000", width=3)
    text_center(d, (215, 735, 1585, 855), "Điểm quan trọng: rule-engine không tự đọc PDF, không tự cộng Invoice và không tự biết nghiệp vụ. Python Handler chuẩn bị context; rule-engine chỉ đánh giá expression; RuleResult mới là kết quả chuẩn để API-AI lưu.", img_font(25, True), "#7F6000")
    path = EVIDENCE_DIR / "05_Cau_hinh_va_chay_Rule_Engine.png"
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    im.save(path, dpi=(180, 180))
    return path


def set_run_tnr(run):
    run.font.name = base.FONT
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), base.FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT)


def add_figure(doc, path, caption):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(str(path), width=Cm(16.1))
    c=doc.add_paragraph(); c.alignment=WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after=Pt(5)
    r=c.add_run(caption); set_run_tnr(r); r.font.size=Pt(9); r.italic=True


def append_guide():
    doc = Document(REPORT)
    if any("Cách Rule Engine được cấu hình" in p.text for p in doc.paragraphs):
        print("guide already exists")
        return
    path = make_practical_diagram()
    base.heading(doc, "16", "Cách Rule Engine được cấu hình và chạy")
    base.callout(doc, "Hiểu ngắn gọn", "Rule Engine không phải một AI đọc file. Nó là bộ máy nhận các biến đã được Python chuẩn hóa rồi kiểm tra một expression. Nếu expression đúng thì Handler tạo RuleResult OK; nếu sai thì tạo NG; nếu thiếu hoặc mâu thuẫn evidence thì dừng ở REVIEW trước khi chạy expression.", base.GREEN)
    add_figure(doc, path, "Hình 5. Một lần chạy Rule Engine POC từ cấu hình đến RuleResult.")

    base.heading(doc, "16.1", "Cấu hình Rule Engine nằm ở đâu")
    base.rich_para(doc, [("Trong POC hiện tại, Rule chưa nằm trong Excel hoặc DB. Rule được khai báo trực tiếp trong file Python ", 10.5, False, base.BLACK), ("App/rule_engine_poc.py", 10.5, True, base.NAVY), (" để kiểm chứng cách chạy.", 10.5, False, base.BLACK)])
    base.code(doc, "Cấu hình thật trong POC", 'RULE_VERSION = "nvl-poc-2026-10-06"\n\nAMOUNT_EXPRESSION_RULE = _ExpressionRule(\n    criterion="SOTIEN",\n    expression="has_required_evidence and has_required_values and "\n               "not evidence_conflict and all_amounts_match",\n)')
    base.table(doc, ["Dòng cấu hình", "Ý nghĩa với dữ liệu thực tế"], [
        ("RULE_VERSION", "Đánh dấu phiên bản Rule đã dùng. Ví dụ: kết quả ngày hôm nay chạy bằng nvl-poc-2026-10-06."),
        ("criterion = SOTIEN", "Rule này chỉ xử lý tiêu chí Số tiền, không xử lý Loại tiền hoặc NCC."),
        ("has_required_evidence", "Đã có đủ chứng từ bắt buộc theo Mapping, ví dụ có Invoice, tờ khai và Ringi."),
        ("has_required_values", "Các chứng từ đã có giá trị số tiền đọc được, không rỗng và parse được."),
        ("evidence_conflict", "Có hay không trường hợp các chứng từ mâu thuẫn, ví dụ Ringi 1.600 nhưng Invoice 1.500."),
        ("all_amounts_match", "Sau khi Handler cộng và chuẩn hóa, các giá trị có khớp với DNTT hay không."),
    ], [5.1, 11.3], 8.5)

    base.heading(doc, "16.2", "Một lần chạy đi qua hàm nào")
    base.code(doc, "Đường chạy thật trong POC", 'input_context\n    ↓\nevaluate_nvl_amount_rule(input_context)\n    ↓\nPython Handler: _coerce_evidence → tổng hợp Invoice → phát hiện thiếu/mâu thuẫn\n    ↓\nAMOUNT_EXPRESSION_RULE.matches(expression_context)\n    ↓\nrule_engine.Rule(expression).matches(dict(context))\n    ↓\n_ok(...) hoặc _ng(...) hoặc _review(...)\n    ↓\nRuleResult.to_dict()')
    base.table(doc, ["Bước", "Dữ liệu cụ thể", "Hệ thống làm gì"], [
        ("1. input_context", "DNTT amount = 1.500; 2 Invoice; tờ khai; Ringi; bảng kê.", "Nhận dữ liệu đã OCR/trích xuất và đã Mapping."),
        ("2. Handler", "Invoice 01 = 900; Invoice 02 = 600.", "Cộng thành 1.500; Statement chỉ kiểm soát, không cộng trùng."),
        ("3. Context", "has_required_evidence=True; has_required_values=True; evidence_conflict=False; all_amounts_match=True.", "Biến đổi dữ liệu nghiệp vụ thành các biến Boolean rõ ràng."),
        ("4. rule-engine", "Expression nhận 4 biến Boolean.", "Trả True vì cả 4 điều kiện đều đúng."),
        ("5. RuleResult", "criterion=SOTIEN; status=OK; reason; evidence; rule_version.", "Đóng gói kết quả để API-AI lưu và ERP9 hiển thị."),
    ], [2.6, 7.3, 6.5], 8.15)

    base.heading(doc, "16.3", "Dữ liệu POC đầy đủ từ đầu vào đến đầu ra")
    base.code(doc, "Đầu vào của hàm evaluate_nvl_amount_rule", '{\n  "dntt": {"amount": "1500.00"},\n  "evidence": [\n    {"doc_type": "CUSTOMSHEET", "file_name": "To_khai_01.pdf", "amount": "1500.00"},\n    {"doc_type": "RINGI", "file_name": "Ringi_09.pdf", "amount": "1500.00"},\n    {"doc_type": "INVOICE", "file_name": "Invoice_01.pdf", "amount": "900.00"},\n    {"doc_type": "INVOICE", "file_name": "Invoice_02.pdf", "amount": "600.00"},\n    {"doc_type": "STATEMENT", "file_name": "Bang_ke.xlsx", "amount": "1500.00"}\n  ]\n}')
    base.code(doc, "Context sau Python Handler", '{\n  "has_required_evidence": true,\n  "has_required_values": true,\n  "evidence_conflict": false,\n  "all_amounts_match": true\n}')
    base.code(doc, "Đầu ra RuleResult", '{\n  "criterion": "SOTIEN",\n  "status": "OK",\n  "reason": "Số tiền DNTT, tờ khai, Ringi và tổng Invoice khớp nhau.",\n  "evidence": {\n    "dntt_amount": "1500.00",\n    "invoice_total": "1500.00",\n    "invoice_files": ["Invoice_01.pdf", "Invoice_02.pdf"],\n    "ringi_amount": "1500.00"\n  },\n  "rule_version": "nvl-poc-2026-10-06"\n}')

    base.heading(doc, "16.4", "Khi nào Rule Engine không được chạy tiếp")
    base.table(doc, ["Dữ liệu thực tế", "Handler xử lý", "Kết quả"], [
        ("Thiếu Invoice bắt buộc.", "has_required_evidence=False.", "REVIEW; dừng trước expression, không gọi LLM đoán."),
        ("Invoice = 1.500; Ringi = 1.600.", "evidence_conflict=True.", "REVIEW; dừng trước expression vì hồ sơ mâu thuẫn."),
        ("Invoice + tờ khai + Ringi đều = 1.400; DNTT = 1.500.", "Đủ evidence, conflict=False, all_amounts_match=False.", "Expression chạy và trả False → NG."),
        ("Phiếu không áp dụng tiêu chí.", "applicable=False.", "N/A; không cần chạy expression."),
    ], [6.0, 6.2, 4.2], 8.25)
    base.callout(doc, "Điểm cần nhớ", "Rule Engine chỉ biết các biến mà Handler đưa vào. Nếu muốn Rule Engine biết “nhiều Invoice cộng lại”, “Statement không được cộng trùng” hoặc “Ringi đang mâu thuẫn”, các logic đó phải được viết ở Python Handler trước.", base.YELLOW)

    doc.add_page_break()
    base.heading(doc, "16.5", "Khi triển khai thật sẽ cấu hình như thế nào")
    base.rich_para(doc, [("POC đang khai báo Rule trong Python để kiểm chứng. Khi triển khai 9 tiêu chí, nên chuyển phần khai báo sang một Rule Catalog có version, nhưng cách chạy vẫn giữ nguyên:", 10.5, False, base.BLACK)])
    base.table(doc, ["Rule Catalog cần lưu", "Ví dụ SOTIEN"], [
        ("rule_id", "SOTIEN"),
        ("rule_version", "nvl-amount-v1"),
        ("required_evidence", "CUSTOMSHEET, RINGI, INVOICE/COMMERCIALINVOICE"),
        ("handler", "normalize_amount_and_detect_conflict"),
        ("expression + status_policy", "has_required_evidence and has_required_values and not evidence_conflict and all_amounts_match\nKết quả: OK / NG / REVIEW / N/A"),
    ], [5.5, 10.9], 8.45)

    for p in doc.paragraphs:
        for run in p.runs:
            set_run_tnr(run)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_tnr(run)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    append_guide()
