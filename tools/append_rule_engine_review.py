from __future__ import annotations

from pathlib import Path
from textwrap import wrap
import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml.ns import qn

import generate_rule_engine_report as base

ROOT = Path(r"E:\Asoft\AI_BEM")
REPORT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_POC_Rule_Engine_BEM_AI_06102026.docx")
EVIDENCE_DIR = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Minh_chung_POC_Rule_Engine_06102026")
IMG_DIR = EVIDENCE_DIR / "So_do_Rule_Engine"
FONT_REG = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
GRAY = "F2F2F2"
DARK = "17365D"
WHITE = "FFFFFF"
BLACK = "000000"


def fnt(size, bold=False):
    path = FONT_BOLD if bold else FONT_REG
    return ImageFont.truetype(path, size)


def text_width(draw, text, font):
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def centered(draw, box, text, font, fill=BLACK, line_gap=5):
    fill = fill if str(fill).startswith("#") else f"#{fill}"
    x1, y1, x2, y2 = box
    lines = []
    max_width = x2 - x1 - 30
    words = text.split()
    current = ""
    for word in words:
        candidate = (current + " " + word).strip()
        if text_width(draw, candidate, font) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + line_gap * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        x = x1 + (x2 - x1 - text_width(draw, line, font)) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += height + line_gap


def box(draw, xy, fill, title, detail="", title_fill=DARK):
    draw.rounded_rectangle(xy, radius=18, fill=f"#{fill}", outline=f"#{DARK}", width=3)
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle((x1, y1, x2, y1 + 55), radius=18, fill=f"#{title_fill}")
    draw.rectangle((x1, y1 + 38, x2, y1 + 55), fill=f"#{title_fill}")
    centered(draw, (x1 + 10, y1 + 5, x2 - 10, y1 + 52), title, fnt(27, True), WHITE)
    if detail:
        centered(draw, (x1 + 18, y1 + 68, x2 - 18, y2 - 18), detail, fnt(22), BLACK)


def arrow(draw, start, end, color=DARK, width=7):
    draw.line([start, end], fill=f"#{color}", width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        if x2 >= x1:
            pts = [(x2, y2), (x2 - 22, y2 - 14), (x2 - 22, y2 + 14)]
        else:
            pts = [(x2, y2), (x2 + 22, y2 - 14), (x2 + 22, y2 + 14)]
    else:
        if y2 >= y1:
            pts = [(x2, y2), (x2 - 14, y2 - 22), (x2 + 14, y2 - 22)]
        else:
            pts = [(x2, y2), (x2 - 14, y2 + 22), (x2 + 14, y2 + 22)]
    draw.polygon(pts, fill=f"#{color}")


def canvas(title, width=1800, height=980):
    im = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(im)
    d.text((50, 28), title, font=fnt(40, True), fill=f"#{NAVY}")
    d.line((50, 92, width - 50, 92), fill=f"#{NAVY}", width=3)
    return im, d


def make_architecture():
    im, d = canvas("Rule Engine nằm trong bước AI đối chiếu")
    # group bands
    d.rounded_rectangle((45, 125, 410, 885), radius=20, fill="#F5F9FC", outline="#9CC2E5", width=3)
    d.rounded_rectangle((440, 125, 1390, 885), radius=20, fill="#EDF6FB", outline="#5B9BD5", width=3)
    d.rounded_rectangle((1420, 125, 1755, 885), radius=20, fill="#F6FBF2", outline="#70AD47", width=3)
    centered(d, (70, 145, 385, 195), "DỮ LIỆU ĐẦU VÀO", fnt(25, True), NAVY)
    centered(d, (470, 145, 1360, 195), "AI PYTHON - BƯỚC ĐỐI CHIẾU", fnt(25, True), NAVY)
    centered(d, (1450, 145, 1725, 195), "LƯU KẾT QUẢ", fnt(25, True), "548235")
    box(d, (80, 235, 375, 385), PALE, "DNTT", "Dòng hàng, số tiền, loại tiền, NCC, điều khoản")
    box(d, (80, 470, 375, 690), BLUE, "Evidence", "Các file đã OCR và trích xuất; có FileVersion")
    box(d, (80, 745, 375, 835), GRAY, "API-AI", "Truyền job và dữ liệu đã chuẩn hóa")
    box(d, (500, 235, 790, 390), BLUE, "1. Mapping", "Gắn dòng DNTT với đúng Invoice, PO, Ringi, tờ khai, bảng kê")
    box(d, (850, 235, 1135, 390), BLUE, "2. Handler", "Chuẩn hóa, cộng nhiều Invoice, phát hiện thiếu và mâu thuẫn")
    box(d, (1195, 235, 1350, 390), "D9EAF7", "3. Context", "Đưa dữ liệu về các biến rõ ràng")
    arrow(d, (375, 310), (500, 310)); arrow(d, (375, 580), (500, 340))
    arrow(d, (790, 310), (850, 310)); arrow(d, (1135, 310), (1195, 310))
    box(d, (620, 510, 1080, 685), GREEN, "4. rule-engine zeroSteiner", "Chạy expression deterministic trên context đã sạch", title_fill="548235")
    box(d, (1145, 510, 1350, 685), GREEN, "5. RuleResult", "OK / NG / REVIEW / N/A", title_fill="548235")
    arrow(d, (1270, 390), (850, 510)); arrow(d, (1080, 598), (1145, 598))
    box(d, (500, 745, 930, 835), ORANGE, "LLM ngoại lệ", "Chỉ gọi khi RuleResult REVIEW thuộc nhóm được phép; không thay logic deterministic", title_fill="C55A11")
    box(d, (980, 745, 1350, 835), YELLOW, "Re-run Rule", "LLM chỉ bổ sung candidate evidence; Handler và Rule Engine chạy lại", title_fill="BF9000")
    arrow(d, (1270, 685), (715, 745), "C55A11"); arrow(d, (930, 790), (980, 790), "BF9000"); arrow(d, (1165, 745), (1250, 685), "BF9000")
    box(d, (1470, 255, 1705, 400), GREEN, "DB", "Lưu RuleResult, evidence refs, RuleVersion, lỗi và thời gian", title_fill="548235")
    box(d, (1470, 530, 1705, 675), BLUE, "ERP9", "Đọc kết quả cuối để hiển thị cho người dùng")
    arrow(d, (1350, 598), (1470, 330), "548235"); arrow(d, (1585, 400), (1585, 530), "548235")
    d.text((55, 915), "Đường chính: Mapping → Handler → rule-engine → RuleResult. LLM là nhánh ngoại lệ, không phải bước mặc định.", font=fnt(23, True), fill=f"#{NAVY}")
    path = IMG_DIR / "01_Kien_truc_Rule_Engine.png"; im.save(path, dpi=(180, 180)); return path



def make_architecture_clean():
    im, d = canvas("Vị trí Rule Engine trong luồng AI đối chiếu", width=1800, height=1050)
    # Two clean rows: deterministic main path and isolated exception path.
    d.rounded_rectangle((45, 125, 1755, 620), radius=22, fill="#F5F9FC", outline="#9CC2E5", width=3)
    d.rounded_rectangle((45, 665, 1755, 965), radius=22, fill="#FFF9F0", outline="#E6B566", width=3)
    d.text((75, 145), "LUỒNG CHÍNH - KHÔNG GỌI LLM", font=fnt(28, True), fill="#1F4E78")
    d.text((75, 685), "LUỒNG NGOẠI LỆ - CHỈ GỌI KHI RULE RESULT = REVIEW VÀ ĐỦ ĐIỀU KIỆN", font=fnt(25, True), fill="#C55A11")

    main = [
        ((80, 260, 300, 500), "1. DỮ LIỆU", "DNTT\n+ Evidence đã OCR/trích xuất", "D9EAF7", "1F4E78"),
        ((350, 260, 590, 500), "2. MAPPING", "Dòng DNTT\n↔ Invoice / PO / Ringi\n↔ Tờ khai / Bảng kê", "D9EAF7", "1F4E78"),
        ((640, 260, 880, 500), "3. HANDLER", "Chuẩn hóa\nCộng nhiều Invoice\nPhát hiện thiếu/mâu thuẫn", "D9EAF7", "1F4E78"),
        ((930, 260, 1170, 500), "4. RULE ENGINE", "Expression\ndeterministic\ntrên context sạch", "E2F0D9", "548235"),
        ((1220, 260, 1450, 500), "5. RULE RESULT", "OK / NG\nREVIEW / N/A\nLý do + evidence + version", "E2F0D9", "548235"),
        ((1500, 260, 1720, 500), "6. API-AI / DB", "Lưu kết quả\ntrạng thái\nRuleVersion", "E2F0D9", "548235"),
    ]
    for xy, title, detail, fill, head in main:
        box(d, xy, fill, title, detail, title_fill=head)
    for i in range(len(main)-1):
        arrow(d, (main[i][0][2] + 12, 380), (main[i+1][0][0] - 12, 380), "1F4E78", 6)
    d.text((1510, 515), "ERP9 đọc kết quả", font=fnt(21, True), fill="#548235")

    exc = [
        ((320, 760, 590, 900), "CHẠY LẠI", "Handler → Rule Engine\nKhông tự ghi OK/NG", "E2F0D9", "548235"),
        ((620, 760, 890, 900), "KIỂM TRA", "Có evidence_refs\nconfidence đạt ngưỡng?", "FFF2CC", "BF9000"),
        ((920, 760, 1190, 900), "LLM NGOẠI LỆ", "Đọc phần cần làm rõ\ntrả candidate evidence", "FCE4D6", "C55A11"),
        ((1220, 760, 1450, 900), "REVIEW", "Thiếu rõ ràng\nhoặc exception được phép", "FFF2CC", "BF9000"),
    ]
    for xy, title, detail, fill, head in exc:
        box(d, xy, fill, title, detail, title_fill=head)
    # Main result points straight down to REVIEW; exception path then runs left.
    arrow(d, (1335, 500), (1335, 760), "C55A11", 5)
    arrow(d, (1220 - 12, 830), (1190 + 12, 830), "C55A11", 6)
    arrow(d, (920 - 12, 830), (890 + 12, 830), "C55A11", 6)
    arrow(d, (620 - 12, 830), (590 + 12, 830), "BF9000", 6)
    d.text((75, 1000), "Nguyên tắc: Rule Engine là nơi kết luận deterministic; LLM chỉ bổ sung candidate evidence rồi phải chạy lại Handler và Rule Engine.", font=fnt(23, True), fill="#1F4E78")
    path = IMG_DIR / "01_Kien_truc_Rule_Engine.png"
    im.save(path, dpi=(180, 180))
    return path

def make_poc_trace():
    im, d = canvas("POC Số tiền: một phiếu, nhiều chứng từ, một kết quả có thể truy vết")
    box(d, (60, 145, 410, 800), PALE, "1. DNTT", "Phiếu: BEMT09-0001\nDòng NVL: NVL-001\nSố tiền: 1.500,00 USD\nPO: PO-4500123\nThanh toán: 1 lần", title_fill=NAVY)
    box(d, (470, 145, 830, 800), BLUE, "2. Evidence Mapping", "Invoice_01.pdf → 900 USD\nInvoice_02.pdf → 600 USD\nTo_khai_01.pdf → 1.500 USD\nRingi_09.pdf → 1.500 USD\nBang_ke.xlsx → 1.500 USD\n\nStatement chỉ là evidence kiểm soát, không cộng vào Invoice.")
    box(d, (890, 145, 1240, 800), GREEN, "3. Python Handler", "invoice_total = 900 + 600\n= 1.500 USD\n\ncustomsheet = 1.500\nringi = 1.500\nstatement = 1.500\n\nconflict = False\nrequired = True", title_fill="548235")
    box(d, (1300, 145, 1550, 800), YELLOW, "4. Expression", "has_required_evidence\nand has_required_values\nand not evidence_conflict\nand all_amounts_match\n\nKết quả: True", title_fill="BF9000")
    box(d, (1610, 145, 1740, 800), GREEN, "5. Result", "SOTIEN\nOK\n\nLý do:\nTất cả số tiền khớp\n\nVersion:\nnvl-poc-2026-10-06", title_fill="548235")
    for a, b in [((410, 470), (470, 470)), ((830, 470), (890, 470)), ((1240, 470), (1300, 470)), ((1550, 470), (1610, 470))]: arrow(d, a, b)
    d.text((60, 845), "Nếu Ringi = 1.600 USD → conflict = True → RuleResult = REVIEW; nếu tất cả evidence = 1.400 USD → RuleResult = NG.", font=fnt(24, True), fill=f"#{NAVY}")
    path = IMG_DIR / "02_POC_So_tien_chi_tiet.png"; im.save(path, dpi=(180, 180)); return path


def make_llm_gate():
    im, d = canvas("Cổng quyết định: khi nào được gọi LLM và khi nào không được gọi")
    box(d, (70, 250, 360, 640), BLUE, "Input", "Evidence đã Mapping\nContext đã chuẩn hóa\nRuleVersion hiện hành")
    box(d, (480, 250, 820, 640), GREEN, "Rule Engine", "Chạy expression\n+ Handler nghiệp vụ\n\nKết quả RuleResult", title_fill="548235")
    arrow(d, (360, 445), (480, 445))
    box(d, (940, 155, 1300, 365), GREEN, "OK hoặc NG", "Evidence đủ, không conflict\n→ Ghi kết quả\n→ Không gọi LLM", title_fill="548235")
    box(d, (940, 515, 1300, 725), ORANGE, "REVIEW có thể xử lý", "Evidence có nhưng diễn đạt/mapping mơ hồ; thuộc allowlist exception\n→ Gọi prompt ngoại lệ", title_fill="C55A11")
    box(d, (1390, 155, 1730, 365), PALE, "Không gọi LLM", "Thiếu file\nOCR rỗng/không tin cậy\nRule đã kết luận OK/NG", title_fill=NAVY)
    box(d, (1390, 515, 1730, 725), YELLOW, "LLM output", "Chỉ trả candidate evidence, evidence_refs, confidence, reason\nKhông tự ghi OK/NG cuối", title_fill="BF9000")
    arrow(d, (820, 350), (940, 260), "548235"); arrow(d, (820, 540), (940, 620), "C55A11")
    arrow(d, (1300, 260), (1390, 260), "548235"); arrow(d, (1300, 620), (1390, 620), "BF9000")
    box(d, (940, 800, 1730, 900), GREEN, "Sau LLM", "Nếu candidate được chấp nhận → Handler + Rule Engine chạy lại. Nếu chưa đủ cơ sở → giữ REVIEW và yêu cầu người dùng xác nhận.", title_fill="548235")
    arrow(d, (1560, 725), (1560, 800), "BF9000")
    path = IMG_DIR / "03_Cong_quyet_dinh_LLM.png"; im.save(path, dpi=(180, 180)); return path


def make_prompt_change():
    im, d = canvas("Cách 9 prompt NVL thay đổi sau khi áp dụng Rule Engine", width=1800, height=1160)
    # headers
    d.rounded_rectangle((70, 145, 830, 1080), radius=18, fill="#FFF4F0", outline="#C55A11", width=3)
    d.rounded_rectangle((970, 145, 1730, 1080), radius=18, fill="#F2F8EC", outline="#548235", width=3)
    centered(d, (100, 170, 800, 235), "HIỆN TẠI - PROMPT TỰ ĐỌC VÀ KẾT LUẬN", fnt(28, True), "C55A11")
    centered(d, (1000, 170, 1700, 235), "SAU RULE ENGINE - PROMPT HỖ TRỢ, ENGINE KẾT LUẬN", fnt(28, True), "548235")
    current = [
        "Đọc DNTT và các file",
        "Tự chuẩn hóa dữ liệu",
        "Tự gom nhiều chứng từ",
        "Tự xử lý thiếu/mâu thuẫn",
        "Tự trả OK / NG / BLANK",
    ]
    future = [
        "Nhận evidence đã Mapping",
        "Đọc/chuẩn hóa ngoại lệ khi cần",
        "Handler gom và tính toán",
        "Rule Engine kết luận OK / NG",
        "Thiếu/mâu thuẫn → REVIEW",
        "LLM chỉ xử lý allowlist exception",
    ]
    y = 330
    for text in current:
        box(d, (145, y, 755, y + 105), "FCE4D6", text, "", title_fill="C55A11"); y += 135
    y = 300
    for text in future:
        fill = "E2F0D9" if "Rule Engine" in text or "Handler" in text else ("FFF2CC" if "LLM" in text or "REVIEW" in text else "D9EAF7")
        title_fill = "548235" if "Rule Engine" in text or "Handler" in text else ("BF9000" if "LLM" in text or "REVIEW" in text else NAVY)
        box(d, (1045, y, 1655, y + 88), fill, text, "", title_fill=title_fill); y += 115
    d.text((105, 1010), "Nhược điểm: LLM vừa đọc vừa quyết định; khó truy vết và dễ suy đoán.", font=fnt(22, True), fill="#C55A11")
    d.text((1005, 1010), "Ưu điểm: prompt cung cấp evidence; Handler + Rule Engine quyết định; LLM không phá Rules.", font=fnt(22, True), fill="#548235")
    path = IMG_DIR / "04_Thay_doi_9_prompt.png"; im.save(path, dpi=(180, 180)); return path


def set_doc_font(document):
    for p in document.paragraphs:
        for run in p.runs:
            run.font.name = base.FONT
            if run._element.rPr is not None and run._element.rPr.rFonts is not None:
                run._element.rPr.rFonts.set(qn("w:ascii"), base.FONT)
                run._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT)
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = base.FONT
                        run._element.rPr.rFonts.set(qn("w:ascii"), base.FONT)
                        run._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT)
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT)


def add_figure(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(16.1))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(caption)
    r.font.name = base.FONT; r.font.size = Pt(9); r.italic = True
    r._element.rPr.rFonts.set(qn("w:ascii"), base.FONT)
    r._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT)


def append_sections():
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    base.build()
    paths = [make_architecture_clean(), make_poc_trace(), make_llm_gate(), make_prompt_change()]
    doc = Document(REPORT)
    base.heading(doc, "11", "Sơ đồ kiến trúc và cách hoạt động")
    base.rich_para(doc, [("Sơ đồ dưới đây mô tả đúng vị trí của Rule Engine: ", 10.5, False, base.BLACK), ("chỉ nằm trong bước AI đối chiếu", 10.5, True, base.NAVY), (", sau OCR, trích xuất và Evidence Mapping.", 10.5, False, base.BLACK)])
    add_figure(doc, paths[0], "Hình 1. Kiến trúc đề xuất: Mapping → Handler → rule-engine → RuleResult; LLM là nhánh ngoại lệ.")
    base.table(doc, ["Thành phần", "Nhiệm vụ trong sơ đồ", "Không được làm"], [
        ("Evidence Mapping", "Xác định file và trường dữ liệu nào thuộc từng dòng DNTT.", "Không tự kết luận OK/NG."),
        ("Python Handler", "Chuẩn hóa, tổng hợp, tính toán, phát hiện thiếu/mâu thuẫn.", "Không thay Rule nghiệp vụ bằng suy đoán."),
        ("rule-engine", "Đánh giá expression deterministic trên context đã sạch.", "Không đọc file thô hoặc gọi OCR."),
        ("LLM", "Chỉ xử lý exception được cho phép, trả candidate evidence.", "Không tự ghi đè OK/NG của Rule Engine."),
        ("DB/API-AI", "Lưu RuleResult, evidence refs, version, lỗi và thời gian.", "Không biến LLM thành nguồn quyết định duy nhất."),
    ], [3.1, 8.1, 5.3], 8.5)

    base.heading(doc, "12", "Ví dụ POC chi tiết theo một phiếu thực tế giả lập")
    base.rich_para(doc, [("Ví dụ này mô phỏng một phiếu NVL có nhiều chứng từ. Mục tiêu là nhìn rõ dữ liệu đi qua từng bước, không chỉ nhìn kết quả cuối.", 10.5, False, base.BLACK)])
    add_figure(doc, paths[1], "Hình 2. Dòng dữ liệu POC SOTIEN từ DNTT, qua Mapping và Handler, đến expression và RuleResult.")
    base.table(doc, ["Bước", "Dữ liệu cụ thể", "Kết quả lưu/truyền"], [
        ("1. DNTT", "BEMT09-0001; dòng NVL-001; 1.500,00 USD; PO-4500123; thanh toán 1 lần.", "Dữ liệu gốc của dòng cần đối chiếu."),
        ("2. File chứng từ", "Invoice_01.pdf = 900; Invoice_02.pdf = 600; To_khai_01.pdf = 1.500; Ringi_09.pdf = 1.500; Bang_ke.xlsx = 1.500.", "Mỗi file có FileVersionID và các trường đã trích xuất."),
        ("3. Mapping", "Hai Invoice cùng map vào dòng NVL-001; tờ khai, Ringi và bảng kê cùng nhóm chứng từ.", "Evidence list có file_name, doc_type, field_name, value, source_version."),
        ("4. Handler", "Cộng 900 + 600 = 1.500; Statement chỉ kiểm soát; conflict = False; required = True.", "Context: has_required_evidence=True, all_amounts_match=True."),
        ("5. Expression", "has_required_evidence and has_required_values and not evidence_conflict and all_amounts_match.", "True."),
        ("6. RuleResult", "SOTIEN = OK; reason nêu các nguồn khớp; evidence ghi lại 5 file; version nvl-poc-2026-10-06.", "API-AI lưu kết quả và ERP9 đọc kết quả cuối."),
    ], [2.5, 9.2, 4.8], 8.1)
    base.callout(doc, "Ba biến thể phải thấy rõ", "Thiếu Invoice → REVIEW; Ringi = 1.600 → REVIEW do mâu thuẫn; tất cả chứng từ = 1.400 nhưng evidence đầy đủ → NG. Không dùng LLM để biến các case này thành OK.", base.YELLOW)

    base.heading(doc, "13", "Quy ước gọi LLM khi Rule Engine chưa đủ cơ sở")
    base.rich_para(doc, [("LLM không phải bước mặc định. Trước khi gọi LLM, Handler phải kiểm tra đủ evidence, chất lượng OCR, trạng thái Mapping và loại exception.", 10.5, False, base.BLACK)])
    add_figure(doc, paths[2], "Hình 3. Cổng quyết định LLM: chỉ gọi theo allowlist exception và luôn chạy lại Rule Engine sau khi bổ sung candidate evidence.")
    base.table(doc, ["Trường hợp", "Gọi LLM?", "Kết quả bắt buộc"], [
        ("Rule Engine đã trả OK hoặc NG với evidence đầy đủ, không mâu thuẫn.", "Không.", "Giữ nguyên kết quả deterministic."),
        ("Thiếu file bắt buộc, OCR rỗng hoặc trường bắt buộc không có giá trị.", "Không.", "REVIEW và yêu cầu bổ sung/xử lý lại file."),
        ("Evidence có đủ nhưng tên chứng từ, cách viết hoặc quan hệ file chưa map được.", "Có, nếu nằm trong allowlist.", "LLM trả candidate mapping/evidence_refs; không tự kết luận cuối."),
        ("Mẫu chứng từ mới hoặc cách diễn đạt ngoại lệ chưa có trong Rules.", "Có, để hỗ trợ phân tích.", "Nếu chưa có Rule được xác nhận thì giữ REVIEW; không tự mở Rule."),
        ("LLM trả confidence thấp, thiếu evidence_refs hoặc JSON sai schema.", "Không chấp nhận kết quả.", "Giữ REVIEW và ghi lỗi LLM."),
    ], [6.0, 2.2, 8.3], 8.35)
    base.code(doc, "Schema đầu ra của LLM ngoại lệ", '{\n  "criterion": "SOTIEN",\n  "candidate_resolution": "Hai Invoice thuộc cùng dòng DNTT",\n  "evidence_refs": ["Invoice_01.pdf", "Invoice_02.pdf"],\n  "confidence": 0.96,\n  "should_rerun_rule": true,\n  "final_status": "REVIEW"\n}')
    base.callout(doc, "Chốt an toàn", "LLM chỉ bổ sung thông tin để Handler chạy lại. Quyền kết luận OK/NG vẫn thuộc Rule Engine. Nếu không đủ cơ sở sau khi chạy lại, kết quả cuối là REVIEW.", base.GREEN)

    base.heading(doc, "14", "Chín prompt NVL thay đổi như thế nào")
    base.rich_para(doc, [("Hiện tại mỗi prompt Warehouse đang nhận DNTT + dataFiles và tự làm cả đọc, chuẩn hóa, đối chiếu, trả OK/NG/BLANK. Sau khi áp dụng Rule Engine, không xóa 9 prompt; chỉ tách lại vai trò.", 10.5, False, base.BLACK)])
    add_figure(doc, paths[3], "Hình 4. So sánh vai trò 9 prompt trước và sau khi Rule Engine được đưa vào bước AI đối chiếu.")
    base.table(doc, ["Tiêu chí", "Prompt sau khi áp dụng Rule Engine", "Handler + Rule Engine", "LLM"], [
        ("NCC", "Đọc tên, MST, alias và evidence nguồn.", "Chuẩn hóa alias, so sánh NCC.", "Chỉ ngoại lệ tên/mẫu mới."),
        ("Số hóa đơn", "Đọc danh sách số và cách viết tắt.", "Tách số, map theo dòng và kiểm tra tập hợp.", "Chỉ khi không xác định được quan hệ file."),
        ("Ngày hóa đơn", "Đọc ngày và trả giá trị chuẩn hóa nếu rõ.", "Kiểm tra quan hệ ngày.", "Chỉ khi mẫu ngày/ý nghĩa ngày mơ hồ."),
        ("Số tiền", "Đọc từng số tiền theo file và dòng.", "Cộng Invoice, phân biệt Statement, kiểm tra khớp.", "Chỉ khi quan hệ chứng từ chưa rõ."),
        ("Loại tiền", "Đọc mã tiền tệ trên DNTT/chứng từ.", "Chuẩn hóa và so sánh.", "Thường không cần nếu evidence đủ."),
        ("Điều kiện giao hàng", "Đọc Incoterm và địa điểm giao hàng.", "Chuẩn hóa và kiểm tra điều kiện.", "Khi câu chữ không chuẩn hoặc thiếu ngữ cảnh."),
        ("Hạn thanh toán", "Đọc payment term và các ngày mốc.", "Tính deadline, kiểm tra ngày và lịch nghỉ.", "Không được tự tính thay handler."),
        ("Ngày hoàn thành kiểm tra", "Đọc các ngày liên quan.", "Chọn mốc theo Rules và kiểm tra quan hệ.", "Chỉ khi chứng từ có nhiều ngày không rõ ý nghĩa."),
        ("Chữ ký/con dấu", "Nhận diện vùng/ký hiệu cần kiểm tra.", "Áp dụng danh sách chứng từ bắt buộc và kết luận có/không.", "Chỉ khi hình ảnh/mẫu mới cần hỗ trợ nhận diện."),
    ], [2.8, 6.3, 4.5, 3.0], 7.75)
    base.table(doc, ["Cách chạy", "9 prompt có được gọi không?", "Kết quả cuối"], [
        ("Case chuẩn, evidence đủ và Mapping rõ", "Không cần gọi prompt đối chiếu; dùng dữ liệu trích xuất + Mapping + Rule Engine.", "Nhanh hơn, ít suy đoán, dễ trace."),
        ("Case cần đọc/chuẩn hóa ngoại lệ", "Chỉ gọi prompt tương ứng với tiêu chí đó, không gửi lại toàn bộ hồ sơ nếu không cần.", "Prompt trả candidate evidence; chạy lại Rule Engine."),
        ("Case thiếu/mâu thuẫn chứng từ", "Không gọi LLM để đoán.", "REVIEW và yêu cầu bổ sung/xác nhận."),
    ], [4.0, 8.5, 4.1], 8.2)

    base.heading(doc, "15", "Ví dụ minh họa các quy ước")
    base.rich_para(doc, [("Các quy ước trong báo cáo được minh họa bằng dữ liệu cụ thể dưới đây. Mục tiêu là người đọc nhìn vào dữ liệu đầu vào có thể hiểu vì sao hệ thống trả kết quả tương ứng.", 10.5, False, base.BLACK)])

    base.heading(doc, "15.1", "Ví dụ quy ước OK NG REVIEW N A")
    base.table(doc, ["Trạng thái", "Dữ liệu đầu vào minh họa", "Cách kết luận", "Có gọi LLM?"], [
        ("OK", "DNTT = 1.500 USD; Invoice 01 = 900; Invoice 02 = 600; Ringi = 1.500; đủ file.", "900 + 600 = 1.500; các nguồn khớp; kết quả deterministic đạt.", "Không."),
        ("NG", "DNTT = 1.500 USD; tờ khai = 1.400; Ringi = 1.400; Invoice = 1.400; đủ file.", "Evidence đầy đủ nhưng không khớp DNTT; vi phạm Rule.", "Không."),
        ("REVIEW", "DNTT = 1.500 USD; có Ringi và tờ khai nhưng thiếu Invoice bắt buộc.", "Thiếu evidence để kết luận; không được suy đoán.", "Không; yêu cầu bổ sung file."),
        ("REVIEW", "Invoice = 1.500 USD nhưng Ringi = 1.600 USD.", "Các chứng từ mâu thuẫn; cần kiểm tra nghiệp vụ/hồ sơ.", "Không; giữ REVIEW."),
        ("N/A", "Phiếu thuộc nghiệp vụ đã xác nhận không áp dụng tiêu chí Loại tiền.", "Tiêu chí được đánh dấu không áp dụng.", "Không."),
    ], [2.1, 7.2, 5.4, 2.0], 7.9)
    base.callout(doc, "Cách nhớ nhanh", "Đủ và khớp = OK; đủ nhưng sai = NG; thiếu hoặc mâu thuẫn = REVIEW; không thuộc phạm vi = N/A.", base.GREEN)

    base.heading(doc, "15.2", "Ví dụ RuleResult được lưu")
    base.code(doc, "Trường hợp OK - SOTIEN", '{\n  "criterion": "SOTIEN",\n  "status": "OK",\n  "reason": "DNTT, tờ khai, Ringi và tổng Invoice khớp nhau.",\n  "evidence": {\n    "dntt_amount": "1500.00",\n    "invoice_files": ["Invoice_01.pdf", "Invoice_02.pdf"],\n    "invoice_total": "1500.00",\n    "ringi_file": "Ringi_09.pdf",\n    "ringi_amount": "1500.00"\n  },\n  "rule_version": "nvl-poc-2026-10-06"\n}')
    base.rich_para(doc, [("Ý nghĩa: ", 10.5, True, base.NAVY), ("API-AI có thể lưu nguyên kết quả này; khi người dùng hỏi vì sao OK, hệ thống biết đã dùng file nào, giá trị nào và RuleVersion nào.", 10.5, False, base.BLACK)])

    base.heading(doc, "15.3", "Ví dụ quy ước gọi LLM")
    base.table(doc, ["Tình huống thực tế", "Xử lý", "Kết quả sau cùng"], [
        ("DNTT = 1.500; Invoice + Ringi + tờ khai đều = 1.400.", "Rule Engine tự kết luận vì đủ dữ liệu và sai rõ ràng.", "NG; không gọi LLM."),
        ("Thiếu Invoice hoặc OCR của Invoice rỗng.", "Không có đủ evidence; LLM không được đoán số tiền.", "REVIEW; yêu cầu bổ sung/chạy lại OCR."),
        ("CUSTOMSHEET ghi `ABCXYZ01-02-03`, còn có Invoice `ABCXYZ01`, `ABCXYZ02`, `ABCXYZ03`; chưa chắc cách viết tắt có cùng nhóm hay không.", "Gọi prompt ngoại lệ Số hóa đơn để đề xuất candidate mapping và trả evidence_refs.", "Nếu được chấp nhận → chạy lại Handler + Rule Engine; nếu chưa rõ → REVIEW."),
        ("Một mẫu chứng từ mới có đủ nội dung nhưng tên trường/ngữ cảnh chưa xác định.", "Gọi prompt đúng tiêu chí để giải thích candidate field; không cho LLM ghi OK/NG trực tiếp.", "Sau khi xác nhận và chạy lại Rule Engine mới có thể OK/NG."),
    ], [6.7, 6.2, 3.8], 8.05)
    base.code(doc, "Ví dụ output LLM ngoại lệ - không phải kết quả cuối", '{\n  "criterion": "SOHOADON",\n  "candidate_resolution": "ABCXYZ01-02-03 tương ứng 3 Invoice riêng",\n  "evidence_refs": ["CUSTOMSHEET_01.pdf", "Invoice_01.pdf", "Invoice_02.pdf", "Invoice_03.pdf"],\n  "confidence": 0.96,\n  "should_rerun_rule": true,\n  "final_status": "REVIEW"\n}')

    base.heading(doc, "15.4", "Ví dụ prompt thay đổi")
    base.table(doc, ["Nội dung", "Prompt hiện tại", "Sau khi áp dụng Rule Engine"], [
        ("Đầu vào", "DNTT + toàn bộ dataFiles.", "Evidence đã Mapping + dữ liệu chuẩn hóa + RuleResult REVIEW nếu có."),
        ("Xử lý Số tiền", "Prompt tự đọc, tự cộng, tự đối chiếu và trả CriteriaStatus.", "Prompt chỉ đọc/giải thích phần chưa rõ; Handler cộng Invoice; Rule Engine kết luận."),
        ("Kết quả", "OK / NG / BLANK do prompt trả.", "RuleResult OK / NG / REVIEW / N/A do Python + Rule Engine trả."),
        ("Ví dụ", "Prompt nhận Invoice 01 = 900, Invoice 02 = 600 và tự kết luận.", "Handler tạo `invoice_total=1500`; expression kiểm tra `all_amounts_match`; prompt chỉ được gọi nếu quan hệ chứng từ mơ hồ."),
    ], [2.5, 6.9, 7.3], 8.05)
    base.callout(doc, "Điểm thay đổi quan trọng", "Không phải thay toàn bộ 9 prompt bằng 9 Rule Engine expression. Mỗi prompt vẫn cung cấp dữ liệu hoặc xử lý ngoại lệ; phần kết luận nghiệp vụ xác định được chuyển sang Handler + Rule Engine.", base.YELLOW)

    # keep Times New Roman for all appended content and save
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = base.FONT
            if run._element.rPr is not None and run._element.rPr.rFonts is not None:
                run._element.rPr.rFonts.set(qn("w:ascii"), base.FONT)
                run._element.rPr.rFonts.set(qn("w:hAnsi"), base.FONT)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), base.FONT)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    append_sections()
