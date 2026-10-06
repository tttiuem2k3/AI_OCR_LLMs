from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generate_rule_engine_report as base


REPORT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_POC_Rule_Engine_BEM_AI_06102026.docx")
EVIDENCE_DIR = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Minh_chung_POC_Rule_Engine_06102026\So_do_Rule_Engine")
FONT_REG = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"


def image_font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def wrap_lines(draw, text, font, max_width):
    result = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            result.append("")
            continue
        current = ""
        for word in words:
            candidate = (current + " " + word).strip()
            if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
                current = candidate
            else:
                if current:
                    result.append(current)
                current = word
        if current:
            result.append(current)
    return result


def centered_text(draw, box, text, font, fill="#000000", gap=5):
    x1, y1, x2, y2 = box
    lines = wrap_lines(draw, text, font, x2 - x1 - 30)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + gap * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=font, fill=fill)
        y += height + gap


def node(draw, box, fill, header_fill, title, detail):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill=fill, outline="#1F4E78", width=3)
    draw.rounded_rectangle((x1, y1, x2, y1 + 55), radius=18, fill=header_fill)
    draw.rectangle((x1, y1 + 35, x2, y1 + 55), fill=header_fill)
    centered_text(draw, (x1 + 8, y1 + 4, x2 - 8, y1 + 50), title, image_font(24, True), "#FFFFFF")
    centered_text(draw, (x1 + 14, y1 + 68, x2 - 14, y2 - 14), detail, image_font(20), "#000000")


def arrow(draw, start, end, color="#1F4E78"):
    draw.line([start, end], fill=color, width=6)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        points = [(x2, y2), (x2 - 20, y2 - 13), (x2 - 20, y2 + 13)]
    else:
        points = [(x2, y2), (x2 + 20, y2 - 13), (x2 + 20, y2 + 13)]
    draw.polygon(points, fill=color)


def make_real_case_diagram() -> Path:
    image = Image.new("RGB", (1900, 1080), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (55, 28),
        "Ví dụ thật tháng 09: NVL/09/2026/0003 đi qua Rule Engine",
        font=image_font(39, True),
        fill="#1F4E78",
    )
    draw.line((55, 92, 1845, 92), fill="#1F4E78", width=3)

    boxes = [
        ((55, 175, 340, 800), "#EAF3F8", "#1F4E78", "1. DNTT + FILE", "Phiếu: NVL/09/2026/0003\n\n8 file đính kèm\n\nDNTT cần thanh toán:\nMK202607041 = 1.210 USD\nMK202607045 = 21.636 USD\n\nTổng DNTT = 22.846 USD"),
        ((420, 175, 710, 800), "#D9EAF7", "#1F4E78", "2. TRÍCH XUẤT", "OCR/LLM đọc:\n\n2 Invoice\n2 Packing list\n2 PO\n2 Tờ khai\n\nMỗi file có:\nloại chứng từ\ntên file\nmã hóa đơn/mã thanh toán\nsố tiền"),
        ((790, 175, 1080, 800), "#E2F0D9", "#548235", "3. MAPPING", "Chọn evidence\nđúng dòng DNTT:\n\nMK202607041\n↔ Invoice + list + PO + tờ khai\n\nMK202607045\n↔ Invoice + list + PO + tờ khai\n\nKhông lấy MK202607039\nhoặc MK202607044"),
        ((1160, 175, 1450, 800), "#E2F0D9", "#548235", "4. HANDLER", "Chuẩn hóa và cộng\nđúng evidence đã map:\n\n1.210 + 21.636\n= 22.846 USD\n\nrequired = True\nconflict = False\nall_match = True"),
        ((1530, 175, 1845, 800), "#FCE4D6", "#C55A11", "5. RULE RESULT", "Expression:\nrequired\nand values\nand not conflict\nand all_match\n\n→ True\n\nSOTIEN = OK\n\nLưu reason + evidence\n+ rule_version"),
    ]
    for box_data in boxes:
        node(draw, *box_data)
    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[0][2] + 10, 490), (right[0][0] - 10, 490))

    draw.rounded_rectangle((180, 875, 1720, 1010), radius=18, fill="#FFF2CC", outline="#BF9000", width=3)
    centered_text(
        draw,
        (205, 895, 1695, 990),
        "Điểm mấu chốt: Rule Engine không tự đọc 8 file và không tự biết file nào thuộc dòng nào. Mapping và Python Handler làm phần nghiệp vụ; rule-engine chỉ kiểm tra expression trên context đã chuẩn hóa.",
        image_font(24, True),
        "#7F6000",
    )
    EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    path = EVIDENCE_DIR / "06_Vi_du_thuc_te_NVLT09_0003.png"
    image.save(path, dpi=(180, 180))
    return path


def set_times_new_roman(run):
    run.font.name = base.FONT
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run._element.rPr.rFonts.set(qn(key), base.FONT)


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=base.Cm(17.2))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.font(cap.add_run(caption), 8.5, False, "7F7F7F")


def append_real_case():
    doc = Document(REPORT)
    if any("Ví dụ dữ liệu tháng 09 đi qua Rule Engine" in p.text for p in doc.paragraphs):
        print("real case already exists")
        return

    doc.add_page_break()
    base.heading(doc, "17", "Ví dụ dữ liệu tháng 09 đi qua Rule Engine")
    base.callout(
        doc,
        "Đây là dữ liệu thật dùng để minh họa",
        "Ví dụ lấy từ phiếu NVL/09/2026/0003 trong dữ liệu tháng 09. Kết quả hiện tại bị NG tiêu chí Số tiền vì tổng hợp nhầm cả chứng từ của mã thanh toán khác. Phần dưới mô tả dữ liệu cần được Mapping và xử lý như thế nào khi áp dụng thiết kế Rule Engine.",
        base.YELLOW,
    )
    path = make_real_case_diagram()
    add_figure(doc, path, "Hình 6. Dữ liệu thật đi qua Mapping, Python Handler và rule-engine.")

    base.heading(doc, "17.1", "Dữ liệu đầu vào của phiếu")
    base.table(
        doc,
        ["Nguồn", "Dữ liệu thực tế", "Ý nghĩa khi chạy Rule"],
        [
            ("Phiếu ĐNTT", "NVL/09/2026/0003; USD; 8 file", "Xác định phạm vi của một lần đối chiếu."),
            ("Dòng thanh toán 1", "MK202607041 = 1.210 USD", "Chỉ lấy evidence thuộc mã này."),
            ("Dòng thanh toán 2", "MK202607045 = 21.636 USD", "Chỉ lấy evidence thuộc mã này."),
            ("Tổng cần đối chiếu", "1.210 + 21.636 = 22.846 USD", "Đây là giá trị DNTT cần so sánh."),
            ("File đính kèm", "2 Invoice, 2 Packing list, 2 PO, 2 Tờ khai", "Tất cả file phải có loại, tên và trường đã trích xuất."),
        ],
        [3.0, 7.0, 6.0],
        8.4,
    )

    base.heading(doc, "17.2", "Mapping đúng và Mapping sai khác nhau ở đâu")
    base.table(
        doc,
        ["Cách xử lý", "Dữ liệu được cộng", "Kết quả"],
        [
            ("Đang bị sai", "Lấy cả số 4.493,2 và 93.155 USD trên tờ khai/bảng kê, dù có mã MK202607039 và MK202607044 không thuộc phiếu.", "AI trả NG Số tiền; kết quả phiếu 88,89%."),
            ("Cách Rule Engine cần nhận", "Chỉ map MK202607041 và MK202607045; tổng Invoice thuộc hai mã = 22.846 USD.", "Đủ evidence, không mâu thuẫn, SOTIEN có thể trả OK."),
        ],
        [3.3, 9.4, 3.3],
        8.2,
    )
    base.callout(
        doc,
        "Ý nghĩa của ví dụ",
        "Nếu chỉ viết expression “các số tiền có khớp không” thì Rule Engine không thể tự biết phải bỏ MK202607039 và MK202607044. Quy tắc chọn đúng file và đúng dòng thanh toán phải nằm ở Transaction/Evidence Mapping và Python Handler.",
        base.GREEN,
    )

    base.heading(doc, "17.3", "Context mà rule-engine thực sự nhận")
    base.code(
        doc,
        "Sau khi Mapping và Handler xử lý xong",
        '{\n  "criterion": "SOTIEN",\n  "dntt_amount": 22846.00,\n  "invoice_total": 22846.00,\n  "has_required_evidence": true,\n  "has_required_values": true,\n  "evidence_conflict": false,\n  "all_amounts_match": true\n}',
    )
    base.code(
        doc,
        "Rule được cấu hình cho tiêu chí Số tiền",
        'expression = "has_required_evidence and has_required_values and "\n             "not evidence_conflict and all_amounts_match"\n\nrule_engine.Rule(expression).matches(context)\n# True',
    )
    base.table(
        doc,
        ["Thành phần", "Làm gì với dữ liệu trên"],
        [
            ("Mapping", "Gắn từng Invoice/list/PO/tờ khai vào đúng mã MK202607041 hoặc MK202607045."),
            ("Python Handler", "Chuẩn hóa số tiền, cộng 1.210 + 21.636, kiểm tra thiếu file và mâu thuẫn."),
            ("rule-engine", "Chỉ đánh giá 4 điều kiện Boolean trong expression và trả True/False."),
            ("RuleResult", "Đóng gói SOTIEN, trạng thái, lý do, evidence và rule_version để API-AI lưu."),
        ],
        [3.5, 13.0],
        8.5,
    )

    base.heading(doc, "17.4", "Khi bổ sung một Rule mới sẽ làm như thế nào")
    base.flow(
        doc,
        [
            ("1. Xác định", "Tiêu chí + evidence"),
            ("2. Viết Handler", "Map / chuẩn hóa / tính"),
            ("3. Khai báo", "expression + version"),
            ("4. Test", "đủ / thiếu / mâu thuẫn"),
            ("5. Lưu", "RuleResult"),
        ],
        2,
    )
    base.table(
        doc,
        ["Cấu hình bắt buộc", "Ví dụ với SOTIEN"],
        [
            ("rule_id", "SOTIEN"),
            ("required_evidence", "DNTT, Invoice, Packing list, PO, Tờ khai theo nghiệp vụ đã xác nhận"),
            ("handler", "map_payment_line_and_sum_invoice"),
            ("expression", "has_required_evidence and has_required_values and not evidence_conflict and all_amounts_match"),
            ("rule_version", "nvl-amount-v1"),
        ],
        [5.0, 11.5],
        8.45,
    )
    base.callout(
        doc,
        "Giới hạn của POC",
        "POC chứng minh cách khai báo và chạy expression với dữ liệu đã chuẩn hóa. Việc map theo từng dòng thanh toán của phiếu NVL/09/2026/0003 là yêu cầu nghiệp vụ cần đưa vào Handler/Mapping khi triển khai thật; ví dụ này chưa sửa production.",
        base.YELLOW,
    )

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_times_new_roman(run)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_times_new_roman(run)
    doc.save(REPORT)
    print(REPORT)


if __name__ == "__main__":
    append_real_case()
