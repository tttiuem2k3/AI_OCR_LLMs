from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
REPORT = ROOT / "Bao_cao_Rule_Engine_BEM_AI_06102026.docx"
BACKUP_DIR = ROOT / "Archive"
IMG_DIR = ROOT / "Minh_chung_POC_Rule_Engine_06102026" / "So_do_Rule_Engine"
FONT = "Times New Roman"
FONT_REG = r"C:\Windows\Fonts\times.ttf"
FONT_BOLD = r"C:\Windows\Fonts\timesbd.ttf"

NAVY = "1F4E78"
BLUE = "D9EAF7"
PALE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
GRAY = "F2F2F2"
WHITE = "FFFFFF"
BLACK = "000000"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, size=10.5, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = rgb(color)
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run._element.rPr.rFonts.set(qn(key), FONT)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), fill)
    node.set(qn("w:val"), "clear")


def border(cell, color="D9E2F3", size="5"):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def margins(cell, top=80, side=90, bottom=80):
    pr = cell._tc.get_or_add_tcPr()
    tc_mar = pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        pr.append(tc_mar)
    for edge, value in (("top", top), ("start", side), ("bottom", bottom), ("end", side)):
        item = tc_mar.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            tc_mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def paragraph(container, text="", size=10.5, bold=False, color=BLACK, align=None, before=0, after=5):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    if text:
        set_run_font(p.add_run(text), size, bold, color)
    return p


def rich_paragraph(container, parts, align=None, before=0, after=5):
    p = paragraph(container, before=before, after=after, align=align)
    for text, size, bold, color in parts:
        set_run_font(p.add_run(text), size, bold, color)
    return p


def heading(doc, number, text):
    p = paragraph(doc, before=10, after=5)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(f"{number}. {text}"), 15, True, NAVY)
    return p


def callout(doc, title, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    shade(cell, fill)
    border(cell, "9CC2E5", "7")
    margins(cell, 120, 140, 120)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(title + ": "), 11, True, NAVY)
    set_run_font(p.add_run(text), 11)
    paragraph(doc, after=0)


def table(doc, headers, rows, widths=None, size=8.8):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for cell, value in zip(t.rows[0].cells, headers):
        shade(cell, NAVY)
        border(cell, WHITE, "4")
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(value), size, True, WHITE)
    for row_index, values in enumerate(rows):
        row = t.add_row()
        for cell, value in zip(row.cells, values):
            shade(cell, WHITE if row_index % 2 == 0 else "F8FBFE")
            border(cell)
            margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for line_index, line in enumerate(str(value).split("\n")):
                if line_index:
                    p.add_run().add_break()
                set_run_font(p.add_run(line), size)
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    paragraph(doc, after=0)
    return t


def code_box(doc, label, text):
    p = paragraph(doc, label, 9.8, True, NAVY, after=2)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0)
    shade(cell, GRAY)
    border(cell)
    margins(cell, 100, 110, 100)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.splitlines()):
        if index:
            p.add_run().add_break()
        set_run_font(p.add_run(line), 8.7)
    paragraph(doc, after=0)


def flow(doc, blocks, highlight=-1):
    t = doc.add_table(rows=1, cols=len(blocks) * 2 - 1)
    t.autofit = False
    for index, (title, detail) in enumerate(blocks):
        cell = t.cell(0, index * 2)
        active = index == highlight
        shade(cell, NAVY if active else BLUE)
        border(cell, "9CC2E5", "7")
        margins(cell, 110, 45, 110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        set_run_font(p.add_run(title), 8.5, True, WHITE if active else NAVY)
        p2 = paragraph(cell, detail, 7.7, False, WHITE if active else NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=0)
        if index < len(blocks) - 1:
            arrow_cell = t.cell(0, index * 2 + 1)
            arrow_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = arrow_cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_run_font(p.add_run("→"), 16, True, NAVY)
    paragraph(doc, after=0)


def image_font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def wrap_draw(draw, text, font, max_width):
    lines = []
    for paragraph_text in text.split("\n"):
        words = paragraph_text.split()
        if not words:
            lines.append("")
            continue
        current = ""
        for word in words:
            candidate = (current + " " + word).strip()
            if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines


def centered(draw, xy, text, font, fill="#000000", gap=5):
    x1, y1, x2, y2 = xy
    lines = wrap_draw(draw, text, font, x2 - x1 - 28)
    heights = [draw.textbbox((0, 0), line, font=font)[3] for line in lines]
    total = sum(heights) + gap * max(0, len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=font)[2]
        draw.text((x1 + (x2 - x1 - width) / 2, y), line, font=font, fill=fill)
        y += height + gap


def img_node(draw, box, fill, title, detail, title_fill="#1F4E78"):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=18, fill=fill, outline="#1F4E78", width=3)
    draw.rounded_rectangle((x1, y1, x2, y1 + 58), radius=18, fill=title_fill)
    draw.rectangle((x1, y1 + 36, x2, y1 + 58), fill=title_fill)
    centered(draw, (x1 + 8, y1 + 4, x2 - 8, y1 + 53), title, image_font(25, True), "#FFFFFF")
    centered(draw, (x1 + 15, y1 + 70, x2 - 15, y2 - 15), detail, image_font(21), "#000000")


def img_arrow(draw, start, end, color="#1F4E78"):
    draw.line([start, end], fill=color, width=6)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        pts = [(x2, y2), (x2 - 22, y2 - 14), (x2 - 22, y2 + 14)]
    else:
        pts = [(x2, y2), (x2 + 22, y2 - 14), (x2 + 22, y2 + 14)]
    draw.polygon(pts, fill=color)


def make_main_diagram():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1900, 980), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 30), "Rule Engine trong BEM AI: hiểu theo 1 luồng dữ liệu", font=image_font(40, True), fill="#1F4E78")
    draw.line((55, 96, 1845, 96), fill="#1F4E78", width=3)
    boxes = [
        ((70, 205, 360, 650), "#EAF3F8", "1. DỮ LIỆU", "DNTT + file đã OCR\nvà đã trích xuất\n\nVí dụ:\nInvoice, Ringi, tờ khai\nsố tiền, loại tiền, ngày"),
        ((445, 205, 735, 650), "#D9EAF7", "2. MAPPING", "Chọn đúng chứng từ\ncho đúng dòng DNTT\n\nTrả lời câu hỏi:\nfile nào dùng để kiểm\ndòng nào?"),
        ((820, 205, 1110, 650), "#E2F0D9", "3. HANDLER", "Chuẩn hóa và tính toán\n\nCộng nhiều Invoice\nkiểm tra thiếu file\nkiểm tra mâu thuẫn"),
        ((1195, 205, 1485, 650), "#E2F0D9", "4. RULE ENGINE", "Chạy công thức IF\ntrên context sạch\n\nKhông đọc file\nkhông đoán nghiệp vụ"),
        ((1570, 205, 1830, 650), "#FCE4D6", "5. KẾT QUẢ", "RuleResult\n\nOK / NG\nREVIEW / N/A\n\nCó lý do + evidence\n+ version"),
    ]
    for box in boxes:
        img_node(draw, *box)
    for left, right in zip(boxes, boxes[1:]):
        img_arrow(draw, (left[0][2] + 12, 430), (right[0][0] - 12, 430))
    draw.rounded_rectangle((180, 760, 1720, 890), radius=18, fill="#FFF2CC", outline="#BF9000", width=3)
    centered(
        draw,
        (210, 780, 1690, 870),
        "Cách nhớ nhanh: Mapping chọn đúng chứng từ; Handler chuẩn bị dữ liệu; Rule Engine kiểm tra công thức; RuleResult là kết quả lưu về API-AI/DB.",
        image_font(27, True),
        "#7F6000",
    )
    path = IMG_DIR / "00_Rule_Engine_Luong_De_Hieu.png"
    image.save(path, dpi=(180, 180))
    return path


def make_real_case_diagram():
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1900, 1040), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 30), "Ví dụ thật: phiếu NVL/09/2026/0003 sai ở bước Mapping", font=image_font(39, True), fill="#1F4E78")
    draw.line((55, 96, 1845, 96), fill="#1F4E78", width=3)
    boxes = [
        ((75, 190, 405, 675), "#EAF3F8", "1. PHIẾU", "NVL/09/2026/0003\n8 file\n\nDNTT cần thanh toán:\nMK202607041 = 1.210\nMK202607045 = 21.636\n\nTổng = 22.846 USD"),
        ((485, 190, 815, 675), "#FCE4D6", "2. SAI HIỆN TẠI", "AI lấy nhầm thêm\nmã không thuộc phiếu:\n\nMK202607039\nMK202607044\n\nNên tổng bị lệch\nvà trả NG Số tiền"),
        ((895, 190, 1225, 675), "#E2F0D9", "3. CÁCH ĐÚNG", "Mapping chỉ chọn file\nthuộc 2 dòng DNTT:\n\nMK202607041\nMK202607045\n\nKhông cộng file\nkhác dòng"),
        ((1305, 190, 1635, 675), "#E2F0D9", "4. HANDLER", "Cộng đúng:\n1.210 + 21.636\n= 22.846 USD\n\nrequired=True\nconflict=False\nall_match=True"),
    ]
    for box in boxes:
        img_node(draw, *box)
    for left, right in zip(boxes, boxes[1:]):
        img_arrow(draw, (left[0][2] + 12, 430), (right[0][0] - 12, 430))
    img_node(draw, (685, 760, 1215, 930), "#D9EAF7", "5. RULE ENGINE", "Expression trả True → RuleResult SOTIEN = OK\nKết quả có lý do, evidence và rule_version", "#1F4E78")
    img_arrow(draw, (1470, 690), (1220, 845))
    path = IMG_DIR / "00_Vi_du_Thuc_Te_NVLT09_0003_De_Hieu.png"
    image.save(path, dpi=(180, 180))
    return path


def add_image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(17.2))
    cap = paragraph(doc, caption, 8.5, False, "7F7F7F", WD_ALIGN_PARAGRAPH.CENTER, after=5)
    return cap


def set_defaults(doc):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.45)
    sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            style._element.rPr.rFonts.set(qn(key), FONT)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("BEM AI | Rule Engine | 06/10/2026"), 7.5, False, "7F7F7F")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("Bản trình bày dễ hiểu để chốt cách áp dụng Rule Engine"), 7.5, False, "7F7F7F")


def build_report():
    if REPORT.exists():
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup = BACKUP_DIR / f"Bao_cao_Rule_Engine_BEM_AI_06102026_backup_before_easy_{stamp}.docx"
        shutil.copy2(REPORT, backup)
    else:
        backup = None

    main_diagram = make_main_diagram()
    real_case_diagram = make_real_case_diagram()

    doc = Document()
    set_defaults(doc)

    paragraph(doc, "BÁO CÁO RULE ENGINE CHO BEM AI", 22, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, before=20, after=3)
    paragraph(doc, "Bản dễ hiểu để chốt cách áp dụng vào bước AI đối chiếu", 12.5, True, "5B9BD5", WD_ALIGN_PARAGRAPH.CENTER, after=10)
    callout(
        doc,
        "Kết luận ngắn",
        "Rule Engine phù hợp để kiểm tra các điều kiện nghiệp vụ đã rõ. Tuy nhiên Rule Engine chỉ nên chạy sau khi dữ liệu đã được OCR, trích xuất, Mapping và Python Handler chuẩn hóa. Không dùng Rule Engine để đọc file, không dùng LLM để thay logic đã xác định được.",
        GREEN,
    )
    table(
        doc,
        ["Người đọc cần nhớ", "Diễn giải ngắn"],
        [
            ("Rule Engine", "Giống công thức IF trong Excel: nhận dữ liệu sạch rồi trả đúng/sai."),
            ("Python Handler", "Chuẩn bị dữ liệu cho Rule Engine: chọn file, chuẩn hóa, cộng nhiều Invoice, phát hiện thiếu/mâu thuẫn."),
            ("Evidence Mapping", "Xác định file/chứng từ nào thuộc dòng DNTT nào."),
            ("LLM", "Chỉ hỗ trợ ngoại lệ khi dữ liệu có thật nhưng cách hiểu/mapping còn mơ hồ; không tự kết luận OK/NG."),
            ("RuleResult", "Kết quả chuẩn để API-AI lưu: tiêu chí, trạng thái, lý do, evidence, rule_version."),
        ],
        [5.0, 11.5],
        9.0,
    )
    add_image(doc, main_diagram, "Hình 1. Luồng dễ hiểu: dữ liệu → mapping → handler → rule-engine → kết quả.")

    heading(doc, "1", "Rule Engine dùng để làm gì")
    table(
        doc,
        ["Rule Engine làm", "Rule Engine không làm"],
        [
            ("Kiểm tra điều kiện đã rõ, ví dụ đủ chứng từ và số tiền khớp.", "Không OCR, không đọc file PDF/Excel."),
            ("Chạy expression ngắn, có version, dễ test lại.", "Không tự biết file nào thuộc dòng DNTT nào."),
            ("Trả True/False để Handler đổi thành OK/NG.", "Không tự cộng nhiều Invoice hoặc xử lý nghiệp vụ phức tạp."),
            ("Giúp tách logic deterministic khỏi LLM.", "Không thay DB, không thay API-AI và không tự lưu kết quả."),
        ],
        [8.2, 8.2],
        8.8,
    )
    callout(doc, "Cách nói đơn giản", "Rule Engine không phải AI mới. Nó là lớp kiểm tra công thức nghiệp vụ sau khi AI đã đọc và hệ thống đã chuẩn hóa dữ liệu.", YELLOW)

    heading(doc, "2", "Rule Engine nằm ở bước nào")
    flow(
        doc,
        [
            ("OCR", "lấy text"),
            ("Trích xuất", "ra trường dữ liệu"),
            ("Mapping", "file ↔ dòng DNTT"),
            ("Handler", "chuẩn hóa/tính"),
            ("Rule Engine", "kiểm công thức"),
            ("RuleResult", "lưu kết quả"),
        ],
        4,
    )
    table(
        doc,
        ["Bước", "Đầu vào", "Đầu ra"],
        [
            ("OCR", "File scan/PDF/ảnh", "Text theo từng file."),
            ("Trích xuất", "Text OCR hoặc file có sẵn text", "Số hóa đơn, số tiền, loại tiền, ngày, NCC, điều kiện giao hàng..."),
            ("Mapping", "DNTT + danh sách dữ liệu đã trích xuất", "Evidence list: dòng DNTT nào dùng file nào."),
            ("Python Handler", "Evidence list", "Context sạch cho từng tiêu chí."),
            ("Rule Engine", "Context sạch + expression", "True/False."),
            ("RuleResult", "Kết quả Handler + Rule Engine", "OK/NG/REVIEW/N/A kèm lý do và evidence."),
        ],
        [3.1, 6.6, 6.8],
        8.4,
    )

    heading(doc, "3", "Cấu hình Rule như thế nào")
    table(
        doc,
        ["Trường cấu hình", "Ý nghĩa", "Ví dụ SOTIEN"],
        [
            ("rule_id", "Tiêu chí cần kiểm tra", "SOTIEN"),
            ("rule_version", "Phiên bản rule để trace lại kết quả", "nvl-amount-v1"),
            ("required_evidence", "Chứng từ/trường bắt buộc phải có", "DNTT, Invoice, Ringi, tờ khai..."),
            ("handler", "Hàm Python chuẩn bị dữ liệu", "map_payment_line_and_sum_invoice"),
            ("expression", "Công thức rule-engine kiểm tra", "required and values and not conflict and all_match"),
            ("status_policy", "Cách đổi kết quả thành trạng thái", "True→OK; False→NG; thiếu/mâu thuẫn→REVIEW"),
        ],
        [3.4, 6.3, 6.8],
        8.2,
    )
    code_box(
        doc,
        "Ví dụ expression dễ hiểu",
        "has_required_evidence\nAND has_required_values\nAND NOT evidence_conflict\nAND all_amounts_match",
    )
    callout(doc, "Điểm quan trọng", "Expression càng ngắn càng tốt. Các việc như cộng nhiều Invoice, bỏ file không thuộc dòng DNTT, phát hiện chứng từ mâu thuẫn phải làm ở Python Handler trước.", GREEN)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    heading(doc, "4", "Ví dụ thật để hình dung")
    callout(
        doc,
        "Phiếu minh họa",
        "Dùng phiếu NVL/09/2026/0003 để giải thích. Vấn đề hiện tại là AI lấy nhầm cả dữ liệu của mã thanh toán không thuộc phiếu, nên tiêu chí Số tiền bị NG.",
        YELLOW,
    )
    add_image(doc, real_case_diagram, "Hình 2. Ví dụ thực tế: sai ở Mapping thì Rule Engine không thể tự sửa nếu Handler chưa chuẩn hóa đúng.")
    table(
        doc,
        ["Nội dung", "Dữ liệu"],
        [
            ("Dòng DNTT cần kiểm", "MK202607041 = 1.210 USD; MK202607045 = 21.636 USD."),
            ("Tổng đúng", "22.846 USD."),
            ("Sai hiện tại", "Cộng nhầm thêm dữ liệu của MK202607039 và MK202607044 không thuộc phiếu."),
            ("Cách xử lý đúng", "Mapping chỉ chọn file/evidence thuộc 2 dòng DNTT cần thanh toán."),
            ("Context đưa vào Rule Engine", "required=True; values=True; conflict=False; all_match=True."),
            ("Kết quả mong muốn", "RuleResult tiêu chí SOTIEN có thể trả OK khi evidence đã map đúng."),
        ],
        [5.1, 11.3],
        8.5,
    )

    heading(doc, "5", "RuleResult lưu những gì")
    table(
        doc,
        ["Trường", "Ý nghĩa", "Ví dụ"],
        [
            ("criterion", "Tiêu chí đang kiểm", "SOTIEN"),
            ("status", "Kết quả cuối của tiêu chí", "OK / NG / REVIEW / N/A"),
            ("reason", "Lý do dễ đọc cho kế toán", "Số tiền DNTT và tổng Invoice khớp."),
            ("evidence", "File/trường/giá trị đã dùng", "Invoice_01, Invoice_02, Ringi, tờ khai..."),
            ("rule_version", "Phiên bản rule đã chạy", "nvl-amount-v1"),
        ],
        [3.5, 7.0, 6.0],
        8.6,
    )
    table(
        doc,
        ["Trạng thái", "Khi nào dùng", "Có gọi LLM không"],
        [
            ("OK", "Dữ liệu đủ, không mâu thuẫn và thỏa Rule.", "Không."),
            ("NG", "Dữ liệu đủ, không mâu thuẫn nhưng vi phạm Rule.", "Không."),
            ("REVIEW", "Thiếu file, OCR/trích xuất thiếu, hoặc chứng từ mâu thuẫn.", "Chỉ gọi nếu thuộc ngoại lệ được phép."),
            ("N/A", "Tiêu chí không áp dụng cho nghiệp vụ đó.", "Không."),
        ],
        [2.7, 10.2, 3.6],
        8.5,
    )

    heading(doc, "6", "Khi nào được gọi LLM")
    table(
        doc,
        ["Tình huống", "Cách xử lý"],
        [
            ("Rule Engine đã đủ cơ sở trả OK hoặc NG", "Không gọi LLM."),
            ("Thiếu file hoặc OCR rỗng", "Không cho LLM đoán; trả REVIEW/chờ bổ sung hoặc OCR lại."),
            ("Chứng từ mâu thuẫn rõ", "Giữ REVIEW để kiểm tra nghiệp vụ/hồ sơ."),
            ("Có evidence thật nhưng cách viết/mapping mơ hồ", "Có thể gọi LLM để đề xuất candidate evidence, sau đó chạy lại Handler + Rule Engine."),
            ("Mẫu chứng từ mới nhưng đủ dữ liệu", "LLM chỉ hỗ trợ nhận diện trường/candidate; kết quả cuối vẫn do Rule Engine quyết định."),
        ],
        [6.0, 10.5],
        8.5,
    )
    callout(doc, "Nguyên tắc chốt", "LLM không được thay Rule Engine để kết luận OK/NG cho logic đã xác định được. LLM chỉ bổ sung candidate evidence cho trường hợp ngoại lệ.", ORANGE)

    heading(doc, "7", "Kết quả POC và cách áp dụng 9 tiêu chí")
    table(
        doc,
        ["Nội dung POC", "Kết quả"],
        [
            ("Thư viện", "Đã kiểm tra cách khai báo, load và chạy expression bằng rule-engine zeroSteiner."),
            ("Tiêu chí POC", "Loại tiền và Số tiền nhiều Invoice."),
            ("Case đã test", "Đủ evidence, thiếu evidence, nhiều evidence và chứng từ mâu thuẫn."),
            ("Kết luận kỹ thuật", "Dùng được làm expression layer, nhưng phải đi kèm Python Handler và Evidence Mapping."),
            ("Phạm vi hiện tại", "POC/chốt thiết kế; chưa triển khai production toàn bộ 9 Rule."),
        ],
        [5.2, 11.2],
        8.6,
    )
    table(
        doc,
        ["Thứ tự triển khai", "Việc cần làm"],
        [
            ("1", "Chốt Evidence Contract cho 9 tiêu chí NVL."),
            ("2", "Viết Handler cho từng tiêu chí: NCC, số hóa đơn, ngày hóa đơn, số tiền, loại tiền, điều kiện giao hàng, hạn thanh toán, ngày hoàn thành kiểm tra, chữ ký/con dấu."),
            ("3", "Khai báo Rule Catalog gồm rule_id, version, evidence, handler, expression, status_policy."),
            ("4", "Tạo test pack đủ/thiếu/mâu thuẫn cho từng tiêu chí."),
            ("5", "Chạy shadow mode trên bộ NVL cũ trước khi bật kết quả chính thức."),
        ],
        [3.0, 13.5],
        8.4,
    )

    doc.save(REPORT)
    return backup


if __name__ == "__main__":
    backup_path = build_report()
    print(REPORT)
    if backup_path:
        print(backup_path)
