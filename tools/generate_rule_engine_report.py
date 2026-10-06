from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09\Bao_cao_POC_Rule_Engine_BEM_AI_06102026.docx")
FONT = "Times New Roman"
NAVY, BLUE, PALE, GRAY = "1F4E78", "D9EAF7", "EAF3F8", "F2F2F2"
GREEN, YELLOW, WHITE, BLACK = "E2F0D9", "FFF2CC", "FFFFFF", "000000"


def font(run, size=10.5, bold=False, color=BLACK, name=FONT):
    run.font.name = name
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        run._element.rPr.rFonts.set(qn(key), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), fill)
    node.set(qn("w:val"), "clear")


def border(cell, color="D9E2F3", size="5"):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:color"), color)


def margins(cell, top=75, side=90, bottom=75):
    pr = cell._tc.get_or_add_tcPr()
    tc_mar = pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        pr.append(tc_mar)
    for edge, value in (("top", top), ("start", side), ("bottom", bottom), ("end", side)):
        item = tc_mar.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            tc_mar.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def para(container, text="", size=10.5, bold=False, color=BLACK, align=None, before=0, after=4):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.03
    if align is not None:
        p.alignment = align
    if text:
        font(p.add_run(text), size, bold, color)
    return p


def rich_para(container, parts, align=None, before=0, after=4):
    p = para(container, before=before, after=after, align=align)
    for text, size, bold, color in parts:
        font(p.add_run(text), size, bold, color)
    return p


def heading(doc, number, title):
    p = para(doc, before=10, after=5)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(f"{number}. {title}"), 14, True, NAVY)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(text), 10)
    return p


def callout(doc, title, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, fill); border(c, "9CC2E5", "7"); margins(c, 110, 130, 110)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(title + ": "), 10.5, True, NAVY)
    font(p.add_run(text), 10.5)
    para(doc, after=0)


def table(doc, headers, rows, widths=None, size=8.6):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"; t.autofit = False
    for c, value in zip(t.rows[0].cells, headers):
        shade(c, NAVY); border(c, WHITE, "4"); margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        font(p.add_run(value), size, True, WHITE)
    for index, values in enumerate(rows):
        row = t.add_row()
        for c, value in zip(row.cells, values):
            shade(c, WHITE if index % 2 == 0 else "F8FBFE"); border(c); margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            for line_index, line in enumerate(str(value).split("\n")):
                if line_index: p.add_run().add_break()
                font(p.add_run(line), size)
    if widths:
        for row in t.rows:
            for c, width in zip(row.cells, widths): c.width = Cm(width)
    para(doc, after=0)
    return t


def code(doc, label, text):
    p = para(doc, label, 9.5, True, NAVY, after=2)
    p.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0); shade(c, GRAY); border(c); margins(c, 100, 110, 100)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    for index, line in enumerate(text.splitlines()):
        if index: p.add_run().add_break()
        font(p.add_run(line), 8.25, False, BLACK, FONT)


def flow(doc, blocks, highlight=-1):
    t = doc.add_table(rows=1, cols=len(blocks) * 2 - 1); t.autofit = False
    for index, (title, detail) in enumerate(blocks):
        c = t.cell(0, index * 2); active = index == highlight
        shade(c, NAVY if active else BLUE); border(c, "9CC2E5", "7"); margins(c, 100, 45, 100)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(1)
        font(p.add_run(title), 8.3, True, WHITE if active else NAVY)
        p = para(c, detail, 7.4, False, WHITE if active else NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=0)
        if index < len(blocks) - 1:
            a = t.cell(0, index * 2 + 1); a.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = a.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
            font(p.add_run("→"), 15, True, NAVY)
    para(doc, after=0)


def defaults(doc):
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(1.55), Cm(1.45)
    sec.left_margin, sec.right_margin = Cm(1.65), Cm(1.65)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet", "List Bullet 2"):
        style = doc.styles[style_name]; style.font.name = FONT
        for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            style._element.rPr.rFonts.set(qn(key), FONT)
    doc.styles["Normal"].font.size = Pt(10.5)
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(h.add_run("BEM AI  |  POC Rule Engine  |  06/10/2026"), 7.5, False, "7F7F7F")
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run("Tài liệu POC kỹ thuật dùng để chốt thiết kế trước khi triển khai"), 7.5, False, "7F7F7F")


def cover(doc):
    for _ in range(4): para(doc, after=0)
    para(doc, "RULE ENGINE CHO AI BEM", 24, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para(doc, "POC và thiết kế áp dụng trong bước AI đối chiếu", 13, False, BLACK, WD_ALIGN_PARAGRAPH.CENTER, after=24)
    callout(doc, "Mục tiêu chốt", "Dùng rule-engine zeroSteiner để thực thi điều kiện nghiệp vụ đã xác định; không dùng LLM thay thế Rules deterministic.", GREEN)
    para(doc, after=16)
    flow(doc, [("OCR", "đã hoàn tất"), ("Trích xuất", "dữ liệu cấu trúc"), ("AI đối chiếu", "Mapping và Rules"), ("Kết quả", "OK NG REVIEW N A")], 2)
    para(doc, "Phạm vi: Nguyên vật liệu  |  POC: Số tiền và Loại tiền  |  Chưa thay đổi production", 9.5, False, "595959", WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()


def build():
    doc = Document(); defaults(doc); cover(doc)
    heading(doc, "1", "rule engine zeroSteiner là gì")
    rich_para(doc, [("rule-engine ", 10.5, True, NAVY), ("là thư viện Python cung cấp ngôn ngữ expression nhẹ. Thư viện nhận dict/object context, đánh giá biểu thức và trả Boolean. Trong POC dùng phiên bản 5.0.2.", 10.5, False, BLACK)])
    table(doc, ["Rule Engine làm", "Rule Engine không làm"], [
        ("Đánh giá điều kiện đã chuẩn hóa, ví dụ đủ chứng từ và các loại tiền khớp nhau.", "Không OCR file, không phân loại chứng từ, không tự Mapping file vào dòng DNTT."),
        ("Giữ expression ngắn, minh bạch, có thể test theo từng rule version.", "Không tự cộng nhiều Invoice, không tính Deadline phức tạp, không giải thích nghiệp vụ."),
        ("Trả Boolean cho handler chuyển thành OK hoặc NG.", "Không tự lưu DB, không quản lý queue và không thay API AI."),
    ], [8.1, 8.1], 9)
    callout(doc, "Nguyên tắc", "Python handler phải chặn thiếu evidence hoặc evidence mâu thuẫn trước. Các trường hợp này trả REVIEW; Rule Engine không được ép ra OK hoặc NG khi dữ liệu chưa đáng tin cậy.", YELLOW)

    heading(doc, "2", "Vị trí và kiến trúc trong AI đối chiếu")
    rich_para(doc, [("Rule Engine chỉ được gọi ở bước ", 10.5, False, BLACK), ("AI đối chiếu", 10.5, True, NAVY), (". OCR và trích xuất phải hoàn tất trước để tạo dữ liệu cấu trúc; Mapping phải xác định evidence nào thuộc dòng DNTT cần kiểm tra.", 10.5, False, BLACK)])
    flow(doc, [("OCR", "text theo file"), ("Trích xuất", "trường dữ liệu"), ("Evidence Mapping", "dòng DNTT ↔ file"), ("Python Handler", "chuẩn hóa / tổng hợp"), ("rule-engine", "expression"), ("RuleResult", "OK NG REVIEW N A")], 4)
    table(doc, ["Lớp", "Nhiệm vụ", "Đầu vào", "Đầu ra"], [
        ("Evidence Mapping", "Chọn đúng chứng từ cho mỗi dòng DNTT.", "Dữ liệu DNTT, loại chứng từ, trường đã trích xuất.", "Evidence list có file và trường sử dụng."),
        ("Python handler", "Chuẩn hóa, cộng/tách nhiều chứng từ, tính toán, phát hiện thiếu hoặc mâu thuẫn.", "Evidence đã Mapping.", "Context đơn giản cho expression hoặc REVIEW."),
        ("rule-engine", "Đánh giá điều kiện deterministic.", "Context Boolean hoặc giá trị đã chuẩn hóa.", "True / False."),
        ("RuleResult", "Chuẩn hóa kết quả để API AI lưu và ERP đọc.", "Kết quả handler + expression.", "Trạng thái, lý do, evidence, version."),
        ("LLM ngoại lệ", "Chỉ hỗ trợ khi evidence đủ nhưng chưa có Rule rõ hoặc diễn giải đặc biệt.", "Evidence và RuleResult REVIEW có kiểm soát.", "Gợi ý xử lý; không thay kết quả Rule deterministic."),
    ], [2.55, 4.3, 4.65, 4.7], 8.35)

    heading(doc, "3", "Khai báo load và thực thi Rule")
    rich_para(doc, [("POC đã chứng minh đủ 3 bước: ", 10.5, True, NAVY), ("khai báo expression → load thư viện và rule → thực thi với context.", 10.5, False, BLACK)])
    table(doc, ["Bước", "POC đang làm", "Khi triển khai 9 Rule"], [
        ("1. Khai báo", "Khai báo criterion và expression trong Python.", "Mỗi RuleDefinition có id, version, evidence contract, handler và expression."),
        ("2. Load", "Import rule_engine; tạo Rule từ expression khi chạy.", "Load Rule Catalog khi AI service khởi động; cache theo RuleVersion."),
        ("3. Thực thi", "Handler lập context, Rule.matches(context) trả True/False.", "Lưu context/evidence và RuleResult để trace lại từng lần chạy."),
    ], [2.3, 6.9, 7.0], 8.7)
    code(doc, "Minh chứng POC: khai báo expression và thực thi rule-engine", 'CURRENCY_EXPRESSION_RULE = _ExpressionRule(\n    criterion="LOAITIEN",\n    expression="has_required_evidence and has_required_values and all_currencies_match",\n)\n\nreturn bool(rule_engine.Rule(self.expression).matches(dict(context)))')
    code(doc, "Context sau khi Python handler đã chuẩn hóa", 'context = {\n  "has_required_evidence": True,\n  "has_required_values": True,\n  "all_currencies_match": True\n}\nRule(...).matches(context)  # True -> handler trả OK')

    heading(doc, "4", "Cấu trúc RuleResult và quy ước trạng thái")
    code(doc, "Cấu trúc bắt buộc cho mọi tiêu chí", 'RuleResult(\n  criterion="SOTIEN",\n  status="OK",\n  reason="Số tiền DNTT, tờ khai, Ringi và tổng Invoice khớp nhau.",\n  evidence={...file, trường, giá trị đã dùng...},\n  rule_version="nvl-poc-2026-10-06"\n)')
    table(doc, ["Trạng thái", "Khi nào dùng", "Ví dụ"], [
        ("OK", "Evidence đầy đủ, không mâu thuẫn và thỏa điều kiện Rule.", "Tổng 2 Invoice = DNTT = Tờ khai = Ringi."),
        ("NG", "Evidence đầy đủ, đáng tin cậy nhưng vi phạm Rule.", "Loại tiền DNTT USD, Invoice JPY."),
        ("REVIEW", "Thiếu evidence, OCR/trích xuất thiếu trường hoặc chứng từ mâu thuẫn.", "Ringi 1.600 nhưng DNTT và Invoice 1.500."),
        ("N/A", "Tiêu chí không áp dụng theo nghiệp vụ đã xác nhận.", "Phiếu không thuộc phạm vi kiểm tra tiêu chí."),
    ], [2.2, 7.25, 6.5], 8.7)
    callout(doc, "Quy tắc an toàn", "REVIEW không được biến thành OK hoặc NG bằng suy đoán của LLM. Thiếu hoặc mâu thuẫn file phải trả về để bổ sung hoặc xác nhận nghiệp vụ.", YELLOW)

    heading(doc, "5", "POC 1 Loại tiền")
    rich_para(doc, [("Mục tiêu POC: ", 10.5, True, NAVY), ("kiểm chứng rule-engine xử lý tốt tiêu chí có dữ liệu đơn giản sau chuẩn hóa. Dữ liệu POC là dữ liệu kiểm thử có cấu trúc, không phải phiếu production.", 10.5, False, BLACK)])
    table(doc, ["Nguồn evidence", "Tên file POC", "Loại tiền"], [
        ("DNTT", "Thông tin dòng DNTT", "USD"),
        ("Tờ khai", "To_khai_01.pdf", "USD"),
        ("PO", "PO_4500123.pdf", "USD"),
        ("Ringi", "Ringi_2026_09.pdf", "USD"),
        ("Invoice", "Invoice_INV001.pdf", "USD"),
    ], [3.3, 7.6, 5.2], 8.7)
    flow(doc, [("Evidence", "4 chứng từ"), ("Handler", "chuẩn hóa USD"), ("Expression", "3 điều kiện true"), ("RuleResult", "OK")], 2)
    table(doc, ["Case test", "Context / evidence", "Kết quả thực tế"], [
        ("Đủ evidence", "DNTT, tờ khai, PO, Ringi, Invoice đều USD.", "OK - Loại tiền trên DNTT và chứng từ đều là USD."),
        ("Thiếu evidence", "Thiếu Ringi và Invoice.", "REVIEW - ghi rõ chứng từ còn thiếu; không gọi LLM kết luận."),
        ("Sai loại tiền", "Invoice là JPY, các nguồn còn lại USD.", "NG - evidence đủ nhưng có chứng từ khác loại tiền."),
        ("Không áp dụng", "Context applicable = False.", "N/A - Rule không chạy cho case này."),
    ], [3.1, 6.6, 6.4], 8.45)

    heading(doc, "6", "POC 2 Số tiền nhiều Invoice")
    rich_para(doc, [("Mục tiêu POC: ", 10.5, True, NAVY), ("chứng minh phần tổng hợp nghiệp vụ phải do Python handler xử lý, sau đó rule-engine mới đánh giá điều kiện cuối.", 10.5, False, BLACK)])
    table(doc, ["Evidence", "Giá trị POC", "Cách dùng"], [
        ("DNTT", "1.500,00", "Giá trị cần kiểm tra."),
        ("Tờ khai", "1.500,00", "Chứng từ bắt buộc."),
        ("Ringi", "1.500,00", "Chứng từ bắt buộc."),
        ("Invoice 01 + 02", "900,00 + 600,00", "Handler cộng thành 1.500,00."),
        ("Bảng kê Statement", "1.500,00", "Chứng từ kiểm soát; không cộng thêm vào Invoice để tránh double count."),
    ], [3.2, 4.2, 8.7], 8.55)
    flow(doc, [("Evidence", "5 chứng từ"), ("Python Handler", "900 + 600 = 1.500"), ("Expression", "đủ và không mâu thuẫn"), ("RuleResult", "OK")], 1)
    code(doc, "Expression chỉ kết luận sau khi handler đã cộng Invoice và phát hiện mâu thuẫn", 'AMOUNT_EXPRESSION_RULE = _ExpressionRule(\n  criterion="SOTIEN",\n  expression="has_required_evidence and has_required_values and "\n             "not evidence_conflict and all_amounts_match"\n)')
    table(doc, ["Case test", "Xử lý handler", "Kết quả"], [
        ("Hai Invoice khớp", "Cộng 900 + 600; Statement dùng đối chiếu kiểm soát.", "OK"),
        ("Ringi mâu thuẫn", "DNTT/tờ khai/Invoice 1.500; Ringi 1.600.", "REVIEW - không kết luận NG vì hồ sơ đang mâu thuẫn."),
        ("Statement mâu thuẫn", "Invoice 1.500; Statement 1.600.", "REVIEW - cần kiểm tra hồ sơ."),
        ("Đủ nhưng sai", "DNTT 1.500; tờ khai/Ringi/Invoice 1.400.", "NG - evidence đầy đủ nhưng không khớp."),
    ], [3.35, 7.3, 5.45], 8.35)

    heading(doc, "7", "Minh chứng kiểm thử POC")
    rich_para(doc, [("Lệnh đã chạy ngày 06/10/2026: ", 10.5, True, NAVY), ("python -m unittest tests.test_rule_engine_poc -v", 10.5, False, BLACK)])
    code(doc, "Kết quả test thực tế", 'test_amount_rule_does_not_double_count_statement_and_invoice_total ... ok\ntest_amount_rule_returns_ng_when_evidence_is_complete_but_amount_does_not_match ... ok\ntest_amount_rule_returns_ok_when_multiple_invoices_sum_to_dntt_amount ... ok\ntest_amount_rule_returns_review_when_evidence_is_conflicting ... ok\ntest_amount_rule_returns_review_when_statement_conflicts_with_invoice_total ... ok\ntest_currency_rule_returns_ng_when_complete_evidence_has_wrong_currency ... ok\ntest_currency_rule_returns_ok_when_required_evidence_matches_dntt_currency ... ok\ntest_currency_rule_returns_review_when_required_evidence_is_missing ... ok\ntest_poc_exposes_clear_error_if_rule_engine_library_is_not_installed ... ok\ntest_rule_returns_not_applicable_when_criterion_is_disabled_for_case ... ok\n\nRan 10 tests in 0.007s\nOK')
    table(doc, ["Nhóm test", "Case đã có", "Mục đích"], [
        ("Đủ evidence", "Loại tiền đúng; tổng nhiều Invoice đúng.", "Chứng minh Rule trả OK khi evidence hoàn chỉnh."),
        ("Thiếu evidence", "Thiếu Ringi / Invoice / trường dữ liệu bắt buộc.", "Chứng minh trả REVIEW, không suy đoán."),
        ("Nhiều evidence", "2 Invoice + Statement.", "Kiểm tra handler tổng hợp, không double count."),
        ("Mâu thuẫn", "Ringi hoặc Statement khác tổng Invoice.", "Nhận diện evidence conflict và trả REVIEW."),
        ("Kết quả sai", "Evidence đủ nhưng số tiền / loại tiền không khớp.", "Chứng minh trả NG."),
        ("Không áp dụng và lỗi thư viện", "N/A và thiếu package.", "Chứng minh contract status và lỗi kỹ thuật rõ ràng."),
    ], [3.35, 6.7, 6.05], 8.45)

    heading(doc, "8", "Ranh giới Expression Python handler và LLM")
    rich_para(doc, [("Nguyên tắc chốt: ", 10.5, True, NAVY), ("mọi logic có thể xác định từ dữ liệu đã Mapping phải được xử lý bằng Python handler và rule-engine. LLM không được thay thế logic này.", 10.5, False, BLACK)])
    table(doc, ["Tiêu chí NVL", "Python handler bắt buộc", "rule-engine expression", "LLM"], [
        ("NCC", "Chuẩn hóa tên, MST, alias; chọn chứng từ nguồn.", "So sánh giá trị đã chuẩn hóa.", "Chỉ ngoại lệ tên/mẫu mới đã đủ evidence."),
        ("Số hóa đơn", "Tách nhiều số, map Invoice đúng dòng DNTT.", "Kiểm tra tập hợp số hóa đơn.", "Không thay Rules."),
        ("Ngày hóa đơn", "Chuẩn hóa ngày, gắn đúng Invoice.", "Kiểm tra điều kiện ngày.", "Chỉ diễn giải ngoại lệ."),
        ("Số tiền", "Cộng Invoice, tiền tệ, phân biệt Statement kiểm soát.", "Kết luận khớp/không khớp khi không conflict.", "Không tự cộng hoặc suy đoán."),
        ("Loại tiền", "Chuẩn hóa USD, JPY, VND và chọn evidence.", "So sánh bằng nhau.", "Không cần nếu Rules đủ."),
        ("Điều kiện giao hàng", "Tách Incoterm và địa điểm từ chứng từ.", "So sánh điều kiện đã chuẩn hóa.", "Chỉ khi cách diễn đạt không chuẩn."),
        ("Deadline thanh toán", "Xác định ngày mốc, điều khoản, lịch nghỉ.", "Kiểm tra đạt/không đạt deadline.", "Không tính thay Python."),
        ("Ngày hoàn thành kiểm tra", "Chọn mốc ngày và chuẩn hóa theo nghiệp vụ.", "Kiểm tra quan hệ ngày.", "Không thay Rules."),
        ("Chữ ký con dấu", "Xác định chứng từ bắt buộc và tín hiệu ký/dấu.", "Kiểm tra có/không theo nghiệp vụ.", "Chỉ nhận diện trường hợp đặc biệt đã đủ ảnh."),
    ], [2.35, 5.1, 4.45, 4.1], 7.55)
    callout(doc, "Không tạo nguồn Rules thứ tư", "Excel Rules NVL là nguồn xác nhận nghiệp vụ; Rule Catalog trong AI Python là bản thực thi đã version hóa; API AI chỉ điều phối và lưu kết quả; Prompt chỉ dùng cho LLM ngoại lệ, không giữ logic deterministic cuối cùng.", YELLOW)

    heading(doc, "9", "Cấu trúc triển khai 9 Rule không cần nghiên cứu lại")
    table(doc, ["Thành phần", "Nội dung bắt buộc", "Mục đích kiểm soát"], [
        ("Rule Catalog", "rule_id, rule_version, evidence contract, handler, expression, applicable condition.", "Biết chính xác Rule nào chạy cho mỗi kết quả."),
        ("Evidence contract", "Loại chứng từ, trường dữ liệu, quan hệ Mapping và dữ liệu bắt buộc cho từng tiêu chí.", "Không chạy Rule khi thiếu đầu vào."),
        ("Python handler", "Chuẩn hóa, tổng hợp, tính toán và phát hiện mâu thuẫn.", "Giữ logic nghiệp vụ phức tạp trong code có test."),
        ("Expression", "Điều kiện ngắn, deterministic, chỉ nhận context đã sạch.", "Dễ đọc, dễ test và không phụ thuộc LLM."),
        ("RuleResult", "criterion, status, reason, evidence, rule_version.", "API AI/DB trace được kết quả theo file và version."),
        ("Test pack", "Case OK, NG, REVIEW, N/A; case nhiều evidence và regression bộ NVL cũ.", "Không giảm kết quả hiện có khi thêm Rule mới."),
    ], [3.0, 7.3, 5.8], 8.45)
    table(doc, ["Thứ tự", "Việc cần làm sau POC", "Điều kiện hoàn thành"], [
        ("1", "Chốt Evidence Contract cho 9 tiêu chí từ Rules NVL đã xác nhận.", "Mỗi tiêu chí có chứng từ/trường bắt buộc và N/A condition."),
        ("2", "Viết 9 handler và expression theo Rule Catalog.", "Mỗi Rule có RuleVersion và RuleResult thống nhất."),
        ("3", "Bổ sung POC test pack cho từng tiêu chí.", "Có OK, NG, REVIEW, N/A; có case nhiều file/mâu thuẫn khi phù hợp."),
        ("4", "Chạy shadow mode trên bộ NVL đã xác nhận.", "So sánh trước/sau, phân loại chênh lệch và sửa trước khi bật."),
        ("5", "API AI lưu RuleResult, evidence và version; bật kết quả chính thức theo version được duyệt.", "Trace được phiếu, file, rule và lý do kết quả."),
    ], [1.2, 8.4, 6.5], 8.35)

    heading(doc, "10", "Nội dung đề nghị chốt")
    table(doc, ["Nội dung", "Đề nghị chốt"], [
        ("Công nghệ", "Dùng rule-engine zeroSteiner làm expression layer trong AI Python."),
        ("Vị trí chạy", "Chỉ chạy sau Evidence Mapping trong bước AI đối chiếu; không thuộc OCR/trích xuất."),
        ("Vai trò các lớp", "Python handler xử lý nghiệp vụ phức tạp; rule-engine kết luận expression; LLM chỉ xử lý ngoại lệ có kiểm soát."),
        ("Kết quả chuẩn", "Dùng RuleResult và 4 trạng thái OK / NG / REVIEW / N/A."),
        ("Cách triển khai", "Triển khai 9 Rule theo Rule Catalog + test pack + shadow mode, không cập nhật trực tiếp vào production."),
        ("Nguyên tắc chất lượng", "Rule mới phải regression trên bộ NVL cũ; không để LLM thay logic deterministic."),
    ], [3.8, 12.3], 8.9)
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        node = doc.paragraphs[-1]._element
        node.getparent().remove(node)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
