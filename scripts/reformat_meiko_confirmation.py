from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
SOURCE = ROOT / "Huong_dan_kiem_tra_thu_cong_Hang_nhap_kho_MEIKO.docx"
TARGETS = [
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026.docx",
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026_hoan_chinh.docx",
]

NAVY = "1F4E79"
BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def has_drawing(paragraph):
    return bool(paragraph._p.xpath(".//w:drawing"))


def clear_and_write(paragraph, text, size=10.5, bold=False, color=None):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run(run, size=size, bold=bold, color=color)


def set_run(run, size=10.5, bold=None, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    element = tcpr.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        tcpr.append(element)
    element.set(qn("w:fill"), color)


def set_cell_width(cell, cm):
    tcpr = cell._tc.get_or_add_tcPr()
    width = tcpr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tcpr.append(width)
    width.set(qn("w:w"), str(int(cm * 567)))
    width.set(qn("w:type"), "dxa")
    cell.width = Cm(cm)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def dont_split(row):
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    properties.append(node)


def repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    properties.append(node)


def normalize_paragraph(paragraph, kind="body"):
    if has_drawing(paragraph):
        return
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_together = True
    for run in paragraph.runs:
        set_run(run, 10.5)
    if kind == "heading1":
        paragraph.style = "Heading 1"
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run(run, 13, True, NAVY)
    elif kind == "heading2":
        paragraph.style = "Heading 2"
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            set_run(run, 11.5, True, NAVY)
    elif kind == "bullet":
        paragraph.style = "List Bullet"
        paragraph.paragraph_format.left_indent = Cm(0.55)
        paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        paragraph.paragraph_format.space_after = Pt(2)


def move_after(element, paragraph):
    paragraph._p.addnext(element)


def add_response_box(document, paragraph):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    move_after(table._tbl, paragraph)
    entries = [
        ("Đúng", "Ghi “Đúng” nếu nội dung mô tả đúng cách kế toán đang kiểm."),
        ("Chưa đúng", "Sửa trực tiếp bảng ghép chứng từ, thứ tự bước và câu trả lời."),
        ("Chưa rõ", "Ghi “Cần kiểm tra thêm” hoặc gửi mã ĐNTT/file mẫu."),
    ]
    for idx, (title, body) in enumerate(entries):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, 5.8)
        set_cell_margins(cell)
        shade(cell, PALE_BLUE if idx != 2 else YELLOW)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        first = p.add_run(title + "\n")
        set_run(first, 10.5, True, NAVY)
        second = p.add_run(body)
        set_run(second, 9.5)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    properties = table._tbl.tblPr
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    columns = len(table.columns)
    if columns == 1:
        widths = [17.4]
    elif columns == 2:
        widths = [5.4, 12.0]
    else:
        widths = [4.0, 5.3, 8.1]
    for row_index, row in enumerate(table.rows):
        dont_split(row)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run(run, 9.5)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(index, len(widths) - 1)])
        if row_index == 0 and len(table.rows) > 1:
            repeat_header(row)
            for cell in row.cells:
                shade(cell, NAVY)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run(run, 9.5, True, WHITE)
        if columns == 1:
            for cell in row.cells:
                shade(cell, YELLOW)


def build():
    document = Document(SOURCE)
    section = document.sections[0]
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.45)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    title_lines = [
        "BẢNG XÁC NHẬN CÁCH ĐỐI CHIẾU HỒ SƠ",
        "NGUYÊN VẬT LIỆU HÀNG NHẬP KHO",
        "MEIKO vui lòng xác nhận cách ghép chứng từ và thứ tự kiểm thực tế. Nội dung chưa được xác nhận được ghi rõ là Cần kiểm tra thêm.",
        "Cách phản hồi theo từng mục",
    ]
    for index, text in enumerate(title_lines):
        paragraph = document.paragraphs[index]
        clear_and_write(paragraph, text, 17 if index == 0 else 13 if index == 1 else 10.5, index < 2, NAVY if index < 2 else None)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index < 2 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(4 if index < 2 else 5)
    add_response_box(document, document.paragraphs[3])

    section_titles = {
        "Trước hết, chọn đúng loại hồ sơ",
        "Tên gọi dùng trong tài liệu",
        "Cách ghi kết quả",
        "Cách so theo nội dung đã ghi",
        "MEIKO ghi thêm các trường hợp chưa được nói rõ",
        "Một hồ sơ mẫu, nếu có",
    }
    for paragraph in list(document.paragraphs[4:]):
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        if " | " in text:
            remove_paragraph(paragraph)
            continue
        if text in section_titles or text.startswith("Ví dụ số tiền trong file"):
            normalize_paragraph(paragraph, "heading2")
        elif re.match(r"^\d+\.\s", text):
            normalize_paragraph(paragraph, "heading1")
        elif text.startswith("Xem:"):
            normalize_paragraph(paragraph)
            for run in paragraph.runs:
                run.bold = True
        elif text.startswith("- "):
            clear_and_write(paragraph, text[2:])
            normalize_paragraph(paragraph, "bullet")
        else:
            normalize_paragraph(paragraph)

    for table in document.tables:
        style_table(table)

    header = section.header.paragraphs[0]
    clear_and_write(header, "MEIKO | Hàng nhập kho", 8.5, False, "7F7F7F")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    clear_and_write(footer, "Tài liệu xác nhận quy tắc đối chiếu", 8.5, False, "7F7F7F")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.core_properties.title = "Xác nhận cách đối chiếu hồ sơ nguyên vật liệu hàng nhập kho"
    document.core_properties.subject = "Tài liệu MEIKO xác nhận cách ghép chứng từ và thứ tự kiểm"
    return document


def main():
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for target in TARGETS:
        if target.exists():
            backup = target.with_name(f"{target.name}.bak_{stamp}_before_professional_format")
            shutil.copy2(target, backup)
            print(f"BACKUP={backup}")
    document = build()
    for target in TARGETS:
        document.save(target)
        print(f"SAVED={target}")


if __name__ == "__main__":
    main()
