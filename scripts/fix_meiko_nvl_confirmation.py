from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
SOURCE = ROOT / "Huong_dan_kiem_tra_thu_cong_Hang_nhap_kho_MEIKO.docx"
TARGETS = [
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026.docx",
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026_hoan_chinh.docx",
]


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def insert_before(paragraph, text="", size=11, bold=False, color=None, align=None):
    new_paragraph = paragraph.insert_paragraph_before()
    new_paragraph.paragraph_format.space_after = Pt(4)
    new_paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        new_paragraph.alignment = align
    run = new_paragraph.add_run(text)
    set_font(run, size, bold, color)
    return new_paragraph


def add_table_before(paragraph, rows):
    table = paragraph._parent.add_table(rows=1, cols=3, width=Cm(18))
    paragraph._p.addprevious(table._tbl)
    table.style = "Table Grid"
    headers = ["Nội dung kiểm tra", "Kết luận", "Việc đã sửa / cần MEIKO xác nhận"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(header), 10.5, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            set_font(cells[index].paragraphs[0].add_run(value), 10.5, index == 1)


def main():
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for target in TARGETS:
        if target.exists():
            backup = target.with_name(f"{target.name}.bak_{stamp}_before_fix_confirm")
            shutil.copy2(target, backup)
            print(f"BACKUP={backup}")
    document = Document(SOURCE)
    first = document.paragraphs[0]
    insert_before(first, "BẢNG XÁC NHẬN CÁCH ĐỐI CHIẾU HỒ SƠ", 16, True, RGBColor(31, 78, 121), WD_ALIGN_PARAGRAPH.CENTER)
    insert_before(first, "NGUYÊN VẬT LIỆU – HÀNG NHẬP KHO", 14, True, RGBColor(31, 78, 121), WD_ALIGN_PARAGRAPH.CENTER)
    insert_before(first, "Kết luận rà soát bản trước: Chưa đúng. Đã sửa lại theo Bo_Rules.xlsx và tài liệu hướng dẫn thủ công; các điểm chưa được MEIKO xác nhận đều giữ là “Cần kiểm tra thêm”.", 11.5, True)
    insert_before(first, "Cách MEIKO phản hồi: Đúng: ghi “Đúng”. Chưa đúng: sửa bảng ghép chứng từ, thứ tự bước và trả lời câu hỏi. Chưa rõ: ghi “Cần kiểm tra thêm” hoặc gửi phiếu/file mẫu.", 11)
    add_table_before(first, [
        ("Bảng ghép chứng từ", "Chưa đúng", "Đã thay bằng bản bám Bo_Rules: tách theo từng tiêu chí, không tự coi ví dụ hoặc câu hỏi là rule."),
        ("Thứ tự bước", "Chưa đúng", "Đã sửa theo luồng thủ công: chọn loại hồ sơ trước, rồi kiểm số hóa đơn/bảng kê/tờ khai và từng tiêu chí."),
        ("Câu hỏi chưa rõ", "Chưa đúng", "Đã chuyển thành ô MEIKO xác nhận thêm; Không tự trả lời thay MEIKO."),
        ("Điểm chưa được MEIKO xác nhận", "Cần kiểm tra thêm", "Giữ nguyên trạng thái chờ xác nhận hoặc yêu cầu gửi mã ĐNTT/file mẫu."),
    ])
    page_break = insert_before(first, "", 1)
    page_break.add_run().add_break(WD_BREAK.PAGE)
    for target in TARGETS:
        document.save(target)
        print(f"SAVED={target}")


if __name__ == "__main__":
    main()
