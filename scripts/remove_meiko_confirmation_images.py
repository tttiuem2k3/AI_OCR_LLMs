from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document


ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
TARGETS = [
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026.docx",
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026_hoan_chinh.docx",
]

REMOVE_TEXT_MARKERS = [
    "Ảnh tờ khai dưới đây lấy trực tiếp",
    "Vùng tổng: có số 100.975,0589",
    "Dòng “Chi tiết khai trị giá”: ghi",
    "Ví dụ có sẵn trong Bo_Rules.xlsx",
    "PO trong ảnh ghi FOB HONGKONG",
]
REPLACE_TEXT = {
    "2) So cả tổng tiền và tiền chi tiết từng hóa đơn với tờ khai/bảng kê. Tờ khai có dòng “Chi tiết khai trị giá” thì lấy trị giá nêu ở dòng đó; không chỉ nhìn con số tổng phía trên. Có ảnh thật ở trang 4.":
    "2) So cả tổng tiền và tiền chi tiết từng hóa đơn với tờ khai/bảng kê. Nếu tờ khai có dòng “Chi tiết khai trị giá” thì lấy trị giá nêu ở dòng đó; không chỉ nhìn con số tổng phía trên.",
}


def remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def remove_table(table) -> None:
    table._tbl.getparent().remove(table._tbl)


def has_drawing(element) -> bool:
    return bool(element.xpath(".//w:drawing"))


def clean_document(path: Path) -> tuple[int, int]:
    document = Document(path)
    removed_paragraphs = 0
    removed_tables = 0

    for paragraph in list(document.paragraphs):
        text = " ".join(paragraph.text.split())
        if text in REPLACE_TEXT:
            paragraph.text = REPLACE_TEXT[text]
        if has_drawing(paragraph._p) or any(marker in text for marker in REMOVE_TEXT_MARKERS):
            remove_paragraph(paragraph)
            removed_paragraphs += 1

    for table in list(document.tables):
        table_text = " ".join(cell.text for row in table.rows for cell in row.cells)
        if has_drawing(table._tbl) or any(marker in table_text for marker in REMOVE_TEXT_MARKERS):
            remove_table(table)
            removed_tables += 1

    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs[:-1]):
        is_blank_page_break = not paragraph.text.strip() and bool(
            paragraph._p.xpath('.//w:br[@w:type="page"]')
        )
        if is_blank_page_break and paragraphs[index + 1].text.startswith("6. Điều kiện giao hàng"):
            remove_paragraph(paragraph)
            removed_paragraphs += 1

    document.save(path)
    purge_orphaned_media(path)
    return removed_paragraphs, removed_tables


def purge_orphaned_media(path: Path) -> None:
    temporary = path.with_name(f"{path.stem}.cleaning{path.suffix}")
    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename.startswith("word/media/"):
                continue
            data = source.read(item.filename)
            if item.filename.endswith(".rels"):
                data = re.sub(
                    rb'\s*<Relationship\b(?=[^>]*/image)(?=[^>]*Target="media/)[^>]*/>',
                    b"",
                    data,
                )
            target.writestr(item, data)
    temporary.replace(path)


def main() -> None:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    for target in TARGETS:
        if target.exists():
            backup = target.with_name(f"{target.name}.bak_{stamp}_before_remove_images")
            shutil.copy2(target, backup)
            print(f"BACKUP={backup}")
            removed_paragraphs, removed_tables = clean_document(target)
            print(f"SAVED={target}; removed_paragraphs={removed_paragraphs}; removed_tables={removed_tables}")


if __name__ == "__main__":
    main()
