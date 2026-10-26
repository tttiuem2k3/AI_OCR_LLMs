from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"E:\Asoft\AI_BEM\AI_BEM_Check_T08_09")
REFERENCE = ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_v1.docx"
TARGETS = [
    REFERENCE,
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026.docx",
    ROOT / "Khung_xac_nhan_Rules_NVL_gui_MEIKO_08102026_hoan_chinh.docx",
]
NAVY = "1F4E79"
LIGHT_BLUE = "D9EAF7"
YELLOW = "FFF2CC"
GREEN = "E2F0D9"
RED = "FCE4D6"
WHITE = "FFFFFF"
GRAY = "F2F2F2"

DOC_TYPES = [
    ("ĐNTT", "Phiếu đề nghị thanh toán, nơi kế toán chọn dòng tiền cần kiểm.", "Xác nhận dòng nào thanh toán cho hóa đơn/chứng từ nào."),
    ("Hóa đơn / VAT / Invoice", "Chứng từ bán hàng hoặc invoice nhà cung cấp phát hành.", "Xác nhận số hóa đơn, ngày, số tiền, loại tiền, chữ ký nếu có."),
    ("Tờ khai hải quan", "Tờ khai của đúng lô hàng đang thanh toán.", "Xác nhận số invoice/bảng kê, trị giá, điều kiện giao hàng, ngày hoàn thành kiểm tra."),
    ("PO", "Đơn đặt hàng có điều kiện mua hàng và thanh toán.", "Xác nhận Payment Term, Shipping Terms, loại tiền và quan hệ với invoice."),
    ("COM", "Hóa đơn thương mại / Commercial Invoice.", "Xác nhận số hóa đơn, ngày, tiền, loại tiền, điều kiện giao hàng."),
    ("Bảng kê HĐTM", "Bảng kê hóa đơn thương mại hoặc bảng kê chi tiết đi kèm.", "Xác nhận dòng invoice nào thuộc khoản thanh toán và số bảng kê nối sang tờ khai."),
    ("Ringi", "Phiếu phê duyệt nội bộ.", "Xác nhận số Ringi và số tiền duyệt cho từng nhóm dòng ĐNTT."),
    ("PL", "Phiếu đóng gói / Packing List.", "Xác nhận số invoice, số kiện, mô tả hàng và quan hệ với hóa đơn/tờ khai."),
]

MATRIX = [
    ("Tên nhà cung cấp", "Tên NCC / mã NCC nếu có", "ĐNTT, hóa đơn, tờ khai, PO, COM, bảng kê, Ringi, PL", "Tên tiếng Việt/Anh/Nhật, mã NCC, MST, địa chỉ nếu cần", "Cần MEIKO chốt khi tên khác ngôn ngữ/cách viết."),
    ("Số hóa đơn", "Số hóa đơn trên ĐNTT", "Hóa đơn VAT hoặc IV/IV_PL", "Bảng kê HĐTM, tờ khai", "Không bảng kê: ĐNTT → hóa đơn → tờ khai. Có bảng kê: hóa đơn → dòng bảng kê → số bảng kê trên tờ khai."),
    ("Ngày hóa đơn", "Không lấy ngày ĐNTT", "Hóa đơn, tờ khai, COM, bảng kê, PL", "Số hóa đơn/số bảng kê đã ghép đúng", "Cần MEIKO chốt ngày nào là ngày ưu tiên khi các giấy khác nhau."),
    ("Số tiền", "Số tiền trên ĐNTT", "Hóa đơn, COM, bảng kê, tờ khai", "Ringi nếu có", "Chọn đúng invoice thuộc ĐNTT; không cộng invoice ngoài khoản thanh toán."),
    ("Loại tiền", "Loại tiền trên ĐNTT", "Hóa đơn, tờ khai, PO, COM, bảng kê, Ringi, PL", "Tỷ giá/chứng từ quy đổi nếu có", "Cần MEIKO chốt trường hợp khác loại tiền."),
    ("Điều kiện giao hàng", "Không lấy trực tiếp từ ĐNTT", "PO, tờ khai, COM, PL/IV", "Địa điểm giao hàng", "DAP và DAP Meiko được ghi nhận là cùng điều kiện; trường hợp khác cần xác nhận."),
    ("Hạn thanh toán", "Hạn thanh toán trên ĐNTT", "Payment Term trên PO", "Ngày hàng đến/ngày hóa đơn/ngày vận đơn/tờ khai", "Cần MEIKO chốt ngày mốc cho AMS và 90 after B/L."),
    ("Ngày hoàn thành kiểm tra", "Không áp dụng cho trả trước", "Tờ khai", "File rõ/thiếu trang", "Có ngày thì OK; ô trống cần MEIKO chốt NG hay REVIEW."),
    ("Chữ ký, con dấu", "Không lấy từ ĐNTT", "VAT, PO, COM, PL/IV", "Địa chỉ PO để phân biệt trong nước/nước ngoài", "Cần MEIKO chốt yêu cầu cụ thể với IV/PL/COM."),
]

CRITERIA = [
    {"title":"Tên nhà cung cấp","basis":"Bo_Rules đánh dấu kiểm tên NCC trên nhiều chứng từ. Ghi chú hiện có nói tên trên tờ khai và ASOFT theo tên MEIKO dùng ở Koubai; tên trên hóa đơn do NCC phát hành.","steps":["Chọn đúng bộ chứng từ thuộc khoản thanh toán.","Đọc tên NCC trên ĐNTT và chứng từ liên quan.","Nếu tên khác ngôn ngữ/cách viết, dùng thông tin MEIKO xác nhận để nhận diện cùng NCC."],"questions":["Tên trên hóa đơn có dùng để kết luận không?", "Khi tên tiếng Việt/Anh/Nhật khác nhau, kế toán dựa vào mã NCC, MST, địa chỉ hay thông tin nào?"],"status":"Cần kiểm tra thêm"},
    {"title":"Số hóa đơn","basis":"Không có bảng kê: ĐNTT ↔ hóa đơn ↔ tờ khai. Có bảng kê: ĐNTT ↔ hóa đơn ↔ dòng bảng kê ↔ số bảng kê trên tờ khai. Trả trước bỏ qua.","steps":["Xác định hồ sơ có bảng kê hay không.","So số hóa đơn trên ĐNTT với VAT/IV/IV_PL.","Nếu có bảng kê, tìm đúng dòng hóa đơn trên bảng kê rồi lấy số bảng kê đối chiếu tờ khai.","Nếu không có bảng kê, so số hóa đơn trực tiếp với tờ khai."],"questions":["Nếu bảng kê có thêm hóa đơn C nhưng ĐNTT chỉ trả A, B thì chọn A, B bằng thông tin nào?", "Các khác biệt định dạng như 000123/123 hoặc AB-01/AB01 xử lý thế nào?"],"status":"Cần kiểm tra thêm"},
    {"title":"Ngày hóa đơn","basis":"Bo_Rules đánh dấu hóa đơn, tờ khai, COM, bảng kê và PL; không đánh dấu ngày trên ĐNTT.","steps":["Ghép chứng từ theo số hóa đơn hoặc số bảng kê đã xác định.","Đọc đúng ngày hóa đơn/ngày phát hành trên từng chứng từ.","Không tự thay bằng ngày giao hàng hoặc ngày đăng ký tờ khai.","Khi các ngày khác nhau, cần cách xử lý MEIKO xác nhận."],"questions":["Hóa đơn ghi 02/09, bảng kê ghi 03/09 thì dùng ngày nào?", "Khi tờ khai ghi số bảng kê thì so ngày bảng kê hay ngày hóa đơn lẻ?"],"status":"Cần kiểm tra thêm"},
    {"title":"Số tiền","basis":"Bo_Rules yêu cầu so tổng và chi tiết từng hóa đơn với tờ khai/bảng kê; nếu có Ringi thì tiền Ringi phải lớn hơn hoặc bằng tổng các dòng cùng Ringi.","steps":["Chọn đúng các invoice thuộc dòng ĐNTT.","Cộng tiền của invoice thuộc khoản thanh toán, không cộng invoice ngoài khoản thanh toán.","So số tiền chi tiết và tổng giữa hóa đơn, bảng kê, COM và tờ khai.","Nếu có Ringi, nhóm dòng ĐNTT cùng số Ringi và so với số tiền duyệt."],"questions":["Thanh toán một phần hóa đơn thì lấy dòng/số tiền nào?", "Chênh lệch do thuế, phí hoặc làm tròn được chấp nhận đến mức nào?"],"status":"Cần kiểm tra thêm"},
    {"title":"Loại tiền","basis":"Bo_Rules đánh dấu ĐNTT, hóa đơn, tờ khai, PO, COM, bảng kê, Ringi và PL nhưng chưa nói rõ ngoại lệ quy đổi.","steps":["Đọc loại tiền trên ĐNTT.","Đọc loại tiền trên từng chứng từ thuộc khoản thanh toán.","Nếu loại tiền khác nhau, chỉ kết luận sau khi có quy tắc quy đổi MEIKO xác nhận."],"questions":["PO ghi USD, hóa đơn/ĐNTT ghi VND thì khi nào được chấp nhận?", "Nếu quy đổi, lấy tỷ giá ở chứng từ nào và theo ngày nào?"],"status":"Cần kiểm tra thêm"},
    {"title":"Điều kiện giao hàng","basis":"Bo_Rules đánh dấu tờ khai, PO, COM, PL; phần diễn giải hướng dẫn đọc điều kiện trong phần chi tiết của IV/INV. DAP và DAP Meiko được xem là cùng điều kiện.","steps":["Chọn PO, tờ khai, COM và PL/IV của cùng lô hàng.","Đọc điều kiện và địa điểm trên từng chứng từ.","Chấp nhận DAP và DAP Meiko là cùng điều kiện.","Các cặp điều kiện khác hoặc địa điểm khác cần MEIKO xác nhận."],"questions":["PO ghi DAP nhưng tờ khai ghi DDU/CIP/FCA/FOB/CIF thì xử lý thế nào?", "Tên tỉnh/nước khác nhau thì khi nào OK, khi nào NG?"],"status":"Cần kiểm tra thêm"},
    {"title":"Hạn thanh toán","basis":"AMS30/60/90: lấy ngày hàng đến + số ngày rồi lấy ngày cuối tháng. 90 after B/L trong diễn giải đang ghi lấy ngày hóa đơn + 90 ngày; điểm này chưa thống nhất với tên B/L.","steps":["Đọc Payment Term trên PO.","Chọn ngày mốc đúng theo từng điều khoản.","Cộng số ngày và lấy ngày cuối tháng; nếu rơi vào ngày nghỉ thì lùi về ngày làm việc trước.","So ngày trên ĐNTT với hạn tính được theo điều kiện không thanh toán sớm hơn."],"questions":["AMS lấy Ngày hàng đến hay Ngày hoàn thành kiểm tra?", "90 after B/L lấy ngày vận đơn hay ngày hóa đơn?", "Ngày ĐNTT đúng bằng hạn tính ra có chấp nhận không?"],"status":"Cần kiểm tra thêm"},
    {"title":"Ngày hoàn thành kiểm tra","basis":"Kiểm trường Ngày hoàn thành kiểm tra trên tờ khai. Có ngày thì OK; để trống ghi không OK nhưng chưa tách NG hay REVIEW. Trả trước bỏ qua.","steps":["Ghép đúng tờ khai với khoản thanh toán.","Tìm đúng ô Ngày hoàn thành kiểm tra.","Không thay bằng Ngày hàng đến, ngày đăng ký hoặc trạng thái thông quan.","Nếu ô trống hoặc file không rõ, xử lý theo cách MEIKO xác nhận."],"questions":["Ô ngày nhìn rõ nhưng trống thì NG hay yêu cầu bổ sung?", "File mờ/thiếu trang xử lý khác ô thực sự trống như thế nào?"],"status":"Cần kiểm tra thêm"},
    {"title":"Chữ ký, con dấu","basis":"Trong nước: VAT có chữ ký điện tử bên bán; PO có chữ ký và con dấu/xác nhận hai bên. Nước ngoài: PO có chữ ký hoặc con dấu hai bên. IV/PL/COM còn cần xác nhận rõ.","steps":["Xem địa chỉ trên PO để phân loại trong nước/nước ngoài.","Mở đủ trang ký/xác nhận của VAT, PO, COM, PL/IV.","Kiểm PO theo yêu cầu trong nước/nước ngoài.","IV/PL/COM cần MEIKO xác nhận rõ chữ ký/con dấu và bên ký."],"questions":["IV, PL, COM cần chữ ký, con dấu hay một trong hai?", "Cần của bên mua, bên bán hay cả hai?", "File mờ/thiếu trang ký thì REVIEW hay NG?"],"status":"Cần kiểm tra thêm"},
]

SELF_CHECKS = [
    ("Tôi có hiểu tài liệu yêu cầu mình làm gì không?", "Đạt", "Trang đầu nói rõ cần xác nhận chứng từ, trường dữ liệu, thứ tự kiểm và ngoại lệ."),
    ("Có còn ảnh minh họa gây rối không?", "Đạt", "Bản này không nhúng ảnh; chỉ giữ bảng và câu hỏi nghiệp vụ."),
    ("Có biết tiêu chí nào đã rõ, tiêu chí nào cần hỏi lại không?", "Đạt", "Mỗi tiêu chí có trạng thái và câu hỏi MEIKO xác nhận."),
    ("Có chỗ nào ASOFT tự chốt thay MEIKO không?", "Đạt", "Các điểm chưa chốt đều ghi Cần kiểm tra thêm."),
    ("Kế toán có thể ghi phản hồi trực tiếp không?", "Đạt", "Mỗi dòng có ô Đúng/Chưa đúng/Cần kiểm tra thêm và dòng ghi ví dụ phiếu/file."),
]

WORD_DOCUMENT_TYPES = [
    ("DNTT", "Đề nghị thanh toán\n(PAYMENT REQUEST)", "VAT", "Hóa đơn giá trị gia tăng\n(VAT INVOICE)"),
    ("INVOICE", "Hóa đơn/Invoice dùng để thanh toán\n(INVOICE)", "INVOICE/PL", "Hóa đơn kiêm Phiếu đóng gói\n(INVOICE/PACKING LIST)"),
    ("COM", "Hóa đơn thương mại dùng trong hồ sơ nhập khẩu\n(COMMERCIAL INVOICE)", "STATEMENT", "Bảng kê\n(STATEMENT)"),
    ("TKHQ", "Tờ khai hải quan\n(CUSTOMS DECLARATION)", "PO", "Đơn đặt hàng\n(PURCHASE ORDER)"),
    ("RINGI", "Phiếu phê duyệt nội bộ\n(RINGI)", "PL", "Phiếu đóng gói\n(PACKING LIST)"),
]

WORD_MATRIX = [
    ("Tên nhà cung cấp", "Tên NCC / mã NCC", "DNTT; VAT/INVOICE/INVOICE-PL; TKHQ; PO; COM; STATEMENT; RINGI; PL", "Đọc tên NCC trên các chứng từ. Nếu tên khác ngôn ngữ/cách viết thì dùng tên chuẩn nội bộ, mã NCC, MST hoặc địa chỉ theo cách MEIKO xác nhận.", "Cần MEIKO chốt thông tin nào là khóa nhận diện cùng NCC."),
    ("Số hóa đơn", "Số hóa đơn trên DNTT", "DNTT; VAT/INVOICE/INVOICE-PL; COM; TKHQ; STATEMENT nếu có", "Không có STATEMENT: nối DNTT ↔ hóa đơn ↔ TKHQ. Có STATEMENT: nối hóa đơn ↔ dòng STATEMENT ↔ số STATEMENT trên TKHQ.", "Cần MEIKO xác nhận cách chọn đúng dòng STATEMENT và cách phân biệt INVOICE/COM trong file mẫu."),
    ("Ngày hóa đơn", "Không lấy ngày DNTT", "VAT/INVOICE; TKHQ; COM; STATEMENT; PL", "Chỉ kiểm sau khi đã ghép đúng hóa đơn hoặc STATEMENT. Cách kiểm hiện tại chưa chốt nguồn ngày ưu tiên khi mỗi chứng từ ghi khác ngày.", "Cần MEIKO chốt nguồn ngày ưu tiên."),
    ("Số tiền", "Số tiền trên DNTT", "DNTT; VAT/INVOICE; COM; STATEMENT; TKHQ; RINGI", "So từng hóa đơn và tổng tiền. Nếu TKHQ có dòng Chi tiết khai trị giá thì dùng số tiền tại dòng đó. Nếu có RINGI, tiền RINGI không nhỏ hơn tổng DNTT cùng RINGI.", "Cần MEIKO chốt thanh toán một phần, thuế, phí và làm tròn."),
    ("Loại tiền", "Loại tiền trên DNTT", "DNTT; VAT/INVOICE; TKHQ; PO; COM; STATEMENT; RINGI; PL", "Đối chiếu loại tiền giữa các chứng từ. Nếu khác loại tiền thì chưa tự quy đổi khi chưa có tỷ giá/ngày tỷ giá được MEIKO xác nhận.", "Cần MEIKO chốt có cho phép quy đổi không."),
    ("Điều kiện giao hàng", "Không lấy trực tiếp từ DNTT", "TKHQ; PO; COM; PL; INVOICE", "Đọc điều kiện giao hàng trên PO, INVOICE/COM và TKHQ. DAP và DAP Meiko đang được xem là cùng điều kiện.", "Cần MEIKO chốt các điều kiện/địa điểm giao hàng khác nhau."),
    ("Hạn thanh toán", "Ngày thanh toán trên DNTT", "DNTT; PO; TKHQ", "Dùng điều khoản thanh toán trên PO để tính ngày cuối tháng phải thanh toán; sau đó so với ngày thanh toán trên DNTT.", "Cần MEIKO chốt ngày mốc và cách xử lý ngày nghỉ."),
    ("Ngày hoàn thành kiểm tra", "Không áp dụng cho trả trước", "TKHQ", "Chỉ kiểm trường Ngày hoàn thành kiểm tra trên TKHQ. Có ngày thì ghi nhận đạt; để trống hoặc file không rõ thì để Cần kiểm tra thêm cho đến khi MEIKO chốt NG hay REVIEW.", "Cần MEIKO chốt NG hay REVIEW khi để trống/file không rõ."),
    ("Chữ ký, con dấu", "Không lấy từ DNTT", "VAT/INVOICE/INVOICE-PL; PO; COM; PL", "Kiểm theo hồ sơ trong nước/nước ngoài. Dùng địa chỉ trên PO để phân loại hồ sơ trước khi kiểm chữ ký/con dấu.", "Cần MEIKO chốt yêu cầu chữ ký/con dấu với INVOICE, COM và PL."),
]

SUMMARY_RULES = [
    "So tên NCC giữa các file; nếu khác cách viết/ngôn ngữ thì dùng danh mục nội bộ theo MEIKO xác nhận.",
    "Không có STATEMENT: DNTT ↔ hóa đơn ↔ TKHQ. Có STATEMENT: hóa đơn ↔ dòng STATEMENT ↔ TKHQ.",
    "Chỉ so sau khi ghép đúng hóa đơn/STATEMENT; chưa chốt nguồn ngày ưu tiên khi chứng từ khác ngày.",
    "So từng hóa đơn và tổng tiền; TKHQ dùng dòng Chi tiết khai trị giá nếu có.",
    "So loại tiền giữa chứng từ; không tự quy đổi khi MEIKO chưa chốt tỷ giá và ngày tỷ giá.",
    "Đọc điều kiện trên PO, INVOICE/COM và TKHQ; DAP và DAP Meiko đang được xem là cùng điều kiện.",
    "Dùng điều khoản thanh toán trên PO để tính hạn, sau đó so với ngày thanh toán trên DNTT.",
    "Chỉ kiểm Ngày hoàn thành kiểm tra trên TKHQ; có ngày thì ghi nhận đạt, để trống/file không rõ thì Cần kiểm tra thêm.",
    "Kiểm theo hồ sơ trong nước/nước ngoài; dùng địa chỉ PO để phân loại trước khi kiểm chữ ký/con dấu.",
]

WORD_CRITERIA = [
    {
        "title": "Tên nhà cung cấp",
        "basis": "Kiểm tên NCC trên DNTT, VAT/INVOICE/INVOICE-PL, TKHQ, PO, COM, STATEMENT, RINGI và PL. Tên trên RINGI thường là tiếng Nhật hoặc tiếng Anh; tên chuẩn trên hệ thống/TKHQ cần theo danh mục nội bộ. Tên do nhà cung cấp phát hành trên hóa đơn có thể khác cách chuẩn hóa nội bộ.",
        "steps": ["Chọn đúng bộ hồ sơ của khoản thanh toán.", "Đối chiếu tên NCC trên DNTT và TKHQ với tên chuẩn nội bộ.", "Dùng PO, RINGI, COM, STATEMENT hoặc PL để kiểm chéo khi tên khác ngôn ngữ/cách viết.", "Khi chưa xác định cùng NCC, kiểm tra thêm mã NCC, MST hoặc địa chỉ."],
        "questions": ["Khóa nào dùng để nhận diện cùng một NCC khi tên Việt/Anh/Nhật khác nhau?", "Tên trên hóa đơn do nhà cung cấp phát hành có được dùng để kết luận không?"],
    },
    {
        "title": "Số hóa đơn",
        "basis": "Hồ sơ trong nước dùng VAT; hồ sơ nước ngoài dùng INVOICE hoặc INVOICE/PL. Không có STATEMENT: số hóa đơn trên DNTT, hóa đơn và TKHQ phải nối đúng nhau. Có STATEMENT: phải nối đúng hóa đơn với dòng STATEMENT, sau đó nối số STATEMENT với TKHQ.",
        "steps": ["Xác định hồ sơ dùng VAT hay INVOICE/INVOICE-PL.", "Xác định có STATEMENT hay không.", "Không có STATEMENT: đối chiếu số hóa đơn DNTT ↔ hóa đơn ↔ TKHQ.", "Có STATEMENT: đối chiếu hóa đơn ↔ dòng STATEMENT ↔ số STATEMENT trên TKHQ."],
        "questions": ["Nếu STATEMENT có thêm hóa đơn ngoài khoản thanh toán, chọn đúng hóa đơn bằng khóa nào?", "Khác biệt định dạng như 000123/123 hoặc AB-01/AB01 xử lý thế nào?"],
    },
    {
        "title": "Ngày hóa đơn",
        "basis": "Đối chiếu ngày trên VAT/INVOICE, TKHQ, COM, STATEMENT và PL; không lấy ngày trên DNTT. Cách kiểm hiện tại chưa quy định nguồn ngày ưu tiên khi chứng từ hiển thị ngày khác nhau.",
        "steps": ["Ghép đúng số hóa đơn hoặc số STATEMENT trước khi kiểm tra ngày.", "Đọc ngày hóa đơn/ngày phát hành trên từng chứng từ đã ghép.", "Không thay bằng ngày giao hàng, ngày đăng ký tờ khai hoặc ngày khác.", "Nếu ngày khác nhau, ghi nguồn ngày kế toán đã dùng trước đây."],
        "questions": ["VAT/INVOICE ghi 02/09 nhưng STATEMENT ghi 03/09 thì dùng ngày nào?", "TKHQ ghi số STATEMENT thì so ngày STATEMENT hay ngày của từng hóa đơn?"],
    },
    {
        "title": "Số tiền",
        "basis": "Trên TKHQ, dùng số tiền tổng hóa đơn ở trang 1; nếu có dòng Chi tiết khai trị giá thì dùng số tiền tại dòng đó. STATEMENT lấy tổng các dòng chi tiết. VAT/INVOICE cộng tổng các hóa đơn để so với TKHQ và STATEMENT; đồng thời so từng hóa đơn. Nếu có RINGI, số tiền RINGI phải không nhỏ hơn tổng các dòng DNTT cùng RINGI.",
        "steps": ["Chọn đúng các hóa đơn thuộc dòng DNTT.", "So số tiền từng hóa đơn với STATEMENT và TKHQ.", "So tổng tiền của nhóm hóa đơn với STATEMENT và TKHQ theo đúng trường tiền được nêu.", "Nếu có RINGI, cộng các dòng DNTT cùng RINGI và so với số tiền phê duyệt."],
        "questions": ["Thanh toán một phần hóa đơn thì dùng số tiền/dòng chi tiết nào?", "Chênh lệch do thuế, phí hoặc làm tròn được chấp nhận đến mức nào?"],
    },
    {
        "title": "Loại tiền",
        "basis": "Đối chiếu loại tiền trên DNTT, VAT/INVOICE, TKHQ, PO, COM, STATEMENT, RINGI và PL. Cách kiểm hiện tại chưa nêu cách quy đổi khi loại tiền giữa các chứng từ khác nhau.",
        "steps": ["Đọc loại tiền trên DNTT.", "Đọc loại tiền trên từng chứng từ của đúng khoản thanh toán.", "Nếu khác loại tiền, ghi rõ chứng từ và loại tiền đang khác.", "Chỉ kết luận sau khi có quy tắc quy đổi được MEIKO xác nhận."],
        "questions": ["PO ghi USD nhưng hóa đơn/DNTT ghi VND thì khi nào được chấp nhận?", "Nếu quy đổi, dùng tỷ giá từ chứng từ nào và theo ngày nào?"],
    },
    {
        "title": "Điều kiện giao hàng",
        "basis": "Đọc điều kiện giao hàng tại dòng tổng trị giá hóa đơn trên TKHQ, Shipping Terms trên PO và phần chi tiết trên INVOICE; COM và PL dùng kiểm chéo. DAP và DAP Meiko được coi là cùng điều kiện. Quy ước hiện tại: giữa tên nước và tỉnh/thành của Việt Nam hoặc Nhật Bản thì đạt; giữa các tỉnh/thành khác nhau thì không đạt.",
        "steps": ["Chọn PO, TKHQ, COM, PL và INVOICE của cùng lô hàng.", "Đọc Shipping Terms trên PO và điều kiện giao hàng tại vị trí tương ứng trên các chứng từ khác.", "Chấp nhận DAP và DAP Meiko là cùng điều kiện.", "Ghi rõ Incoterm/địa điểm khi các chứng từ khác nhau."],
        "questions": ["PO ghi DAP nhưng TKHQ ghi DDU/CIP/FCA/FOB/CIF thì xử lý thế nào?", "Tên nước/tỉnh/thành khác nhau thì trường hợp nào đạt, trường hợp nào không đạt?"],
    },
    {
        "title": "Hạn thanh toán",
        "basis": "Dùng điều khoản thanh toán trên PO. Với AMS30/60/90: lấy Ngày hàng đến cộng số ngày, sau đó lấy ngày cuối tháng của tháng tính được. Với 90 after B/L: cách kiểm hiện tại lấy ngày hóa đơn cộng 90 ngày, sau đó lấy ngày cuối tháng. Nếu ngày cuối tháng rơi vào thứ bảy, chủ nhật hoặc ngày nghỉ lễ thì dùng ngày làm việc ngay trước đó. Ngày thanh toán trên DNTT không được sớm hơn ngày đã tính.",
        "steps": ["Đọc điều khoản thanh toán trên PO.", "Xác định điều khoản AMS30/60/90 hoặc 90 after B/L.", "Lấy ngày mốc theo cách kiểm hiện tại, cộng số ngày và xác định ngày cuối tháng.", "Nếu trùng ngày nghỉ, lùi về ngày làm việc trước; so với ngày thanh toán trên DNTT."],
        "questions": ["Ngày hàng đến lấy ở chứng từ/trường nào?", "Vì sao điều khoản 90 after B/L đang lấy ngày hóa đơn làm ngày mốc; MEIKO có xác nhận không?", "Ngày DNTT đúng bằng hạn đã tính có được chấp nhận không?"],
    },
    {
        "title": "Ngày hoàn thành kiểm tra",
        "basis": "Chỉ kiểm tra trường Ngày hoàn thành kiểm tra trên TKHQ; có ngày thì ghi nhận đạt. Nếu để trống hoặc file không rõ thì để Cần kiểm tra thêm cho đến khi MEIKO chốt NG hay REVIEW. Hồ sơ Đặt cọc/Trả trước không áp dụng tiêu chí này.",
        "steps": ["Ghép đúng TKHQ với khoản thanh toán.", "Tìm trường Ngày hoàn thành kiểm tra trên TKHQ.", "Không thay bằng Ngày hàng đến, ngày đăng ký tờ khai hoặc trạng thái thông quan.", "Nếu trường trống hoặc file không rõ, ghi rõ thực trạng để MEIKO chốt cách xử lý."],
        "questions": ["Trường ngày nhìn rõ nhưng để trống thì kết luận NG hay yêu cầu bổ sung?", "File mờ/thiếu trang khác gì với trường thực sự trống?"],
    },
    {
        "title": "Chữ ký, con dấu",
        "basis": "Hồ sơ trong nước: VAT có chữ ký điện tử của bên bán; PO có chữ ký và con dấu/xác nhận của cả bên mua và bên bán; INVOICE/PL có chữ ký điện tử hoặc chữ ký/con dấu. Hồ sơ nước ngoài: PO có chữ ký hoặc con dấu của cả hai bên; INVOICE/PL có chữ ký điện tử, dấu hoặc chữ ký. Dùng địa chỉ trên PO để phân biệt trong nước/nước ngoài.",
        "steps": ["Dùng địa chỉ trên PO để phân loại hồ sơ trong nước hoặc nước ngoài.", "Mở đủ trang ký/xác nhận của VAT/INVOICE, PO, COM và PL.", "Kiểm PO theo yêu cầu tương ứng với loại hồ sơ.", "Ghi rõ loại chữ ký/con dấu và bên ký khi INVOICE, COM hoặc PL chưa rõ."],
        "questions": ["COM, INVOICE và PL cần chữ ký, con dấu hay một trong hai?", "Cần của bên mua, bên bán hay cả hai?", "File mờ/thiếu trang ký thì REVIEW hay NG?"],
    },
]

# helper functions omitted? no

def display_text(value):
    return str(value).replace("DNTT", "ĐNTT")


def set_run(run, size=11, bold=None, color=None):
    run.font.name = "Times New Roman"
    for key in ("ascii", "hAnsi", "eastAsia"):
        run._element.rPr.rFonts.set(qn(f"w:{key}"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=11, bold=False, color=None, align=None, before=0, after=4):
    text = display_text(text)
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size, bold, color)
    return p


def add_bullet(doc, text):
    text = display_text(text)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, 10.5)
    return p


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)


def width(cell, cm):
    cell.width = Cm(cm)
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(cm * 567)))
    tcw.set(qn("w:type"), "dxa")


def margins(cell, value=90):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name in ["top", "start", "bottom", "end"]:
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_text(cell, text, size=10, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(display_text(text).split("\n")):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run(r, size, bold, color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def table(doc, headers, rows, widths, header_color=NAVY, font_size=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = t._tbl.tblPr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        t._tbl.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for i, h in enumerate(headers):
        set_text(t.rows[0].cells[i], h, font_size, True, WHITE)
        shade(t.rows[0].cells[i], header_color)
        width(t.rows[0].cells[i], widths[i])
        margins(t.rows[0].cells[i])
        t.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_properties = t.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    for row in rows:
        cells = t.add_row().cells
        row_properties = t.rows[-1]._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        for i, val in enumerate(row):
            set_text(cells[i], val, font_size)
            width(cells[i], widths[i])
            margins(cells[i])
    return t


def response_text():
    return "☐ Đúng   ☐ Chưa đúng   ☐ Cần kiểm tra thêm\nPhản hồi: ........................................\nPhiếu/file mẫu: ................................."


def current_rule_response():
    return "MEIKO chọn: ☐ Đúng   ☐ Chưa đúng   ☐ Cần kiểm tra thêm\nNếu chưa đúng/chưa rõ, ghi sửa tại đây:\n................................................"


def mapping_response():
    return "File thực tế MEIKO dùng: ....................\nTrường cần đọc: ..............................\nThông tin ghép/ưu tiên khi khác nhau: ........"


def manual_process_response():
    return "Thứ tự MEIKO đã làm (ghi 1 → 2 → 3): ........\n................................................\nNgoại lệ/bước bỏ qua: ........................"


def add_page_number(paragraph):
    run = paragraph.add_run("Trang ")
    set_run(run, 8.5, False, "7F7F7F")
    field_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.append(begin)
    field_run._r.append(instruction)
    field_run._r.append(end)
    set_run(field_run, 8.5, False, "7F7F7F")
    run = paragraph.add_run(" / 8  |  Không tự kết luận khi thiếu chứng từ hoặc cách xử lý chưa được MEIKO xác nhận")
    set_run(run, 8.5, False, "7F7F7F")


def setup_styles(doc):
    for name in ["Normal", "List Bullet"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"; h1.font.size = Pt(15); h1.font.bold = True; h1.font.color.rgb = RGBColor.from_string(NAVY)
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"; h2.font.size = Pt(12.5); h2.font.bold = True; h2.font.color.rgb = RGBColor.from_string(NAVY)
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def build_doc_legacy():
    doc = Document()
    setup_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(1.4)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.7)
    header = sec.header.paragraphs[0]
    header.text = ""
    run = header.add_run("MEIKO | Xác nhận cách đối chiếu NVL")
    set_run(run, 8.5, False, "7F7F7F")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = sec.footer.paragraphs[0]
    footer.text = ""
    run = footer.add_run("Không để AI tự đoán khi thiếu chứng từ hoặc rules chưa được chốt")
    set_run(run, 8.5, False, "7F7F7F")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, "KHUNG XÁC NHẬN QUY TẮC ĐỐI CHIẾU NGUYÊN VẬT LIỆU", size=18, bold=True, color=NAVY, after=8)
    add_bullet(doc, "Phạm vi: Bộ Đề nghị thanh toán Nguyên vật liệu; quy tắc đối chiếu hiện tại theo Bo_Rules.xlsx - sheet Hàng nhập kho.")
    add_bullet(doc, "Mục đích: MEIKO ghi rõ cách kế toán đang kiểm tra hồ sơ để ASOFT cấu hình lại mapping và rule engine, giảm nhầm chứng từ và giảm sai kết quả.")
    add_bullet(doc, "Lưu ý: các điểm chưa được MEIKO xác nhận chỉ ghi là Cần kiểm tra thêm; không tự coi là quy tắc đã chốt.")

    add_para(doc, "1. MEIKO cần làm gì trong tài liệu này?", style="Heading 1")
    add_para(doc, "Tài liệu này là mẫu để MEIKO ghi lại cách kế toán đang kiểm tra một bộ hồ sơ Nguyên vật liệu. Người đọc chỉ cần xác nhận: chứng từ nào cần mở, dữ liệu nào cần đọc, ghép chứng từ bằng thông tin nào, và kết luận OK/NG/REVIEW trong trường hợp nào.")
    for b in [
        "Với mỗi tiêu chí, MEIKO xác nhận chứng từ nào là căn cứ kiểm tra: ĐNTT, VAT/Invoice/Com Invoice, PO, Ringi, tờ khai, bảng kê...",
        "MEIKO xác nhận dữ liệu nào trên chứng từ cần đọc và so với dữ liệu nào trên mặt ĐNTT.",
        "MEIKO chốt cách kết luận: OK khi đủ căn cứ và khớp, NG khi đủ căn cứ nhưng sai, REVIEW khi thiếu chứng từ, mâu thuẫn hoặc rules chưa rõ.",
        "Nếu có ngoại lệ hoặc ví dụ thực tế, MEIKO ghi trực tiếp vào ô phản hồi để ASOFT cấu hình lại hệ thống.",
    ]:
        add_bullet(doc, b)
    table(doc, ["MEIKO cần chuẩn bị", "Cần ghi vào file Word", "Ví dụ cách ghi"], [
        ("Quan hệ ghép chứng từ", "Xác nhận dòng ĐNTT liên quan đến hóa đơn, tờ khai, PO/Ringi hoặc bảng kê nào; dùng dữ liệu nào để ghép và kiểm theo thứ tự nào.", "Dòng ĐNTT ↔ Hóa đơn A/B ↔ Bảng kê BK01 ↔ Tờ khai X; ghép theo số hóa đơn + NCC + lần thanh toán."),
        ("Quy tắc đối chiếu 9 tiêu chí", "Xác nhận trường dữ liệu trên ĐNTT, chứng từ cần dùng, cách so sánh và khi nào kết luận OK/NG/REVIEW.", "Số tiền: chỉ cộng các hóa đơn thuộc ĐNTT; không cộng hóa đơn khác dù nằm chung tờ khai/bảng kê."),
        ("Tình huống chưa rõ", "Ghi Cần kiểm tra thêm hoặc gửi mã ĐNTT/file mẫu.", "PO ghi USD, ĐNTT ghi VND: cần nêu có quy đổi không, lấy tỷ giá ở đâu."),
    ], [4.0, 7.7, 5.5])

    add_para(doc, "2. Cách hiểu kết quả OK NG REVIEW", style="Heading 1")
    table(doc, ["OK", "NG", "REVIEW / Cần kiểm tra thêm"], [
        ("Đủ chứng từ và dữ liệu khớp theo rule đã chốt.", "Đủ căn cứ xác định sai hoặc không đạt.", "Thiếu chứng từ, file mờ, dữ liệu mâu thuẫn hoặc rule nghiệp vụ chưa rõ."),
        ("Tiêu chí không áp dụng được bỏ qua và tính OK.", "Không dùng NG khi chỉ thiếu dữ liệu hoặc chưa có rule rõ.", "Dùng REVIEW khi cần MEIKO bổ sung hồ sơ hoặc xác nhận ngoại lệ."),
    ], [5.7, 5.7, 5.8])
    for cell in doc.tables[-1].rows[1].cells:
        shade(cell, GREEN)
    shade(doc.tables[-1].rows[1].cells[1], RED)
    shade(doc.tables[-1].rows[1].cells[2], YELLOW)

    add_para(doc, "3. Tên chứng từ dùng trong tài liệu", style="Heading 1")
    table(doc, ["Từ dùng trong tài liệu", "Kế toán hiểu đơn giản là", "MEIKO cần cung cấp gì"], DOC_TYPES, [4.2, 6.4, 6.6])

    add_para(doc, "4. Ma trận ghép chứng từ và trình tự đối chiếu MEIKO cần xác nhận", style="Heading 1")
    add_para(doc, "Phần này giúp MEIKO chỉ rõ dòng ĐNTT đang được chứng minh bởi chứng từ nào, trường dữ liệu nào, ghép bằng khóa nào và kiểm theo thứ tự nào. Nếu bước ghép chứng từ sai, kết quả đối chiếu có thể sai dù quy tắc nghiệp vụ đúng.")
    table(doc, ["Tiêu chí", "Trường ĐNTT cần lấy", "Chứng từ chính + trường cần đối chiếu", "Chứng từ hỗ trợ / kiểm chéo", "Khóa ghép & thứ tự MEIKO cần chốt"], MATRIX, [2.4, 3.5, 4.2, 3.3, 3.8], font_size=8.5)

    add_para(doc, "5. Phiếu xác nhận chi tiết từng tiêu chí", style="Heading 1")
    add_para(doc, "Mỗi bảng dưới đây là mẫu để MEIKO xác nhận hoặc chỉnh lại trình tự đối chiếu thực tế. Nếu nội dung đúng, đánh dấu Đúng. Nếu chưa đúng, sửa bảng ghép chứng từ, thứ tự bước và trả lời câu hỏi ngay trong ô phản hồi. Nếu chưa rõ, ghi Cần kiểm tra thêm hoặc gửi mã ĐNTT/file mẫu.")
    for idx, item in enumerate(CRITERIA, 1):
        add_para(doc, f"5.{idx}. {item['title']}", style="Heading 2", before=6)
        rows = [
            ("Căn cứ đang ghi nhận", item["basis"], item["status"]),
            ("Thứ tự kế toán kiểm", "\n".join(f"{i+1}) {s}" for i, s in enumerate(item["steps"])), "MEIKO xác nhận/chỉnh sửa"),
            ("Câu hỏi cần trả lời", "\n".join(f"- {q}" for q in item["questions"]), response_text()),
        ]
        table(doc, ["Nội dung", "ASOFT ghi nhận để MEIKO kiểm tra", "MEIKO xác nhận/chỉnh sửa"], rows, [3.3, 8.4, 5.5], font_size=8.8)
        for row in doc.tables[-1].rows[1:]:
            shade(row.cells[2], YELLOW if "Cần" in row.cells[2].text or "MEIKO" in row.cells[2].text or "☐" in row.cells[2].text else LIGHT_BLUE)

    add_para(doc, "6. Câu hỏi chung cần MEIKO chốt trước khi cấu hình", style="Heading 1")
    common_rows = [
        ("1", "Nhóm hồ sơ", "MEIKO phân loại hồ sơ NVL thành các nhóm nào trước khi đối chiếu: trả trước; công nợ không có bảng kê; công nợ có bảng kê; hay còn nhóm khác?", response_text()),
        ("2", "Chứng từ bắt buộc", "Với từng nhóm hồ sơ, chứng từ nào là bắt buộc và chứng từ nào chỉ dùng kiểm chéo?", response_text()),
        ("3", "Khóa ghép", "Có thể dùng bộ khóa Số hóa đơn + Nhà cung cấp + PO/Ringi + lần thanh toán để ghép dòng ĐNTT với file không?", response_text()),
        ("4", "File dùng chung", "Nếu một tờ khai dùng chung cho nhiều hóa đơn hoặc một file dùng cho nhiều dòng ĐNTT, hệ thống phải liên kết một-nhiều như thế nào?", response_text()),
        ("5", "Tên file sai", "Nếu người dùng đặt sai tên file nhưng nội dung bên trong đúng, hệ thống được dùng nội dung file để phân loại lại hay yêu cầu người dùng sửa tên/đính lại file?", response_text()),
    ]
    table(doc, ["STT", "Nhóm câu hỏi", "Nội dung cần MEIKO xác nhận", "MEIKO phản hồi"], common_rows, [1.1, 3.0, 8.2, 4.9], font_size=8.7)
    for row in doc.tables[-1].rows[1:]:
        shade(row.cells[3], YELLOW)

    add_para(doc, "7. Checklist kế toán tự kiểm trước khi gửi lại", style="Heading 1")
    table(doc, ["Câu hỏi tự kiểm", "Kết quả", "Ghi chú"], SELF_CHECKS, [7.0, 2.4, 7.8])
    for row in doc.tables[-1].rows[1:]:
        shade(row.cells[1], GREEN)

    add_para(doc, "Người phản hồi: ........................................    Ngày: ........................................", after=0)
    add_para(doc, "ASOFT sẽ chỉ cập nhật rule sau khi MEIKO ghi Đúng hoặc chỉnh sửa rõ nội dung.", bold=True, color=NAVY)
    return doc


def build_doc():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run("MEIKO | Xác nhận đối chiếu NVL")
    set_run(run, 8.5, False, "7F7F7F")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    add_page_number(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, "KHUNG XÁC NHẬN QUY TẮC ĐỐI CHIẾU NGUYÊN VẬT LIỆU", size=18, bold=True, color=NAVY, after=8)
    add_bullet(doc, "Phạm vi: Bộ Đề nghị thanh toán Nguyên vật liệu; trình bày cách đối chiếu, cách ghép chứng từ và quy trình kế toán đang thực hiện.")
    add_bullet(doc, "Mục đích: MEIKO xác nhận nội dung kiểm tra, chứng từ sử dụng và cách kế toán đã làm trước đây để nội dung đối chiếu được ghi nhận đúng.")
    add_bullet(doc, "Nguyên tắc: nội dung chưa được MEIKO xác nhận luôn để trạng thái Cần kiểm tra thêm; không tự coi là cách xử lý đã chốt.")

    add_para(doc, "1. Đọc và phản hồi theo 3 bước", style="Heading 1", size=13, bold=True)
    table(doc, ["Bước", "MEIKO đọc nội dung gì", "MEIKO cần làm"], [
        ("1. Xem cách đối chiếu", "Đọc bảng tóm tắt 9 tiêu chí, dữ liệu trên DNTT và chứng từ đang được sử dụng.", "Chọn Đúng / Chưa đúng / Cần kiểm tra thêm cho từng nội dung cần xác nhận."),
        ("2. Xác nhận cách làm thực tế", "Đọc file cần mở, cách ghép chứng từ và thứ tự kiểm đang được ghi nhận.", "Ghi lại file, thông tin cần đọc và các bước kế toán đã làm trước đây."),
        ("3. Trả lời điểm chưa rõ", "Đọc các trường hợp còn mâu thuẫn, thiếu cách xử lý hoặc cần ví dụ thực tế.", "Trả lời trực tiếp hoặc gửi mã DNTT, phiếu và file mẫu."),
    ], [4.2, 7.0, 6.0], font_size=9.5)

    add_para(doc, "2. Cách ghi câu trả lời", style="Heading 1", size=13, bold=True)
    table(doc, ["Đúng", "Chưa đúng", "Cần kiểm tra thêm"], [
        ("Nội dung đối chiếu, cách ghép chứng từ hoặc quy trình đang ghi nhận đúng với thực tế.", "Sửa trực tiếp nội dung, thứ tự bước hoặc cách ghép chứng từ.", "Chưa đủ căn cứ; ghi câu hỏi hoặc gửi phiếu/file mẫu."),
        ("Chỉ đánh dấu sau khi đã kiểm tra chứng từ thực tế.", "Nêu rõ nội dung đúng thay thế, không chỉ ghi chung chung.", "Không dùng nội dung chưa rõ làm cách kiểm chính thức."),
    ], [5.7, 5.7, 5.8], font_size=9.3)
    for cell in doc.tables[-1].rows[1].cells:
        shade(cell, GREEN)
    shade(doc.tables[-1].rows[1].cells[1], RED)
    shade(doc.tables[-1].rows[1].cells[2], YELLOW)
    add_para(doc, "Cách hiểu kết quả đối chiếu: OK = đủ chứng từ và khớp; NG = đủ căn cứ xác định không đạt; REVIEW = thiếu, mờ, mâu thuẫn hoặc chưa có cách xử lý được MEIKO xác nhận.", size=8.5, after=1)

    add_para(doc, "Quy ước hiện tại cho hồ sơ Đặt cọc/Trả trước", style="Heading 2", size=11, bold=False, color="000000", before=5, after=2)
    table(doc, ["Phạm vi chứng từ", "Cách đối chiếu hiện tại"], [
        ("DNTT; PO; VAT/INVOICE; RINGI nếu có", "Không kiểm tra số hóa đơn, ngày hóa đơn và Ngày hoàn thành kiểm tra. Bỏ qua thông tin từ COM và TKHQ; các tiêu chí còn lại kiểm tra như hồ sơ thông thường."),
    ], [5.2, 12.0], font_size=8.3)

    add_para(doc, "Tên chứng từ sử dụng thống nhất", style="Heading 2", size=11, bold=False, color="000000", before=5, after=2)
    table(doc, ["Viết tắt", "Tên chuẩn", "Viết tắt", "Tên chuẩn"], WORD_DOCUMENT_TYPES, [1.8, 7.2, 2.1, 6.1], font_size=8.0)
    add_para(doc, "Từ trang 2 trở đi, các viết tắt được dùng theo bảng này. INVOICE và COM đều liên quan hóa đơn thương mại; MEIKO cần ghi rõ nếu thực tế đang dùng cùng một file hoặc dùng khác nhau theo từng bộ hồ sơ.", size=8.5, after=0)

    doc.add_page_break()
    add_para(doc, "3. Cách đối chiếu hiện tại", style="Heading 1", size=13, bold=True)
    add_para(doc, "Bảng dưới đây trình bày cách đối chiếu hiện tại bằng tên chứng từ đã chuẩn hóa. Toàn bộ 9 tiêu chí vẫn ở trạng thái Cần kiểm tra thêm cho đến khi MEIKO phản hồi.")
    rule_rows = [
        (item[0], item[1], item[2], SUMMARY_RULES[index])
        for index, item in enumerate(WORD_MATRIX)
    ]
    table(doc, ["Tiêu chí", "Dữ liệu DNTT", "File/chứng từ đối chiếu", "Cách kiểm hiện tại"], rule_rows, [2.5, 2.9, 4.8, 7.0], font_size=8.2)

    def add_criterion(criterion_index, heading_style="Heading 2"):
        item = WORD_CRITERIA[criterion_index - 1]
        mapping = WORD_MATRIX[criterion_index - 1]
        add_para(doc, f"4.{criterion_index}. {item['title']}", style=heading_style, size=11, bold=True, color="000000", before=4, after=3)
        rows = [
            ("Cách kiểm hiện tại", item["basis"], current_rule_response()),
            ("File cần mở và cách ghép", f"Trên DNTT cần đọc: {mapping[1]}\nFile cần mở: {mapping[2]}\nCách kiểm hiện tại: {mapping[3]}\nMEIKO cần chốt: {mapping[4]}", mapping_response()),
            ("Quy trình đối chiếu thủ công trước đây", "\n".join(f"{step_number + 1}) {step}" for step_number, step in enumerate(item["steps"])), manual_process_response()),
            ("Câu hỏi nghiệp vụ cần phản hồi", "\n".join(f"- {question}" for question in item["questions"]), response_text()),
        ]
        table(doc, ["Nội dung cần xác nhận", "Nội dung đang ghi nhận để MEIKO kiểm tra", "MEIKO phản hồi"], rows, [3.4, 8.3, 5.5], font_size=8.1)
        for row in doc.tables[-1].rows[1:]:
            shade(row.cells[2], YELLOW)

    detail_groups = [(1, 2), (3, 4), (5, 6), (7, 8)]
    for group_number, group in enumerate(detail_groups, start=1):
        doc.add_page_break()
        add_para(doc, f"4. Xác nhận cách kiểm, file ghép và quy trình thủ công - Nhóm {group_number}/5", style="Heading 1", size=13, bold=True, after=2)
        if group_number == 1:
            add_para(doc, "MEIKO chọn Đúng nếu phù hợp; nếu Chưa đúng hoặc Cần kiểm tra thêm, sửa trực tiếp file dùng, thứ tự kiểm và câu trả lời nghiệp vụ.", size=9.5, after=3)
        for criterion_index in group:
            add_criterion(criterion_index)

    doc.add_page_break()
    add_para(doc, "4. Xác nhận cách kiểm, file ghép và quy trình thủ công - Nhóm 5/5", style="Heading 1", size=13, bold=True, after=2)
    add_criterion(9)

    add_para(doc, "5. Cách ghép chứng từ dùng chung cần MEIKO mô tả", style="Heading 1", size=13, bold=True, before=5, after=2)
    shared_mapping_rows = [
        ("Một TKHQ - nhiều hóa đơn", "Kế toán xác định từng hóa đơn/dòng DNTT thuộc TKHQ nào bằng thông tin nào?", "☐ Đúng   ☐ Chưa đúng   ☐ Cần kiểm tra thêm\nThông tin dùng để ghép + file mẫu: ............"),
        ("Một STATEMENT - nhiều hóa đơn", "Dùng số STATEMENT, số hóa đơn hay thứ tự dòng để chọn đúng INVOICE?", "☐ Đúng   ☐ Chưa đúng   ☐ Cần kiểm tra thêm\nCách làm cũ + file mẫu: ......................"),
        ("Một hóa đơn - nhiều lần thanh toán", "Phân biệt lần thanh toán bằng số tiền, PO/RINGI hay thông tin nào?", "☐ Đúng   ☐ Chưa đúng   ☐ Cần kiểm tra thêm\nThông tin kế toán dùng để ghép: .............."),
        ("Tên file sai nhưng nội dung đúng", "Kế toán trước đây dùng nội dung file hay yêu cầu đặt/đính lại file?", response_text()),
    ]
    table(doc, ["Tình huống ghép chứng từ", "Nội dung cần MEIKO xác nhận", "MEIKO phản hồi"], shared_mapping_rows, [4.4, 7.4, 5.4], font_size=8.2)
    for row in doc.tables[-1].rows[1:]:
        shade(row.cells[2], YELLOW)

    doc.add_page_break()
    add_para(doc, "6. Câu hỏi nghiệp vụ cần phản hồi trước khi chốt cách đối chiếu", style="Heading 1", size=13, bold=True)
    common_rows = [
        ("1", "Nhóm hồ sơ", "MEIKO phân loại hồ sơ NVL thành các nhóm nào: trả trước; công nợ không có STATEMENT; công nợ có STATEMENT; hay nhóm khác?", response_text()),
        ("2", "Chứng từ bắt buộc", "Với từng nhóm hồ sơ, chứng từ nào bắt buộc và chứng từ nào chỉ dùng kiểm chéo?", response_text()),
        ("3", "Ưu tiên dữ liệu", "Khi cùng một thông tin khác nhau giữa DNTT, VAT/INVOICE, PO, TKHQ, COM, PL hoặc STATEMENT thì ưu tiên nguồn nào?", response_text()),
        ("4", "Thiếu/mờ chứng từ", "Trường hợp nào kết luận NG, trường hợp nào REVIEW hoặc yêu cầu bổ sung file?", response_text()),
        ("5", "Khác loại tiền", "Có cho phép quy đổi không; dùng tỷ giá từ chứng từ nào và ngày nào?", response_text()),
        ("6", "File dùng chung", "Một file dùng cho nhiều dòng DNTT: kế toán xác định mỗi dòng thuộc file nào bằng các thông tin nào?", response_text()),
        ("7", "Kết quả đối chiếu", "MEIKO chốt rõ điều kiện OK, NG và REVIEW cho các điểm chưa có cách xử lý cụ thể.", response_text()),
    ]
    table(doc, ["STT", "Nhóm câu hỏi", "Câu hỏi nghiệp vụ", "MEIKO phản hồi"], common_rows, [1.0, 2.9, 7.8, 5.5], font_size=8.1)
    for row in doc.tables[-1].rows[1:]:
        shade(row.cells[3], YELLOW)

    add_para(doc, "Hồ sơ mẫu cần gửi kèm khi phản hồi", style="Heading 2", size=11, bold=True, color="000000", before=5, after=2)
    add_para(doc, "Với mỗi bộ hồ sơ, MEIKO ghi mã DNTT; chứng từ đã dùng; thông tin dùng để ghép; chứng từ được ưu tiên khi số liệu khác nhau; và kết luận mong muốn cho từng tiêu chí. Nếu không phát sinh nhóm nào, ghi rõ Không phát sinh.", size=8.5, after=2)
    sample_rows = [
        ("Đặt cọc/Trả trước", "01 bộ đã xử lý đúng; 01 bộ có ngoại lệ nếu có."),
        ("Hồ sơ trong nước", "01 bộ đã xử lý đúng; 01 bộ có ngoại lệ nếu có."),
        ("Hồ sơ nước ngoài không có STATEMENT", "01 bộ đã xử lý đúng; 01 bộ có ngoại lệ nếu có."),
        ("Hồ sơ nước ngoài có STATEMENT", "01 bộ đã xử lý đúng; 01 bộ có ngoại lệ nếu có."),
    ]
    table(doc, ["Nhóm hồ sơ cần có mẫu", "Mẫu MEIKO cần gửi"], sample_rows, [9.0, 8.2], font_size=8.3)
    add_para(doc, "Trước khi gửi lại: đã chọn trạng thái cho đủ 9 tiêu chí; các mục Chưa đúng/Cần kiểm tra thêm đã ghi nội dung cần sửa; mỗi file mẫu đã ghi rõ mã DNTT và nhóm hồ sơ.", size=8.3, bold=True, after=1)

    add_para(doc, "Người phản hồi: ........................................    Bộ phận: ........................    Ngày: ........................", before=5, after=1)
    add_para(doc, "Nội dung đối chiếu và cách ghép chứng từ chỉ được cập nhật sau khi MEIKO ghi Đúng hoặc chỉnh sửa rõ nội dung.", bold=True, color=NAVY, after=0)
    return doc


def remove_media_if_any(path: Path):
    with ZipFile(path) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    if media:
        raise RuntimeError(f"Unexpected embedded media in {path}: {media}")


def main():
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    doc = build_doc()
    for target in TARGETS:
        if target.exists():
            backup = target.with_name(f"{target.name}.bak_{stamp}_before_v1_format_rebuild")
            shutil.copy2(target, backup)
            print(f"BACKUP={backup}")
        doc.save(target)
        remove_media_if_any(target)
        print(f"SAVED={target}")


if __name__ == "__main__":
    main()
